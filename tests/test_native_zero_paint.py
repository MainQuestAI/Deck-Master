"""Compile approved SVG paint boundaries into actual DrawingML, without render mocks."""
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
import zipfile
import pytest
sys.path[:0] = [str(Path(__file__).resolve().parents[1] / 'scripts'), str(Path(__file__).resolve().parent)]
from test_high_density_builder_v2 import _prepared_fixture
from native_pptx.pptx import compile_pptx, pptx_path
from high_density.svg import svg_path, validate_approved_svg


@pytest.mark.parametrize('opacity,stop_alpha', [('0', '1'), ('1', '0'), ('1', '0.5')])
def test_zero_and_normal_gradient_alpha_survive_real_compile(tmp_path, opacity, stop_alpha):
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, 'P001')
    doc = ET.fromstring(svg.read_text())
    ns = '{http://www.w3.org/2000/svg}'
    defs = ET.SubElement(doc, ns + 'defs')
    gradient = ET.SubElement(defs, ns + 'linearGradient', {'id': 'boundary.gradient', 'x1':'0%', 'y1':'0%', 'x2':'100%', 'y2':'0%'})
    ET.SubElement(gradient, ns+'stop', {'offset':'0%', 'stop-color':'#123456', 'stop-opacity':stop_alpha})
    ET.SubElement(gradient, ns+'stop', {'offset':'100%', 'stop-color':'#abcdef', 'stop-opacity':stop_alpha})
    shape = next(e for e in doc.iter() if e.get('id') == 'block.03')
    shape.set('fill', 'url(#boundary.gradient)')
    shape.set('opacity', opacity)
    svg.write_text(ET.tostring(doc, encoding='unicode'))
    validate_approved_svg(svg, scene, lock)
    compile_pptx(run, [scene], {'P001':lock}, validate_approved=validate_approved_svg)
    with zipfile.ZipFile(pptx_path(run)) as z:
        slide = ET.fromstring(z.read('ppt/slides/slide1.xml'))
    alpha = [int(n.get('val')) for n in slide.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill//{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')]
    expected = round(float(opacity)*float(stop_alpha)*100000)
    assert len(alpha) == 2 and alpha == [expected, expected]


@pytest.mark.parametrize('attrs,expected_alpha,zero_stroke', [({'opacity':0},0,False), ({'fill_opacity':0},0,False), ({'fill_opacity':0.5},50000,False), ({'stroke_width':0},100000,True)])
def test_solid_paint_zero_values_are_serialized(attrs, expected_alpha, zero_stroke):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from native_pptx.pptx import _set_shape_fill
    p=Presentation(); s=p.slides.add_slide(p.slide_layouts[6]).shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,100000,100000)
    style={'fill':'#123456', 'stroke':'#abcdef', 'stroke_width':attrs.get('stroke_width',1)}
    paint={'fill':{'kind':'solid','color':'#123456'},'stroke':{'kind':'solid','color':'#abcdef'},'opacity':attrs.get('opacity',1),'fill_opacity':attrs.get('fill_opacity',1),'stroke_opacity':1}
    _set_shape_fill(s,style,paint)
    alpha=s._element.xpath('./p:spPr/a:solidFill/a:srgbClr/a:alpha')
    assert len(alpha)==1 and int(alpha[0].get('val'))==expected_alpha
    if zero_stroke: assert s._element.xpath('./p:spPr/a:ln/a:noFill')


@pytest.mark.parametrize('value', ['0', '0.00', '0e0'])
def test_zero_text_alpha_never_satisfies_content_lock(tmp_path, value):
    from high_density.svg import SvgVisualError
    run,lock,scene = _prepared_fixture(tmp_path)
    svg=svg_path(run,'P001'); doc=ET.fromstring(svg.read_text())
    title=next(e for e in doc.iter() if e.get('id')=='title.main')
    title.set('fill-opacity',value)
    svg.write_text(ET.tostring(doc,encoding='unicode'))
    with pytest.raises(SvgVisualError,match='hidden.*SVG text'):
        validate_approved_svg(svg,scene,lock)
