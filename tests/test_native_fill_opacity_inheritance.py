"""Actual SVG renderer and approved-content validator agree on inherited paint."""
import sys
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET
import pytest
from PIL import Image
sys.path[:0]=[str(Path(__file__).resolve().parents[1]/'scripts'),str(Path(__file__).parent)]
from high_density.svg import validate_approved_svg,SvgVisualError,svg_path
from test_high_density_builder_v2 import _prepared_fixture

@pytest.fixture(scope='module')
def sample(tmp_path_factory):
    root,lock,scene=_prepared_fixture(tmp_path_factory.mktemp('opacity_inheritance'))
    return root,lock,scene,svg_path(root,'P001').read_text()

@pytest.mark.parametrize('scope',['group','tspan'])
def test_explicit_child_fill_alpha_overrides_inherited_zero(sample,scope):
    root,lock,scene,original=sample
    doc=ET.fromstring(original)
    title=next(n for n in doc.iter() if n.get('id')=='title.main')
    if scope=='group':
        group=next(n for n in doc.iter() if n.get('id')=='page.P001');group.set('fill-opacity','0')
        for n in doc.iter():
            if n.tag.rsplit('}',1)[-1] in {'text','rect','circle','ellipse','path','polygon','polyline'}:n.set('fill-opacity','1')
    else:
        title.set('fill-opacity','0')
        text=''.join(title.itertext());title.text=None
        for child in list(title):title.remove(child)
        ET.SubElement(title,'{http://www.w3.org/2000/svg}tspan',{'fill-opacity':'1'}).text=text
    mutated=root/f'{scope}-inheritance.svg';ET.ElementTree(doc).write(mutated,encoding='utf-8')
    # Compare the actual rendered title against the original visible title;
    # background shapes may inherit zero, but this crop contains text only.
    source=root/f'{scope}-original.svg';source.write_text(original)
    pngs=[]
    for svg in [source,mutated]:
        png=svg.with_suffix('.png');subprocess.run(['rsvg-convert','-w','1672','-h','941','-o',str(png),str(svg)],check=True);pngs.append(Image.open(png).convert('RGBA'))
    x,y,w,h=[float(v) for v in title.get('data-pptx-bounds').split(',')]
    box=(int(x),int(y),int(x+w),int(y+h))
    # Composite both on the same white canvas so transparent backgrounds do
    # not obscure the check that actual glyph pixels remain nonempty.
    def dark_pixels(image):
        white=Image.new('RGBA',image.size,'white');white.alpha_composite(image)
        return sum(value < 100 for value in white.convert('L').crop(box).tobytes())
    reference_dark=dark_pixels(pngs[0]);assert reference_dark>100
    assert dark_pixels(pngs[1])>=reference_dark*0.95
    validate_approved_svg(mutated,scene,lock)
    from native_pptx.pptx import compile_pptx
    from pptx import Presentation
    pptx,_=compile_pptx(root,[scene],{'P001':lock},validate_approved=validate_approved_svg,
                       svg_paths={'P001':mutated},output_root=root/f'{scope}-compiled',canvas_mode='native')
    shape=next(s for s in Presentation(pptx).slides[0].shapes if s.name=='title.main')
    alphas=shape._element.xpath('.//a:rPr/a:solidFill/a:srgbClr/a:alpha')
    assert alphas and all(int(a.get('val'))==100000 for a in alphas)


def test_group_and_child_opacity_still_compose_below_visibility_threshold(sample):
    root,lock,scene,original=sample;doc=ET.fromstring(original)
    group=next(n for n in doc.iter() if n.get('id')=='page.P001');group.set('opacity','0.1')
    title=next(n for n in doc.iter() if n.get('id')=='title.main');title.set('opacity','0.1')
    svg=root/'composited-opacity.svg';ET.ElementTree(doc).write(svg,encoding='utf-8')
    png=svg.with_suffix('.png');subprocess.run(['rsvg-convert','-w','1672','-h','941','-o',str(png),str(svg)],check=True)
    rendered=Image.open(png).convert('RGBA');white=Image.new('RGBA',rendered.size,'white');white.alpha_composite(rendered)
    x,y,w,h=[float(v) for v in title.get('data-pptx-bounds').split(',')]
    assert not any(value<100 for value in white.convert('L').crop((int(x),int(y),int(x+w),int(y+h))).tobytes())
    with pytest.raises(SvgVisualError,match='hidden.*SVG text'):
        validate_approved_svg(svg,scene,lock)


def test_parent_fill_zero_cannot_hide_required_tail_text(sample):
    root,lock,scene,original=sample;doc=ET.fromstring(original)
    title=next(n for n in doc.iter() if n.get('id')=='title.main')
    content=''.join(title.itertext());title.set('fill-opacity','0');title.text=None
    for child in list(title):title.remove(child)
    span=ET.SubElement(title,'{http://www.w3.org/2000/svg}tspan',{'fill-opacity':'1'})
    span.text=content[:len(content)//2];span.tail=content[len(content)//2:]
    svg=root/'hidden-tail.svg';ET.ElementTree(doc).write(svg,encoding='utf-8')
    with pytest.raises(SvgVisualError,match='hidden.*SVG text'):
        validate_approved_svg(svg,scene,lock)
