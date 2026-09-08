"""Actual P006 pattern: absolute same-baseline tspan columns remain separated."""
from xml.etree import ElementTree as ET
import pytest
from pptx import Presentation
from native_pptx.pptx import _svg_text_lines, _add_text, _NATIVE_CANVAS

def emit(*, reverse=False):
    node=ET.fromstring('<text id="P006.e006" x="100" y="352" font-size="40" font-family="Noto Sans SC" fill="#092442" data-pptx-text="库存准确性｜盘点一致数量与样本口径"><tspan x="140" font-weight="700">库存准确性</tspan><tspan x="562">｜</tspan><tspan x="636">盘点一致数量与样本口径</tspan></text>')
    if reverse:
        list(node)[1].set('x','120')
    element={'element_id':'P006.e006','kind':'text','bbox':{'x':100,'y':312,'w':1460,'h':70},'text':node.get('data-pptx-text'),'style':{},'_text_lines':_svg_text_lines(node,{'gradients':{},'effects':{}})}
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);token=_NATIVE_CANVAS.set(True)
    trace=[]
    try:_add_text(slide,element,trace)
    finally:_NATIVE_CANVAS.reset(token)
    return slide.shapes[0], trace

def test_explicit_same_baseline_x_is_not_serialized_as_contiguous_text():
    shape,trace=emit()
    assert [s.text for s in shape.shapes]==['库存准确性','｜','盘点一致数量与样本口径']
    margins=[s.left/12700 for s in shape.shapes]
    assert margins==pytest.approx([140*960/1672,562*960/1672,636*960/1672],abs=.001)

    assert trace[0]['element_id']=='P006.e006'
    assert ''.join(c['text'] for c in trace[0]['children'])=='库存准确性｜盘点一致数量与样本口径'
    assert len(set(c['element_id'] for c in trace[0]['children']))==3


def test_reverse_absolute_x_keeps_declared_position_instead_of_tab_reflow():
    shape,trace=emit(reverse=True)
    assert shape.shapes[1].left/12700==pytest.approx(120*960/1672,abs=.001)
    assert shape.shapes[0].left > shape.shapes[1].left
