from __future__ import annotations

import json
import shutil
import sys
import threading
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pytest
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from build.manifest import build_manifest_v2
from high_density.blueprint import BlueprintInvalid, _default_slide_frame, build_blueprint_prompt, build_blueprint_prompt_artifact, ensure_blueprint_manifest, record_provider_host_result
from high_density.blueprint_content_review import BlueprintContentReviewRequired, archive_rejected_blueprint, load_blueprint_content_review, next_attempt_index, write_blueprint_content_review
from high_density.capability import REQUIRED_SCHEMAS, inspect_high_density_capability
from high_density.content import (
    build_content_lock,
    build_mbb_page,
    build_mbb_plan,
    enrich_selected_mbb_plan,
    load_content_lock,
    load_mbb_plan,
    load_page_packages,
    seal_mbb_plan,
    select_mbb_storyline,
    write_mbb_plan,
)
from high_density.contracts import ContractError, assert_valid, read_json, sha256_file, sha256_json, write_json
from high_density.engine import (
    build_high_density_status,
    prepare_high_density,
    retry_high_density,
    run_high_density,
    watch_high_density_status,
)
from high_density.migration import MIGRATION_REQUIRED_CODE, assert_current_mbb_artifact, retired_method_token
from high_density.pptx import PptxEditabilityError, _render_pptx_page, compile_pptx, pptx_path, readback_pptx, trace_path
from high_density.scene import _fixture_background, build_fixture_scene, load_scene, validate_scene_content, write_scene
from high_density.style import write_style_lock
from high_density.svg import SvgVisualError, _font_path, compile_svg, load_visual_review, main_review_receipt_path, preview_path, render_preview, review_path, svg_path, validate_approved_svg, validate_svg
from high_density.visibility import build_visibility_policy, visible_text_violation
from production.page_package import PageContent, PagePackageIndex, build_page_package
from runtime.run_state import create_run


FIXTURE_DIR = ROOT / "tests" / "fixtures" / "high_density"
FIXTURE = json.loads((FIXTURE_DIR / "fixture.json").read_text(encoding="utf-8"))


def _package(run_id: str, page: dict) -> dict:
    return build_page_package(
        run_id=run_id,
        content=PageContent(
            page_id=str(page["page_id"]),
            order=int(page["order"]),
            title=str(page["title"]),
            subtitle=str(page["subtitle"]),
            body_blocks=list(page["body_blocks"]),
            labels=list(page["labels"]),
            footnotes=list(page["footnotes"]),
            speaker_notes=str(page["speaker_notes"]),
            claim_bindings=list(page["claim_bindings"]),
            evidence_bindings=list(page["evidence_bindings"]),
            visual_spec={"page_type": page["page_class"]},
        ),
        status="ready_for_build",
        now=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )


def _make_run(tmp_path: Path, *, mode: str = "fixture", page_count: int = 1, project_name: str = "v2 fixture") -> tuple[Path, PagePackageIndex]:
    run = create_run(tmp_path / "runs", {"project_name": project_name, "run_mode": mode}, run_id="hd-v2-test", force=True)
    index = PagePackageIndex(run)
    for page in FIXTURE["pages"][:page_count]:
        index.write(_package(str(run.name), page))
    return run, index


def _blueprint(run: Path, page_id: str = "P001", *, fill: str = "#f7f9fb") -> Path:
    path = run / "high_density_build" / "blueprints" / f"{page_id}.svg"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text((FIXTURE_DIR / "blueprint.svg").read_text(encoding="utf-8").replace("#f7f9fb", fill), encoding="utf-8")
    return path


def _prepared_fixture(tmp_path: Path) -> tuple[Path, dict, dict]:
    run, _ = _make_run(tmp_path)
    _blueprint(run)
    prepare_high_density(run)
    result = run_high_density(run)
    assert result["status"] == "completed"
    lock = read_json(run / "high_density_build/content_locks/P001.json")
    scene = load_scene(run, "P001")
    return run, lock, scene


def _approved_page_plan(package: dict) -> tuple[dict, dict]:
    plan = build_mbb_plan([package], run_id=str(package["run_id"]))
    plan["selection"].update(
        {
            "status": "selected_pending_enrichment",
            "selected_storyline_id": "storyline.decision",
            "selected_by": "test",
            "selected_at": "2026-01-01T00:00:00+00:00",
        }
    )
    plan["storyline_audit"].update({"status": "selected_pending_enrichment", "selected_id": "storyline.decision"})
    plan = enrich_selected_mbb_plan(plan, [package])
    return plan, plan["pages"][0]


def test_retired_content_plan_requests_deck_rebuild_without_deleting_artifact(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="fixture")
    retired_dir = run / "high_density_build" / retired_method_token()
    retired_dir.mkdir(parents=True)
    marker = retired_dir / "plan.json"
    marker.write_text("{}\n", encoding="utf-8")

    result = run_high_density(run)

    assert result["status"] == "awaiting_agent_build"
    assert result["current_stage"] == "content_lock"
    assert result["next_action"]["kind"] == "agent_mbb_candidates"
    assert result["next_action"]["rebuild_scope"] == "deck"
    assert MIGRATION_REQUIRED_CODE in result["next_action"]["reason"]
    assert marker.is_file()
    assert not (run / "high_density_build/mbb/mbb_plan.json").exists()


def test_retired_schema_in_current_artifact_requires_mbb_regeneration(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="fixture")
    retired_schema = "deck_" + retired_method_token() + "_plan.v1"
    write_json(
        run / "high_density_build/mbb/mbb_plan.json",
        {"schema_version": retired_schema, "run_id": run.name},
    )

    with pytest.raises(ContractError, match=MIGRATION_REQUIRED_CODE):
        load_mbb_plan(run, expected_run_id=run.name, require_approved=False)


def test_retired_content_lock_lineage_requires_mbb_regeneration(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="fixture")
    retired_hash_field = retired_method_token() + "_plan_sha256"
    write_json(
        run / "high_density_build/content_locks/P001.content_lock.json",
        {"schema_version": "deck_content_lock.v2", "page_id": "P001", retired_hash_field: "a" * 64},
    )

    with pytest.raises(ContractError, match=MIGRATION_REQUIRED_CODE):
        load_content_lock(run, "P001", expected_run_id=run.name)


@pytest.mark.parametrize(
    ("relative_path", "payload"),
    [
        (
            "high_density_build/mbb/selection_receipt.json",
            {"schema_version": "deck_" + retired_method_token() + "_selection_receipt.v1"},
        ),
        (
            "high_density_build/mbb/runtime_seal.json",
            {"schema_version": "deck_" + retired_method_token() + "_runtime_seal.v1"},
        ),
        (
            "high_density_build/prompts/P001.blueprint_prompt.json",
            {retired_method_token() + "_plan_sha256": "a" * 64},
        ),
        (
            "high_density_build/blueprints/P001.blueprint_manifest.json",
            {retired_method_token() + "_plan_sha256": "a" * 64},
        ),
    ],
)
def test_retired_receipt_prompt_and_manifest_require_mbb_regeneration(
    tmp_path: Path,
    relative_path: str,
    payload: dict[str, str],
) -> None:
    run, _ = _make_run(tmp_path, mode="fixture")
    write_json(run / relative_path, payload)

    with pytest.raises(ContractError, match=MIGRATION_REQUIRED_CODE):
        assert_current_mbb_artifact(run)


def test_active_high_density_surface_has_no_retired_content_plan_terminology() -> None:
    active_roots = [
        ROOT / "scripts" / "high_density",
        ROOT / "scripts" / "deck_master.py",
        ROOT / "scripts" / "skills",
        ROOT / "docs" / "contracts",
        ROOT / "skills" / "deck-builder-high-density",
        ROOT / "skills" / "manifest.json",
        ROOT / "tests",
    ]
    retired = retired_method_token()
    forbidden = (retired, retired.upper(), retired + "_", "/" + retired + "/")
    offenders: list[str] = []
    for active_root in active_roots:
        paths = [active_root] if active_root.is_file() else active_root.rglob("*")
        for path in paths:
            if not path.is_file() or path.suffix not in {".py", ".json", ".md"}:
                continue
            text = path.read_text(encoding="utf-8")
            if any(token in text for token in forbidden):
                offenders.append(str(path.relative_to(ROOT)))
    assert offenders == []


def _write_approved_mbb_plan(run: Path) -> dict:
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    return seal_mbb_plan(run)


def _write_selected_enriched_mbb_plan(run: Path, packages: list[dict]) -> dict:
    plan = build_mbb_plan(packages, run_id=run.name)
    write_mbb_plan(run, plan)
    plan = select_mbb_storyline(run, "storyline.decision", selected_by="test")
    plan = enrich_selected_mbb_plan(plan, packages)
    write_mbb_plan(run, plan)
    return plan


def _approve_blueprint(run: Path, page_id: str = "P001") -> None:
    lock = read_json(run / f"high_density_build/content_locks/{page_id}.json")
    style_lock = read_json(run / "high_density_build/style/style_lock.json")
    mbb_plan = read_json(run / "high_density_build/mbb/mbb_plan.json")
    write_blueprint_content_review(run, page_id, lock, findings=[], reviewer_id="test", action_id="test-content-review")
    ensure_blueprint_manifest(
        run,
        page_id,
        lock,
        style_lock=style_lock,
        mbb_plan_sha256=str(mbb_plan["mbb_plan_sha256"]),
        approval={
            "status": "approved",
            "source": "explicit_user",
            "approved_by": "test",
            "approved_at": "2026-01-01T00:00:00+00:00",
        },
    )


def _rehash_page_and_plan(plan: dict, page: dict) -> None:
    page["page_plan_sha256"] = sha256_json(
        {key: value for key, value in page.items() if key not in {"page_plan_sha256", "created_at", "updated_at"}}
    )
    plan["mbb_plan_sha256"] = sha256_json(
        {key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}}
    )


def test_prepare_rejects_malformed_page_package(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    (run / "page_packages/P001.json").write_text("{broken", encoding="utf-8")

    with pytest.raises(ValueError, match="invalid page package"):
        prepare_high_density(run)
    assert build_high_density_status(run)["status"] == "blocked"


def test_prepare_rejects_non_object_page_package_index(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    (run / "page_packages/index.json").write_text("[]", encoding="utf-8")

    with pytest.raises(ValueError, match="index must be an object"):
        prepare_high_density(run)


def test_contract_rejects_non_string_schema_version() -> None:
    with pytest.raises(ContractError):
        assert_valid("page_package", {"schema_version": 1})


def test_prepare_rejects_cross_run_page_package(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    package["run_id"] = "another-run"
    (run / "page_packages/P001.json").write_text(json.dumps(package), encoding="utf-8")

    with pytest.raises(ValueError, match="run_id mismatch"):
        prepare_high_density(run)


def test_retry_rejects_unknown_or_unsafe_page_id(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    prepare_high_density(run)
    outside = tmp_path / "outside.txt"
    outside.write_text("untouched", encoding="utf-8")

    with pytest.raises(ValueError, match="not registered"):
        retry_high_density(run, page_id="../outside", stage="svg")
    assert outside.read_text(encoding="utf-8") == "untouched"


def test_watch_waits_until_stable_end_state(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    status_path = run / "high_density_build/status.json"
    write_json(
        status_path,
        {
            "schema_version": "deck_high_density_status.v2",
            "run_id": run.name,
            "builder_profile": "high_density",
            "status": "building",
            "current_stage": "blueprint",
            "next_action": {"kind": "agent_imagegen"},
            "updated_at": "2026-01-01T00:00:00+00:00",
        },
    )

    def finish() -> None:
        time.sleep(0.05)
        write_json(
            status_path,
            {
                "schema_version": "deck_high_density_status.v2",
                "run_id": run.name,
                "builder_profile": "high_density",
                "status": "completed",
                "current_stage": "handback",
                "next_action": {"kind": "none"},
                "updated_at": "2026-01-01T00:00:01+00:00",
            },
        )

    worker = threading.Thread(target=finish)
    worker.start()
    watched = watch_high_density_status(run, timeout_seconds=1, poll_seconds=0.01)
    worker.join()
    assert watched["status"] == "completed"
    assert watched["watch"]["timed_out"] is False
    assert len(watched["watch"]["events"]) >= 2


def test_watch_returns_immediately_for_actionable_agent_state(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    status_path = run / "high_density_build/status.json"
    write_json(
        status_path,
        {
            "schema_version": "deck_high_density_status.v2",
            "run_id": run.name,
            "builder_profile": "high_density",
            "status": "awaiting_agent_build",
            "current_stage": "blueprint",
            "next_action": {"kind": "agent_imagegen"},
            "updated_at": "2026-01-01T00:00:00+00:00",
        },
    )

    started = time.monotonic()
    watched = watch_high_density_status(run, timeout_seconds=30, poll_seconds=1)

    assert watched["status"] == "awaiting_agent_build"
    assert watched["watch"]["timed_out"] is False
    assert time.monotonic() - started < 1


def test_cli_high_density_runtime_exposes_watch_action() -> None:
    from deck_master import _high_density_runtime

    assert callable(_high_density_runtime()["watch"])


def test_capability_blocks_missing_visual_dependency(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import high_density.capability as capability

    original_which = capability.shutil.which
    monkeypatch.setattr(capability.shutil, "which", lambda name: "" if name == "rsvg-convert" else original_which(name))
    report = inspect_high_density_capability(ROOT)
    assert report["ready"] is False
    assert report["status"] == "blocked_runtime_dependency"
    svg_check = next(item for item in report["checks"] if item["name"] == "renderers")
    assert svg_check["ready"] is False


def test_capability_accepts_release_contract_directory(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import high_density.capability as capability

    contract_dir = tmp_path / "contracts"
    contract_dir.mkdir()
    for name in REQUIRED_SCHEMAS:
        shutil.copy2(ROOT / "docs" / "contracts" / name, contract_dir / name)
    monkeypatch.setattr(capability.shutil, "which", lambda name: f"/usr/bin/{name}")
    monkeypatch.setattr(capability, "_font_ready", lambda: (True, "fixture-font"))
    monkeypatch.setattr(capability.importlib.util, "find_spec", lambda name: object())

    report = inspect_high_density_capability(tmp_path)

    schema_check = next(item for item in report["checks"] if item["name"] == "v2_schemas")
    assert schema_check["ready"] is True
    assert report["ready"] is True


def test_production_requires_approved_style_lock(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="production", project_name="production example")
    prepared = prepare_high_density(run)

    assert prepared["status"] == "awaiting_user_decision"
    assert build_high_density_status(run)["current_stage"] == "style_lock"
    assert (run / "high_density_build/style/style_options.json").exists()
    style_options = read_json(run / "high_density_build/style/style_options.json")
    assert len(style_options["styles"]) == 8
    assert all((run / item["sample_ref"]).is_file() for item in style_options["styles"])
    assert len({item["sample_sha256"] for item in style_options["styles"]}) == 8
    waiting = run_high_density(run)
    assert waiting["status"] == "awaiting_user_decision"
    watched = watch_high_density_status(run, timeout_seconds=0.01, poll_seconds=0.01)
    assert watched["status"] == "awaiting_user_decision"
    assert watched["watch"]["timed_out"] is False
    assert "build select-style" in watched["next_action"]["approval_command"]


def test_build_cli_exposes_style_blueprint_and_provider_commands() -> None:
    import subprocess

    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "deck_master.py"), "build", "--help"],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "select-style" in result.stdout
    assert "approve-blueprint" in result.stdout
    assert "import-provider-result" in result.stdout


def test_style_lock_change_invalidates_blueprints(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    write_style_lock(run, run.name, "cyber-02", approved=True, approver="test")

    result = run_high_density(run)

    assert result["status"] == "awaiting_agent_build"
    assert result["current_stage"] == "blueprint"
    assert not (run / "high_density_build/blueprints/P001.blueprint_manifest.json").exists()


def test_mbb_blocks_low_density_without_evidence() -> None:
    package = _package("mbb-run", FIXTURE["pages"][0])
    package["customer_visible"]["body_blocks"] = [{"type": "text", "text": "One short point"}]
    package["customer_visible"]["callouts"] = []
    package["evidence_bindings"] = []

    with pytest.raises(ContractError, match="too sparse"):
        build_mbb_page(package)


def test_structural_page_allows_low_density_without_business_evidence() -> None:
    package = _package("mbb-run", FIXTURE["pages"][0])
    package["visual_spec"]["page_type"] = "cover"
    package["customer_visible"]["body_blocks"] = []
    package["customer_visible"]["callouts"] = []
    package["evidence_bindings"] = []
    package["claim_bindings"] = []

    result = build_mbb_page(package)

    assert result["analysis"]["structural_page"] is True
    assert result["analysis"]["page_role"] == "cover"


def test_mbb_rejects_unsupported_factual_claim() -> None:
    package = _package("mbb-run", FIXTURE["pages"][0])
    package["customer_visible"]["body_blocks"] = [{"type": "text", "text": "42% unsupported claim"} for _ in range(6)]
    package["evidence_bindings"] = []

    with pytest.raises(ContractError, match="unsupported factual values"):
        build_mbb_page(package)


def test_mbb_plan_contains_content_specific_candidates_and_precise_page_refs(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, page_count=2)
    packages = load_page_packages(run, expected_run_id=run.name)

    plan = build_mbb_plan(packages, run_id=run.name)
    write_mbb_plan(run, plan)

    assert plan["selection"]["status"] == "pending_user_decision"
    assert len(plan["storyline_candidates"]) == 3
    conclusions = {candidate["management_conclusion"] for candidate in plan["storyline_candidates"]}
    assert len(conclusions) == 3
    assert "Synthetic framework page" in plan["storyline_candidates"][0]["management_conclusion"]
    assert plan["scr"] is None
    assert plan["pages"] == []
    assert all(candidate["claim_bindings"] for candidate in plan["storyline_candidates"])
    assert load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)["mbb_plan_sha256"] == plan["mbb_plan_sha256"]


def test_mbb_rejects_duplicate_evidence_ids_across_page_packages(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, page_count=2)
    packages = load_page_packages(run, expected_run_id=run.name)
    packages[1]["evidence_bindings"] = [
        {"evidence_id": "E001", "source_ref": "fixture.duplicate", "meaning": "Conflicting page evidence."}
    ]

    with pytest.raises(ContractError, match="globally unique across Page Packages"):
        build_mbb_plan(packages, run_id=run.name)


def test_mbb_loader_rejects_duplicate_evidence_ids_in_agent_plan(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, page_count=2)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = build_mbb_plan(packages, run_id=run.name)
    plan["evidence_ledger"].append(dict(plan["evidence_ledger"][0]))
    plan["mbb_plan_sha256"] = sha256_json(
        {key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}}
    )
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="globally unique across MBB plan"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_build_mbb_plan_cannot_auto_enrich_selected_storyline() -> None:
    package = _package("mbb-run", FIXTURE["pages"][0])

    with pytest.raises(TypeError):
        build_mbb_plan(
            [package],
            run_id="mbb-run",
            selected_storyline_id="storyline.decision",
            approved_by="runtime",
        )


def test_content_lock_requires_approved_mbb_page_plan(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")

    with pytest.raises(ContractError, match="approved MBB page plan is required"):
        build_content_lock(package, mbb_plan_sha256="a" * 64)


def test_production_records_storyline_confirmation_without_host_key(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="production", project_name="production mbb")
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    prepare_high_density(run)
    assert run_high_density(run)["next_action"]["kind"] == "agent_mbb_candidates"

    packages = load_page_packages(run, expected_run_id=run.name)
    pending = build_mbb_plan(packages, run_id=run.name)
    write_mbb_plan(run, pending)
    waiting = run_high_density(run)

    assert waiting["status"] == "awaiting_user_decision"
    assert waiting["current_stage"] == "content_lock"
    assert waiting["next_action"]["recommended_storyline_id"] == "storyline.decision"
    assert len(waiting["next_action"]["storyline_candidates"]) == 3

    selected = retry_high_density(run, page_id="", stage="content_lock", storyline_id="storyline.risk")

    assert selected["status"] == "awaiting_agent_build"
    assert selected["next_action"]["kind"] == "agent_mbb_enrich_selected"
    assert not (run / "high_density_build/mbb/user_decision_receipt.json").exists()
    assert not (run / "high_density_build/content_locks/P001.json").exists()
    selected_plan = load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)
    assert selected_plan["selection"]["status"] == "selected_pending_enrichment"
    enriched = enrich_selected_mbb_plan(selected_plan, packages)
    write_mbb_plan(run, enriched)

    resumed = run_high_density(run)

    assert resumed["status"] == "awaiting_agent_build"
    assert resumed["current_stage"] == "blueprint"
    lock = read_json(run / "high_density_build/content_locks/P001.json")
    assert lock["lineage"]["selected_storyline_id"] == "storyline.risk"
    assert lock["lineage"]["mbb_page_plan_sha256"]
    context = lock["enrichment"]["storyline_context"]
    assert context["storyline_id"] == "storyline.risk"
    assert context["management_conclusion"].startswith("Reduce the execution risk")
    assert lock["enrichment"]["conclusion"].startswith("Synthetic framework page: Reduce the execution risk")
    assert "risk route" in lock["enrichment"]["so_what"]
    prompt = read_json(run / "high_density_build/prompts/P001.blueprint_prompt.json")
    assert context["management_conclusion"] in prompt["prompt_text"]


def test_storyline_selection_does_not_rewrite_agent_content(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="production", project_name="agent content preservation")
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    prepare_high_density(run)
    packages = load_page_packages(run, expected_run_id=run.name)
    write_mbb_plan(run, build_mbb_plan(packages, run_id=run.name))

    before = read_json(run / "high_density_build/mbb/mbb_plan.json")
    selected = select_mbb_storyline(run, "storyline.decision", selected_by="user")

    assert selected["storyline_candidates"] == before["storyline_candidates"]
    assert selected["scr"] is None
    assert selected["pages"] == []
    enriched = enrich_selected_mbb_plan(selected, packages)
    page = enriched["pages"][0]
    page["conclusion"] = "AGENT RICH CONCLUSION PRESERVE ME for Synthetic framework page"
    page["supporting_arguments"][0] = "AGENT ARGUMENT PRESERVE ME with Evidence, constraints, and decision context"
    page["so_what"] = "AGENT SO WHAT PRESERVE ME for Synthetic framework page"
    page["business_implication"] = page["so_what"]
    next(item for item in page["required_text_refs"] if item["ref"] == "content_lock.enrichment.business_implication")["value"] = page["business_implication"]
    for binding in page["claim_bindings"]:
        if binding["target"] == "conclusion":
            binding["text_sha256"] = sha256_json(page["conclusion"])
            binding["origin"] = "derived"
        elif binding["target"] == "supporting_arguments.0":
            binding["text_sha256"] = sha256_json(page["supporting_arguments"][0])
            binding["origin"] = "derived"
        elif binding["target"] == "so_what":
            binding["text_sha256"] = sha256_json(page["so_what"])
            binding["origin"] = "derived"
        elif binding["target"] == "business_implication":
            binding["text_sha256"] = sha256_json(page["business_implication"])
            binding["origin"] = "derived"
        elif binding["target"].endswith(".value"):
            binding["text_sha256"] = sha256_json(page["business_implication"])
            binding["origin"] = "derived"
    _rehash_page_and_plan(enriched, page)
    write_mbb_plan(run, enriched)

    result = run_high_density(run)

    assert result["current_stage"] == "blueprint"
    lock = read_json(run / "high_density_build/content_locks/P001.json")
    assert lock["enrichment"]["conclusion"] == "AGENT RICH CONCLUSION PRESERVE ME for Synthetic framework page"
    assert lock["enrichment"]["supporting_arguments"][0] == "AGENT ARGUMENT PRESERVE ME with Evidence, constraints, and decision context"
    assert lock["enrichment"]["so_what"] == "AGENT SO WHAT PRESERVE ME for Synthetic framework page"


def test_selected_storyline_requires_agent_enrichment_before_lock(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path, mode="production", project_name="selected enrichment gate")
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    prepare_high_density(run)
    packages = load_page_packages(run, expected_run_id=run.name)
    write_mbb_plan(run, build_mbb_plan(packages, run_id=run.name))

    waiting = retry_high_density(run, page_id="", stage="content_lock", storyline_id="storyline.decision")

    assert waiting["next_action"]["kind"] == "agent_mbb_enrich_selected"
    assert not (run / "high_density_build/content_locks/P001.json").exists()


def test_unbound_mbb_fact_and_numeric_claim_are_rejected(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    page["conclusion"] = "Synthetic framework page reports 42% without source support."
    binding = next(item for item in page["claim_bindings"] if item["target"] == "conclusion")
    binding["text_sha256"] = sha256_json(page["conclusion"])
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="unsupported factual values"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)

    page["conclusion"] = "Evidence-bound editorial conclusion."
    page["claim_bindings"] = [item for item in page["claim_bindings"] if item["target"] != "conclusion"]
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)
    with pytest.raises(ContractError, match="claim binding coverage failed"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_unrelated_nonnumeric_mbb_claim_is_rejected(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    packages[0]["evidence_bindings"] = []
    packages[0]["citations"] = [
        {"evidence_id": "E001", "source_ref": "fixture.metric", "meaning": "Evidence, constraints, and decision context."}
    ]
    write_json(run / "page_packages/P001.json", packages[0])
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    page["conclusion"] = "Synthetic framework page has dominant market share and guaranteed profitability."
    binding = next(item for item in page["claim_bindings"] if item["target"] == "conclusion")
    binding["text_sha256"] = sha256_json(page["conclusion"])
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="does not match evidence content"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_numeric_claim_must_exist_in_its_cited_evidence(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    packages[0]["speaker_notes"] += " An unrelated note mentions 42%."
    write_json(run / "page_packages/P001.json", packages[0])
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    page["conclusion"] = "Synthetic framework page records 42%."
    binding = next(item for item in page["claim_bindings"] if item["target"] == "conclusion")
    binding["text_sha256"] = sha256_json(page["conclusion"])
    binding["origin"] = "derived"
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="unsupported factual values"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_agent_approved_mbb_plan_without_runtime_seal_is_rejected(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    plan["selection"]["status"] = "approved"
    plan["selection"]["sealed_at"] = "2026-01-01T00:00:00+00:00"
    plan["storyline_audit"]["status"] = "approved"
    plan["mbb_plan_sha256"] = sha256_json(
        {key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}}
    )
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="missing its Runtime seal"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=True)


def test_agent_selected_mbb_plan_without_runtime_receipt_is_rejected(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = build_mbb_plan(packages, run_id=run.name)
    plan["selection"].update(
        {
            "status": "selected_pending_enrichment",
            "selected_storyline_id": "storyline.decision",
            "selected_by": "forged-agent",
            "selected_at": "2026-01-01T00:00:00+00:00",
        }
    )
    plan["storyline_audit"].update({"status": "selected_pending_enrichment", "selected_id": "storyline.decision"})
    plan = enrich_selected_mbb_plan(plan, packages)
    write_json(run / "high_density_build/mbb/mbb_plan.json", plan)

    with pytest.raises(ContractError, match="selection_receipt"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_forged_mbb_runtime_seal_signature_is_rejected(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    plan["selection"]["status"] = "approved"
    plan["selection"]["sealed_at"] = "2026-01-01T00:00:00+00:00"
    plan["storyline_audit"]["status"] = "approved"
    plan["mbb_plan_sha256"] = sha256_json(
        {key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}}
    )
    write_mbb_plan(run, plan)
    receipt_path = run / "high_density_build/mbb/selection_receipt.json"
    fake_seal = {
        "schema_version": "deck_mbb_runtime_seal.v1",
        "run_id": run.name,
        "selected_storyline_id": "storyline.decision",
        "mbb_plan_sha256": plan["mbb_plan_sha256"],
        "selection_receipt_sha256": sha256_file(receipt_path),
        "page_package_sha256": {package["page_id"]: sha256_json(package) for package in packages},
        "sealed_at": plan["selection"]["sealed_at"],
        "integrity": {
            "algorithm": "hmac-sha256",
            "key_id": "0" * 16,
            "payload_sha256": "0" * 64,
            "signature": "0" * 64,
        },
    }
    write_json(run / "high_density_build/mbb/runtime_seal.json", fake_seal)

    with pytest.raises(ContractError, match="Runtime integrity"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=True)


def test_runtime_seal_is_invalidated_when_approved_plan_changes(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    sealed = seal_mbb_plan(run)
    assert (run / "high_density_build/mbb/runtime_seal.json").is_file()

    sealed["storyline_audit"]["notes"] = "mutated-after-seal"
    sealed["mbb_plan_sha256"] = sha256_json(
        {key: value for key, value in sealed.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}}
    )
    write_mbb_plan(run, sealed)

    assert not (run / "high_density_build/mbb/runtime_seal.json").exists()
    with pytest.raises(ContractError, match="missing its Runtime seal"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=True)


def test_numeric_claim_uses_exact_token_match_in_cited_evidence(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    package["speaker_notes"] += " Source notes mention 42%."
    package["evidence_bindings"] = []
    package["citations"] = [
            {"evidence_id": "E001", "source_ref": "fixture.metric", "meaning": "Synthetic framework page has a cited metric of 142%."}
    ]
    write_json(run / "page_packages/P001.json", package)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    page["conclusion"] = "Synthetic framework page records 42%."
    binding = next(item for item in page["claim_bindings"] if item["target"] == "conclusion")
    binding["text_sha256"] = sha256_json(page["conclusion"])
    binding["origin"] = "derived"
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="unsupported factual values"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_wide_and_tall_frames_remain_inside_source_canvas() -> None:
    for width, height in ((2000, 1000), (1000, 2000)):
        frame = _default_slide_frame(width, height)
        assert frame["x"] >= 0 and frame["y"] >= 0
        assert frame["x"] + frame["w"] <= width
        assert frame["y"] + frame["h"] <= height
        assert frame["w"] / frame["h"] == pytest.approx(1672 / 941, abs=0.02)


def test_prepare_invalidates_downstream_when_page_package_changes(tmp_path: Path) -> None:
    run, index = _make_run(tmp_path)
    _blueprint(run)
    prepare_high_density(run)
    assert run_high_density(run)["status"] == "completed"

    package = read_json(run / "page_packages/P001.json")
    package["customer_visible"]["title"] = "Changed after the completed build"
    index.write(package)

    prepared = prepare_high_density(run)

    assert prepared["status"] == "prepared"
    assert (run / "high_density_build/content_locks/P001.json").exists()
    assert not (run / "high_density_build/blueprints/P001.svg").exists()
    assert not (run / "high_density_build/page_scenes/P001.json").exists()
    waiting = run_high_density(run)
    assert waiting["status"] == "awaiting_agent_build"
    assert waiting["current_stage"] == "blueprint"


def test_mbb_plan_rejects_unknown_page_evidence_ref(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    plan["pages"][0]["evidence_refs"] = ["E999"]
    plan["pages"][0]["page_plan_sha256"] = sha256_json({key: value for key, value in plan["pages"][0].items() if key != "page_plan_sha256"})
    plan["mbb_plan_sha256"] = sha256_json({key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}})
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="evidence refs are invalid"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_mbb_plan_rejects_missing_required_component(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    plan["pages"][0]["components"] = plan["pages"][0]["components"][1:]
    plan["pages"][0]["page_plan_sha256"] = sha256_json({key: value for key, value in plan["pages"][0].items() if key != "page_plan_sha256"})
    plan["mbb_plan_sha256"] = sha256_json({key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}})
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="components are outside the Runtime registry"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_mbb_plan_rejects_stale_storyline_material_pool(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    plan["pages"][0]["material_pool"]["storyline_id"] = "storyline.risk"
    plan["pages"][0]["page_plan_sha256"] = sha256_json({key: value for key, value in plan["pages"][0].items() if key != "page_plan_sha256"})
    plan["mbb_plan_sha256"] = sha256_json({key: value for key, value in plan.items() if key not in {"mbb_plan_sha256", "created_at", "updated_at"}})
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="material pool is outside the Runtime registry"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_mbb_plan_blocks_agent_added_material_component_and_required_text(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    page["components"].append(
        {"component_id": "component.unverified", "kind": "callout", "priority": "P0", "region": "insight"}
    )
    page["required_text_refs"].append(
        {"ref": "content_lock.enrichment.unverified", "priority": "P0", "required": True, "value": "Unverified growth reached 999%."}
    )
    _rehash_page_and_plan(plan, page)
    write_mbb_plan(run, plan)

    with pytest.raises(ContractError, match="components are outside the Runtime registry"):
        load_mbb_plan(run, packages=packages, expected_run_id=run.name, require_approved=False)


def test_mbb_claim_bindings_cover_components_and_required_text_registry(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    packages = load_page_packages(run, expected_run_id=run.name)
    plan = _write_selected_enriched_mbb_plan(run, packages)
    page = plan["pages"][0]
    targets = {str(binding["target"]) for binding in page["claim_bindings"]}

    for index, component in enumerate(page["components"]):
        for field in component:
            assert f"components.{index}.{field}" in targets
    for index, text_ref in enumerate(page["required_text_refs"]):
        for field in text_ref:
            assert f"required_text_refs.{index}.{field}" in targets


def test_pending_mbb_plan_invalidates_previous_downstream(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    _blueprint(run)
    prepare_high_density(run)
    assert run_high_density(run)["status"] == "completed"

    packages = load_page_packages(run, expected_run_id=run.name)
    pending = build_mbb_plan(packages, run_id=run.name)
    write_mbb_plan(run, pending)
    waiting = run_high_density(run)

    assert waiting["status"] == "awaiting_user_decision"
    assert not (run / "high_density_build/content_locks/P001.json").exists()
    assert not (run / "high_density_build/svg/P001.svg").exists()
    assert not (run / "high_density_build/pptx/deck_high_density.pptx").exists()


def test_content_lock_contains_required_components_and_text_refs(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="a" * 64)

    assert lock["schema_version"] == "deck_content_lock.v2"
    assert lock["required_component_ids"]
    assert {item["ref"] for item in lock["required_text_refs"]} >= {"content_lock.customer_visible.title", "content_lock.enrichment.business_implication"}
    assert lock["lineage"]["mbb_plan_sha256"] == "a" * 64


def test_blueprint_prompt_changes_with_locked_content(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="b" * 64)
    style = {"style_id": "cyber-01", "name": "Ink Cobalt", "palette": {"primary": "#419BFD"}, "grid": {"system": "12-column"}}
    changed = json.loads(json.dumps(lock))
    changed["customer_visible"]["title"] = "Changed locked title"

    first = build_blueprint_prompt(lock, style, mbb_plan_sha256="b" * 64)
    second = build_blueprint_prompt(changed, style, mbb_plan_sha256="b" * 64)
    assert first != second
    assert "Changed locked title" in second


def test_prompt_contains_full_mbb_and_style_lock(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="b" * 64)
    style = {
        "style_id": "cyber-01",
        "name": "Ink Cobalt",
        "palette": {"primary": "#419BFD"},
        "grid": {"system": "12-column"},
        "typography": {"family": "Arial"},
        "chart_language": {"axis": "hairline"},
        "table_language": {"cell_padding_px": 14},
        "surface_system": {"radius_px": 8},
        "density_rules": {"minimum_information_regions": 3},
    }

    prompt = build_blueprint_prompt(lock, style, mbb_plan_sha256="b" * 64)

    assert lock["enrichment"]["supporting_arguments"][0] in prompt
    for key in ("typography", "chart_language", "table_language", "surface_system", "density_rules"):
        assert f"{key}=" in prompt


def test_prompt_excludes_internal_analysis_fields(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="e" * 64)
    prompt = build_blueprint_prompt(lock, {"style_id": "cyber-01", "name": "Ink Cobalt", "palette": {}, "grid": {}}, mbb_plan_sha256="e" * 64)

    assert lock["enrichment"]["conclusion"] in prompt
    assert lock["enrichment"]["supporting_arguments"][0] in prompt
    for internal_field in ("Evidence ID", "Evidence hierarchy", "Evidence assessment", "Derived claim lineage", "SO WHAT", "Caveat", "MBB", "SCR"):
        assert internal_field not in prompt


def test_page_number_is_hard_forbidden_and_internal_labels_need_allowlist(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    policy = build_visibility_policy(package, page_id="P001")

    assert visible_text_violation(policy, "1 / 12", page_id="P001") == "page_number"
    assert visible_text_violation(policy, "P001", page_id="P001") == "page_number"
    assert visible_text_violation(policy, "来源：", page_id="P001") == "source_marker"
    policy["allowed_visible_terms"] = [{"term": "SWOT", "content_ref": "content_lock.customer_visible.body_blocks.0", "approved_by": "client", "reason": "client-facing framework"}]
    assert visible_text_violation(policy, "SWOT", page_id="P001") is None


def test_blueprint_requires_content_review_before_approval(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    _blueprint(run)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="f" * 64)
    build_blueprint_prompt_artifact(run, "P001", lock, style_lock={"style_lock_sha256": "0" * 64})

    with pytest.raises(BlueprintContentReviewRequired, match="content review is required"):
        load_blueprint_content_review(run, "P001", lock)

    write_blueprint_content_review(
        run,
        "P001",
        lock,
        findings=[{"category": "page_number", "bbox": {"x": 1600, "y": 900, "w": 20, "h": 20}, "description": "visible page number", "observed_text": "1", "disposition": "allowed", "allowlist_term": "1"}],
        reviewer_id="test",
        action_id="test-review",
    )
    with pytest.raises(BlueprintContentReviewRequired, match="requires regeneration"):
        load_blueprint_content_review(run, "P001", lock)


def test_rejected_blueprint_attempts_are_archived_and_counted(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="a" * 64)
    source = FIXTURE_DIR / "blueprint.svg"
    (run / "high_density_build/blueprints").mkdir(parents=True, exist_ok=True)

    for attempt in range(1, 4):
        (run / "high_density_build/blueprints/P001.svg").write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
        build_blueprint_prompt_artifact(run, "P001", lock, style_lock={"style_lock_sha256": "0" * 64})
        write_blueprint_content_review(
            run,
            "P001",
            lock,
            findings=[{"category": "page_number", "bbox": {"x": 1600, "y": 900, "w": 20, "h": 20}, "description": "visible page number", "observed_text": "1 / 3", "disposition": "blocked"}],
            reviewer_id="test",
            action_id=f"test-review-{attempt}",
        )
        assert archive_rejected_blueprint(run, "P001") == attempt
        archive = run / f"high_density_build/blueprints/rejected/P001/attempt-{attempt}"
        assert (archive / "P001.svg").is_file()
        assert (archive / "P001.blueprint_content_review.json").is_file()

    assert next_attempt_index(run, "P001") == 3


def test_blueprint_prompt_preserves_structured_content(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    package["customer_visible"]["body_blocks"][0] = {"title": "Metric table", "rows": [{"metric": "conversion", "value": "42%"}], "type": "table"}
    package["evidence_bindings"] = [
        {"evidence_id": "E001", "source_ref": "fixture.metric_table", "source_position": "row.conversion", "meaning": "conversion is 42%"}
    ]
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="c" * 64)
    style = {"style_id": "cyber-01", "name": "Ink Cobalt", "palette": {}, "grid": {}}

    prompt = build_blueprint_prompt(lock, style, mbb_plan_sha256="c" * 64)

    assert "conversion" in prompt
    assert "42%" in prompt


def test_blueprint_manifest_requires_prompt_before_image(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    image = _blueprint(run)
    package = read_json(run / "page_packages/P001.json")
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="d" * 64)

    assert image.exists()
    with pytest.raises(BlueprintInvalid, match="prompt must be written"):
        ensure_blueprint_manifest(run, "P001", lock)


def test_production_blueprint_rejects_self_declared_provider_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    run, _ = _make_run(tmp_path, mode="production")
    _blueprint(run)
    prepare_high_density(run)
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    packages = load_page_packages(run, expected_run_id=run.name)
    _write_selected_enriched_mbb_plan(run, packages)
    plan = seal_mbb_plan(run)
    waiting = run_high_density(run)

    assert waiting["status"] == "awaiting_agent_build"
    assert waiting["next_action"]["kind"] == "agent_imagegen"
    prompt = read_json(run / "high_density_build/prompts/P001.blueprint_prompt.json")
    challenge = prompt["provider_challenge"]
    provider = {
        "tool": "image_gen.imagegen",
        "model": "self-declared-model",
        "request_id": "self-declared-request",
        "challenge_nonce": challenge["nonce"],
        "prompt_sha256": prompt["prompt_sha256"],
        "requested_at": challenge["issued_at"],
        "responded_at": challenge["issued_at"],
    }
    provider["request_sha256"] = sha256_json(
        {key: provider[key] for key in ("tool", "model", "request_id", "challenge_nonce", "prompt_sha256", "requested_at")}
    )
    write_json(
        run / "high_density_build/blueprints/P001.blueprint_manifest.json",
        {"schema_version": "deck_blueprint_manifest.v2", "provider": provider},
    )
    lock = read_json(run / "high_density_build/content_locks/P001.json")
    style = read_json(run / "high_density_build/style/style_lock.json")

    with pytest.raises(BlueprintInvalid, match="Host-managed provider receipt"):
        ensure_blueprint_manifest(
            run,
            "P001",
            lock,
            style_lock=style,
            mbb_plan_sha256=str(plan["mbb_plan_sha256"]),
            approval={
                "status": "approved",
                "source": "explicit_user",
                "approved_by": "test",
                "approved_at": challenge["issued_at"],
            },
        )


def test_distinct_blueprints_produce_distinct_svg(tmp_path: Path) -> None:
    run, lock, _ = _prepared_fixture(tmp_path)
    first = build_fixture_scene(lock, "1" * 64)
    second = build_fixture_scene(lock, "2" * 64)
    first_path = run / "high_density_build/svg/first.svg"
    second_path = run / "high_density_build/svg/second.svg"
    compile_svg(first, first_path)
    compile_svg(second, second_path)

    assert first_path.read_text(encoding="utf-8") != second_path.read_text(encoding="utf-8")


def test_fixture_background_ignores_malformed_rect_attributes(tmp_path: Path) -> None:
    blueprint = tmp_path / "malformed.svg"
    blueprint.write_text('<svg xmlns="http://www.w3.org/2000/svg"><rect x="invalid" y="0" width="1672" height="941" fill="#123456"/></svg>', encoding="utf-8")

    assert _fixture_background(blueprint) == "#f7f9fb"


def test_production_uses_agent_svg_without_overwriting_it(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    run, _ = _make_run(tmp_path, mode="production", project_name="production visual")
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    _blueprint(run)
    prepare_high_density(run)
    _write_approved_mbb_plan(run)
    waiting = run_high_density(run)
    assert waiting["current_stage"] == "blueprint"
    from PIL import Image

    provider_root = tmp_path / "provider-results"
    provider_root.mkdir()
    monkeypatch.setenv("DECK_MASTER_PROVIDER_RESULT_ROOTS", str(provider_root))
    provider_image = provider_root / "exec-00000000-0000-0000-0000-000000000001.png"
    Image.new("RGB", (1672, 941), "#f7f9fb").save(provider_image)
    record_provider_host_result(run, "P001", provider_image)
    (run / "high_density_build/blueprints/P001.svg").unlink()
    _approve_blueprint(run)
    assert run_high_density(run)["current_stage"] == "page_scene"

    lock = read_json(run / "high_density_build/content_locks/P001.json")
    blueprint_manifest = read_json(run / "high_density_build/blueprints/P001.blueprint_manifest.json")
    scene = build_fixture_scene(
        lock,
        str(blueprint_manifest["image_sha256"]),
        blueprint_path=FIXTURE_DIR / "blueprint.svg",
    )
    write_scene(run, scene)
    svg_file = svg_path(run, "P001")
    compile_svg(scene, svg_file)
    custom_svg = svg_file.read_text(encoding="utf-8").replace("#f7f9fb", "#123456", 1)
    svg_file.write_text(custom_svg, encoding="utf-8")

    result = run_high_density(run)

    assert result["status"] == "awaiting_agent_build"
    assert result["current_stage"] == "svg"
    assert result["next_action"]["kind"] == "agent_svg_repair"
    assert "#123456" in svg_file.read_text(encoding="utf-8")


def test_scene_rejects_missing_required_content(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    scene["elements"] = [element for element in scene["elements"] if element["element_id"] != "title.main"]

    with pytest.raises(ContractError, match="missing required text refs"):
        validate_scene_content(scene, lock)


def test_scene_required_sets_match_content_lock(tmp_path: Path) -> None:
    _, lock, scene = _prepared_fixture(tmp_path)
    scene["required_component_ids"] = scene["required_component_ids"][:-1]
    scene["required_text_refs"] = scene["required_text_refs"][:-1]

    with pytest.raises(ContractError, match="does not exactly match content lock"):
        validate_scene_content(scene, lock)


def test_callout_removal_blocks_scene(tmp_path: Path) -> None:
    run, _ = _make_run(tmp_path)
    package = read_json(run / "page_packages/P001.json")
    package["customer_visible"]["callouts"] = [{"text": "Decision gate requires explicit evidence."}]
    _, page_plan = _approved_page_plan(package)
    lock = build_content_lock(package, page_plan, mbb_plan_sha256="c" * 64)
    blueprint = _blueprint(run)
    scene = build_fixture_scene(lock, sha256_file(blueprint), blueprint_path=blueprint)
    assert "component.callouts" in scene["required_component_ids"]
    scene["elements"] = [element for element in scene["elements"] if element.get("text_ref") != "content_lock.customer_visible.callouts"]

    with pytest.raises(ContractError, match="missing required text refs"):
        validate_scene_content(scene, lock)


def test_svg_rejects_missing_required_component(tmp_path: Path) -> None:
    run, _, scene = _prepared_fixture(tmp_path)
    scene["required_component_ids"].append("component.missing")

    with pytest.raises(SvgVisualError, match="missing required components"):
        compile_svg(scene, run / "high_density_build/svg/missing.svg")


def test_svg_rejects_missing_scene_element(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    document = ElementTree.parse(svg)
    parents = {child: parent for parent in document.getroot().iter() for child in list(parent)}
    missing = next(node for node in document.getroot().iter() if node.get("id") == "header.rule")
    parents[missing].remove(missing)
    document.write(svg, encoding="utf-8", xml_declaration=True)

    with pytest.raises(SvgVisualError, match="missing scene element: header.rule"):
        validate_approved_svg(svg, scene, lock)


def test_visual_metrics_are_computed_from_artifacts(tmp_path: Path) -> None:
    run, _, scene = _prepared_fixture(tmp_path)
    from high_density.visual import compute_visual_metrics

    output = render_preview(svg_path(run, "P001"), run / "high_density_build/previews/metrics.png")
    metrics = compute_visual_metrics(run, scene, output, output)

    assert metrics["status"] == "pass"
    assert metrics["inputs"]["reference_sha256"] == sha256_file(output)
    assert metrics["values"]["text_masked_ssim"] == pytest.approx(1.0)


def test_visual_bbox_metrics_read_actual_svg_geometry(tmp_path: Path) -> None:
    run, _, scene = _prepared_fixture(tmp_path)
    reference = run / "high_density_build/previews/reference.png"
    shutil.copy2(preview_path(run, "P001"), reference)
    svg = svg_path(run, "P001")
    original = svg.read_text(encoding="utf-8")
    svg.write_text(original.replace('id="block.01" x="80.00"', 'id="block.01" x="120.00"', 1), encoding="utf-8")
    candidate = run / "high_density_build/previews/mutated.png"
    render_preview(svg, candidate)

    from high_density.visual import compute_visual_metrics

    metrics = compute_visual_metrics(run, scene, reference, candidate)

    assert metrics["status"] == "failed"
    assert metrics["values"]["bbox_max_delta_px"] >= 40
    assert any(item["code"] == "p0_p1_bbox_drift" for item in metrics["findings"])


def test_visual_review_rejects_handwritten_metrics(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    metrics_file = run / "high_density_build/reviews/P001.metrics.json"
    review_file = review_path(run, "P001")
    metrics = read_json(metrics_file)
    metrics["geometry"]["pairs"][0]["target"]["x"] += 1
    write_json(metrics_file, metrics)
    review = read_json(review_file)
    review["metrics_sha256"] = sha256_file(metrics_file)
    write_json(review_file, review)

    with pytest.raises(SvgVisualError, match="Runtime challenge metrics_sha256 is stale"):
        load_visual_review(run, "P001")


def test_visual_review_requires_independent_reviewer_ids(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    review_file = review_path(run, "P001")
    review = read_json(review_file)
    review["main_review"]["reviewer_id"] = review["self_review"]["reviewer_id"]
    write_json(review_file, review)

    with pytest.raises(SvgVisualError, match="independent reviewer IDs"):
        load_visual_review(run, "P001")


def test_main_review_requires_separate_host_attestation(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    from high_density.svg import _load_main_review_receipt

    review = read_json(review_path(run, "P001"))
    with pytest.raises(SvgVisualError, match="independent main visual review attestation is required"):
        _load_main_review_receipt(run, "P001", review)


def test_default_production_visual_review_does_not_require_host_attestation(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)

    review = load_visual_review(run, "P001", require_external_receipt=False)

    assert review["visual_status"] == "pass"
    assert not main_review_receipt_path(run, "P001").exists()


def test_near_full_image_is_blocked(tmp_path: Path) -> None:
    run, _, scene = _prepared_fixture(tmp_path)
    scene["elements"].append({"element_id": "image.near_full", "component_id": "component.proof", "kind": "image", "role": "proof", "priority": "P2", "bbox": {"x": 20, "y": 20, "w": 900, "h": 800}, "asset_ref": "proof", "asset_sha256": "a" * 64, "editability_target": "registered_asset", "asset_policy": "registered"})

    with pytest.raises(SvgVisualError, match="exceeds 35%"):
        compile_svg(scene, run / "high_density_build/svg/near-full.svg")


def test_image_cannot_cover_p0_text(tmp_path: Path) -> None:
    run, _, scene = _prepared_fixture(tmp_path)
    scene["elements"].append({"element_id": "image.cover", "component_id": "component.proof", "kind": "image", "role": "proof", "priority": "P2", "z_index": 30, "bbox": {"x": 80, "y": 56, "w": 300, "h": 100}, "asset_ref": "proof", "asset_sha256": "a" * 64, "editability_target": "registered_asset", "asset_policy": "registered"})

    with pytest.raises(SvgVisualError, match="covers P0/P1 text"):
        compile_svg(scene, run / "high_density_build/svg/cover.svg")


def test_production_svg_blocks_near_full_image_and_p0_p1_coverage(tmp_path: Path) -> None:
    from PIL import Image

    run, lock, scene = _prepared_fixture(tmp_path)
    asset = run / "assets/proof.png"
    asset.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (32, 32), "#419bfd").save(asset)
    scene["elements"].append(
        {
            "element_id": "image.proof",
            "component_id": "component.proof",
            "kind": "image",
            "role": "proof",
            "priority": "P2",
            "z_index": 30,
            "bbox": {"x": 1400, "y": 866, "w": 120, "h": 24},
            "asset_ref": "proof",
            "asset_sha256": sha256_file(asset),
            "editability_target": "registered_asset",
            "asset_policy": "registered",
        }
    )
    output = run / "high_density_build/svg/with-proof.svg"
    compile_svg(scene, output, assets={"proof": asset})
    validate_approved_svg(output, scene, lock, {"proof": asset})
    source = output.read_text(encoding="utf-8")
    near_full = source.replace('x="1400.00" y="866.00" width="120.00" height="24.00"', 'x="20.00" y="20.00" width="900.00" height="800.00"')
    near_full = near_full.replace('data-pptx-bounds="1400.00,866.00,120.00,24.00"', 'data-pptx-bounds="20.00,20.00,900.00,800.00"')
    output.write_text(near_full, encoding="utf-8")

    with pytest.raises(SvgVisualError, match="exceeds 35%"):
        validate_approved_svg(output, scene, lock, {"proof": asset})

    coverage = source.replace('x="1400.00" y="866.00" width="120.00" height="24.00"', 'x="100.00" y="240.00" width="300.00" height="100.00"')
    coverage = coverage.replace('data-pptx-bounds="1400.00,866.00,120.00,24.00"', 'data-pptx-bounds="100.00,240.00,300.00,100.00"')
    output.write_text(coverage, encoding="utf-8")

    with pytest.raises(SvgVisualError, match="covers P0/P1 text"):
        validate_approved_svg(output, scene, lock, {"proof": asset})


def test_image_heavy_page_role_allows_large_registered_image(tmp_path: Path) -> None:
    from PIL import Image

    run, lock, scene = _prepared_fixture(tmp_path)
    scene["page_role"] = "visual"
    lock["enrichment"]["analysis"]["page_role"] = "visual"
    asset = run / "assets/visual.png"
    asset.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (32, 32), "#419bfd").save(asset)
    scene["elements"].append(
        {
            "element_id": "image.hero",
            "component_id": "component.visual",
            "kind": "image",
            "role": "visual",
            "priority": "P2",
            "z_index": 2,
            "bbox": {"x": 40, "y": 220, "w": 1100, "h": 720},
            "asset_ref": "visual",
            "asset_sha256": sha256_file(asset),
            "editability_target": "registered_asset",
            "asset_policy": "registered",
        }
    )

    output = run / "high_density_build/svg/image-heavy.svg"
    compile_svg(scene, output, assets={"visual": asset})

    assert 'data-pptx-page-role="visual"' in output.read_text(encoding="utf-8")


def test_explicit_provider_import_records_hash_without_host_root(tmp_path: Path) -> None:
    from PIL import Image

    run, _ = _make_run(tmp_path, mode="production", project_name="production provider import")
    write_style_lock(run, run.name, "cyber-01", approved=True, approver="test")
    _blueprint(run)
    prepare_high_density(run)
    _write_approved_mbb_plan(run)
    waiting = run_high_density(run)
    assert waiting["current_stage"] == "blueprint"
    imported = tmp_path / "outside-provider.png"
    Image.new("RGB", (1672, 941), "#f7f9fb").save(imported)

    record_provider_host_result(run, "P001", imported, source_type="explicit_import")

    receipt = read_json(run / "high_density_build/blueprints/P001.provider_host_receipt.json")
    assert receipt["source_type"] == "explicit_import"
    assert receipt["source_file_sha256"] == sha256_file(imported)


def test_approved_svg_blocks_unresolvable_font_and_text_overflow(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    source = svg.read_text(encoding="utf-8")
    svg.write_text(source.replace('font-family="Arial"', 'font-family="Definitely Missing Font 123"', 1), encoding="utf-8")

    with pytest.raises(SvgVisualError, match="unresolvable SVG font"):
        validate_approved_svg(svg, scene, lock, {})

    svg.write_text(source.replace('font-size="42.00px"', 'font-size="200.00px"', 1), encoding="utf-8")
    with pytest.raises(SvgVisualError, match="text overflow"):
        validate_approved_svg(svg, scene, lock, {})


def test_required_text_same_as_rendered_background_fails_closed(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    document = ElementTree.parse(svg)
    title = next(node for node in document.getroot().iter() if node.get("id") == "title.main")
    title.set("fill", "#f7f9fb")
    document.write(svg, encoding="utf-8", xml_declaration=True)
    render_preview(svg, preview_path(run, "P001"))

    with pytest.raises(SvgVisualError, match="contrast"):
        validate_approved_svg(svg, scene, lock, {})

    from high_density.visual import compute_visual_metrics, normalize_blueprint

    manifest = read_json(run / "high_density_build/blueprints/P001.blueprint_manifest.json")
    metrics = compute_visual_metrics(run, scene, normalize_blueprint(run, manifest), preview_path(run, "P001"))
    assert metrics["status"] == "failed"
    assert any(finding["code"] == "p0_p1_text_contrast_below_threshold" for finding in metrics["findings"])


def test_arial_accepts_metric_compatible_fontconfig_alias(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import high_density.svg as svg_module

    font_file = tmp_path / "LiberationSans-Regular.ttf"
    font_file.write_bytes(b"font-fixture")
    monkeypatch.setattr(svg_module.shutil, "which", lambda name: "/usr/bin/fc-match" if name == "fc-match" else None)
    monkeypatch.setattr(
        svg_module.subprocess,
        "run",
        lambda *args, **kwargs: type("Result", (), {"returncode": 0, "stdout": f"Liberation Sans|{font_file}\n"})(),
    )

    assert _font_path("Arial", "P001", "title.main") == font_file


def test_blueprint_geometry_mutation_changes_scene_svg_and_metrics(tmp_path: Path) -> None:
    run, lock, first_scene = _prepared_fixture(tmp_path)
    original_svg = svg_path(run, "P001")
    original_svg_sha256 = sha256_file(original_svg)
    original_preview = preview_path(run, "P001")
    original_metrics = read_json(run / "high_density_build/reviews/P001.metrics.json")
    blueprint = run / "high_density_build/blueprints/P001.svg"
    mutated = blueprint.read_text(encoding="utf-8").replace('x="80" y="200" width="744"', 'x="120" y="200" width="704"', 1)
    blueprint.write_text(mutated, encoding="utf-8")
    second_scene = build_fixture_scene(lock, sha256_file(blueprint), blueprint_path=blueprint)
    second_svg = run / "high_density_build/svg/P001.svg"
    compile_svg(second_scene, second_svg)
    second_preview = run / "high_density_build/previews/P001.mutated.png"
    render_preview(second_svg, second_preview)
    from high_density.visual import compute_visual_metrics

    metrics = compute_visual_metrics(run, second_scene, original_preview, second_preview)

    assert first_scene["scene_signature"] != second_scene["scene_signature"]
    assert original_svg_sha256 != sha256_file(second_svg)
    assert metrics["inputs"]["candidate_sha256"] != original_metrics["inputs"]["candidate_sha256"]
    assert metrics["values"] != original_metrics["values"]


def test_svg_mutation_changes_pptx_trace_and_render(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    original_pptx_hash = sha256_file(pptx_path(run))
    original_trace = read_json(trace_path(run))
    original_preview = sha256_file(run / "high_density_build/previews/P001.pptx.png")
    svg = svg_path(run, "P001")
    svg.write_text(svg.read_text(encoding="utf-8").replace("#f7f9fb", "#fff3e8", 1), encoding="utf-8")
    render_preview(svg, preview_path(run, "P001"))
    compile_pptx(run, [scene], {"P001": lock})
    new_preview = _render_pptx_page(run, pptx_path(run), "P001", 0)
    new_trace = read_json(trace_path(run))

    assert sha256_file(pptx_path(run)) != original_pptx_hash
    assert new_trace["pages"][0]["svg_sha256"] != original_trace["pages"][0]["svg_sha256"]
    assert sha256_file(new_preview) != original_preview


def test_svg_gradient_shadow_compile_to_drawingml_and_readback(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    original = svg.read_text(encoding="utf-8")
    defs = '<defs><linearGradient id="gradient.primary" x1="0%" y1="0%" x2="100%" y2="0%"><stop offset="0%" stop-color="#fff3e8"/><stop offset="100%" stop-color="#fff3e8"/></linearGradient><filter id="effect.primary"><feDropShadow dx="0.5" dy="0.5" stdDeviation="0.5" flood-color="#102030" flood-opacity="0.02"/></filter></defs>'
    mutated = original.replace('<g id="page.P001"', defs + '<g id="page.P001"', 1).replace('id="block.03" x="80.00" y="511.00" width="744.00" height="289.00" rx="0.00" fill="#fff3e8"', 'id="block.03" x="80.00" y="511.00" width="744.00" height="289.00" rx="0.00" fill="url(#gradient.primary)" filter="url(#effect.primary)"', 1)
    svg.write_text(mutated, encoding="utf-8")
    render_preview(svg, preview_path(run, "P001"))

    compile_pptx(run, [scene], {"P001": lock})
    trace = read_json(trace_path(run))
    entry = next(item for item in trace["elements"] if item["element_id"] == "block.03")
    assert entry["paint"]["fill"]["gradient_id"] == "gradient.primary"
    assert entry["effect"]["effect_id"] == "effect.primary"
    readback_pptx(run, [scene], {"P001": lock}, pptx_path(run))
    report = read_json(run / "high_density_build/readback/readback_report.json")
    assert report["drawingml_paint"]["expected_gradient_count"] == 1
    assert report["drawingml_paint"]["expected_effect_count"] == 1
    assert report["drawingml_paint"]["status"] == "pass"


def test_svg_gradient_stroke_with_opacity_stays_gradient(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    original = svg.read_text(encoding="utf-8")
    defs = '<defs><linearGradient id="gradient.stroke" x1="0%" y1="0%" x2="100%" y2="0%"><stop offset="0%" stop-color="#419bfd"/><stop offset="100%" stop-color="#ff8c42"/></linearGradient></defs>'
    mutated = original.replace('<g id="page.P001"', defs + '<g id="page.P001"', 1).replace('id="block.03" x="80.00" y="511.00" width="744.00" height="289.00" rx="0.00" fill="#fff3e8" stroke="#d5dde5"', 'id="block.03" x="80.00" y="511.00" width="744.00" height="289.00" rx="0.00" fill="#fff3e8" stroke="url(#gradient.stroke)" stroke-opacity="0.5"', 1)
    svg.write_text(mutated, encoding="utf-8")
    render_preview(svg, preview_path(run, "P001"))

    compile_pptx(run, [scene], {"P001": lock})
    trace = read_json(trace_path(run))
    entry = next(item for item in trace["elements"] if item["element_id"] == "block.03")
    assert entry["paint"]["stroke"]["gradient_id"] == "gradient.stroke"
    with zipfile.ZipFile(pptx_path(run)) as package:
        slide_xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
    assert slide_xml.count("<a:gradFill") >= 1


def test_svg_gradient_inheritance_is_blocked(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    original = svg.read_text(encoding="utf-8")
    defs = '<defs><linearGradient id="gradient.bad" href="#gradient.other"><stop offset="0%" stop-color="#123456"/><stop offset="100%" stop-color="#abcdef"/></linearGradient></defs>'
    svg.write_text(original.replace('<g id="page.P001"', defs + '<g id="page.P001"', 1), encoding="utf-8")

    with pytest.raises(SvgVisualError, match="inheritance or transform"):
        validate_svg(svg, page_id="P001")


def test_readback_batches_full_deck_render(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    run, _ = _make_run(tmp_path, page_count=2)
    _blueprint(run, "P001")
    _blueprint(run, "P002")
    prepare_high_density(run)
    assert run_high_density(run)["status"] == "completed"
    packages = load_page_packages(run, expected_run_id=run.name)
    scenes = [load_scene(run, str(package["page_id"])) for package in packages]
    locks = {str(package["page_id"]): read_json(run / f"high_density_build/content_locks/{package['page_id']}.json") for package in packages}

    import high_density.pptx as pptx_module

    original_run = pptx_module.subprocess.run
    commands: list[list[str]] = []

    def counted_run(command: list[str], *args: Any, **kwargs: Any) -> Any:
        commands.append([str(item) for item in command])
        return original_run(command, *args, **kwargs)

    monkeypatch.setattr(pptx_module.subprocess, "run", counted_run)
    readback_pptx(run, scenes, locks, pptx_path(run))

    soffice_commands = [command for command in commands if command and command[0].endswith("soffice")]
    assert len(soffice_commands) == 1
    assert any(argument.startswith("-env:UserInstallation=file:") for argument in soffice_commands[0])
    assert sum(command and command[0].endswith("pdftoppm") for command in commands) == 1


def test_readback_rejects_unregistered_shape(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    presentation = Presentation(pptx_path(run))
    extra = presentation.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    extra.name = ""
    presentation.save(pptx_path(run))
    trace = read_json(trace_path(run))
    trace["pptx_sha256"] = sha256_file(pptx_path(run))
    write_json(trace_path(run), trace)

    with pytest.raises(PptxEditabilityError, match="PPTX readback failed"):
        readback_pptx(run, [scene], {"P001": lock}, pptx_path(run))


def test_scene_mutation_without_svg_change_does_not_change_pptx(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    original_svg_hash = sha256_file(svg_path(run, "P001"))
    original_pptx = _render_pptx_page(run, pptx_path(run), "P001", 0)
    with zipfile.ZipFile(pptx_path(run)) as archive:
        original_slide_xml = archive.read("ppt/slides/slide1.xml")
    title = next(element for element in scene["elements"] if element["element_id"] == "title.main")
    title["bbox"] = {"x": 160, "y": 80, "w": 700, "h": 120}
    title["style"] = {"fill": "#ff0000", "font_family": "Arial", "font_weight": "400"}
    title["text_fit"] = {"preferred_size_px": 12, "min_size_px": 9, "max_lines": 5, "line_height": 2.0}
    title["z_index"] = 999
    write_scene(run, scene)
    compile_pptx(run, [scene], {"P001": lock})
    mutated_pptx = _render_pptx_page(run, pptx_path(run), "P001", 0)
    readback_pptx(run, [scene], {"P001": lock}, pptx_path(run))
    with zipfile.ZipFile(pptx_path(run)) as archive:
        mutated_slide_xml = archive.read("ppt/slides/slide1.xml")

    assert sha256_file(svg_path(run, "P001")) == original_svg_hash
    assert mutated_slide_xml == original_slide_xml
    assert sha256_file(mutated_pptx) == sha256_file(original_pptx)


def test_tspan_runs_preserve_text_and_style(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    document = ElementTree.fromstring(svg.read_text(encoding="utf-8"))
    title = next(node for node in document.iter() if node.get("id") == "title.main")
    for child in list(title):
        title.remove(child)
    title.text = None
    namespace = str(title.tag).split("}")[0].removeprefix("{")
    first = ElementTree.SubElement(title, f"{{{namespace}}}tspan", {"x": str(title.get("x")), "dy": "0", "fill": "#c65c42", "font-weight": "400"})
    first.text = "Synthetic "
    second = ElementTree.SubElement(title, f"{{{namespace}}}tspan", {"dy": "0", "fill": "#1f6fd1", "font-weight": "700"})
    second.text = "framework page"
    ElementTree.ElementTree(document).write(svg, encoding="utf-8", xml_declaration=True)
    render_preview(svg, preview_path(run, "P001"))

    compile_pptx(run, [scene], {"P001": lock})

    presentation = Presentation(pptx_path(run))
    shape = next(shape for shape in presentation.slides[0].shapes if shape.name == "title.main")
    runs = list(shape.text_frame.paragraphs[0].runs)
    trace = read_json(trace_path(run))
    title_trace = next(element for element in trace["elements"] if element["element_id"] == "title.main")
    assert "".join(run.text for run in runs) == lock["customer_visible"]["title"]
    assert len(runs) == 2
    assert runs[0].font.bold is False
    assert runs[1].font.bold is True
    assert [run["paint"]["color"] for run in title_trace["runs"]] == ["#c65c42", "#1f6fd1"]


def test_visible_tspan_text_cannot_hide_behind_declared_metadata(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    document = ElementTree.fromstring(svg.read_text(encoding="utf-8"))
    title = next(node for node in document.iter() if node.get("id") == "title.main")
    title.text = "Visible drift"
    ElementTree.ElementTree(document).write(svg, encoding="utf-8", xml_declaration=True)

    with pytest.raises(SvgVisualError, match="visible SVG text drift"):
        validate_approved_svg(svg, scene, lock)
    with pytest.raises(PptxEditabilityError, match="visible SVG text drift"):
        compile_pptx(run, [scene], {"P001": lock})


@pytest.mark.parametrize("paint", ['fill="none"', 'fill-opacity="0"'])
def test_required_tspan_cannot_be_hidden_by_run_paint(tmp_path: Path, paint: str) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    svg.write_text(svg.read_text(encoding="utf-8").replace("<tspan ", f"<tspan {paint} ", 1), encoding="utf-8")

    with pytest.raises(SvgVisualError, match="hidden SVG tspan"):
        validate_approved_svg(svg, scene, lock)


def test_required_text_cannot_be_hidden_by_ancestor_group(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    source = svg.read_text(encoding="utf-8").replace('<g id="page.P001"', '<g id="page.P001" opacity="0"', 1)
    svg.write_text(source, encoding="utf-8")

    with pytest.raises(SvgVisualError, match="ancestor"):
        validate_approved_svg(svg, scene, lock)


def test_tspan_font_size_participates_in_overflow_measurement(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    svg.write_text(svg.read_text(encoding="utf-8").replace("<tspan ", '<tspan font-size="200px" ', 1), encoding="utf-8")

    with pytest.raises(SvgVisualError, match="text overflow"):
        validate_approved_svg(svg, scene, lock)


def test_excessive_text_mask_fails_closed() -> None:
    import numpy as np
    from PIL import Image
    from high_density.visual import _ssim

    image = Image.new("RGB", (64, 64), "white")
    mask = np.ones((64, 64), dtype=bool)

    with pytest.raises(ContractError, match="too few valid pixels"):
        _ssim(image, image, mask)


def test_svg_text_drift_blocks_drawingml_compile(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    title = str(lock["customer_visible"]["title"])
    svg = svg_path(run, "P001")
    svg.write_text(svg.read_text(encoding="utf-8").replace(title, "drifted text"), encoding="utf-8")

    with pytest.raises(PptxEditabilityError, match="SVG text drift"):
        compile_pptx(run, [scene], {"P001": lock})


def test_circle_is_compiled_as_native_shape(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    scene["elements"].append(
        {
            "element_id": "proof.circle",
            "component_id": "component.proof",
            "kind": "circle",
            "role": "proof_marker",
            "priority": "P2",
            "bbox": {"x": 1460, "y": 100, "w": 48, "h": 48},
            "editability_target": "native_shape",
            "asset_policy": "none",
            "style": {"fill": "#419BFD", "stroke": "#1f6fd1", "stroke_width": 1},
        }
    )
    write_scene(run, scene)
    (run / "high_density_build/reviews/P001.visual_review.json").unlink()

    result = run_high_density(run)

    assert result["status"] == "completed"
    trace = read_json(run / "high_density_build/traces/P001.json")
    circle_trace = next(item for item in trace["trace"]["elements"] if item["element_id"] == "proof.circle")
    assert circle_trace["object_type"] == "shape"


def test_unsupported_svg_element_blocks_compile(tmp_path: Path) -> None:
    run, lock, scene = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    original = svg.read_text(encoding="utf-8")
    svg.write_text(original.replace("</g></svg>", '<mask id="unsupported.mask"><rect x="0" y="0" width="10" height="10" /></mask></g></svg>'), encoding="utf-8")

    with pytest.raises(PptxEditabilityError, match="unsupported SVG element mask"):
        compile_pptx(run, [scene], {"P001": lock})


def test_svg_rejects_unsafe_event_attribute(tmp_path: Path) -> None:
    run, _, _ = _prepared_fixture(tmp_path)
    svg = svg_path(run, "P001")
    svg.write_text(svg.read_text(encoding="utf-8").replace("<svg ", '<svg onload="alert(1)" ', 1), encoding="utf-8")

    with pytest.raises(SvgVisualError, match="event attribute"):
        validate_svg(svg, page_id="P001")
