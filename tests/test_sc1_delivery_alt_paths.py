from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from pptx import Presentation
from pptx.util import Inches
from quality.semantic_checks import scan_delivery_pptx


def test_recursive_alt_path_leak_is_found_without_blocking_business_qualifier(tmp_path):
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]);group=slide.shapes.add_group_shape();shape=group.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1));shape.text='建议方案，客户事实待确认'
    shape._element.xpath('.//p:cNvPr')[0].set('descr','/Users/private-client/internal/credentials.md')
    for field in ('author','last_modified_by','comments','keywords','category'):setattr(p.core_properties,field,'')
    path=tmp_path/'delivery.pptx';p.save(path)
    findings=scan_delivery_pptx(path)
    assert any(f['check']=='internal_path_leak' and f['slide']==1 for f in findings)
    shape._element.xpath('.//p:cNvPr')[0].set('descr','业务限定说明')
    p.save(path)
    assert scan_delivery_pptx(path)==[]
