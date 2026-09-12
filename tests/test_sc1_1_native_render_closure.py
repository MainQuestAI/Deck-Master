"""Real renderer output, isolated roots and capability evidence regressions."""
from __future__ import annotations
import json
import subprocess
import sys
from pathlib import Path
from unittest import mock

import pytest
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from native_pptx.api import NativeCompileError, NativeCompileResult, render_pptx
from native_pptx.probe import _probe_compilable


def result_for(tmp_path, count=2):
    pptx = tmp_path / 'deck.pptx'
    deck = Presentation()
    for _ in range(count):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(pptx)
    trace = tmp_path / 'trace.json'
    trace.write_text(json.dumps({'pages': [{'page_id': f'P{i+1:03d}'} for i in range(count)]}))
    return NativeCompileResult('compiled', pptx, trace)


def fake_renderer(cmd, **kwargs):
    assert kwargs['timeout'] > 0
    if '--convert-to' in cmd:
        profile = next(arg for arg in cmd if arg.startswith('-env:UserInstallation='))
        assert profile.endswith('/lo-profile')
        out = Path(cmd[cmd.index('--outdir') + 1])
        (out / (Path(cmd[-1]).stem + '.pdf')).write_bytes(b'%PDF-1.4\nprobe')
    else:
        prefix = Path(cmd[-1])
        # Poppler zero pads filenames in decks with >=10 pages. Exact
        # numeric parsing must not select page 10 when page 1 is missing.
        for i in (1, 2):
            Image.new('RGB', (160,90), 'white').save(prefix.parent / f'{prefix.name}-{i:02d}.png')
    return subprocess.CompletedProcess(cmd, 0, '', '')


def test_render_emits_all_pages_and_uses_isolated_profile(tmp_path):
    result = result_for(tmp_path)
    with mock.patch('shutil.which', side_effect=lambda name: '/bin/' + name), mock.patch('subprocess.run', side_effect=fake_renderer):
        output = render_pptx(result, {'output_root': tmp_path/'rendered', 'timeout_seconds': 7})
    assert output['status'] == 'rendered'
    assert [p['page_id'] for p in output['pages']] == ['P001', 'P002']
    assert all(Path(p['path']).is_file() and len(p['sha256']) == 64 for p in output['pages'])
    assert Path(output['pdf_path']).is_file()


def test_render_rejects_missing_or_extra_page_without_publishing(tmp_path):
    result = result_for(tmp_path, 1)
    with mock.patch('shutil.which', side_effect=lambda name: '/bin/'+name), mock.patch('subprocess.run', side_effect=fake_renderer):
        with pytest.raises(NativeCompileError, match='page set'):
            render_pptx(result, {'output_root': tmp_path/'rendered'})
    assert not (tmp_path/'rendered'/'P001.png').exists()


def test_render_timeout_is_structured_failure(tmp_path):
    result = result_for(tmp_path)
    with mock.patch('shutil.which', return_value='/bin/tool'), mock.patch('subprocess.run', side_effect=subprocess.TimeoutExpired('soffice', 1)):
        with pytest.raises(NativeCompileError) as error:
            render_pptx(result, {'timeout_seconds': 1})
    assert error.value.code == 'NDC_RENDER_TIMEOUT'


def test_compile_probe_must_emit_a_pptx_not_only_parse_svg():
    # An emission failure must defeat readiness even though imports and
    # SVG parsing succeed.
    with mock.patch('pptx.presentation.Presentation.save', side_effect=OSError('save failed')):
        probe = _probe_compilable()
    assert probe['status'] == 'blocked'
    assert 'save failed' in probe['error']


def test_native_contain_keeps_uniform_geometry_font_and_stroke(tmp_path):
    from native_pptx.pptx import _contain_elements
    svg = tmp_path / 'square.svg'
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="10 20 900 900"/>')
    style = {'font_size': '24px', 'stroke_width': 3}
    element = {'element_id': 'text', 'bbox': {'x':10,'y':20,'w':90,'h':90}, 'style':style,
               '_text_lines': [[{'style':style}]], '_native_commands':[{'op':'M','x':10,'y':20}]}
    mapped = _contain_elements([element], svg)[0]
    scale = (1672*9/16)/900
    assert mapped['bbox']['w'] == mapped['bbox']['h'] == 90*scale
    assert mapped['bbox']['x'] == pytest.approx((1672-1672*9/16)/2)
    assert mapped['bbox']['y'] == pytest.approx(0)
    assert float(mapped['style']['font_size'].removesuffix('px')) == pytest.approx(24*scale)
    assert mapped['style']['stroke_width'] == 3*scale
    assert mapped['_native_commands'][0]['x'] == mapped['bbox']['x']
    assert element['bbox']['x'] == 10  # input immutable


def test_explicit_roots_compile_native_without_hd_output_layout(tmp_path):
    sys.path.insert(0, str(ROOT/'tests'))
    import test_high_density_builder as fixtures
    from native_pptx.api import NativeCompileRequest, compile_svg_deck, readback_pptx
    from high_density.svg import validate_approved_svg
    from high_density.scene import load_scene
    from high_density.content import load_content_lock
    from native_pptx.contracts import sha256_file
    import shutil

    run, _ = fixtures._make_run(tmp_path, mode='fixture', page_count=1)
    fixtures._blueprint(run, 'P001')
    fixtures.prepare_high_density(run)
    fixtures.run_high_density(run)
    source = run/'approved'/'P001.svg'
    source.parent.mkdir()
    shutil.copyfile(run/'high_density_build/svg/P001.svg', source)
    scene, lock = load_scene(run, 'P001'), load_content_lock(run, 'P001')
    output = tmp_path/'revision-output'
    result = compile_svg_deck(NativeCompileRequest(root=run, scenes=[scene], locks={'P001':lock},
            validate_approved=validate_approved_svg, svg_paths={'P001':source},
            expected_sha256={'P001':sha256_file(source)}, output_root=output, canvas_mode='native'))
    assert result.pptx_path == output/'deck.pptx'
    assert result.trace_path == output/'pptx_trace.json'
    deck = Presentation(result.pptx_path)
    assert deck.slide_width*9 == deck.slide_height*16
    report = readback_pptx(result, [scene], {'P001':lock})
    assert report['status'] == 'pass'
    assert report['scope'] == 'structural' and report['visual_review'] == 'not_run'
    # The input must not be reread from an unrelated mutable HD projection.
    (run/'high_density_build/svg/P001.svg').unlink()
    assert readback_pptx(result, [scene], {'P001':lock})['status'] == 'pass'
    # A non-16:9 source must compile through the same validator and keep
    # shapes isotropic, rather than stretching x/y independently.
    source.write_text(source.read_text().replace('viewBox="0 0 1672 941"', 'viewBox="0 0 2000 1600"'))
    second = compile_svg_deck(NativeCompileRequest(root=run, scenes=[scene], locks={'P001':lock},
            validate_approved=validate_approved_svg, svg_paths={'P001':source},
            expected_sha256={'P001':sha256_file(source)}, output_root=tmp_path/'contain-output', canvas_mode='native'))
    assert readback_pptx(second, [scene], {'P001':lock})['status'] == 'pass'
    from xml.etree import ElementTree
    rect = next(node for node in ElementTree.parse(source).getroot().iter() if node.tag.endswith('}rect') and node.get('id') and float(node.get('height','0')) > 0)
    emitted = next(shape for shape in Presentation(second.pptx_path).slides[0].shapes if shape.name == rect.get('id'))
    assert emitted.width/emitted.height == pytest.approx(float(rect.get('width'))/float(rect.get('height')), rel=1e-4)


def test_compile_request_rejects_symlink_input_escape(tmp_path):
    from native_pptx.api import NativeCompileRequest
    root = tmp_path/'run'
    root.mkdir()
    outside = tmp_path/'outside.svg'
    outside.write_text('<svg/>')
    (root/'approved.svg').symlink_to(outside)
    request = NativeCompileRequest(root, [{'page_id':'P001'}], {}, svg_paths={'P001':root/'approved.svg'})
    with pytest.raises(NativeCompileError, match='escapes'):
        request.validated()
