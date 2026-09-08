"""Prevent LibreOffice reflow of explicitly positioned native CJK/Latin lines."""
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from xml.etree import ElementTree as ET
import pytest
from pptx import Presentation
from native_pptx.pptx import _add_text,_svg_text_lines,_NATIVE_CANVAS

@pytest.mark.parametrize('native',[True,False])
def test_explicit_native_lines_disable_template_shape_autofit(native):
 node=ET.fromstring('<text x="326" y="281" font-family="Noto Sans SC" font-size="31" fill="#06163e" data-pptx-text="现有ERP与WMS：提供库存事实"><tspan x="326" dy="0" font-size="40" font-weight="700">现有ERP与WMS：</tspan><tspan x="326" dy="56">提供库存事实</tspan></text>')
 e={'element_id':'P004.body0','text':'现有ERP与WMS：提供库存事实','bbox':{'x':326,'y':243,'w':368,'h':140},'style':{'font_size':'31px'},'_text_lines':_svg_text_lines(node,{'gradients':{},'effects':{}},preserve_positions=native)}
 deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);token=_NATIVE_CANVAS.set(native)
 try:_add_text(slide,e,[])
 finally:_NATIVE_CANVAS.reset(token)
 shape=slide.shapes[0];body=shape.text_frame._txBody.bodyPr
 if native:
  assert body.get('wrap')=='none'
  assert body.xpath('./a:noAutofit')
  assert not body.xpath('./a:spAutoFit')
  assert [p.text for p in shape.text_frame.paragraphs]==['现有ERP与WMS：','提供库存事实']
  assert shape.text_frame.paragraphs[0].line_spacing is None
  assert shape.text_frame.paragraphs[1].line_spacing.pt==pytest.approx(56*960/1672,abs=.01)
 else:
  assert body.get('wrap')=='square'
  assert body.xpath('./a:spAutoFit')
