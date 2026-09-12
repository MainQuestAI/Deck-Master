"""Locked logical newlines must not discard SVG positions or mixed styles."""
from xml.etree import ElementTree as ET
from pathlib import Path
import sys

import pytest
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from native_pptx.pptx import (
    _NATIVE_CANVAS, _add_text, _svg_elements, _svg_text_lines, PptxEditabilityError,
)


def example(tmp_path, *, changed=False):
    text = '库存准确率\n数量一致的门店记录数；同刻、同单位比较。'
    node = ET.Element('text', {
        'id': 'P006.metric', 'x': '140', 'y': '358', 'font-family': 'Noto Sans SC',
        'font-size': '36', 'font-weight': '400', 'fill': '#071f45',
        'data-pptx-text': text, 'data-pptx-text-ref': 'content_lock.customer_visible.body_blocks.0.text',
        'data-pptx-bounds': '139,276,1457,132',
    })
    for copy, x, dy, size, weight in [
        ('库存准确率', 141, 0, 54, 700),
        ('数量一致的门店记录数；' if not changed else '数量不一致的门店记录数；', 638, -26, 36, 400),
        ('同刻、同单位比较。', 638, 49, 36, 400),
    ]:
        ET.SubElement(node, 'tspan', {'x': str(x), 'dy': str(dy), 'font-size': str(size), 'font-weight': str(weight)}).text = copy
    root = ET.Element('svg', {'width': '1672', 'height': '941', 'viewBox': '0 0 1672 941'})
    root.append(node)
    path = tmp_path / 'page.svg'
    ET.ElementTree(root).write(path, encoding='unicode')
    element = {'element_id': 'P006.metric', 'kind': 'text', 'text': text,
               'bbox': {'x': 139, 'y': 276, 'w': 1457, 'h': 132}, 'priority': 'P1'}
    return node, path, {'page_id': 'P006', 'elements': [element], 'visual_registry': []}


def test_multiline_parser_preserves_positions_and_tspan_style(tmp_path):
    node, _, _ = example(tmp_path)
    runs = [run for line in _svg_text_lines(node, {'gradients': {}, 'effects': {}}) for run in line]
    assert [run['position']['x'] for run in runs] == [141, 638, 638]
    assert [run['position']['y'] for run in runs] == [358, 332, 381]
    assert runs[0]['style']['font_weight'] == '700'


def test_real_consumer_emits_positioned_multiline_chunks(tmp_path):
    _, path, scene = example(tmp_path)
    token = _NATIVE_CANVAS.set(True)
    try:
        element = _svg_elements(path, scene)[0]
        assert len(element['_text_lines']) == 3
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        trace = []
        _add_text(slide, element, trace)
        group = slide.shapes[0]
        assert [shape.text for shape in group.shapes] == ['库存准确率', '数量一致的门店记录数；', '同刻、同单位比较。']
        assert [shape.left / 12700 for shape in group.shapes] == pytest.approx([141 * 960 / 1672, 638 * 960 / 1672, 638 * 960 / 1672], abs=.001)
        assert trace[0]['object_type'] == 'group'
    finally:
        _NATIVE_CANVAS.reset(token)


def test_native_consumer_never_silently_replaces_invalid_positioned_text(tmp_path):
    _, path, scene = example(tmp_path, changed=True)
    token = _NATIVE_CANVAS.set(True)
    try:
        with pytest.raises(PptxEditabilityError, match='visible SVG text'):
            _svg_elements(path, scene)
    finally:
        _NATIVE_CANVAS.reset(token)
