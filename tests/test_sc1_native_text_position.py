from xml.etree import ElementTree as ET
from native_pptx.pptx import _svg_text_lines

def test_locked_semantic_text_does_not_flatten_explicit_visual_lines():
    node=ET.fromstring('<text id="P004.body1" x="471" y="653" font-family="Noto Sans SC" font-size="33" fill="#06163e" data-pptx-text="库存服务：可售量计算与预占"><tspan x="551.5" dy="0">库存服务：</tspan><tspan x="502" dy="44.55">可售量计算与预占</tspan></text>')
    lines=_svg_text_lines(node,{'gradients':{},'effects':{}})
    assert [''.join(r['text'] for r in line) for line in lines]==['库存服务：','可售量计算与预占']
    assert lines[0][0]['position']=={'x':551.5,'y':653.0,'anchor':'start'}
    assert lines[1][0]['position']=={'x':502.0,'y':697.55,'anchor':'start'}

import pytest
from pptx import Presentation
from pptx.util import Inches
from native_pptx.pptx import _add_text, _contain_elements, _NATIVE_CANVAS

@pytest.mark.parametrize('anchor,x,alignment,margin',[('start',150,'l',50),('middle',200,'ctr',0),('end',250,'r',50)])
def test_anchor_and_explicit_baseline_survive_uniform_mapping(tmp_path,anchor,x,alignment,margin):
    node=ET.fromstring(f'<text id="caption" x="{x}" y="130" text-anchor="{anchor}" font-family="Arial" font-size="20" fill="#000000" data-pptx-text="firstsecond"><tspan x="{x}" dy="0">first</tspan><tspan x="{x}" dy="30">second</tspan></text>')
    element={'element_id':'caption','kind':'text','bbox':{'x':100,'y':110,'w':200,'h':80},'style':{'font_size':'20px'},'text':'firstsecond','_text_lines':_svg_text_lines(node,{'gradients':{},'effects':{}})}
    svg=tmp_path/'canvas.svg';svg.write_text('<svg viewBox="0 0 800 800"/>')
    mapped=_contain_elements([element],svg)[0]
    scale=940.5/800;dx=(1672-800*scale)/2
    assert mapped['_text_lines'][0][0]['position']['x']==pytest.approx(x*scale+dx)
    deck=Presentation();deck.slide_width=Inches(40/3);slide=deck.slides.add_slide(deck.slide_layouts[6])
    token=_NATIVE_CANVAS.set(True)
    try:_add_text(slide,mapped,[])
    finally:_NATIVE_CANVAS.reset(token)
    shape=slide.shapes[0];paragraphs=shape.text_frame.paragraphs
    assert [p.text for p in paragraphs]==['first','second']
    assert shape.text_frame.word_wrap is False
    assert paragraphs[0]._p.pPr.get('algn')==alignment
    key='marR' if anchor=='end' else 'marL'
    assert int(paragraphs[0]._p.pPr.get(key))==pytest.approx(margin*scale*960/1672*12700,abs=1)
    assert paragraphs[0].line_spacing.pt==pytest.approx(30*scale*960/1672,abs=.01)

def test_legacy_character_flattening_is_unchanged():
    node=ET.fromstring('<text id="legacy" font-size="20" fill="#000000" data-pptx-text="甲乙"><tspan dy="0">甲</tspan><tspan dy="30">乙</tspan></text>')
    lines=_svg_text_lines(node,{'gradients':{},'effects':{}},preserve_positions=False)
    assert len(lines)==1
    assert ''.join(r['text'] for r in lines[0])=='甲乙'

from pathlib import Path
import json
import shutil
from native_pptx.pptx import _svg_elements

def test_actual_p004_reconstruction_keeps_four_card_line_positions(tmp_path):
    fixture=Path(__file__).parent/'fixtures/native_text_fidelity'
    target=tmp_path/'high_density_build/svg/P004.svg';target.parent.mkdir(parents=True)
    shutil.copyfile(fixture/'P004.svg',target)
    scene=json.loads((fixture/'P004.scene.json').read_text())
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
    token=_NATIVE_CANVAS.set(True)
    try:
        elements=_contain_elements(_svg_elements(target,scene),target)
        for element in elements:
            if element['element_id'].startswith('P004.body'):
                _add_text(slide,element,[])
    finally:_NATIVE_CANVAS.reset(token)
    assert [s.text for s in slide.shapes]==['现有ERP与WMS：\n提供库存事实','库存服务：\n可售量计算与预占','订单服务：\n承诺校验与履约编排','门店与配送：\n执行回传与异常处理']
    assert all(not s.text_frame.word_wrap for s in slide.shapes)
    assert int(slide.shapes[1].text_frame.paragraphs[0]._p.pPr.get('marL')) > int(slide.shapes[1].text_frame.paragraphs[1]._p.pPr.get('marL'))
