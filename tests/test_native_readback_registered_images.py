"""Native readback must classify real registered picture objects consistently."""
import sys
from pathlib import Path
from xml.etree import ElementTree
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'scripts'), str(ROOT/'tests')]
import test_high_density_builder as fixtures
from high_density.content import load_content_lock
from high_density.scene import load_scene
from high_density.svg import compile_svg, validate_approved_svg
from native_pptx.api import NativeCompileRequest, compile_svg_deck, readback_pptx
from native_pptx.contracts import sha256_file
from native_pptx.pptx import _flatten_shape_objects


@pytest.mark.parametrize('extra_attributes', [{}, {'rx': '10', 'ry': '10'}])
def test_registered_picture_is_counted_as_image_not_shape(tmp_path, extra_attributes):
    run, _ = fixtures._make_run(tmp_path, mode='fixture', page_count=1)
    fixtures._blueprint(run, 'P001')
    fixtures.prepare_high_density(run)
    fixtures.run_high_density(run)
    scene, lock = load_scene(run, 'P001'), load_content_lock(run, 'P001')
    asset = run/'registered.png'
    Image.new('RGB', (100, 50), '#8b532f').save(asset)
    scene['elements'].append({
        'element_id':'asset.proof', 'kind':'image', 'role':'proof_image', 'priority':'P2',
        'bbox':{'x':1450,'y':850,'w':100,'h':50}, 'asset_ref':'proof',
        'asset_sha256':sha256_file(asset), 'editability_target':'registered_asset', 'asset_policy':'registered',
    })
    svg = run/'registered.svg'
    compile_svg(scene, svg, assets={'proof':asset})
    document = ElementTree.parse(svg)
    for node in document.getroot().iter():
        if node.tag.rsplit('}', 1)[-1] == 'image':
            node.attrib.update(extra_attributes)
    document.write(svg, encoding='unicode')
    result = compile_svg_deck(NativeCompileRequest(root=run, scenes=[scene], locks={'P001':lock},
        asset_paths_by_page={'P001':{'proof':asset}}, validate_approved=validate_approved_svg,
        svg_paths={'P001':svg}, expected_sha256={'P001':sha256_file(svg)},
        output_root=tmp_path/'native', canvas_mode='native'))
    report = readback_pptx(result, [scene], {'P001':lock})['pages'][0]
    shapes = _flatten_shape_objects(Presentation(result.pptx_path).slides[0].shapes)
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert [shape.name for shape in pictures] == ['asset.proof']
    assert report['image_objects'] == len(pictures) == 1
    assert sum(report[key] for key in ('image_objects','shape_objects','text_objects','group_objects')) == len(shapes)
