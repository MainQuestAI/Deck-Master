import sys
from pathlib import Path
from xml.etree import ElementTree as ET
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from native_pptx.svg_native import parse_svg_native, SvgNativeError

def parse(attrs='',group=''):
 return parse_svg_native(ET.fromstring(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200"><g {group}><rect id="panel" x="20" y="30" width="200" height="100" fill="#ffffff" stroke="#000000" stroke-width="2" {attrs}/></g></svg>'))['elements'][0]

def test_rect_radius_is_geometry_not_default_office_radius():
 e=parse('rx="10"')
 assert e['normalized_tag']=='path'
 assert sum(c['op']=='C' for c in e['commands'])==4
 assert e['commands'][0]=={'op':'M','x':30.0,'y':30.0}
 assert e['bbox']=={'x':20.0,'y':30.0,'w':200.0,'h':100.0}

def test_elliptical_radius_clamped_and_transformed():
 e=parse('rx="150" ry="20"','transform="translate(5 7)"')
 assert e['commands'][0]=={'op':'M','x':125.0,'y':37.0}
 assert sum(c['op']=='C' for c in e['commands'])==4

def test_dash_array_inherits_and_emits_exact_stroke_width_ratios():
 from native_pptx.pptx import _native_style,_apply_line_style
 from pptx import Presentation
 from pptx.enum.shapes import MSO_SHAPE
 e=parse('rx="10"','stroke-dasharray="11 7"')
 assert e['style']['stroke-dasharray']=='11 7'
 s=Presentation().slides.add_slide(Presentation().slide_layouts[6]).shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,100000,100000)
 _apply_line_style(s,_native_style(e['style']))
 ds=s._element.xpath('.//a:custDash/a:ds')
 assert len(ds)==1 and ds[0].get('d')=='550000' and ds[0].get('sp')=='350000'

@pytest.mark.parametrize('attrs',['rx="-1"','stroke-dasharray="11 -7"','stroke-dasharray="10% 5%"','stroke-dasharray="11 7" stroke-dashoffset="3"'])
def test_unsupported_or_invalid_style_not_silently_dropped(attrs):
 with pytest.raises(SvgNativeError):parse(attrs)
