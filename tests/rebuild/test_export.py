"""T15 export evidence: review/delivery purposes (AC-R09), honest status
reporting (T15.03), and editability claims (AC-K16 export side)."""
import copy
import json
import uuid
from pathlib import Path

import pytest

import deck_master.cli as cli
import deck_master.service as service
from deck_master.editing import export_project, review_status
from deck_master.models import bump_revision, sha256_bytes
from deck_master import collaboration, delivery_eligibility, exports
from rc3_fixtures import authorize_scope, evidence_for, public_project
from deck_master.store import Store, StoreError

REQUIRED = ('content', 'blueprint_content', 'blueprint_fidelity', 'conversion',
            'readability', 'privacy')


def export_scope(*, layers=None):
    return {'subject': 'deck', 'page_ids': ['p1'], 'page_manifest_ref': None,
            'external_artifact_ref': None,
            'layers': list(layers or ['narrative', 'plan', 'page_copy', 'speech', 'design',
                                      'blueprint', 'svg', 'review', 'output']),
            'allow_page_membership_change': False, 'allow_page_order_change': False,
            'purpose': 'working'}


def _authorize_delivery(store):
    """Explicit L1 test instruction, scoped execute grant and business consent.

    This is authorization for synthetic test operations, never human review.
    Refresh it explicitly after a policy change, whose basis invalidates consent.
    """
    project = store.project_root
    evidence = evidence_for(project, text=(
        'L1 synthetic instruction: review this one-page fixture and deliver its '
        'selected output only when current checks and configured professional '
        'policy are satisfied. No human or native execution is asserted.'),
        locator='synthetic-l1:test_export:scoped-delivery')
    authorize_scope(project, actions=['export.delivery', 'handoff.delivery',
        'delivery.historical', 'studio.deliver', 'dispatch.host_task', 'review.record',
        'review.import', 'change.commit', 'candidate.adopt'], scope=export_scope(),
        max_external_calls=0, max_candidate_count=0,
        adoption_policy='within_scope', evidence=evidence)
    document = store.load_document()
    collaboration.record_decision(project, input={
        'topic': 'delivery_acceptance', 'scope': export_scope(layers=['output']),
        'subject_refs': [],
        'value': {'scenario': 'L1 synthetic scoped delivery only'},
        'disposition': 'accepted', 'authority_mode': 'direct_user',
        'instruction_evidence': evidence}, base_revision=document['revision_id'],
        operation_id='synthetic-consent-' + uuid.uuid4().hex)


def projection_subject(store, document, *, notes_policy=None):
    """Materialize the exact sanitized public .pptx that the review names."""
    packet = exports.create(store.project_root, revision=document['revision_id'],
                            purpose='review', notes_policy=notes_policy)
    data = exports.download(store.project_root, export_id=packet['export_id'],
                            filename='deck.pptx')[0]
    digest = sha256_bytes(data)
    subject = {'path': f'.deckmaster/objects/{digest[:2]}/{digest}.pptx', 'sha256': digest}
    assert store.read_object_bytes(subject) == data
    return subject


def selected_professional_policy(store, document=None, *, notes_policy=None):
    document = document or store.load_document()
    return delivery_eligibility.professional_policy(store, document,
        selected_output=delivery_eligibility.selected_output(store, document,
                                                             notes_policy=notes_policy))


def _deck(tmp_path: Path, *, notes='') -> tuple[Path, Store]:
    # Explicit current public intake and exact L1 scope. Seeded engineering
    # artifacts below are synthetic, never claimed as actual production.
    fixture = public_project(tmp_path, create_input={'brief':
        'L1 synthetic export fixture: one fixed page, scoped review and delivery; no image generation',
        'draft': {'pages': [{
            'schema_version': 'deck_page_package.v2', 'page_id': 'p1',
            'speaker_notes': notes,
            'customer_visible': {'title': '导出一页', 'body_blocks': [
                {'id': 'b1', 'type': 'paragraph', 'text': '正文。'}]},
            'visual_spec': {'intent': 'export', 'reference_mode': 'new_design'}}]}},
        request_input={'intent': 'full_deck', 'content_intent': None, 'scope': export_scope()})
    project = fixture['project']
    store = Store(project)
    work = store.staging_dir / 'pptx'
    work.mkdir(parents=True, exist_ok=True)
    pptx_file = work / 'deck.pptx'
    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(100000, 100000, 5000000, 1000000)
    box.text_frame.text = '导出一页\n正文。'
    presentation.save(pptx_file)
    from deck_master.compiler.native import attach_notes
    page = store.read_object_json(store.load_document()['pages'][0]['page'])
    pptx_file.write_bytes(attach_notes(pptx_file.read_bytes(),
        [{'page_id': 'p1', 'notes': page.get('speaker_notes', '')}]))
    svg_file = work / 'page.svg'
    svg_file.write_text('<svg viewBox="0 0 10 10"/>')
    report_file = work / 'readback.json'
    report_file.write_text(json.dumps({'status': 'pass', 'findings': [], 'pages': []}, ensure_ascii=False))
    from deck_master.pipeline import artifact as adopt_artifact
    document = store.load_document()
    bumped = bump_revision(document, {'operation_id': 'attach-pptx', 'kind': 'task_update',
                                      'description': 'outputs', 'read_set': []})
    bumped['outputs']['pptx'] = adopt_artifact(store, pptx_file, 'pptx')
    bumped['outputs']['render_report'] = adopt_artifact(store, report_file, 'render_report')
    bumped['pages'][0]['svg'] = adopt_artifact(store, svg_file, 'svg', page_id='p1')
    bumped['pages'][0]['ppt_preview'] = adopt_artifact(store, svg_file, 'ppt_preview', page_id='p1')
    store.commit_change(base_revision=document['revision_id'], document=bumped, operation_id='attach-pptx')
    _authorize_delivery(store)
    assert store.load_document()['compatibility']['minimum_writer'] == 'rc3-entry.v1'
    return project, Store(project)


def qualify_synthetic_manuscript(store):
    """L1 host_self fidelity review via the actual dispatch/claim/accept API."""
    from deck_master import content_review
    project = store.project_root
    dispatched = content_review.dispatch_preproduction_review(project)
    assert dispatched['status'] == 'review_dispatched'
    assert dispatched['requirement'] == 'fidelity_only'
    task = dispatched['task']
    service.task_start(project, task_id=task['task_id'],
        execution_ref='synthetic-l1-fidelity-host',
        supported_protocols=[task['protocol_version']], capabilities=task['required_capabilities'])
    document = store.load_document()
    subjects = [document['content_plan'], *[entry['page'] for entry in document['pages']]]
    deps = [{'kind': 'content', 'identity': 'content_plan',
             'sha256': document['content_plan']['sha256']}]
    deps.extend({'kind': 'content', 'identity': 'page:' + entry['page_id'],
                 'sha256': entry['page']['sha256']} for entry in document['pages'])
    evidence = store.put_json_object({'scenario': 'L1 synthetic host_self fidelity only'})
    reviews = [{'schema_version': 'deck_review.v2', 'review_id': 'l1-fidelity-' + kind,
        'kind': kind, 'review_stage': 'preproduction', 'status': 'pass',
        'subjects': subjects, 'dependencies': deps,
        'reviewer': {'type': 'host_self', 'id': 'synthetic-l1-fidelity-host',
            'execution_ref': 'synthetic-l1-fidelity-host', 'independence_confirmed': False,
            'identity_receipt': None, 'evidence_ref': evidence, 'assurance': 'host_attested'},
        'observations': ['L1 synthetic fidelity check, no independent or human review.'],
        'findings': [], 'created_at': '2026-10-15T00:00:00Z', 'replaces': None}
        for kind in content_review.required_units(store, document)]
    result = service.accept_result(project,
        **{key: task[key] for key in ('task_id', 'operation_id', 'produced_against')},
        result_payload={'kind': 'review', 'files': [], 'reviews': reviews})
    assert result['status'] == 'accepted'
    assert content_review.qualify_for_production(store, store.load_document())['allowed']


def professional_dependencies(store, document, *, subject=None):
    """The runtime minimum dependency set for a professional review: the
    exact selected public PPTX file plus every page content digest. An empty array is
    never a complete dependency list (AC-P02)."""
    deps = [{'kind': 'content', 'identity': f"page:{entry['page_id']}",
             'sha256': entry['page']['sha256']}
            for entry in document['pages'] if entry.get('page')]
    pptx_ref = document['outputs'].get('pptx')
    if pptx_ref:
        selected = subject or delivery_eligibility.selected_output(store, document)['subject_ref']
        deps.append({'kind': 'artifact', 'identity': 'pptx', 'sha256': selected['sha256']})
    return deps


def human_professional_reviewer(document, *, event_id='evt-human-review-1',
                                reviewer_id='reviewer-chen',
                                adapter_id='synthetic-unit-fixture', subjects=None):
    """Controlled L1 receipt fixture (CONTRACTS §12.1): proves the predicate
    combination only. This is a synthetic/unit fixture and must never be
    presented as a real human having reviewed real files."""
    subjects = subjects if subjects is not None else [document['outputs']['pptx']] + [
        entry['page'] for entry in document['pages'] if entry.get('page')]
    return {
        'type': 'human_external', 'id': reviewer_id, 'execution_ref': None,
        'independence_confirmed': True,
        'identity_receipt': {
            'adapter_id': adapter_id, 'event_id': event_id,
            'issued_at': '2026-09-21T08:00:00Z', 'reviewer_id': reviewer_id,
            'project_id': document['project_id'], 'subjects': subjects,
        },
    }


def _review(store, document, kind, status, *, seed, findings=(), reviewer=None,
            observations=None, dependencies=None, subjects=None):
    """Callers state who reviewed and what they executed. A seed only names the
    explicit L1 synthetic record: seed="human" on a host_self record was the old false positive and
    is gone (T01). Defaults name an explicit host_self reviewer."""
    page_ref = document['pages'][0]['page']
    subject = projection_subject(store, document) if kind == 'professional_use' else None
    selected_subjects = list(subjects if subjects is not None else ([subject, page_ref] if subject else [document['outputs']['pptx'], page_ref]))
    reviewer = copy.deepcopy(reviewer)
    if reviewer and reviewer.get('identity_receipt') and subjects is None:
        reviewer['identity_receipt']['subjects'] = selected_subjects
    review = {
        'schema_version': 'deck_review.v1', 'review_id': f'rv-{seed}', 'kind': kind, 'status': status,
        'subjects': selected_subjects,
        'dependencies': list(dependencies if dependencies is not None
                             else professional_dependencies(store, document, subject=subject or
                                 store.read_object_json(document['outputs']['pptx'])['file'])),
        'reviewer': reviewer or {'type': 'host_self', 'id': f'host-{seed}', 'execution_ref': None,
                                 'independence_confirmed': False},
        'observations': list(observations if observations is not None else ['实际检查记录']),
        'findings': list(findings),
        'created_at': '2026-09-21T00:00:00Z', 'replaces': None,
    }
    ref = store.put_json_object(review)
    bumped = bump_revision(store.load_document(), {'operation_id': f'rv-{seed}', 'kind': 'task_update',
                                                   'description': 'review', 'read_set': []})
    bumped['reviews'] = list(bumped.get('reviews') or []) + [ref]
    store.commit_change(base_revision=document['revision_id'], document=bumped, operation_id=f'rv-{seed}')


def _must_fix():
    return [{'finding_id': 'f-1', 'kind': 'conversion', 'impact': 'must_fix', 'page_id': 'p1',
             'element_refs': ['atom:p1:block:b1:text'], 'message': '正文缺失', 'expected': '有',
             'actual': '无', 'evidence': [], 'resolution': 'open'}]


def _passing_deck(tmp_path, *, notes=''):
    project, store = _deck(tmp_path, notes=notes)
    for kind in REQUIRED:
        _review(store, store.load_document(), kind, 'pass', seed=f'pass-{kind}')
    return project, store


# ---------------------------------------------------------------------------
# AC-R09: review shows failing work honestly; delivery refuses unresolved fixes.


def test_review_export_of_failing_deck_marks_unresolved(tmp_path):
    project, store = _deck(tmp_path)
    document = store.load_document()
    _review(store, document, 'conversion', 'fail', seed='fail', findings=_must_fix())

    outcome = export_project(project, output_dir=tmp_path / 'review-out', purpose='review')
    assert outcome['status'] == 'exported'
    report = json.loads((tmp_path / 'review-out' / 'delivery.json').read_text('utf-8'))
    assert report['purpose'] == 'review'
    assert report['revision_id'] == store.load_document()['revision_id']
    assert report['review_status'] == 'fail'
    assert report['unresolved']['failed_dimensions'] == ['conversion:p1']
    assert sorted(report['unresolved']['missing_dimensions']) == sorted(
        f'{kind}:p1' for kind in REQUIRED if kind != 'conversion')
    # Exported bytes are the real current objects.
    selected = delivery_eligibility.selected_output(store, store.load_document())
    assert sha256_bytes((tmp_path / 'review-out' / 'deck.pptx').read_bytes()) == selected['file_sha256']
    assert report['editability'] == 'editable_shapes_and_text'
    assert report['evidence_level'] == 'engineering'


def test_delivery_export_rejected_with_specific_reason(tmp_path):
    project, store = _deck(tmp_path)
    document = store.load_document()
    _review(store, document, 'conversion', 'fail', seed='fail', findings=_must_fix())
    with pytest.raises(StoreError, match='delivery requires every required check'):
        export_project(project, output_dir=tmp_path / 'delivery-out', purpose='delivery')
    assert not (tmp_path / 'delivery-out').exists(), 'a refused export leaves no partial package'


def test_working_alias_matches_review_semantics(tmp_path):
    project, store = _deck(tmp_path)
    document = store.load_document()
    _review(store, document, 'readability', 'fail', seed='fail', findings=_must_fix())
    export_project(project, output_dir=tmp_path / 'alias-out', purpose='working')
    report = json.loads((tmp_path / 'alias-out' / 'delivery.json').read_text('utf-8'))
    assert report['purpose'] == 'review', 'working is an accepted alias, canonically reported as review'
    assert report['review_status'] == 'fail'
    assert report['unresolved']['failed_dimensions'] == ['readability:p1']


def test_review_and_delivery_of_passing_current_deck(tmp_path):
    project, store = _passing_deck(tmp_path)
    assert review_status(store, store.load_document()) == 'pass'
    review_out = export_project(project, output_dir=tmp_path / 'pass-review', purpose='review')
    delivery_out = export_project(project, output_dir=tmp_path / 'pass-delivery', purpose='delivery')
    for outcome in (review_out, delivery_out):
        report = json.loads((Path(outcome['output_dir']) / 'delivery.json').read_text('utf-8'))
        assert report['review_status'] == 'pass'
        assert report['unresolved']['missing_dimensions'] == []
        assert report['unresolved']['failed_dimensions'] == []
    assert (tmp_path / 'pass-delivery' / 'deck.pptx').is_file()
    assert not (tmp_path / 'pass-delivery' / 'project').exists(), 'delivery packs the deck, not the portable project'
    assert not (tmp_path / 'pass-review' / 'project').exists()
    export_project(project, output_dir=tmp_path/'internal', purpose='engineering')
    assert Store(tmp_path/'internal'/'project').load_document()==store.load_document()


def test_adopting_an_existing_svg_candidate_invalidates_delivery_with_named_gaps(tmp_path):
    """The audit's before/after pair: a passing deck exports, then adoption refuses."""
    import uuid as uuid_mod
    from deck_master import candidates as candidates_mod, exports
    from candidate_fixtures import accept, dispatch, plan as plan_adoption
    project, store = _passing_deck(tmp_path)
    # The candidate/changes machinery works on explicit workbench.v3 projects, and
    # an SVG trial needs the current original image to exist.
    from PIL import Image
    from deck_master.pipeline import artifact as adopt_artifact
    document = store.load_document()
    work = store.staging_dir / 'workbench-basis'; work.mkdir(parents=True, exist_ok=True)
    image_file = work / 'page.png'; Image.new('RGB', (960, 540), 'white').save(image_file)
    bumped = bump_revision(document, {'operation_id': 'workbench-basis', 'kind': 'task_update',
                                      'description': 'workbench basis', 'read_set': []})
    bumped['pages'][0]['blueprint'] = adopt_artifact(store, image_file, 'blueprint', page_id='p1')
    store.commit_change(base_revision=document['revision_id'], document=bumped, operation_id='workbench-basis')
    assert exports.create(project, purpose='delivery')['export_id']
    qualify_synthetic_manuscript(store)
    authorize_scope(project, actions=['change.commit'], scope=export_scope(),
        max_external_calls=0, max_candidate_count=1, adoption_policy='candidate_only',
        evidence=evidence_for(project,
            text='L1 explicit instruction: create one SVG comparison candidate without changing current artifacts.',
            locator='L1:test_export:single-svg-trial'))
    task = dispatch(store, 'p1')
    service.task_start(project, task_id=task['task_id'], execution_ref='synthetic-l1-candidate-host',
        supported_protocols=[task['protocol_version']], capabilities=task['required_capabilities'])
    candidate_id = accept(store, task)['candidate_ids'][0]
    value = plan_adoption(store, [candidate_id])
    candidates_mod.adopt(store.project_root, input=value, base_revision=value['base_revision'],
                         operation_id=str(uuid_mod.uuid4()))
    adopted = store.load_document()
    assert adopted['pages'][0]['svg'] == value['selections'][0]['result_ref'], 'adoption replaces the page SVG'
    with pytest.raises(exports.ExportError) as failure:
        exports.create(project, purpose='delivery')
    gaps = failure.value.payload()['error']['gaps']
    reasons = {(gap['layer'], gap['reason']) for gap in gaps}
    assert any(layer in ('deck_outputs', 'ppt', 'pptx') for layer, _ in reasons), reasons
    assert any(reason in ('not_recorded', 'basis_changed') for _, reason in reasons), reasons
    assert exports.create(project, purpose='review')['export_id'], 'a review copy stays available'


def test_missing_human_reviews_reported_not_evaluated(tmp_path):
    # Six required checks pass but no human/professional review exists: the
    # export must say not_evaluated, never a fabricated human pass.
    project, store = _passing_deck(tmp_path)
    export_project(project, output_dir=tmp_path / 'human-missing', purpose='delivery')
    report = json.loads((tmp_path / 'human-missing' / 'delivery.json').read_text('utf-8'))
    assert report['professional_evidence'] == {
        'professional_use': 'not_evaluated', 'desktop_editing': 'not_evaluated'}
    assert report['desktop_editing'] == 'not_evaluated'


def test_recorded_human_review_status_is_reported_verbatim(tmp_path):
    project, store = _passing_deck(tmp_path)
    _review(store, store.load_document(), 'professional_use', 'pass', seed='recorded-professional')
    export_project(project, output_dir=tmp_path / 'human-present', purpose='delivery')
    report = json.loads((tmp_path / 'human-present' / 'delivery.json').read_text('utf-8'))
    assert report['professional_evidence']['professional_use'] == 'pass'
    assert report['professional_evidence']['desktop_editing'] == 'not_evaluated'


def test_cli_export_exit_zero_is_not_professional_pass(tmp_path, capsys):
    project, store = _deck(tmp_path)
    document = store.load_document()
    _review(store, document, 'conversion', 'fail', seed='fail', findings=_must_fix())
    exit_code = cli.main(['export', '--project', str(project),
                          '--out', str(tmp_path / 'cli-out'), '--purpose', 'review'])
    assert exit_code == 0, 'exporting a failing deck for review is allowed'
    report = json.loads((tmp_path / 'cli-out' / 'delivery.json').read_text('utf-8'))
    assert report['review_status'] == 'fail'
    assert report['evidence_level'] == 'engineering'
    assert report['unresolved']['failed_dimensions'], 'exit 0 carries the real unresolved list'
    capsys.readouterr()


def test_export_reports_unknown_editability_for_unverified_legacy_pptx(tmp_path):
    project, store = _deck(tmp_path)
    legacy_bytes = store.read_object_bytes(store.read_object_json(store.load_document()['outputs']['pptx'])['file'])
    file_ref = store.put_blob(legacy_bytes, ext='pptx')
    legacy_artifact = {
        'schema_version': 'deck_artifact.v1', 'artifact_id': 'legacy-1', 'page_id': None,
        'role': 'pptx', 'file': file_ref, 'media_type':
            'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'created_at': '2026-09-21T00:00:00Z', 'dependencies': [], 'derived_from': [],
        'provenance': {'source_type': 'user_supplied'}, 'limitations': [],
        'editability': 'unknown',
    }
    ref = store.put_json_object(legacy_artifact)
    document = store.load_document()
    bumped = bump_revision(document, {'operation_id': 'legacy-out', 'kind': 'task_update',
                                      'description': 'legacy', 'read_set': []})
    bumped['outputs']['pptx'] = ref
    store.commit_change(base_revision=document['revision_id'], document=bumped, operation_id='legacy-out')
    export_project(project, output_dir=tmp_path / 'legacy-out', purpose='review')
    report = json.loads((tmp_path / 'legacy-out' / 'delivery.json').read_text('utf-8'))
    assert report['editability'] == 'unknown', \
        'an unverified legacy file is exported with its real unknown editability'


def _set_policy(store, **updates):
    document = store.load_document()
    operation = 'synthetic-policy-' + uuid.uuid4().hex
    bumped = bump_revision(document, {'operation_id': operation, 'kind': 'policy_update',
                                      'description': 'Explicit L1 synthetic policy change', 'read_set': []})
    bumped['policy'] = {**bumped['policy'], **updates}
    store.commit_change(base_revision=document['revision_id'], document=bumped, operation_id=operation)
    _authorize_delivery(store)


def test_delivery_honors_professional_review_required_policy(tmp_path, synthetic_human_adapter):
    # Positive: six passing engineering checks are not enough when the policy
    # demands professional review, and the audit's old false positive — a
    # host_self professional_use pass whose seed merely said "human" — must
    # stay blocked (AC-P02). Only a receipt-carrying human review unlocks it.
    project, store = _passing_deck(tmp_path)
    _set_policy(store, professional_review_required_for_delivery=True)
    with pytest.raises(StoreError, match='verified human professional review'):
        export_project(project, output_dir=tmp_path / 'needs-human', purpose='delivery')
    assert not (tmp_path / 'needs-human').exists()

    _review(store, store.load_document(), 'professional_use', 'pass', seed='self-named-human')
    with pytest.raises(StoreError, match='verified human professional review') as failure:
        export_project(project, output_dir=tmp_path / 'self-named', purpose='delivery')
    assert 'professional_review_required' in failure.value.payload()['error']['reason_codes']
    reasons = set(selected_professional_policy(store)['gap_reasons'])
    assert 'reviewer_not_human' in reasons
    assert not (tmp_path / 'self-named').exists()

    document = store.load_document()
    _review(store, document, 'professional_use', 'pass', seed='human-verified',
            reviewer=human_professional_reviewer(document))
    outcome = export_project(project, output_dir=tmp_path / 'with-human', purpose='delivery')
    assert outcome['status'] == 'exported'
    report = json.loads((tmp_path / 'with-human' / 'delivery.json').read_text('utf-8'))
    assert report['professional']['required'] is True
    assert report['professional']['satisfied'] is True
    assert report['professional']['human_signoff'] is True
    assert report['professional_evidence']['professional_use'] == 'pass'

    # A failing professional review blocks delivery even earlier: the shared
    # interpretation already fails the deck (refused, destination cleaned).
    project2, store2 = _passing_deck(tmp_path / 'second')
    _set_policy(store2, professional_review_required_for_delivery=True)
    _review(store2, store2.load_document(), 'professional_use', 'fail', seed='recorded-fail',
            findings=_must_fix())
    with pytest.raises(StoreError, match='verified human professional review'):
        export_project(project2, output_dir=tmp_path / 'second' / 'fail-human', purpose='delivery')
    assert not (tmp_path / 'second' / 'fail-human').exists()

    # The policy only gates delivery; review export stays open and says so (07.7).
    review_outcome = export_project(project2, output_dir=tmp_path / 'second' / 'review-ok', purpose='review')
    review_report = json.loads((Path(review_outcome['output_dir']) / 'delivery.json').read_text('utf-8'))
    assert review_report['professional']['required'] is True
    assert review_report['professional']['satisfied'] is False
    assert review_report['professional']['human_signoff'] is False


@pytest.mark.render
def test_cli_needs_tool_maps_to_exit_3(tmp_path, capsys):
    from deck_master.cli import _fail
    from deck_master.pipeline import NeedsTool
    assert _fail(NeedsTool('soffice unavailable')) == 3
    payload = json.loads(capsys.readouterr().err)
    assert payload['error']['code'] == 'needs_tool'
