"""D08 selected historical public-file delivery, using real sanitized PPTX bytes.

Engineering artifacts and the registered human adapter are explicit L1
synthetic fixtures. Public requests/grants, review dispatch/claim/accept,
external byte readback, and delivery gates are real runtime calls. No test
here is evidence of an actual human, native image, or desktop session.
"""
import io
import json
import uuid
from pathlib import Path

import pytest

from deck_master import collaboration, delivery_eligibility, eligibility, exports, external_review
from deck_master.compiler.native import attach_notes, notes_records
from deck_master.models import bump_revision, canonical_json_bytes, sha256_bytes
from deck_master.pipeline import artifact as adopt_artifact
from rc3_fixtures import evidence_for, source_snapshot
from test_delivery_professional_gate import _accept_review, _professional_review
from test_export import (_passing_deck, _set_policy, export_scope,
                         human_professional_reviewer, professional_dependencies)

NOTES = 'L1 synthetic speaker cue: retain the qualification, then ask for the next owner.'


def _fixed_public_file(store, root, *, revision=None, notes_policy='pptx_notes', name='selected'):
    artifact = store.load_document(revision) if revision else store.load_document()
    packet = exports.create(store.project_root, purpose='review',
        revision=artifact['revision_id'], notes_policy=notes_policy)
    data, _ = exports.download(store.project_root, export_id=packet['export_id'], filename='deck.pptx')
    selected = delivery_eligibility.selected_output(store, artifact, notes_policy=notes_policy)
    assert sha256_bytes(data) == selected['file_sha256'] == selected['subject_ref']['sha256']
    assert store.read_object_bytes(selected['subject_ref']) == data
    assert selected['subject_ref']['path'].endswith('.pptx')
    assert packet['manifest']['notes_policy'] == selected['notes_policy']
    assert sha256_bytes(canonical_json_bytes(packet['manifest']['notes_policy'])) == sha256_bytes(
        canonical_json_bytes(selected['notes_policy']))
    source = root / (name + '.pptx')
    source.write_bytes(data)
    before = source_snapshot(source.parent)
    scope = {**export_scope(layers=['review']), 'page_ids': [], 'purpose': 'review'}
    readback = external_review.evaluate(source, source_sha256=selected['file_sha256'], scope=scope)
    report = readback['external_report']['report']
    assert report['source_unchanged'] is True
    assert report['source_manifest']['source_kind'] == 'file'
    assert report['subjects'][0]['sha256'] == selected['file_sha256']
    assert report['reviews'] == [] and report['provenance']['reviewer_executed'] is False
    assert report['engineering_observations'][0]['slide_count'] == 1
    assert source_snapshot(source.parent) == before
    assert readback['external_report']['report_sha256'] == sha256_bytes(canonical_json_bytes(report))
    (root / (name + '-external-packet.json')).write_bytes(
        canonical_json_bytes(readback['external_report']))
    (root / (name + '-selection.json')).write_bytes(canonical_json_bytes({
        'evidence_level': 'L1_SYNTHETIC', 'artifact_revision_id': artifact['revision_id'],
        'review_export_id': packet['export_id'], 'selected_output': selected,
        'actual_pptx_sha256': sha256_bytes(data),
        'notes_policy_sha256': sha256_bytes(canonical_json_bytes(selected['notes_policy']))}))
    return {'artifact': artifact, 'selected': selected, 'file': source, 'bytes': data,
            'review_export': packet, 'external_packet': readback['external_report'],
            'notes_policy_sha256': sha256_bytes(canonical_json_bytes(selected['notes_policy']))}


def _record_exact_review(store, fixed, *, tag, status='pass', original=False):
    """Actual final Task and Host claim; receipt remains synthetic L1."""
    artifact = fixed['artifact']
    current = store.load_document()
    subject = artifact['outputs']['pptx'] if original else fixed['selected']['subject_ref']
    # Formal final adoption also checks the actual current engineering output.
    # All named subjects are covered by the explicitly synthetic receipt.
    subjects = [subject]
    if current['outputs']['pptx'] not in subjects:
        subjects.append(current['outputs']['pptx'])
    subjects += [entry['page'] for entry in artifact['pages']]
    dependency_subject = (store.read_object_json(subject)['file'] if original else subject)
    reviewer = human_professional_reviewer(current, event_id='evt-d08-' + tag, subjects=subjects)
    review = _professional_review(store, artifact, review_id='d08-' + tag, status=status,
        reviewer=reviewer, subjects=subjects,
        deps=professional_dependencies(store, artifact, subject=dependency_subject),
        observations=['L1 synthetic predicate check on exactly the named file and its page dependencies.'])
    outcome = _accept_review(store.project_root, store, review,
        task_id='d08-' + tag, operation_id='d08-' + tag)
    assert outcome['status'] == 'accepted'
    recorded = store.read_object_json(store.load_document()['reviews'][-1])
    assert recorded['subjects'] == subjects
    assert recorded['reviewer']['identity_receipt']['subjects'] == subjects
    return store.load_document()['reviews'][-1]


def _refused(project, fixed, destination, *, notes_policy='pptx_notes', code='professional_review_required'):
    with pytest.raises(eligibility.EligibilityBlocked) as blocked:
        exports.create(project, purpose='delivery', revision=fixed['artifact']['revision_id'],
                       notes_policy=notes_policy, output_dir=destination)
    assert code in blocked.value.payload()['error']['reason_codes']
    assert not destination.exists()
    report = delivery_eligibility.describe(project, revision=fixed['artifact']['revision_id'],
        file_path=fixed['file'], notes_policy=notes_policy)
    assert report['status'] == 'blocked' and code in report['reason_codes']
    return report


def _delivered(project, fixed, destination, *, notes_policy='pptx_notes'):
    outcome = exports.create(project, purpose='delivery', revision=fixed['artifact']['revision_id'],
        notes_policy=notes_policy, output_dir=destination)
    assert outcome['status'] == 'exported'
    assert (destination / 'deck.pptx').read_bytes() == fixed['bytes']
    report = json.loads((destination / 'delivery.json').read_text())
    assert report['revision_id'] == fixed['artifact']['revision_id']
    assert sha256_bytes(canonical_json_bytes(report['notes_policy'])) == fixed['notes_policy_sha256']
    readiness = delivery_eligibility.describe(project, revision=fixed['artifact']['revision_id'],
        file_path=fixed['file'], notes_policy=notes_policy)
    assert readiness['status'] == 'ready' and readiness['engineering_ready'] is True
    assert readiness['selected_output'] == fixed['selected']
    return report, readiness


def test_selected_old_policy_true_uses_current_false_without_claiming_human(tmp_path):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _set_policy(store, professional_review_required_for_delivery=True)
    fixed = _fixed_public_file(store, tmp_path)
    assert fixed['artifact']['policy']['professional_review_required_for_delivery'] is True
    assert notes_records(fixed['bytes'])[0]['notes'] == NOTES
    _set_policy(store, professional_review_required_for_delivery=False)
    report, readiness = _delivered(project, fixed, tmp_path / 'delivery')
    assert readiness['control_revision_id'] != fixed['artifact']['revision_id']
    assert report['professional'] == {'required': False, 'satisfied': True,
        'human_signoff': False, 'gap_reasons': []}
    assert store.load_document(fixed['artifact']['revision_id']) == fixed['artifact']


def test_selected_old_policy_false_requires_new_exact_file_review_under_current_true(
        tmp_path, synthetic_human_adapter):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    fixed = _fixed_public_file(store, tmp_path)
    assert fixed['artifact']['policy']['professional_review_required_for_delivery'] is False
    _set_policy(store, professional_review_required_for_delivery=True)
    blocked = _refused(project, fixed, tmp_path / 'before-review')
    assert blocked['professional']['human_signoff'] is False
    assert 'no_professional_review' in blocked['professional']['gap_reasons']
    _record_exact_review(store, fixed, tag='new-exact-public')
    report, readiness = _delivered(project, fixed, tmp_path / 'after-review')
    assert report['professional']['human_signoff'] is True
    assert readiness['professional']['reviews'][0]['subject_current'] is True
    assert store.load_document(fixed['artifact']['revision_id'])['reviews'] == fixed['artifact']['reviews']


def test_original_file_review_cannot_authorize_different_public_projection(tmp_path, synthetic_human_adapter):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _set_policy(store, professional_review_required_for_delivery=True)
    fixed = _fixed_public_file(store, tmp_path)
    original = store.read_object_json(fixed['artifact']['outputs']['pptx'])['file']
    assert original['sha256'] != fixed['selected']['file_sha256']
    _record_exact_review(store, fixed, tag='internal-original-only', original=True)
    report = _refused(project, fixed, tmp_path / 'public-refused')
    assert 'subject_not_current_pptx' in report['professional']['gap_reasons']
    # Exact original-file judgement can pass while the distinct public copy fails.
    original_path = tmp_path / 'internal-original.pptx'
    original_path.write_bytes(store.read_object_bytes(original))
    original_report = delivery_eligibility.describe(project, file_path=original_path)
    assert original_report['selected_output']['projection'] == 'original_internal'
    assert original_report['professional']['satisfied'] is True
    _record_exact_review(store, fixed, tag='public-after-original')
    assert _delivered(project, fixed, tmp_path / 'public-reviewed')[0]['professional']['human_signoff'] is True


def test_notes_projection_change_requires_the_new_exact_public_file(tmp_path, synthetic_human_adapter):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _set_policy(store, professional_review_required_for_delivery=True)
    with_notes = _fixed_public_file(store, tmp_path, name='with-notes')
    reading = _fixed_public_file(store, tmp_path, revision=with_notes['artifact']['revision_id'],
        notes_policy='pure_reading', name='pure-reading')
    assert with_notes['selected']['file_sha256'] != reading['selected']['file_sha256']
    assert with_notes['notes_policy_sha256'] != reading['notes_policy_sha256']
    assert notes_records(with_notes['bytes'])[0]['notes'] == NOTES
    assert notes_records(reading['bytes'])[0]['notes'] == ''
    _record_exact_review(store, with_notes, tag='with-notes')
    _delivered(project, with_notes, tmp_path / 'notes-delivery')
    blocked = _refused(project, reading, tmp_path / 'reading-refused', notes_policy='pure_reading')
    assert 'subject_not_current_pptx' in blocked['professional']['gap_reasons']
    _record_exact_review(store, reading, tag='pure-reading')
    _delivered(project, reading, tmp_path / 'reading-reviewed', notes_policy='pure_reading')


def test_current_appended_fail_for_old_exact_file_overrides_its_earlier_pass(tmp_path, synthetic_human_adapter):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _set_policy(store, professional_review_required_for_delivery=True)
    fixed = _fixed_public_file(store, tmp_path)
    old_reviews = list(fixed['artifact']['reviews'])
    _record_exact_review(store, fixed, tag='historical-pass')
    _delivered(project, fixed, tmp_path / 'before-new-fail')
    _record_exact_review(store, fixed, tag='historical-current-fail', status='fail')
    report = _refused(project, fixed, tmp_path / 'after-new-fail')
    assert 'professional_review_failed' in report['professional']['gap_reasons']
    assert report['engineering_ready'] is True
    assert store.load_document(fixed['artifact']['revision_id'])['reviews'] == old_reviews


def _grant_fixed(store, *, binding='exact_artifact', selection='pptx_notes'):
    evidence = evidence_for(store.project_root,
        text='L1 synthetic instruction: deliver only this exact public file with its fixed notes policy; no production permission.',
        locator='synthetic-l1:D08:exact-file-authorization')
    return collaboration.grant(store.project_root, input={
        'mode': 'execute', 'actions': ['export.delivery'], 'scope': export_scope(layers=['output']),
        'delivery_binding': binding, 'constraints': {'delivery_notes_policy': selection},
        'max_external_calls': 0, 'max_candidate_count': 0, 'adoption_policy': 'candidate_only',
        'instruction_evidence': evidence}, base_revision=store.load_document()['revision_id'],
        operation_id='fixed-grant-' + uuid.uuid4().hex)


def _next_request(store):
    current = store.load_document()
    evidence = evidence_for(store.project_root,
        text='L1 synthetic next request: edit the one-page copy only; do not grant any delivery or image permission.',
        locator='synthetic-l1:D08:next-request')
    return collaboration.start_request(store.project_root, input={'intent': 'edit', 'content_intent': None,
        'scope': export_scope(layers=['page_copy']), 'instruction_evidence': evidence,
        'replace_request_id': current['collaboration']['request']['request_id']},
        base_revision=current['revision_id'], operation_id='next-request-' + uuid.uuid4().hex)


def _seed_second_actual_output(store, root):
    """Explicit L1 state setup with a genuinely different, valid PPTX file."""
    from pptx import Presentation
    document = store.load_document()
    data = store.read_object_bytes(store.read_object_json(document['outputs']['pptx'])['file'])
    deck = Presentation(io.BytesIO(data))
    deck.slides[0].shapes[0].text_frame.text += '\nDifferent visible second-version artifact.'
    path = root / 'second-version.pptx'
    deck.save(path)
    path.write_bytes(attach_notes(path.read_bytes(), [{'page_id': 'p1', 'notes': NOTES}]))
    ref = adopt_artifact(store, path, 'pptx')
    fresh = store.load_document()
    operation = 'l1-second-output-' + uuid.uuid4().hex
    updated = bump_revision(fresh, {'operation_id': operation, 'kind': 'task_update',
        'description': 'Explicit L1 synthetic second actual output', 'read_set': []})
    updated['outputs']['pptx'] = ref
    store.commit_change(base_revision=fresh['revision_id'], document=updated, operation_id=operation)
    return ref


def test_fixed_old_file_permission_survives_next_request_but_current_withdrawal_stops_it(tmp_path):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    materialized = _fixed_public_file(store, tmp_path, name='materialized-before-grant')
    grant = _grant_fixed(store)
    fixed = _fixed_public_file(store, tmp_path, name='frozen-after-grant')
    record = store.read_object_json(grant['authorization_ref'])
    frozen = record['constraints']['fixed_delivery']
    assert frozen == {**{key: fixed['selected'][key] for key in
        ('artifact_ref', 'subject_ref', 'file_sha256', 'notes_policy')},
        'source_request_id': fixed['artifact']['collaboration']['request']['request_id']}
    assert materialized['bytes'] == fixed['bytes']
    _next_request(store)
    _seed_second_actual_output(store, tmp_path)
    current = store.load_document()
    assert current['outputs']['pptx'] != fixed['artifact']['outputs']['pptx']
    assert current['collaboration']['request']['request_id'] != frozen['source_request_id']
    _delivered(project, fixed, tmp_path / 'old-file-still-delivered')
    current_status = delivery_eligibility.describe(project, notes_policy='pptx_notes')
    assert current_status['status'] == 'blocked'
    assert current_status['selected_output']['file_sha256'] != frozen['file_sha256']
    collaboration.revoke_authorization(project, authorization_id=grant['authorization_id'],
        reason='L1 synthetic current withdrawal of the old exact file permission',
        base_revision=store.load_document()['revision_id'], operation_id='withdraw-fixed')
    report = _refused(project, fixed, tmp_path / 'withdrawn-old-file', code='authorization_revoked')
    assert report['professional']['required'] is False
    assert store.load_document(fixed['artifact']['revision_id']) == fixed['artifact']


def test_process_delegation_does_not_become_arbitrary_old_file_permission(tmp_path):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _fixed_public_file(store, tmp_path, name='before-process-grant')
    _grant_fixed(store, binding='process_delegation')
    fixed = _fixed_public_file(store, tmp_path, name='process-selected')
    _next_request(store)
    with pytest.raises(eligibility.EligibilityBlocked):
        exports.create(project, purpose='delivery', revision=fixed['artifact']['revision_id'],
            notes_policy='pptx_notes', output_dir=tmp_path / 'process-old-file')
    assert not (tmp_path / 'process-old-file').exists()


def test_fixed_old_file_permission_does_not_cover_another_notes_projection(tmp_path):
    project, store = _passing_deck(tmp_path, notes=NOTES)
    _fixed_public_file(store, tmp_path, name='pre-grant-notes')
    _grant_fixed(store, selection='pptx_notes')
    with_notes = _fixed_public_file(store, tmp_path, name='authorized-notes')
    reading = _fixed_public_file(store, tmp_path,
        revision=with_notes['artifact']['revision_id'], notes_policy='pure_reading', name='unauthorized-reading')
    _next_request(store)
    _delivered(project, with_notes, tmp_path / 'fixed-notes-delivered')
    assert reading['notes_policy_sha256'] != with_notes['notes_policy_sha256']
    with pytest.raises(eligibility.EligibilityBlocked) as failure:
        exports.create(project, purpose='delivery', revision=reading['artifact']['revision_id'],
            notes_policy='pure_reading', output_dir=tmp_path / 'fixed-reading-refused')
    assert 'authority_insufficient' in failure.value.payload()['error']['reason_codes']
    assert not (tmp_path / 'fixed-reading-refused').exists()
