"""Synthetic counters for v1.2 response-local facts and warm projection reuse."""
from concurrent.futures import ThreadPoolExecutor
import threading
import time

import pytest

from deck_master import content_review, delivery_eligibility, editing, export_sanitize, workbench
from deck_master.store import Store
from test_x3_design_plan_actions import production_project


def test_v2_builds_artifacts_once_and_warm_read_reuses_projection(tmp_path, monkeypatch):
    project, store = production_project(tmp_path, 'full_deck')
    workbench._summary_cache.clear()
    calls = []
    original = editing._current_artifact_digests
    def counted(*args, **kwargs):
        calls.append(args[1]['revision_id'])
        return original(*args, **kwargs)
    monkeypatch.setattr(editing, '_current_artifact_digests', counted)
    monkeypatch.setattr(content_review, '_current_artifact_digests', counted)
    first = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert calls == [store.current_revision_id()]
    assert workbench.workbench_summary(project, protocol='workbench_summary.v2') == first
    assert len(calls) == 1
    workbench._summary_cache.clear()
    assert workbench.workbench_summary(project, protocol='workbench_summary.v2') == first
    assert len(calls) == 2


def test_same_v2_identity_coalesces_builds(tmp_path, monkeypatch):
    project, _ = production_project(tmp_path, 'full_deck')
    workbench._summary_cache.clear()
    entered, release = threading.Event(), threading.Event()
    original = workbench._build_v2_summary
    calls = []
    def blocked(*args, **kwargs):
        calls.append(1)
        entered.set()
        assert release.wait(5)
        return original(*args, **kwargs)
    monkeypatch.setattr(workbench, '_build_v2_summary', blocked)
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = [executor.submit(workbench.workbench_summary, project, protocol='workbench_summary.v2') for _ in range(4)]
        assert entered.wait(5)
        time.sleep(.05)
        release.set()
        results = [future.result(timeout=10) for future in futures]
    assert len(calls) == 1
    assert all(value == results[0] for value in results)
    assert not workbench._summary_inflight


def test_pptx_warm_selection_does_not_read_source_bytes(tmp_path, monkeypatch):
    from test_export import _deck
    _, store = _deck(tmp_path)
    document = store.load_document()
    ref = store.read_object_json(document['outputs']['pptx'])['file']
    delivery_eligibility._projection_cache.clear()
    calls = {'source': 0, 'sanitize': 0}
    original_read, original_pptx = Store.read_object_bytes, export_sanitize.pptx
    def read(self, item):
        if item == ref:
            calls['source'] += 1
        return original_read(self, item)
    def sanitize(*args, **kwargs):
        calls['sanitize'] += 1
        return original_pptx(*args, **kwargs)
    monkeypatch.setattr(Store, 'read_object_bytes', read)
    monkeypatch.setattr(export_sanitize, 'pptx', sanitize)
    first = delivery_eligibility.selected_output(store, document)
    assert calls == {'source': 1, 'sanitize': 1}
    assert delivery_eligibility.selected_output(store, document) == first
    assert calls == {'source': 1, 'sanitize': 1}
    path = store.project_root / ref['path']
    path.write_bytes(path.read_bytes() + b'corrupt')
    with pytest.raises(Exception):
        delivery_eligibility.selected_output(store, document)
    assert calls['source'] == 2


def test_drafts_share_one_history_reader(tmp_path, monkeypatch):
    from deck_master import samples, snapshots, ui_journal
    from test_ui_journal import draft
    project = tmp_path / 'project'
    samples.create_sample(project, page_count=2, readonly=False)
    for index in range(3):
        ui_journal.save(project, draft=draft(project, draft_id=f'd{index}'))
    original = snapshots.FixedSnapshotReader.__init__
    calls = []
    def counted(self, *args, **kwargs):
        calls.append(1)
        original(self, *args, **kwargs)
    monkeypatch.setattr(snapshots.FixedSnapshotReader, '__init__', counted)
    result = ui_journal.list_drafts(project)
    assert len(result['records']) == 3 and not result['errors']
    assert calls == [1]


def test_warm_v2_rechecks_historical_visual_dependency(tmp_path):
    from test_workbench_reads import mixed_project, commit
    project, store = mixed_project(tmp_path)
    document = store.load_document()
    old_ref = document['pages'][0]['page']
    page = store.read_object_json(old_ref)
    page['speaker_notes'] = 'Only narration changed; retained visuals still match.'
    document['pages'][0]['page'] = store.put_json_object(page)
    commit(store, document, 'notes-only-history')
    workbench._summary_cache.clear()
    first = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert workbench._summary_cache, 'fixture must exercise a reusable v2 response'
    assert workbench.workbench_summary(project, protocol='workbench_summary.v2') == first
    # Neither current Head nor its Page changes. Only a Page consumed through
    # the retained visual's historical content binding becomes unreadable.
    (store.project_root / old_ref['path']).write_bytes(b'corrupt historical Page')
    guarded = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    workbench._summary_cache.clear()
    reference = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert guarded == reference
    assert guarded != first


def test_v2_does_not_freeze_synthetic_human_adapter_verdict(tmp_path, synthetic_human_adapter):
    from test_export import _passing_deck, _set_policy, _review, human_professional_reviewer
    project, store = _passing_deck(tmp_path)
    _set_policy(store, professional_review_required_for_delivery=True)
    document = store.load_document()
    _review(store, document, 'professional_use', 'pass', seed='runtime-identity',
            reviewer=human_professional_reviewer(document, event_id='evt-runtime-identity'))
    head = store.current_revision_id()
    workbench._summary_cache.clear()
    def professional(summary):
        check = next(item for item in summary['current_work']['delivery']['checks']
                     if item['kind'] == 'requested_delivery')
        return check['basis']['professional']
    first = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert professional(first)['human_signoff'] is True
    # The controlled verifier now rejects the same immutable receipt. There is
    # no product revision or file mutation to invalidate a static cache key.
    synthetic_human_adapter['evt-runtime-identity'] = 'different-receipt-fingerprint'
    second = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert store.current_revision_id() == head
    assert professional(second)['human_signoff'] is False
    assert 'receipt_event_replayed' in professional(second)['gap_reasons']
    workbench._summary_cache.clear()
    assert second == workbench.workbench_summary(project, protocol='workbench_summary.v2')


def test_v2_owner_rechecks_cache_after_miss_registration_race(tmp_path, monkeypatch):
    project, _ = production_project(tmp_path, 'full_deck')
    workbench._summary_cache.clear()
    original_lookup = workbench._cached_summary
    original_build = workbench._build_v2_summary
    delayed_thread = threading.get_ident()
    delayed_once = False
    builds = []
    published = []

    def counted(*args, **kwargs):
        builds.append(1)
        return original_build(*args, **kwargs)

    with ThreadPoolExecutor(max_workers=1) as executor:
        def interleaved(ctx, key, revision):
            nonlocal delayed_once
            cached = original_lookup(ctx, key, revision)
            if threading.get_ident() == delayed_thread and not delayed_once:
                delayed_once = True
                assert cached is None
                # Freeze this caller after its miss but before registration.
                # Another real caller owns, builds, publishes and exits fully.
                owner = executor.submit(workbench.workbench_summary, project,
                                        protocol='workbench_summary.v2')
                published.append(owner.result(timeout=10))
                assert key not in workbench._summary_inflight
                assert original_lookup(ctx, key, revision) is not None
                return None  # Deliver this caller's original, now stale miss.
            return cached

        monkeypatch.setattr(workbench, '_cached_summary', interleaved)
        monkeypatch.setattr(workbench, '_build_v2_summary', counted)
        result = workbench.workbench_summary(project, protocol='workbench_summary.v2')

    assert result == published[0]
    assert builds == [1]
    assert not workbench._summary_inflight


@pytest.mark.parametrize('planning', [False, True])
def test_v2_reuses_design_completion_including_nonplanning_result(tmp_path, monkeypatch, planning):
    from deck_master import design_completion
    from test_x3_design_plan_actions import plan_project
    project, _ = plan_project(tmp_path, confirmed=False) if planning else production_project(tmp_path, 'full_deck')
    original = design_completion.completion_status
    calls = []
    def counted(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)
    monkeypatch.setattr(design_completion, 'completion_status', counted)
    workbench._summary_cache.clear()
    workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert calls == [1]


def test_v2_reuses_current_input_digest_across_read_consumers(tmp_path, monkeypatch):
    from deck_master import models, content_plan, tasks
    from test_workbench_reads import commit
    project, store = production_project(tmp_path, 'full_deck')
    document = store.load_document()
    document['content_basis'] = {'input_digest': models.compute_input_digest(document),
                                'input_revision_id': None, 'resolved_by_task_id': None}
    commit(store, document, 'bind-current-inputs')
    current = store.load_document()
    original = models.compute_input_digest
    calls = []
    def counted(document):
        if document == current:
            calls.append(1)
        return original(document)
    for module in (models, editing, content_plan, tasks):
        monkeypatch.setattr(module, 'compute_input_digest', counted)
    workbench._summary_cache.clear()
    result = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert result['input_alignment'] == 'current'
    assert calls == [1]


def test_v2_reuses_each_verified_page_visual_digest(tmp_path, monkeypatch):
    from deck_master import models, visual_dependencies
    from test_workbench_reads import mixed_project, commit
    project, store = mixed_project(tmp_path)
    document = store.load_document()
    page = store.read_object_json(document['pages'][0]['page'])
    page['speaker_notes'] = 'Narration only: historical and current visual facts are equal.'
    document['pages'][0]['page'] = store.put_json_object(page)
    commit(store, document, 'visual-digest-reuse')
    original = visual_dependencies.page_visual_digest
    counts = {}
    def counted(page):
        key = models.sha256_bytes(models.canonical_json_bytes(page))
        counts[key] = counts.get(key, 0) + 1
        return original(page)
    monkeypatch.setattr(visual_dependencies, 'page_visual_digest', counted)
    workbench._summary_cache.clear()
    workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert counts
    assert set(counts.values()) == {1}


def test_input_facts_guard_actual_inputs_and_leave_overlays_unchanged(tmp_path, monkeypatch):
    import copy
    from deck_master import models
    project, store = production_project(tmp_path, 'full_deck')
    document = store.load_document()
    before = copy.deepcopy(document)
    other = copy.deepcopy(document)
    other['task']['title'] += ' changed interpretation'
    original = models.compute_input_digest
    calls = []
    def counted(value):
        calls.append(copy.deepcopy(value))
        return original(value)
    monkeypatch.setattr(models, 'compute_input_digest', counted)
    ctx = workbench._ReadContext(store, document)
    try:
        first = ctx.input_digest(document)
        assert ctx.input_digest(copy.deepcopy(document)) == first
        assert len(calls) == 1
        second = ctx.input_digest(other)
        assert second != first
        assert len(calls) == 2
        # Unrelated policy/pages/tasks are not compute_input_digest inputs.
        overlay = copy.deepcopy(other)
        overlay['policy'] = {**overlay['policy'], 'synthetic_read_overlay': True}
        overlay['tasks'] = [{'unrelated': 'task registry'}]
        overlay['pages'] = [{'unrelated': 'page registry'}]
        assert ctx.input_digest(overlay) == second
        assert len(calls) == 2
        source_overlay = copy.deepcopy(other)
        source_overlay['sources'] = [*source_overlay.get('sources', []),
            {'source_id': 'synthetic-input', 'original_sha256': 'a' * 64,
             'usage_note': 'new source semantics'}]
        third = ctx.input_digest(source_overlay)
        assert third != second
        assert len(calls) == 3
        # Revision buckets remain separate even when actual inputs agree.
        revision_overlay = copy.deepcopy(source_overlay)
        revision_overlay['revision_id'] += '-other'
        assert ctx.input_digest(revision_overlay) == third
        assert len(calls) == 4
        # Saved inputs are detached: mutating a caller-owned source cannot
        # silently mutate the guard alongside the requested interpretation.
        source_overlay['sources'][-1]['usage_note'] = 'changed source semantics'
        assert ctx.input_digest(source_overlay) != third
        assert len(calls) == 5
    finally:
        ctx.close()
    assert document == before
    assert 'synthetic_read_overlay' not in other['policy']


def test_pptx_metadata_rechecks_notes_order_and_policy(tmp_path, monkeypatch):
    import copy
    import io
    from pptx import Presentation
    from deck_master.compiler.native import attach_notes
    from deck_master.operations import OperationError
    from test_export import _deck
    _, store = _deck(tmp_path, notes='First page notes')
    document = store.load_document()
    output = store.read_object_json(document['outputs']['pptx'])
    presentation = Presentation(io.BytesIO(store.read_object_bytes(output['file'])))
    presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    second_page = copy.deepcopy(store.read_object_json(document['pages'][0]['page']))
    second_page.update(page_id='p2', speaker_notes='Second page notes')
    document['pages'].append({'page_id': 'p2', 'page': store.put_json_object(second_page)})
    def bind_fixture_plan(selected):
        from deck_master.models import validate_content_plan_semantics
        plan = copy.deepcopy(store.read_object_json(selected['content_plan']))
        order = [entry['page_id'] for entry in selected['pages']]
        data = plan['input']
        data['page_order'] = order
        roles = data.get('page_roles') or []
        if roles and not any(role['page_id'] == 'p2' for role in roles):
            roles.append({**copy.deepcopy(roles[0]), 'page_id': 'p2'})
        if roles:
            data['page_roles'] = sorted(roles, key=lambda role: order.index(role['page_id']))
        chapters = data.get('chapters') or []
        if chapters:
            # The source fixture has one page/chapter. Its added second slide
            # belongs to that same synthetic chapter, in the selected order.
            chapters[0]['page_ids'] = order
        for field in ('exact', 'max'):
            if (data.get('page_budget') or {}).get(field) is not None:
                data['page_budget'][field] = len(order)
        plan['page_links'] = [{'page_id': entry['page_id'], 'page_ref': entry['page'],
                               'plan_role_page_id': entry['page_id']}
                              for entry in selected['pages']]
        validate_content_plan_semantics(plan)
        selected['content_plan'] = store.put_json_object(plan)
    bind_fixture_plan(document)
    output['file'] = store.put_blob(attach_notes(buffer.getvalue(), [
        {'page_id': 'p1', 'notes': 'First page notes'},
        {'page_id': 'p2', 'notes': 'Second page notes'}]), ext='pptx')
    document['outputs']['pptx'] = store.put_json_object(output)
    source_ref = output['file']
    original_read, original_pptx = Store.read_object_bytes, export_sanitize.pptx
    calls = {'source': 0, 'sanitize': 0}
    def read(self, ref):
        if ref == source_ref:
            calls['source'] += 1
        return original_read(self, ref)
    def sanitize(*args, **kwargs):
        calls['sanitize'] += 1
        return original_pptx(*args, **kwargs)
    monkeypatch.setattr(Store, 'read_object_bytes', read)
    monkeypatch.setattr(export_sanitize, 'pptx', sanitize)
    delivery_eligibility._projection_cache.clear()
    notes = delivery_eligibility.selected_output(store, document, notes_policy='pptx_notes')
    assert delivery_eligibility.selected_output(store, document, notes_policy='pptx_notes') == notes
    assert calls == {'source': 1, 'sanitize': 1}
    reading = delivery_eligibility.selected_output(store, document, notes_policy='pure_reading')
    assert reading['file_sha256'] != notes['file_sha256']
    assert delivery_eligibility.selected_output(store, document, notes_policy='pure_reading') == reading
    assert calls == {'source': 2, 'sanitize': 2}
    reordered = copy.deepcopy(document)
    reordered['pages'].reverse()
    bind_fixture_plan(reordered)
    # Same source bytes cannot silently inherit the earlier successful notes
    # order. The real sanitizer must reject its stable Page-identity mismatch.
    with pytest.raises(OperationError, match="actual PPTX notes differ"):
        delivery_eligibility.selected_output(store, reordered, notes_policy='pptx_notes')
    assert calls == {'source': 3, 'sanitize': 3}
    delivery_eligibility._projection_cache.clear()
    with pytest.raises(OperationError, match="actual PPTX notes differ"):
        delivery_eligibility.selected_output(store, reordered, notes_policy='pptx_notes')
    assert calls == {'source': 4, 'sanitize': 4}


def test_v2_paired_delivery_reuses_real_pptx_output_check(tmp_path, monkeypatch):
    import copy
    from deck_master import pipeline
    from test_export import _deck
    from test_workbench_reads import commit
    project, store = _deck(tmp_path)
    document = store.load_document()
    output = store.read_object_json(document['outputs']['pptx'])
    output['dependencies'] = [
        {'kind': 'svg', 'identity': page['page_id'], 'sha256': page['svg']['sha256']}
        for page in document['pages']]
    document['outputs']['pptx'] = store.put_json_object(output)
    commit(store, document, 'bind-real-pptx-svg-dependencies')
    document = store.load_document()
    original = pipeline.output_matches_current_svg
    calls = []
    def counted(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)
    monkeypatch.setattr(pipeline, 'output_matches_current_svg', counted)
    workbench._summary_cache.clear()
    summary = workbench.workbench_summary(project, protocol='workbench_summary.v2')
    assert summary['outputs']['pptx']['applicability'] == 'current'
    assert calls == [1]
    paired = next(item['basis'] for item in summary['current_work']['delivery']['checks']
                  if item['kind'] == 'requested_delivery')
    before = copy.deepcopy(summary)
    # No paired fact is supplied to this standalone path: it must still check
    # actual output dependencies, and arrive at the identical delivery facts.
    standalone = delivery_eligibility.describe_snapshots(store, document, document)
    assert calls == [1, 1]
    assert paired == standalone
    assert summary == before
