"""W11 frozen copies, metadata privacy and recoverable engineering exports."""

import io
import copy
import json
import uuid
import zipfile

from PIL import Image, PngImagePlugin
from pptx import Presentation
import pytest

from deck_master import exports, service
from deck_master.export_sanitize import png, svg, pptx
from deck_master.models import bump_revision, sha256_bytes
from deck_master.operations import OperationError
from deck_master.store import Store

CANARY = "/Users/private-customer/private-prompt-CANARY"


def project(tmp_path):
    root = tmp_path / "project"
    service.create(
        root,
        brief="Internal brief " + CANARY,
        draft={
            "pages": [
                {
                    "schema_version": "deck_page_package.v2",
                    "page_id": "p1",
                    "customer_visible": {
                        "title": "Visible title",
                        "body_blocks": [{"id": "b1", "type": "paragraph", "text": "Visible body"}],
                    },
                    "visual_spec": {"intent": "Public diagram", "reference_mode": "new_design"},
                }
            ]
        },
    )
    return root, Store(root)


def test_png_metadata_removed_pixels_unchanged():
    image = Image.new("RGBA", (7, 9), (20, 50, 80, 170))
    meta = PngImagePlugin.PngInfo()
    meta.add_text("prompt", CANARY)
    meta.add_itxt("XML:com.adobe.xmp", CANARY)
    source = io.BytesIO()
    image.save(source, format="PNG", pnginfo=meta)
    raw = source.getvalue()
    clean = png(raw, "image")
    assert CANARY.encode() in raw and CANARY.encode() not in clean
    assert Image.open(io.BytesIO(clean)).tobytes() == image.tobytes()
    assert source.getvalue() == raw


def test_svg_metadata_removed_links_and_gradients_preserved():
    raw = f"""<svg xmlns="http://www.w3.org/2000/svg"><metadata>{CANARY}</metadata><!--{CANARY}--><defs><linearGradient id="g"><stop offset="0" stop-color="red"/></linearGradient></defs><rect width="20" height="20" fill="url( '#g' )"/><a href="https://example.org/guide"><text x="1" y="2">Normal text</text></a></svg>""".encode()
    clean = svg(raw, "image")
    assert CANARY.encode() not in clean
    assert b"Normal text" in clean and b"https://example.org/guide" in clean and b"linearGradient" in clean
    with pytest.raises(OperationError, match="local path"):
        svg(f"<svg><text>{CANARY}</text></svg>".encode(), "image")
    with pytest.raises(OperationError, match="external SVG"):
        svg(b'<svg><rect fill="url(https://example.org/private)"/></svg>', "image")


def presentation():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(100, 100, 1000000, 1000000)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Normal visible body"
    run.hyperlink.address = "https://example.org/guide"
    box._element.xpath(".//p:cNvPr")[0].set("descr", "PRIVATE GENERATION PROMPT CANARY")
    deck.core_properties.author = CANARY
    slide.notes_slide.notes_text_frame.text = CANARY
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def test_pptx_properties_notes_removed_body_and_link_preserved():
    raw = presentation()
    clean = pptx(raw, "deck")
    read = Presentation(io.BytesIO(clean))
    run = read.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.text == "Normal visible body" and run.hyperlink.address == "https://example.org/guide"
    with zipfile.ZipFile(io.BytesIO(clean)) as archive:
        assert not any(n.startswith(("docProps/", "ppt/notesSlides/")) for n in archive.namelist())
        assert all(
            CANARY.encode() not in archive.read(n) and b"PRIVATE GENERATION PROMPT CANARY" not in archive.read(n)
            for n in archive.namelist()
        )
    assert raw != clean


def test_review_without_ppt_and_engineering_restores_only_fixed_ancestry(tmp_path):
    root, store = project(tmp_path)
    old = store.load_document()
    review = exports.create(root, revision=old["revision_id"], output_dir=tmp_path / "review")
    assert review["purpose"] == "review"
    assert not (tmp_path / "review" / "deck.pptx").exists()
    assert not (tmp_path / "review" / "project").exists()
    for item in review["manifest"]["files"]:
        data = (tmp_path / "review" / item["path"]).read_bytes()
        assert sha256_bytes(data) == item["sha256"] and CANARY.encode() not in data
    updated = bump_revision(old, {"operation_id": "advance", "kind": "task_update", "description": "later", "read_set": []})
    store.commit_change(base_revision=old["revision_id"], document=updated, operation_id="advance")
    engineering = exports.create(root, purpose="engineering", revision=old["revision_id"], output_dir=tmp_path / "engineering")
    restored = Store(tmp_path / "engineering" / "project").load_document()
    assert restored == old
    assert not (tmp_path / "engineering" / "project" / ".deckmaster" / "revisions" / (updated["revision_id"] + ".json")).exists()
    assert CANARY in json.dumps(restored)
    assert engineering["manifest"]["metadata_policy"] == "original_internal"
    assert store.load_document()["revision_id"] == updated["revision_id"]
    again = exports.create(root, revision=old["revision_id"], export_id=review["export_id"])
    assert again["manifest"] == review["manifest"]
    with pytest.raises(OperationError, match="another fixed request"):
        exports.create(root, export_id=review["export_id"])


def test_download_whitelist_cross_project_and_tampering(tmp_path):
    root, store = project(tmp_path)
    out = exports.create(root)
    data, name = exports.download(root, export_id=out["export_id"], filename=out["archive"])
    assert name == out["archive"] and sha256_bytes(data) == out["files"][name]["sha256"]
    with pytest.raises(OperationError):
        exports.download(root, export_id=out["export_id"], filename="../../current.json")
    other, _ = project(tmp_path / "other")
    with pytest.raises(OperationError):
        exports.show(other, export_id=out["export_id"])
    with pytest.raises(OperationError):
        exports.show(root, export_id="export-" + str(uuid.uuid4()))
    cache = store.deck_root / "workbench" / "exports" / out["export_id"]
    (cache / name).write_bytes(b"changed")
    with pytest.raises(OperationError, match="manifest hash"):
        exports.download(root, export_id=out["export_id"], filename=name)
    record = json.loads((cache / "export-record.json").read_text())
    record["files"]["extra.txt"] = {"sha256": "a" * 64, "size": 1}
    (cache / "export-record.json").write_text(json.dumps(record))
    with pytest.raises(OperationError, match="whitelist"):
        exports.show(root, export_id=out["export_id"])


def test_delivery_blocks_no_ppt_with_structured_gaps(tmp_path):
    root, _ = project(tmp_path)
    with pytest.raises(exports.ExportError) as failure:
        exports.create(root, purpose="delivery")
    assert failure.value.exit_code == 3 and failure.value.http_status == 409
    assert any(g["layer"] == "ppt" for g in failure.value.payload()["error"]["gaps"])


def test_restore_plan_cancel_conflict_and_atomic_retry(tmp_path, monkeypatch):
    from deck_master import restoration, operations

    root, store = project(tmp_path)
    old = store.load_document()
    changed = bump_revision(
        copy.deepcopy(old), {"operation_id": "edit", "kind": "policy_update", "description": "new policy", "read_set": []}
    )
    changed["policy"]["user_stop"] = True
    page = store.read_object_json(changed["pages"][0]["page"])
    page["customer_visible"]["title"] = "New title"
    changed["pages"][0]["page"] = store.put_json_object(page)
    store.commit_change(base_revision=old["revision_id"], document=changed, operation_id="edit")
    plan = restoration.plan(root, revision_id=old["revision_id"], base_revision=changed["revision_id"])
    assert plan["plan"]["impact"]["changed_pages"] == ["p1"]
    assert store.load_document() == changed, "preview/cancel never moves the current pointer"
    # Force disposable journal-index failure after the durable commit.
    monkeypatch.setattr(operations, "publish_index", lambda *args: {"code": "receipt_cache_unavailable"})
    op = str(uuid.uuid4())
    result = restoration.commit(root, plan_id=plan["plan_id"], base_revision=changed["revision_id"], operation_id=op)
    current = store.load_document()
    assert current["revision_id"] not in (old["revision_id"], changed["revision_id"])
    assert current["pages"] == old["pages"] and current["policy"]["user_stop"]
    assert result["journal_warning"]["code"] == "receipt_cache_unavailable"
    replay = restoration.commit(root, plan_id=plan["plan_id"], base_revision=changed["revision_id"], operation_id=op)
    assert replay["committed_revision_id"] == current["revision_id"]
    with pytest.raises(OperationError, match="project changed"):
        restoration.commit(root, plan_id=plan["plan_id"], base_revision=changed["revision_id"], operation_id=str(uuid.uuid4()))
    assert store.load_document() == current


def test_export_cli_and_same_origin_download_routes(tmp_path, capsys):
    from deck_master import cli
    from deck_master.web import WorkbenchServer
    from test_web import _get, _get_json, _post_json, _session_token

    root, store = project(tmp_path)
    revision = store.current_revision_id()
    assert cli.main(["export", "--project", str(root), "--revision", revision, "--purpose", "review", "--out", str(tmp_path / "cli")]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["revision_id"] == revision
    assert cli.main(["export", "--project", str(root), "--purpose", "delivery", "--out", str(tmp_path / "delivery")]) == 3
    error = json.loads(capsys.readouterr().err)
    assert error["error"]["code"] == "delivery_blocked"
    instance = WorkbenchServer(root)
    url = instance.start().rstrip("/")
    try:
        token = _session_token(url)
        assert _post_json(url + "/api/exports", {"purpose": "review"}, token=token)[0] == 403
        status, result = _post_json(url + "/api/exports", {"purpose": "review", "revision": revision}, token=token, origin=url)
        assert status == 200
        status, headers, raw = _get(url + result["download_url"])
        assert status == 200 and "attachment;" in headers["Content-Disposition"]
        assert sha256_bytes(raw) == result["files"][result["archive"]]["sha256"]
        assert _get_json(url + "/api/exports/" + result["export_id"] + "/files/unknown")[0] == 404
        assert _post_json(url + "/api/exports", {"output_dir": "/tmp/escape"}, token=token, origin=url)[0] == 422
        status, error = _post_json(url + "/api/exports", {"purpose": "delivery"}, token=token, origin=url)
        assert status == 409 and error["error"]["code"] == "delivery_blocked"
        status, plan = _post_json(
            url + "/api/history/plan-restore", {"revision_id": revision, "base_revision": revision}, token=token, origin=url
        )
        assert status == 200
        status, done = _post_json(
            url + "/api/history/commit-restore",
            {"plan_id": plan["plan_id"], "base_revision": revision, "operation_id": str(uuid.uuid4())},
            token=token,
            origin=url,
        )
        assert status == 200 and done["operation_result"]["restored_from"] == revision
    finally:
        instance.stop()


def test_background_commit_during_export_cannot_mix_snapshots(tmp_path, monkeypatch):
    root, store = project(tmp_path)
    old = store.load_document()
    original = exports._write_public

    def advancing(current_store, doc, target, report):
        changed = bump_revision(
            copy.deepcopy(old), {"operation_id": "arrived", "kind": "task_update", "description": "background result", "read_set": []}
        )
        page = store.read_object_json(changed["pages"][0]["page"])
        page["customer_visible"]["title"] = "Background new title"
        changed["pages"][0]["page"] = store.put_json_object(page)
        store.commit_change(base_revision=old["revision_id"], document=changed, operation_id="arrived")
        return original(current_store, doc, target, report)

    monkeypatch.setattr(exports, "_write_public", advancing)
    result = exports.create(root, output_dir=tmp_path / "out")
    assert result["revision_id"] == old["revision_id"] != store.current_revision_id()
    text = (tmp_path / "out" / "page-001" / "content.txt").read_text()
    assert "Visible title" in text and "Background new title" not in text
    assert json.loads((tmp_path / "out" / "delivery.json").read_text())["revision_id"] == old["revision_id"]


def test_public_copies_clean_immutable_artifacts_and_partial_failure(tmp_path):
    from deck_master.pipeline import artifact

    root, store = project(tmp_path)
    doc = store.load_document()
    raw = presentation()
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(raw)
    image = tmp_path / "blueprint.png"
    meta = PngImagePlugin.PngInfo()
    meta.add_text("prompt", CANARY)
    Image.new("RGB", (10, 10), "red").save(image, pnginfo=meta)
    changed = bump_revision(
        copy.deepcopy(doc), {"operation_id": "artifacts", "kind": "task_update", "description": "artifacts", "read_set": []}
    )
    changed["outputs"]["pptx"] = artifact(store, deck, "pptx")
    changed["pages"][0]["blueprint"] = artifact(store, image, "blueprint", page_id="p1")
    store.commit_change(base_revision=doc["revision_id"], document=changed, operation_id="artifacts")
    result = exports.create(root, output_dir=tmp_path / "clean")
    assert (tmp_path / "clean" / "deck.pptx").read_bytes() != raw
    assert store.read_object_bytes(store.read_object_json(changed["outputs"]["pptx"])["file"]) == raw
    for item in result["manifest"]["files"]:
        if item["path"] in ("deck.pptx", "page-001/blueprint.png"):
            assert item["original_sha256"] and item["original_sha256"] != item["sha256"]
    doc = store.load_document()
    page = store.read_object_json(doc["pages"][0]["page"])
    page["customer_visible"]["title"] = CANARY
    changed = bump_revision(
        copy.deepcopy(doc), {"operation_id": "private-visible", "kind": "task_update", "description": "private visible", "read_set": []}
    )
    changed["pages"][0]["page"] = store.put_json_object(page)
    store.commit_change(base_revision=doc["revision_id"], document=changed, operation_id="private-visible")
    export_id = "export-" + str(uuid.uuid4())
    with pytest.raises(OperationError, match="local path"):
        exports.create(root, export_id=export_id, output_dir=tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()
    assert not (store.deck_root / "workbench" / "exports" / export_id).exists()
    assert not list((store.deck_root / "workbench" / "exports").glob(".export-*"))


def test_engineering_does_not_interpret_material_json_as_project_references(tmp_path):
    source = tmp_path / "source.json"
    source.write_text(json.dumps({"path": ".deckmaster/objects/aa/" + "a" * 64 + ".json", "sha256": "a" * 64}))
    root = tmp_path / "project"
    service.create(root, brief="Keep source bytes", sources=[source])
    store = Store(root)
    (store.deck_root / "private-auth.json").write_text("NEVER EXPORT RUNTIME SECRET")
    result = exports.create(root, purpose="engineering", output_dir=tmp_path / "internal")
    assert not any("private-auth" in f["path"] for f in result["manifest"]["files"])
    recovered = Store(tmp_path / "internal" / "project")
    assert recovered.load_document() == store.load_document()


def test_svg_style_classes_and_internal_paint_preserved():
    raw = b'<svg><style>.label { fill: red; } .bar { fill: url(#g); }</style><defs><linearGradient id="g"/></defs><text class="label">Text</text></svg>'
    clean = svg(raw, "style")
    assert b'class="label"' in clean and b".label { fill: red; }" in clean
    with pytest.raises(OperationError):
        svg(b'<svg><style>@import "https://example.org/style";</style></svg>', "style")


def test_jpeg_metadata_after_scan_and_orientation(tmp_path):
    from deck_master.export_sanitize import jpeg

    raw = io.BytesIO()
    exif = Image.Exif()
    exif[270] = CANARY
    Image.new("RGB", (20, 12), (30, 80, 120)).save(raw, format="JPEG", exif=exif, progressive=True)
    original = raw.getvalue()
    comment = CANARY.encode()
    payload = original[:-2] + b"\xff\xfe" + (len(comment) + 2).to_bytes(2, "big") + comment + b"\xff\xd9" + comment
    cleaned = jpeg(payload, "photo")
    assert CANARY.encode() not in cleaned
    assert Image.open(io.BytesIO(cleaned)).tobytes() == Image.open(io.BytesIO(original)).tobytes()
    exif[274] = 6
    oriented = io.BytesIO()
    Image.new("RGB", (20, 12)).save(oriented, format="JPEG", exif=exif)
    with pytest.raises(OperationError, match="orientation"):
        jpeg(oriented.getvalue(), "photo")


def test_private_xml_namespace_cannot_leak_through_metadata_cleanup():
    with pytest.raises(OperationError, match="local path"):
        svg(f'<svg xmlns:private="{CANARY}"><text>Visible</text></svg>'.encode(), "image")


def test_historical_delivery_uses_its_checks_and_input_alignment(tmp_path):
    from test_export import _passing_deck

    root, store = _passing_deck(tmp_path)
    good = store.load_document()
    service.inputs_update(
        root,
        patch={"reason": "new audience", "task_patch": {"audience": "new audience"}},
        base_revision=good["revision_id"],
        operation_id="new-input",
    )
    with pytest.raises(exports.ExportInputPending):
        exports.create(root, purpose="delivery")
    historical = exports.create(root, purpose="delivery", revision=good["revision_id"])
    assert historical["revision_id"] == good["revision_id"]
    report = json.loads(exports.download(root, export_id=historical["export_id"], filename="delivery.json")[0])
    assert report["input_alignment"] == "current" and report["review_status"] == "pass"
