"""T03 sources tests: material reading across formats (AC-C01, AC-C08).

Every format proves it can locate the relevant copy; missing tools report
``needs_tool`` with the format named — never an empty success. Image pages
dispatch to visual reading (pending_visual).
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest

from deck_master.sources import read_source, tail_constraint


def _write(tmp_path: Path, name: str, content) -> Path:
    path = tmp_path / name
    if isinstance(content, str):
        path.write_text(content, encoding="utf-8")
    else:
        path.write_bytes(content)
    return path


MATERIAL_TEXT = "\n".join(
    [
        "# 季度处理对照",
        "第一季度人工处理 3120 单，条件自动处理 480 单。",
        "",
        "（材料尾部约束，关键）本材料为合成输入，不可回写。",
    ]
)


def test_txt_and_md_reading_keep_positions_and_hash(tmp_path: Path) -> None:
    for name in ("brief.txt", "brief.md"):
        path = _write(tmp_path, name, MATERIAL_TEXT)
        extract = read_source(path)
        assert extract.status == "read"
        assert extract.original_sha256 is not None
        assert len(extract.original_sha256) == 64
        locators = [item["locator"] for item in extract.locators]
        assert "L1" in locators and "L3" in locators
        # AC-C01: the tail constraint is the last content line and stays in the real input.
        assert "不可回写" in tail_constraint(extract)


def test_json_reading_keeps_numbers_zero_and_nesting(tmp_path: Path) -> None:
    payload = {
        "quarters": [
            {"q": "Q1", "manual": 3120, "auto": 480, "ratio": 0.0},
            {"q": "Q2", "manual": 2690, "auto": 910},
        ],
        "unit": "单",
    }
    path = _write(tmp_path, "data.json", json.dumps(payload, ensure_ascii=False))
    extract = read_source(path)
    assert extract.status == "read"
    values = {item["locator"]: item["text"] for item in extract.locators}
    assert values["/quarters/0/manual"] == "3120"
    assert values["/quarters/0/ratio"] == "0.0"
    assert "/quarters/1/auto" in values


def test_invalid_json_is_explicit_not_empty_success(tmp_path: Path) -> None:
    path = _write(tmp_path, "broken.json", b"{not json")
    extract = read_source(path)
    assert extract.status == "needs_tool"
    assert "invalid JSON" in extract.detail
    assert extract.text == ""


def _pdf_stream(text: str) -> bytes:
    return f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET\n".encode("ascii")


def _minimal_pdf(text: str) -> bytes:
    stream = _pdf_stream(text)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer << /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref "
        + str(xref_at).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


def test_pdf_reading(tmp_path: Path) -> None:
    path = _write(tmp_path, "report.pdf", _minimal_pdf("Total 3120"))
    extract = read_source(path)
    if extract.status == "needs_tool":
        assert "pdftotext" in extract.detail
        return
    assert "3120" in extract.text
    locators = [item["locator"] for item in extract.locators]
    assert "page-1" in locators


def test_pdf_without_tool_needs_tool(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import deck_master.sources as sources

    monkeypatch.setattr(sources.shutil, "which", lambda name: None)
    path = _write(tmp_path, "report.pdf", _minimal_pdf("Total 3120"))
    extract = read_source(path)
    assert extract.status == "needs_tool"
    assert "pdftotext" in extract.detail
    assert extract.original_sha256 is not None


def test_docx_reading(tmp_path: Path) -> None:
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{ns}"><w:body>'
        "<w:p><w:r><w:t>季度处理量对照</w:t></w:r></w:p>"
        "<w:tbl>"
        "<w:tr>"
        "<w:tc><w:p><w:r><w:t>季度</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>3120</w:t></w:r></w:p></w:tc>"
        "</w:tr>"
        "</w:tbl>"
        "</w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        "</Types>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("word/document.xml", document_xml)
    path = _write(tmp_path, "note.docx", buffer.getvalue())
    extract = read_source(path)
    assert extract.status == "read"
    texts = [item.get("text", "") for item in extract.locators]
    assert any("季度处理量对照" in text for text in texts)
    assert any("3120" in text for text in texts)
    assert any(item.get("kind") == "table" for item in extract.locators)


def test_pptx_reading(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "幻灯片一：条件自动处理"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "季度"
    table.cell(0, 1).text = "处理量"
    table.cell(1, 0).text = "第一季度"
    table.cell(1, 1).text = "3120单"
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    extract = read_source(path)
    assert extract.status == "read"
    locators = [item["locator"] for item in extract.locators]
    assert "slide-1" in locators
    assert "3120单" in extract.text


def test_image_dispatches_to_visual_reading(tmp_path: Path) -> None:
    path = _write(tmp_path, "photo.png", b"\x89PNG\r\n\x1a\nplaceholder-bytes")
    extract = read_source(path)
    assert extract.status == "pending_visual"
    assert extract.detail
    assert extract.original_sha256 is not None


def test_unknown_format_names_itself(tmp_path: Path) -> None:
    path = _write(tmp_path, "data.xyz", b"whatever")
    extract = read_source(path)
    assert extract.status == "needs_tool"
    assert ".xyz" in extract.detail


def test_tail_constraint_is_last_content(tmp_path: Path) -> None:
    text = "开头说明\n中间正文\n（尾部约束）不得包含真实客户名。"
    path = _write(tmp_path, "brief.txt", text)
    extract = read_source(path)
    assert tail_constraint(extract).endswith("不得包含真实客户名。")


def test_missing_file_raises() -> None:
    with pytest.raises(FileNotFoundError):
        read_source("/nonexistent/material.txt")


def test_hash_matches_file_bytes(tmp_path: Path) -> None:
    path = _write(tmp_path, "brief.txt", MATERIAL_TEXT)
    extract = read_source(path)
    expected = hashlib.sha256(path.read_bytes()).hexdigest()
    assert extract.original_sha256 == expected


def test_pdf_empty_pages_keep_physical_number_and_full_text(tmp_path,monkeypatch):
    import deck_master.sources as sources
    from types import SimpleNamespace
    monkeypatch.setattr(sources.shutil,'which',lambda _: '/tools/pdftotext')
    text='\f'+('A'*4500)+' TAIL READ ONLY\f'
    monkeypatch.setattr(sources.subprocess,'run',lambda *a,**k:SimpleNamespace(returncode=0,stdout=text.encode(),stderr=b''))
    result=read_source(_write(tmp_path,'scanned.pdf',b'pdf'))
    assert result.status=='pending_visual'
    assert result.locators[0]['locator']=='page-1' and result.locators[0]['text']==''
    assert result.locators[1]['locator']=='page-2' and result.locators[1]['text'].endswith('TAIL READ ONLY')
    assert [p['locator'] for p in result.image_pages]==['page-1','page-2']


def test_grouped_ppt_text_notes_and_picture_detection(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    group=slide.shapes.add_group_shape();group.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='Grouped tail constraint'
    slide.notes_slide.notes_text_frame.text='READ ONLY; DO NOT WRITE BACK'
    image=tmp_path/'input.png';Image.new('RGB',(40,30),'red').save(image)
    slide.shapes.add_picture(str(image),0,0,Inches(1),Inches(1))
    p=tmp_path/'input.pptx';prs.save(p);result=read_source(p)
    assert 'Grouped tail constraint' in result.text and 'DO NOT WRITE BACK' in result.text
    assert result.status=='pending_visual' and result.image_pages[0]['locator']=='slide-1'


def test_create_preserves_source_status_and_immutable_original(tmp_path):
    """v1.1 §4: an explicitly named unsupported file fails the whole create
    (exit 2, code source_unsupported) and no Document is left behind; a
    directory-discovered unsupported file is only skipped, with a reason."""
    from deck_master import service
    from deck_master.errors import SourceUnsupported
    source=_write(tmp_path,'unread.xyz',b'original material')
    project=tmp_path/'project'
    with pytest.raises(SourceUnsupported):
        service.create(project,brief='read actual material',sources=[str(source)])
    assert not (project/'.deckmaster').exists()
    directory=tmp_path/'materials';directory.mkdir()
    (directory/'real.md').write_text('真实材料',encoding='utf-8')
    (directory/'unread.xyz').write_bytes(b'original material')
    response=service.create(tmp_path/'project2',brief='read actual material',sources=[directory])
    assert [item['name'] for item in response['sources_adopted']]==['real.md']
    skipped={Path(item['path']).name:item['reason'] for item in response['sources_skipped']}
    assert 'unsupported format' in skipped['unread.xyz']


def test_json_pointer_escapes_keys(tmp_path):
    result=read_source(_write(tmp_path,'escaped.json','{"a/b~c":0}'))
    assert result.locators[0]['locator']=='/a~1b~0c'


# ---------------------------------------------------------------------------
# T12 AC-K15: assets resolve from project-relative object refs after
# relocation; system fonts are fingerprinted by actual file hash and never
# packaged into the release.


def test_logo_reads_from_object_ref_after_relocation(tmp_path: Path, monkeypatch) -> None:
    import shutil
    from PIL import Image
    from deck_master import service
    from deck_master.store import Store

    logo = tmp_path / "logo.png"
    Image.new("RGB", (24, 12), (20, 60, 180)).save(logo)
    project = tmp_path / "proj"
    service.create(project, brief="品牌页", draft={"pages": [{
        "schema_version": "deck_page_package.v2", "page_id": "p1",
        "customer_visible": {"title": "品牌", "body_blocks": []},
        "visual_spec": {"intent": "brand", "reference_mode": "new_design"}}]},
        design={"assets": [{"asset_id": "brand-logo", "kind": "logo", "file": str(logo),
                            "external_use": "allowed"}],
                "allowed_asset_ids": ["brand-logo"]})
    store = Store(project)
    document = store.load_document()
    asset = document["design_context"]["assets"][0]
    assert asset["artifact"]["path"].startswith(".deckmaster/objects/"), \
        "design assets are stored as project-relative object refs, not absolute paths"
    assert store.read_object_bytes(store.read_object_json(asset["artifact"])["file"]) == logo.read_bytes()

    moved = tmp_path / "relocated"
    shutil.copytree(project, moved)
    relocated = Store(moved)
    moved_doc = relocated.load_document()
    moved_asset = moved_doc["design_context"]["assets"][0]
    assert relocated.read_object_bytes(
        relocated.read_object_json(moved_asset["artifact"])["file"]) == logo.read_bytes()


def test_font_fingerprint_tracks_actual_file_hash(tmp_path: Path) -> None:
    import hashlib
    import shutil
    import subprocess
    from deck_master.compiler import CompileOptions, SvgInput, compile_deck

    from hostenv import host_font_file, resolve_host_font
    family = resolve_host_font()
    font_path = Path(host_font_file())
    svg = tmp_path / "page.svg"
    svg.write_text(f'<svg viewBox="0 0 960 720"><text x="20" y="60" '
                   f'font-family="{family}" font-size="24">字体指纹</text></svg>')
    result = compile_deck([SvgInput("p1", svg)],
                          CompileOptions(width_px=960, height_px=720,
                                         fonts={family: str(font_path)}),
                          tmp_path / "out")
    manifest = json.loads(result.manifest_path.read_text())
    recorded = manifest["fonts"][family]["sha256"]
    # The recorded fingerprint is the actual font file hash and re-resolves
    # after relocation; a different file is a detectable font change.
    assert recorded == hashlib.sha256(font_path.read_bytes()).hexdigest()
    copied = shutil.copyfile(font_path, tmp_path / font_path.name)
    assert hashlib.sha256(Path(copied).read_bytes()).hexdigest() == recorded
    other = tmp_path / "other.ttf"
    other.write_bytes(b"not-a-real-font")
    assert hashlib.sha256(other.read_bytes()).hexdigest() != recorded


def test_missing_font_blocks_new_work_but_not_old_media(tmp_path: Path) -> None:
    import pytest
    from PIL import Image
    from deck_master import service
    from deck_master.compiler import CompileOptions, SvgInput, compile_deck
    from deck_master.pipeline import artifact as adopt_artifact
    from deck_master.store import Store

    project = tmp_path / "proj"
    service.create(project, brief="旧媒体", draft={"pages": [{
        "schema_version": "deck_page_package.v2", "page_id": "p1",
        "customer_visible": {"title": "旧页", "body_blocks": []},
        "visual_spec": {"intent": "keep", "reference_mode": "new_design"}}]})
    store = Store(project)
    document = store.load_document()
    preview = tmp_path / "preview.png"
    Image.new("RGB", (16, 9), (250, 250, 250)).save(preview)
    bumped = store.load_document()
    from deck_master.models import bump_revision
    bumped = bump_revision(bumped, {"operation_id": "attach-preview", "kind": "task_update",
                                    "description": "old media", "read_set": []})
    bumped["pages"][0]["ppt_preview"] = adopt_artifact(store, preview, "ppt_preview", page_id="p1")
    store.commit_change(base_revision=document["revision_id"], document=bumped, operation_id="attach-preview")

    # New work that needs the font is refused with a located error...
    svg = tmp_path / "text.svg"
    svg.write_text('<svg viewBox="0 0 960 720"><text x="20" y="60" font-family="Noto Sans SC">新字</text></svg>')
    with pytest.raises(ValueError, match="font file not supplied"):
        compile_deck([SvgInput("p1", svg)], CompileOptions(width_px=960, height_px=720, fonts={}),
                     tmp_path / "denied")

    # ...while previously produced media stays viewable from the store.
    current = store.load_document()
    assert store.read_object_bytes(
        store.read_object_json(current["pages"][0]["ppt_preview"])["file"]) == preview.read_bytes()
@pytest.mark.parametrize('suffix,data', [('json', b'{broken'), ('docx', b'not a ZIP'), ('txt', b'\xff\xfe')])
def test_create_rejects_explicit_unreadable_material_before_document(tmp_path, suffix, data):
    from deck_master.service import create
    from deck_master.errors import SourceUnreadable
    source = tmp_path / f'bad.{suffix}'
    source.write_bytes(data)
    project = tmp_path / 'project'
    with pytest.raises(SourceUnreadable):
        create(project, brief='读取材料', sources=[source])
    assert not (project / '.deckmaster/current.json').exists()


def test_directory_reports_bad_material_and_stops_symlink_cycles(tmp_path):
    from deck_master.service import create
    materials = tmp_path / 'materials'
    materials.mkdir()
    (materials / 'good.md').write_text('可读材料')
    (materials / 'bad.json').write_text('{bad')
    (materials / 'again').symlink_to(materials, target_is_directory=True)
    result = create(tmp_path / 'project', brief='读取材料', sources=[materials])
    assert [entry['name'] for entry in result['sources_adopted']] == ['good.md']
    assert result['sources_errored'][0]['code'] == 'source_unreadable'
    assert any('cycle' in entry['reason'] or 'visited' in entry['reason']
               for entry in result['sources_skipped'])


def test_input_update_uses_staged_bytes_and_rejects_parse_failure(tmp_path, monkeypatch):
    from deck_master import service
    from deck_master.errors import SourceUnreadable
    from deck_master.store import Store
    project = tmp_path / 'project'
    service.create(project, brief='读取材料')
    store = Store(project)
    before = store.load_document()
    source = tmp_path / 'source.json'
    source.write_text('{broken')
    patch = {'reason': '补充材料', 'source_changes': {'add': [{'path': str(source)}]}}
    with pytest.raises(SourceUnreadable):
        service.inputs_update(project, patch=patch, base_revision=before['revision_id'], operation_id='bad-json')
    assert store.load_document() == before
    source.write_text('{"version":1}')
    original_read = service._read_patch_material
    def change_after_staging(path):
        staged = original_read(path)
        source.write_text('{"version":2}')
        return staged
    monkeypatch.setattr(service, '_read_patch_material', change_after_staging)
    service.inputs_update(project, patch=patch, base_revision=before['revision_id'], operation_id='staged-bytes')
    entry = store.load_document()['sources'][0]
    extract = store.read_object_json(entry['extract'])
    assert json.loads(extract['text']) == {'version': 1}
    assert store.read_object_bytes(extract['original_file']) == b'{"version":1}'
