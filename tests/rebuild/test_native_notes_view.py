"""Actual compatible-reader notes pages must contain the complete long notes."""
import io
import subprocess

import pytest
from pptx import Presentation

from deck_master.compiler.native import attach_notes, notes_records
from deck_master.compiler import CompileOptions, SvgInput, compile_deck
from deck_master.export_sanitize import pptx as public_pptx


@pytest.mark.render
@pytest.mark.parametrize('route', ['native_compile', 'attachment', 'public_export'])
def test_long_notes_are_visible_in_compatible_notes_page(tmp_path, route):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    raw = io.BytesIO()
    deck.save(raw)
    notes = '\n\n'.join(
        f'第{i}段：型号和版本未知时保留原因，人员核查许可资料及适用性；'
        '只读引用不能当作故障已解决，54/120=45%也不能当作收益。'
        for i in range(1, 8)
    ) + '\n末段必须可见：完整保留中文、数字与限定。'
    records = [{'page_id': 'notes-page', 'notes': notes}]
    if route == 'native_compile':
        svg = tmp_path / 'page.svg'
        svg.write_text('<svg viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#fff"/></svg>')
        compiled = compile_deck([SvgInput('notes-page', svg, notes)], CompileOptions(), tmp_path / 'compiled')
        data = compiled.pptx_path.read_bytes()
    else:
        data = attach_notes(raw.getvalue(), records)
    if route == 'public_export':
        data = public_pptx(data, 'long-notes-regression', notes=records, include_notes=True)
    source = tmp_path / 'notes.pptx'
    source.write_bytes(data)
    assert notes_records(data)[0]['notes'] == notes
    from deck_master.pipeline import executable
    subprocess.run([
        executable('soffice'), f'-env:UserInstallation={(tmp_path / "lo-profile").as_uri()}',
        '--headless', '--convert-to',
        'pdf:impress_pdf_Export:{"ExportNotesPages":{"type":"boolean","value":"true"},'
        '"ExportOnlyNotesPages":{"type":"boolean","value":"true"}}',
        '--outdir', str(tmp_path), str(source),
    ], check=True, capture_output=True, timeout=60)
    result = subprocess.run(
        ['pdftotext', str(tmp_path / 'notes.pdf'), '-'],
        check=True, capture_output=True, text=True, timeout=20,
    )
    normalize = lambda text: ''.join(text.split())
    assert normalize(notes) in normalize(result.stdout)
