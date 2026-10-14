"""Reports do not initialize a project, mutate sources or promote identities."""

import hashlib
import uuid

import pytest

from test_eligibility_enforcement import _project, _start, _scope
from deck_master import external_review, models


def _hashes(path):
    return {str(p.relative_to(path)): hashlib.sha256(p.read_bytes()).hexdigest() for p in path.rglob("*") if p.is_file()}


def test_standalone_manuscript_report_is_outside_source_and_preserves_all_bytes(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "copy.md").write_text("# 版本核对\n未知版本先取得铭牌，人工核查后提交。")
    before = _hashes(source)
    result = external_review.evaluate(source, report_output=tmp_path / "reports/report.json")
    assert result["status"] == "needs_reviewer"
    assert result["persistence"] == "none"
    assert result["professional_evidence"] == "not_evaluated"
    assert "人工核查后提交" in result["manuscript"][0]["text"]
    assert _hashes(source) == before
    assert not (source / ".deckmaster").exists()
    with pytest.raises(models.ModelError, match="outside"):
        external_review.evaluate(source, report_output=source / "report.json")
    assert _hashes(source) == before
    with pytest.raises(FileExistsError):
        external_review.evaluate(source, report_output=tmp_path / "reports/report.json")


def test_report_source_hash_mismatch_does_not_write_a_report(tmp_path):
    source = tmp_path / "copy.md"
    source.write_text("已改变的稿件")
    output = tmp_path / "out.json"
    with pytest.raises(models.ModelError, match="differ"):
        external_review.evaluate(source, source_sha256="0" * 64, report_output=output)
    assert not output.exists()


def test_external_pptx_selection_reads_actual_notes_and_refuses_unknown_slide(tmp_path):
    from pptx import Presentation

    source = tmp_path / "external.pptx"
    presentation = Presentation()
    for title in ["第一页", "第二页"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = title + "口播限定"
    presentation.save(source)
    before = source.read_bytes()
    report = external_review.evaluate(source, scope={"slide_numbers": [2]})
    assert len(report["manuscript"][0]["slides"]) == 1
    assert report["manuscript"][0]["slides"][0]["speaker_notes"] == "第二页口播限定"
    assert source.read_bytes() == before
    with pytest.raises(models.ModelError, match="not in"):
        external_review.evaluate(source, scope={"slide_numbers": [3]})


def test_explicit_import_retains_needs_review_and_tool_identity_even_when_report_claims_human(tmp_path):
    project, store = _project(tmp_path)
    _start(project, store, intent="review_only", scope=_scope(page_ids=["p1"], layers=["review"]))
    doc = store.load_document()
    page = next(e for e in doc["pages"] if e["page_id"] == "p1")
    report = {
        "schema_version": "external_review_report.v1",
        "status": "pass",
        "reviewer": {"type": "human_external", "id": "unverified-person"},
        "source_hashes": [{"source_uri": "provided:copy.md", "sha256": "1" * 64}],
    }
    payload = {
        "external_report": report,
        "source_bindings": [{"source_uri": "provided:copy.md", "source_sha256": "1" * 64, "page_id": "p1", "page_ref": page["page"]}],
    }
    operation = str(uuid.uuid4())
    result = external_review.import_report(project, input=payload, base_revision=doc["revision_id"], operation_id=operation)
    review = store.read_object_json(result["review_refs"][0])
    assert review["status"] == "needs_review" and review["reviewer"]["type"] == "tool"
    assert result["assurance"] == "unverified"
    assert store.load_document()["pages"] == doc["pages"]
    replay = external_review.import_report(project, input=payload, base_revision=doc["revision_id"], operation_id=operation)
    assert replay["status"] == "already_applied"
