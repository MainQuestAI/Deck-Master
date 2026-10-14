"""Actual notes text, stable ordering, and notes-only visual preservation."""

import copy
import hashlib
import json

import pytest
from pptx import Presentation

from deck_master.compiler import CompileOptions, SvgInput, compile_deck
from deck_master.compiler.svg import parse_svg
from deck_master import editing, models
from deck_master.pipeline import readback, output_matches_current_svg


def _svg(tmp_path):
    path = tmp_path / "slide.svg"
    path.write_text('<svg viewBox="0 0 200 100"><rect width="200" height="100" fill="#fff"/></svg>')
    return path


def test_three_page_notes_are_complete_and_follow_stable_ids_after_reorder(tmp_path):
    path = _svg(tmp_path)
    notes = {
        "p-a": "中文 54/120＝45%\n限定：不是节约率。",
        "p-b": "用户原备注\n\n保留第二行与空行。",
        "p-c": "版本未知时停止检索；人工核对后提交。",
    }
    for order, dirname in [(["p-a", "p-b", "p-c"], "first"), (["p-c", "p-a", "p-b"], "reordered")]:
        result = compile_deck([SvgInput(i, path, notes[i]) for i in order], CompileOptions(), tmp_path / dirname)
        actual = Presentation(result.pptx_path)
        assert [s.notes_slide.notes_text_frame.text for s in actual.slides] == [notes[i] for i in order]
        manifest = json.loads(result.manifest_path.read_text())
        assert [(p["page_id"], p["speaker_notes"]) for p in manifest["pages"]] == [(i, notes[i]) for i in order]
        assert all(p["speaker_notes_sha256"] == hashlib.sha256(p["speaker_notes"].encode()).hexdigest() for p in manifest["pages"])


def test_readback_fails_when_real_pptx_notes_differ_even_when_visible_text_matches(tmp_path):
    path = _svg(tmp_path)
    result = compile_deck([SvgInput("p-a", path, "已核对版本")], CompileOptions(), tmp_path / "compiled")
    page = {"page_id": "p-a", "customer_visible": {}, "speaker_notes": "尚未核对版本", "visual_spec": {}}
    report = readback(result.pptx_path, [parse_svg(path.read_bytes(), page_id="p-a")], [page])
    assert report["status"] == "fail"
    assert any(f["code"] == "speaker_notes_mismatch" for f in report["findings"])


def test_absent_notes_does_not_create_new_notes_and_duplicate_page_ids_are_rejected(tmp_path):
    path = _svg(tmp_path)
    result = compile_deck([SvgInput("p-a", path)], CompileOptions(), tmp_path / "none")
    assert not Presentation(result.pptx_path).slides[0].has_notes_slide
    with pytest.raises(ValueError, match="stable page IDs"):
        compile_deck([SvgInput("p-a", path, "A"), SvgInput("p-a", path, "B")], CompileOptions(), tmp_path / "bad")


def test_notes_only_edit_keeps_visual_slots_and_invalidates_file(created_store):
    store = created_store
    document = store.load_document()
    page = {
        "schema_version": "deck_page_package.v2",
        "page_id": "p-a",
        "customer_visible": {"title": "测试", "body_blocks": []},
        "speaker_notes": "原备注",
        "visual_spec": {"intent": "说明", "layout_hint": "保留正文", "reference_mode": "new_design"},
    }
    # Immutable fixture refs model already completed visual slots; the edit
    # transaction must retain each exact ref and make no Host/image call.
    ref = store.put_json_object(page)
    visual = store.put_json_object({"fixture": "existing visual"})
    updated = models.bump_revision(document, {"operation_id": "setup", "kind": "content_update", "description": "fixture", "read_set": []})
    updated["pages"] = [{"page_id": "p-a", "page": ref, "blueprint": visual, "svg": visual, "svg_preview": visual, "ppt_preview": visual}]
    updated["outputs"] = {key: visual for key in updated["outputs"]}
    updated["content_plan"] = store.put_json_object(
        {"schema_version": "content_plan.v2", "mode": "bound", "page_links": [{"page_id": "p-a", "page_ref": ref}]}
    )
    models.require_writer(updated, "content-plan.v1")
    store.commit_change(base_revision=document["revision_id"], document=updated, operation_id="setup")
    next_page = {**page, "speaker_notes": "修订备注\n限定已保留。"}
    editing.edit_page(
        store.project_root, page=next_page, base_revision=updated["revision_id"], page_hash=ref["sha256"], operation_id="notes-edit"
    )
    now = store.load_document()
    assert all(now["pages"][0][slot] == visual for slot in ["blueprint", "svg", "svg_preview"])
    assert all(value is None for value in now["outputs"].values())
    assert now["tasks"] == updated["tasks"]
    assert store.read_object_json(now["content_plan"])["page_links"][0]["page_ref"] == now["pages"][0]["page"]


def test_old_file_with_same_svg_but_prior_notes_is_not_current(created_store):
    store = created_store
    page = store.put_json_object({"speaker_notes": "当前限定"})
    svg = store.put_json_object({"svg": True})
    old_page = store.put_json_object({"speaker_notes": "旧限定"})
    pptx = store.put_json_object(
        {
            "dependencies": [
                {"kind": "svg", "identity": "p-a", "sha256": svg["sha256"]},
                {"kind": "content", "identity": "p-a", "sha256": old_page["sha256"]},
            ]
        }
    )
    document = {"pages": [{"page_id": "p-a", "page": page, "svg": svg}], "outputs": {"pptx": pptx}}
    assert not output_matches_current_svg(store, document)
    document = copy.deepcopy(document)
    document["pages"][0]["page"] = old_page
    assert output_matches_current_svg(store, document)
