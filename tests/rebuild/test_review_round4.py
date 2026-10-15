"""Migration coverage and enforceable page limits before merge."""
import os
from copy import deepcopy

import pytest

from deck_master import install, service, tasks
from deck_master.editing import check_summary
from test_install_host_skill import _build_legacy_layout, _make_release
from test_flow_quality import _setup_project_with_pages, _raw_commit, _load_fixture
from test_flow_quality import _accept_result
from test_service_flow import _import_current
from rc3_fixtures import claim_fixture_task
from deck_master.models import ModelError


@pytest.mark.parametrize('deferred', [False, True])
def test_migrate_ppt_links_with_ownership_and_optout(tmp_path, monkeypatch, deferred):
    prefix, codex = tmp_path / 'prefix', tmp_path / 'codex'
    monkeypatch.setenv('CODEX_HOME', str(codex))
    _build_legacy_layout(prefix, codex)
    _make_release(prefix, 'new')
    names = ['ppt-' + suffix for suffix in ('master', 'library', 'deck-pro-max', 'quality-gate')]
    root = prefix / '.deck-master'
    for name in names:
        (codex / 'skills' / name).symlink_to(root / 'current/skills' / name)
    foreign = codex / 'skills/ppt-foreign'
    foreign.symlink_to(tmp_path / 'other/current/skills/ppt-foreign')
    (codex / 'skills/ppt-user').mkdir()
    if deferred:
        result = install.activate(prefix, 'new', register_host=False)
        assert result['migration']['removed_links'] == []
        assert all((codex / 'skills' / n).is_symlink() for n in names)
    result = install.activate(prefix, 'new')
    assert len(result['migration']['removed_links']) == 19
    assert all(not (codex / 'skills' / n).is_symlink() for n in names)
    assert os.readlink(foreign) == str(tmp_path / 'other/current/skills/ppt-foreign')
    assert (codex / 'skills/ppt-user').is_dir()
    assert install.activate(prefix, 'new')['migration']['removed_links'] == []


def test_page_limit_rejects_full_import(tmp_path):
    from deck_master.store import Store
    project = tmp_path / 'project'
    draft = _load_fixture('initial-pages.json')
    ids = draft['page_order']
    scope = {'subject': 'deck', 'page_ids': ids, 'page_manifest_ref': None,
        'external_artifact_ref': None, 'layers': ['narrative', 'plan', 'page_copy', 'speech',
            'design', 'blueprint', 'svg', 'review', 'output'],
        'allow_page_membership_change': True, 'allow_page_order_change': True, 'purpose': 'working'}
    service.create(project, brief='Honor the exact two-page maximum', page_limit=2,
        draft={'pages': draft['pages'][:2], 'page_order': ids[:2]}, request_input={
            'intent': 'full_deck', 'content_intent': None, 'scope': scope,
            'instruction_evidence': {'kind': 'host_attested_instruction',
                'text': 'Import these selected pages subject to the two-page maximum.', 'source_locator': 'L1:page-limit'}})
    store = Store(project)
    before = store.load_document()
    with pytest.raises(ModelError, match='hard maximum'):
        _import_current(project, draft_payload=_load_fixture('initial-pages.json'))
    assert store.load_document()['pages'] == before['pages']
    assert store.load_document()['content_basis'] == before['content_basis']


def test_page_limit_rejects_unchanged_revision_then_accepts_reduction(tmp_path):
    project, store, _ = _setup_project_with_pages(tmp_path)
    task = service.inputs_update(project, patch={'reason': 'shorter', 'task_patch': {'page_limit': 2}},
                                 base_revision=store.load_document()['revision_id'], operation_id='shorten')['pending_tasks'][0]
    stored = tasks._lookup_task(store.load_document(), task['task_id'], store)
    claim_fixture_task(project, stored)
    before = store.load_document()
    update = {'input_digest': stored['input_digest'], 'upsert_pages': [], 'remove_page_ids': [],
              'page_order': ['p01','p02','p03'], 'unchanged_reason': 'no change'}
    def accept():
        return _accept_result(project, task_id=task['task_id'], operation_id=task['operation_id'],
                                     produced_against=task['produced_against'], result_payload={'kind': 'compose','content_update': update})
    with pytest.raises(tasks.EnvelopeError, match='page_limit'):
        accept()
    assert store.load_document() == before
    update.update(page_order=['p01','p02'], remove_page_ids=['p03'])
    update['upsert_pages'] = [deepcopy(store.read_object_json(entry['page']))
                             for entry in store.load_document()['pages'][:2]]
    assert accept()['status'] == 'accepted'
    assert len(store.load_document()['pages']) == 2


def test_existing_overlimit_cannot_pass_summary(tmp_path):
    project, store, _ = _setup_project_with_pages(tmp_path)
    _raw_commit(store, lambda d: d['task'].update(page_limit=2))
    summary = check_summary(store, store.load_document())
    assert summary['status'] == 'fail'
    assert 'page_limit' in summary['reason']


def test_ppt_links_restored_after_failed_registration(tmp_path, monkeypatch):
    prefix, codex = tmp_path / 'prefix', tmp_path / 'codex'
    monkeypatch.setenv('CODEX_HOME', str(codex))
    _build_legacy_layout(prefix, codex)
    _make_release(prefix, 'new')
    link = codex / 'skills/ppt-retired'
    target = str(prefix / '.deck-master/current/skills/ppt-retired')
    link.symlink_to(target)
    def fail(*args, **kwargs):
        raise OSError('injected registration')
    monkeypatch.setattr(install, '_sync_host_skill', fail)
    with pytest.raises(OSError, match='injected'):
        install.activate(prefix, 'new')
    assert os.readlink(link) == target
    assert (prefix / '.deck-master/current/companion-manifest.json').is_file()
    assert not list((prefix / '.deck-master').glob('legacy-companion-*'))


@pytest.mark.parametrize('limit', [None, 3, 4])
def test_full_import_at_or_under_limit_remains_accepted(tmp_path, limit):
    project = tmp_path / 'project'
    service.create(project, brief='Explicit import under the actual hard limit', page_limit=limit,
        draft=_load_fixture('initial-pages.json'))
    assert _import_current(project, draft_payload=_load_fixture('initial-pages.json'))['status'] == 'accepted'


def test_initial_manuscript_cannot_exceed_limit(tmp_path):
    project = tmp_path / 'project'
    with pytest.raises(ModelError, match='hard maximum'):
        service.create(project, brief='two pages only', page_limit=2,
                       draft=_load_fixture('initial-pages.json'))
    assert not project.exists()


def test_existing_overlimit_blocks_delivery_and_handoff(tmp_path):
    from deck_master.editing import export_project
    from deck_master.handoff import check_handoff
    from deck_master.store import StoreError
    from deck_master import collaboration
    from deck_master.pipeline import artifact
    from pptx import Presentation
    project, store, _ = _setup_project_with_pages(tmp_path)
    deck = Presentation()
    for entry in store.load_document()['pages']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.name = entry['page_id']
        notes = store.read_object_json(entry['page']).get('speaker_notes') or ''
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    file = tmp_path / 'historical-output.pptx'
    deck.save(file)
    _raw_commit(store, lambda doc: doc['outputs'].update(pptx=artifact(store, file, 'pptx')))
    def historical_overlimit(doc):
        from deck_master.models import compute_input_digest
        doc['task']['page_limit'] = 2
        doc['content_basis']['input_digest'] = compute_input_digest(doc)
    _raw_commit(store, historical_overlimit)
    collaboration.grant(project, input={'mode': 'execute', 'actions': ['export.delivery', 'handoff.delivery'],
        'scope': store.load_document()['collaboration']['request']['scope'],
        'max_external_calls': 0, 'max_candidate_count': 0, 'delivery_binding': 'process_delegation',
        'instruction_evidence': {'kind': 'host_attested_instruction',
            'text': 'Check explicit delivery against the actual hard page limit.', 'source_locator': 'L1:delivery-page-limit'}},
        base_revision=store.current_revision_id(), operation_id='L1-delivery-page-limit')
    doc = store.load_document()
    with pytest.raises(StoreError, match='page_limit'):
        export_project(project, output_dir=tmp_path / 'delivery', purpose='delivery')
    assert not (tmp_path / 'delivery').exists()
    output = tmp_path / 'existing.pptx'
    output.write_bytes(store.read_object_bytes(store.read_object_json(doc['outputs']['pptx'])['file']))
    result = check_handoff(project, file_path=output, purpose='delivery')
    assert result['status'] == 'blocked' and result['review_status'] == 'fail'
