"""RC3-03 / CQ-05: legal Page -> actual SVG -> editable native PPTX.

The manuscript is an explicit synthetic engineering sample, not customer
material, a generated blueprint, model review, human acceptance or E2/E3.
Every compile/font/render path is marked render. No project is created, no
candidate is adopted, and no external image call or authority fixture is used.
"""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
from pptx import Presentation
import pytest

from deck_master.compiler import CompileOptions, SvgInput, compile_deck
from deck_master.compiler.svg import parse_svg
from deck_master.content import validate_page_references, visible_atoms
from deck_master.models import validate_page_semantics
from deck_master.pipeline import executable, readback, render_deck, run
from hostenv import host_fonts, resolve_host_font


pytestmark = pytest.mark.render
PAGE_ID = "synthetic-rc3-relations"
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _atom(kind: str, stable_id: str, field: str = "text") -> str:
    return f"atom:{PAGE_ID}:{kind}:{stable_id}:{field}"


def _page() -> dict:
    page = {
        "schema_version": "deck_page_package.v2", "page_id": PAGE_ID,
        "customer_visible": {
            "title": "人工确认后才提交正式记录",
            "body_blocks": [
                {"id": "model-duty", "type": "paragraph", "text": "生成候选草稿"},
                {"id": "human-duty", "type": "paragraph", "text": "人工核对并确认"},
                {"id": "submit-duty", "type": "paragraph", "text": "提交已确认记录"},
                {"id": "archive-duty", "type": "paragraph", "text": "保留版本与回执"},
                {"id": "exceptions", "type": "table", "columns": [
                    {"id": "owner", "label": "责任主体"},
                    {"id": "action", "label": "异常处理"},
                    {"id": "limit", "label": "边界与单位"},
                ], "rows": [
                    {"id": "human", "cells": [
                        {"column_id": "owner", "display_text": "人工审核员"},
                        {"column_id": "action", "display_text": "事实冲突退回核对"},
                        {"column_id": "limit", "display_text": "确认前不得提交"},
                    ]},
                    {"id": "model", "cells": [
                        {"column_id": "owner", "display_text": "候选生成器"},
                        {"column_id": "action", "display_text": "等待人工确认后重试"},
                        {"column_id": "limit", "display_text": "每次等待 10 ms",
                         "value": 10, "unit": "ms"},
                    ]},
                ]},
            ],
            "labels": [{"id": node, "text": label} for node, label in [
                ("model", "模型起草"), ("human", "人工确认"),
                ("submit", "正式提交"), ("archive", "版本归档"),
            ]],
            "footnotes": [{"id": "qualification", "text": "工程示例；10 ms 是测试参数，不是服务承诺。"}],
        },
        "visual_spec": {
            "intent": "先由人工确认，再提交；异常责任与限制保持各自归属。",
            "reference_mode": "new_design",
            "nodes": [{"node_id": node, "label_ref": _atom("label", node),
                       "status": "existing",
                       "responsibility_refs": [_atom("block", node + "-duty")]} for node in
                      ("model", "human", "submit", "archive")],
            "edges": [{"edge_id": f"{left}-to-{right}", "from": left, "to": right,
                       "direction": "forward", "relationship": relationship} for left, right, relationship in [
                           ("model", "human", "候选交给人工核对"),
                           ("human", "submit", "仅在人工确认后提交"),
                           ("submit", "archive", "保存已提交版本"),
                       ]],
        },
        "speaker_notes": "先核对人工确认，再核对提交箭头。异常处理按责任主体逐行阅读。\n10 ms 只用于工程验证。",
        "citations": [],
    }
    # These are complete legal objects, unlike old geometry-test stubs that
    # inserted illegal Page node x/y fields or omitted required label/status.
    validate_page_semantics(page)
    validate_page_references(page)
    return page


def _text(parent, atom_id: str, text: str, x: float, y: float, *, node=None,
          size=24, anchor="start"):
    attrs = {"id": atom_id.replace(":", "-"), "data-atom-id": atom_id,
             "x": str(x), "y": str(y), "font-size": str(size),
             "font-family": resolve_host_font(), "text-anchor": anchor, "fill": "#20242a"}
    if node:
        attrs["data-node-ref"] = node
    element = ET.SubElement(parent, "text", attrs)
    element.text = text
    return element


def _svg(page: dict, *, transform=None, swap_cells=False, wrong_responsibility=False,
         wrong_responsibility_geometry=False, reverse_edge=False) -> bytes:
    root = ET.Element("svg", {"xmlns": "http://www.w3.org/2000/svg", "viewBox": "0 0 1280 720"})
    ET.SubElement(root, "rect", {"width": "1280", "height": "720", "fill": "#ffffff"})
    _text(root, f"atom:{PAGE_ID}:title", page["customer_visible"]["title"], 68, 70, size=34)
    group = ET.SubElement(root, "g", {"transform": transform} if transform else {})
    labels = {label["id"]: label["text"] for label in page["customer_visible"]["labels"]}
    blocks = {block["id"]: block for block in page["customer_visible"]["body_blocks"]}
    centers = {node: (168 + 300 * index, 190) for index, node in enumerate(labels)}
    for node, (x, y) in centers.items():
        ET.SubElement(group, "rect", {"id": "node-" + node, "data-node-ref": node,
                      "x": str(x - 100), "y": str(y - 64), "width": "200", "height": "128",
                      "fill": "#f3f5f7", "stroke": "#aab3bd", "stroke-width": "2"})
        _text(group, _atom("label", node), labels[node], x, y - 26, node=node, anchor="middle")
        responsibility_node = "model" if node == "human" and wrong_responsibility else node
        responsibility_x = centers["model"][0] if node == "human" and wrong_responsibility_geometry else x
        _text(group, _atom("block", node + "-duty"), blocks[node + "-duty"]["text"],
              responsibility_x, y + 44, node=responsibility_node, size=19, anchor="middle")
    for edge in page["visual_spec"]["edges"]:
        left, right = centers[edge["from"]], centers[edge["to"]]
        reversed_now = reverse_edge and edge["edge_id"] == "human-to-submit"
        start, end = (right, left) if reversed_now else (left, right)
        ET.SubElement(group, "line", {"id": "edge-" + edge["edge_id"], "data-edge-ref": edge["edge_id"],
                      "x1": str(start[0]), "y1": str(start[1]), "x2": str(end[0]), "y2": str(end[1]),
                      "stroke": "#445b78", "stroke-width": "3", "fill": "none"})
        # SVG marker-end is outside the supported subset: draw a real editable
        # triangle, as well as the edge line, rather than a decorative raster.
        sign = -1 if reversed_now else 1
        tip_x = end[0] - sign * 112
        points = [(tip_x, end[1]), (tip_x - sign * 12, end[1] - 7), (tip_x - sign * 12, end[1] + 7)]
        ET.SubElement(group, "polygon", {"id": "arrowhead-" + edge["edge_id"],
                      "points": " ".join(f"{x},{y}" for x, y in points), "fill": "#445b78"})
    table = blocks["exceptions"]
    x_by_col = {"owner": 86, "action": 404, "limit": 802}
    for index, y in enumerate((374, 434, 494)):
        ET.SubElement(group, "rect", {"id": f"table-row-{index}", "x": "68", "y": str(y - 38),
                      "width": "1120", "height": "60", "fill": "#edf1f5" if index == 0 else "#ffffff",
                      "stroke": "#c2c9d0", "stroke-width": "1"})
    for column in table["columns"]:
        _text(group, _atom("column", f'exceptions:{column["id"]}', "label"), column["label"],
              x_by_col[column["id"]], 374)
    for row_index, row in enumerate(table["rows"]):
        for cell in row["cells"]:
            position_index = 1 - row_index if swap_cells and cell["column_id"] == "action" else row_index
            _text(group, _atom("cell", f'exceptions:{row["id"]}:{cell["column_id"]}', "display_text"),
                  cell["display_text"], x_by_col[cell["column_id"]], 434 + position_index * 60)
    footnote = page["customer_visible"]["footnotes"][0]
    _text(group, _atom("footnote", footnote["id"]), footnote["text"], 68, 578, size=20)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _compile(tmp_path: Path, page: dict, *, name="sample", **svg_options):
    svg_path = tmp_path / f"{name}.svg"
    svg_path.write_bytes(_svg(page, **svg_options))
    compiled = compile_deck([SvgInput(PAGE_ID, svg_path, page["speaker_notes"])],
                            CompileOptions(width_px=1280, height_px=720, fonts=host_fonts()),
                            tmp_path / f"{name}-native")
    parsed = parse_svg(svg_path.read_bytes(), page_id=PAGE_ID)
    report = readback(compiled.pptx_path, [parsed], [page])
    (tmp_path / f"{name}-readback.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return svg_path, compiled, parsed, report


def _native_texts(path):
    with zipfile.ZipFile(path) as archive:
        slide = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    return slide, [node.text or "" for node in slide.findall(".//a:t", NS)]


def _located_atom(report, atom_id):
    findings = [finding for finding in report["findings"] if finding.get("atom_id") == atom_id]
    assert report["status"] == "fail", report
    assert findings, f"relation failure did not locate {atom_id}: {report}"
    assert all(finding["page_id"] == PAGE_ID for finding in findings)
    return findings


def test_legal_full_page_has_native_editable_atoms_units_notes_and_reading_order(tmp_path):
    page = _page()
    _, compiled, parsed, report = _compile(tmp_path, page)
    assert report["status"] == "pass", report
    assert report["visual_review"] == report["desktop_editing"] == "not_evaluated"
    slide, texts = _native_texts(compiled.pptx_path)
    assert texts == [shape["text"] for shape in parsed["shapes"] if shape["kind"] == "text"]
    assert Counter(texts) == Counter(atom["text"] for atom in visible_atoms(page))
    assert report["notes"][0]["page_id"] == PAGE_ID
    assert report["notes"][0]["notes"] == page["speaker_notes"]
    assert "每次等待 10 ms" in texts
    assert "工程示例；10 ms 是测试参数，不是服务承诺。" in texts
    assert not slide.findall(".//p:pic", NS)
    native_atoms = {shape.find("p:nvSpPr/p:cNvPr", NS).get("name"): shape
                    for shape in slide.findall(".//p:sp", NS) if shape.findall(".//a:t", NS)}
    assert set(native_atoms) == {atom["atom_id"] for atom in visible_atoms(page)}
    def native_offset(atom_id, axis):
        return int(native_atoms[atom_id].find("p:spPr/a:xfrm/a:off", NS).get(axis))
    # Check the declared sample's actual visible reading sequence, separately
    # from XML text-run order: chain left-to-right, table headings/cells in
    # column order, and human exceptions above model exceptions.
    chain_x = [native_offset(_atom("label", node), "x") for node in ("model", "human", "submit", "archive")]
    assert chain_x == sorted(chain_x) and len(set(chain_x)) == 4
    for row in ("human", "model"):
        row_x = [native_offset(_atom("cell", f"exceptions:{row}:{column}", "display_text"), "x")
                 for column in ("owner", "action", "limit")]
        assert row_x == sorted(row_x) and len(set(row_x)) == 3
    assert native_offset(_atom("cell", "exceptions:human:owner", "display_text"), "y") < native_offset(
        _atom("cell", "exceptions:model:owner", "display_text"), "y")
    # A real native object can be edited and saved without reconstructing a
    # picture. This does not claim an actual desktop editing acceptance.
    editable = Presentation(compiled.pptx_path)
    atom_id = _atom("block", "human-duty")
    shape = next(shape for shape in editable.slides[0].shapes if shape.name == atom_id)
    shape.text_frame.paragraphs[0].runs[0].text = "已人工核对并确认"
    edited_path = tmp_path / "native-edit-proof.pptx"
    editable.save(edited_path)
    reopened = Presentation(edited_path)
    assert next(shape for shape in reopened.slides[0].shapes if shape.name == atom_id).text == "已人工核对并确认"


def test_identical_table_text_at_wrong_row_is_located_at_cell_atom(tmp_path):
    page = _page()
    _, good, _, good_report = _compile(tmp_path, page, name="table-good")
    _, bad, _, report = _compile(tmp_path, page, name="table-swapped", swap_cells=True)
    assert good_report["status"] == "pass", good_report
    assert _native_texts(good.pptx_path)[1] == _native_texts(bad.pptx_path)[1]
    findings = _located_atom(report, _atom("cell", "exceptions:human:action", "display_text"))
    assert any(finding.get("table_id") == "exceptions" and finding.get("row_id") == "human"
               and finding.get("column_id") == "action" for finding in findings)
    assert not any(finding["code"] in ("missing_visible_atom", "svg_ppt_text_mismatch") for finding in report["findings"])


def test_identical_table_text_at_wrong_column_is_located_at_cell_atom(tmp_path):
    page = _page()
    svg = ET.fromstring(_svg(page))
    atom_id = _atom("cell", "exceptions:human:action", "display_text")
    next(element for element in svg.iter() if element.get("data-atom-id") == atom_id).set("x", "802")
    path = tmp_path / "wrong-column.svg"
    path.write_bytes(ET.tostring(svg))
    compiled = compile_deck([SvgInput(PAGE_ID, path, page["speaker_notes"])],
                            CompileOptions(fonts=host_fonts()), tmp_path / "wrong-column-native")
    report = readback(compiled.pptx_path, [parse_svg(path.read_bytes(), page_id=PAGE_ID)], [page])
    findings = _located_atom(report, atom_id)
    assert any(finding.get("table_id") == "exceptions" and finding.get("column_id") == "action" for finding in findings)


def test_native_only_cell_geometry_corruption_is_located(tmp_path):
    page = _page()
    _, compiled, parsed, original_report = _compile(tmp_path, page, name="native-cell-source")
    assert original_report["status"] == "pass", original_report
    human_atom = _atom("cell", "exceptions:human:action", "display_text")
    model_atom = _atom("cell", "exceptions:model:action", "display_text")
    changed = Presentation(compiled.pptx_path)
    human = next(shape for shape in changed.slides[0].shapes if shape.name == human_atom)
    model = next(shape for shape in changed.slides[0].shapes if shape.name == model_atom)
    human.top, model.top = model.top, human.top
    corrupt = tmp_path / "native-only-cell-swap.pptx"
    changed.save(corrupt)
    assert _native_texts(compiled.pptx_path)[1] == _native_texts(corrupt)[1]
    findings = _located_atom(readback(corrupt, [parsed], [page]), human_atom)
    assert any(finding.get("table_id") == "exceptions" and finding.get("row_id") == "human" for finding in findings)


def test_same_text_reversed_business_arrow_is_located(tmp_path):
    page = _page()
    _, good, _, good_report = _compile(tmp_path, page, name="arrow-good")
    _, bad, _, report = _compile(tmp_path, page, name="arrow-reversed", reverse_edge=True)
    assert good_report["status"] == "pass", good_report
    assert _native_texts(good.pptx_path)[1] == _native_texts(bad.pptx_path)[1]
    finding = next(finding for finding in report["findings"] if finding["code"] == "reversed_arrow")
    assert finding["page_id"] == PAGE_ID and finding["edge_id"] == "human-to-submit"
    assert finding["from"] == "human" and finding["to"] == "submit"
    assert finding["element"] == "edge-human-to-submit"


def _native_arrow_ending(source_pptx, output_path, *, reverse):
    """Change only actual DrawingML end decorations on the business edge."""
    with zipfile.ZipFile(source_pptx) as source, zipfile.ZipFile(output_path, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                root = ET.fromstring(content)
                edge = next(shape for shape in root.findall(".//p:sp", NS)
                            if shape.find("p:nvSpPr/p:cNvPr", NS).get("name") == "edge-human-to-submit")
                line = edge.find("p:spPr/a:ln", NS)
                assert line is not None
                for tag in ("headEnd", "tailEnd"):
                    for ending in line.findall("a:" + tag, NS):
                        line.remove(ending)
                    ET.SubElement(line, "{" + NS["a"] + "}" + tag,
                                  {"type": "triangle" if (tag == "headEnd") == reverse else "none",
                                   "w": "lg", "len": "lg"})
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, content)


def test_native_only_arrow_ending_reversal_with_unchanged_endpoints_is_located(tmp_path):
    page = _page()
    svg = ET.fromstring(_svg(page))
    # This edge uses actual native line-end decorations; no second explicit
    # polygon may obscure which direction the native line itself conveys.
    for parent in svg.iter():
        for child in list(parent):
            if child.get("id") == "arrowhead-human-to-submit":
                parent.remove(child)
    svg_path = tmp_path / "native-endings-source.svg"
    svg_path.write_bytes(ET.tostring(svg))
    compiled = compile_deck([SvgInput(PAGE_ID, svg_path, page["speaker_notes"])],
                            CompileOptions(fonts=host_fonts()), tmp_path / "native-endings-compiled")
    parsed = parse_svg(svg_path.read_bytes(), page_id=PAGE_ID)
    correct = tmp_path / "native-ending-forward.pptx"
    reversed_pptx = tmp_path / "native-ending-reversed.pptx"
    _native_arrow_ending(compiled.pptx_path, correct, reverse=False)
    _native_arrow_ending(correct, reversed_pptx, reverse=True)
    correct_xml, correct_text = _native_texts(correct)
    reversed_xml, reversed_text = _native_texts(reversed_pptx)
    assert correct_text == reversed_text
    correct_edge = next(shape for shape in correct_xml.findall(".//p:sp", NS)
                        if shape.find("p:nvSpPr/p:cNvPr", NS).get("name") == "edge-human-to-submit")
    reversed_edge = next(shape for shape in reversed_xml.findall(".//p:sp", NS)
                         if shape.find("p:nvSpPr/p:cNvPr", NS).get("name") == "edge-human-to-submit")
    for path in ("p:nvSpPr/p:cNvPr", "p:spPr/a:xfrm", "p:spPr/a:custGeom"):
        assert ET.tostring(correct_edge.find(path, NS)) == ET.tostring(reversed_edge.find(path, NS))
    good_line, bad_line = (edge.find("p:spPr/a:ln", NS) for edge in (correct_edge, reversed_edge))
    assert good_line.find("a:headEnd", NS).get("type") == bad_line.find("a:tailEnd", NS).get("type") == "none"
    assert good_line.find("a:tailEnd", NS).get("type") == bad_line.find("a:headEnd", NS).get("type") == "triangle"
    # Check every remaining native XML object/attribute, not just text: only
    # these two line-end decorations are allowed to differ in the paired deck.
    for line in (good_line, bad_line):
        for tag in ("headEnd", "tailEnd"):
            for ending in line.findall("a:" + tag, NS):
                line.remove(ending)
    assert ET.tostring(correct_xml) == ET.tostring(reversed_xml)
    good_report = readback(correct, [parsed], [page])
    bad_report = readback(reversed_pptx, [parsed], [page])
    for name, report in (("forward", good_report), ("reversed", bad_report)):
        (tmp_path / f"native-ending-{name}-readback.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    good_png = render_deck(correct, tmp_path / "native-ending-forward-render", fonts=host_fonts())
    bad_png = render_deck(reversed_pptx, tmp_path / "native-ending-reversed-render", fonts=host_fonts())
    assert len(good_png) == len(bad_png) == 1
    with Image.open(good_png[0]) as forward_image, Image.open(bad_png[0]) as reversed_image:
        assert forward_image.size == reversed_image.size
        assert forward_image.convert("RGB").tobytes() != reversed_image.convert("RGB").tobytes()
    assert good_report["status"] == "pass", good_report
    assert bad_report["status"] == "fail", bad_report
    finding = next(finding for finding in bad_report["findings"] if finding["code"] == "reversed_arrow")
    assert finding["page_id"] == PAGE_ID and finding["edge_id"] == "human-to-submit"
    assert finding["from"] == "human" and finding["to"] == "submit"
    assert finding["element"] == "edge-human-to-submit"


@pytest.mark.parametrize("options", [{"wrong_responsibility": True}, {"wrong_responsibility_geometry": True}])
def test_same_text_wrong_responsibility_owner_is_located(tmp_path, options):
    page = _page()
    _, good, _, good_report = _compile(tmp_path, page, name="responsibility-good")
    _, bad, _, report = _compile(tmp_path, page, name="responsibility-wrong", **options)
    assert good_report["status"] == "pass", good_report
    assert _native_texts(good.pptx_path)[1] == _native_texts(bad.pptx_path)[1]
    findings = _located_atom(report, _atom("block", "human-duty"))
    assert any(finding.get("node_id") == "human" for finding in findings)


@pytest.mark.parametrize("transform", ["translate(24 30)", "translate(30 0) scale(0.92)"])
def test_same_page_relations_with_only_real_svg_transform_pass(tmp_path, transform):
    page = _page()
    original_page = deepcopy(page)
    _, good, good_ir, good_report = _compile(tmp_path, page, name="geometry-good")
    _, moved, moved_ir, report = _compile(tmp_path, page, name="geometry-moved", transform=transform)
    assert page == original_page
    assert good_report["status"] == report["status"] == "pass", report
    assert good.pptx_path.read_bytes() != moved.pptx_path.read_bytes()
    assert _native_texts(good.pptx_path)[1] == _native_texts(moved.pptx_path)[1]
    for original, changed in zip(good_ir["shapes"], moved_ir["shapes"], strict=True):
        assert (original.get("atom_id"), original.get("node_ref"), original.get("edge_ref")) == (
            changed.get("atom_id"), changed.get("node_ref"), changed.get("edge_ref"))
    old_node = next(shape for shape in good_ir["shapes"] if shape["id"] == "node-human")
    new_node = next(shape for shape in moved_ir["shapes"] if shape["id"] == "node-human")
    assert old_node["x"] != new_node["x"]
    assert not any(finding.get("code") in ("reversed_arrow", "missing_visible_atom") for finding in report["findings"])


def test_actual_svg_and_native_renders_retain_the_sample_and_hashes(tmp_path):
    page = _page()
    svg, compiled, _, report = _compile(tmp_path, page, name="render-positive")
    assert report["status"] == "pass", report
    svg_png = tmp_path / "source-svg.png"
    run([executable("rsvg-convert"), "-o", str(svg_png), str(svg)])
    native_pngs = render_deck(compiled.pptx_path, tmp_path / "native-render", fonts=host_fonts())
    assert len(native_pngs) == 1
    for path in (svg_png, native_pngs[0]):
        with Image.open(path) as image:
            assert image.width >= 1280 and image.height >= 720
            assert image.convert("RGB").getextrema() != ((255, 255),) * 3
    evidence = {"evidence_kind": "synthetic_engineering_actual_files",
                "authority_and_review": "not_used", "image_calls": 0,
                "visual_review": "not_evaluated", "desktop_editing": "not_evaluated",
                "files": [{"name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                          for path in (svg, compiled.pptx_path, svg_png, native_pngs[0])]}
    (tmp_path / "native-render-evidence.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    assert len({item["sha256"] for item in evidence["files"]}) == 4


@pytest.mark.parametrize("options, expected_code", [
    ({"swap_cells": True}, "table_cell_ownership_mismatch"),
    ({"reverse_edge": True}, "reversed_arrow"),
    ({"wrong_responsibility_geometry": True}, "node_atom_ownership_mismatch"),
])
def test_negative_source_and_native_renders_are_retained_with_located_readback(tmp_path, options, expected_code):
    page = _page()
    good_svg, good, _, good_report = _compile(tmp_path, page, name="paired-good")
    bad_svg, bad, _, bad_report = _compile(tmp_path, page, name="paired-bad", **options)
    assert _native_texts(good.pptx_path)[1] == _native_texts(bad.pptx_path)[1]
    images = []
    for svg, compiled, name in ((good_svg, good, "good"), (bad_svg, bad, "bad")):
        svg_png = tmp_path / f"{name}-svg.png"
        run([executable("rsvg-convert"), "-o", str(svg_png), str(svg)])
        native_pngs = render_deck(compiled.pptx_path, tmp_path / f"{name}-native-render", fonts=host_fonts())
        assert len(native_pngs) == 1
        images.append((svg_png, native_pngs[0]))
    for source_index in (0, 1):
        with Image.open(images[0][source_index]) as original, Image.open(images[1][source_index]) as changed:
            assert original.size == changed.size
            assert original.convert("RGB").tobytes() != changed.convert("RGB").tobytes()
    # Files are kept even if relation checking fails: a renderer result or full
    # character coverage is never used to waive the semantic negative case.
    assert good_report["status"] == "pass", good_report
    assert bad_report["status"] == "fail", bad_report
    assert any(finding["code"] == expected_code for finding in bad_report["findings"])
