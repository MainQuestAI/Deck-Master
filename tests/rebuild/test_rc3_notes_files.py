"""Synthetic L1 pages, actual native PPTX/export bytes; no desktop/human claim."""
import copy
import io
import json
import zipfile
import uuid

import pytest
from pptx import Presentation

from deck_master import exports, service
from deck_master.compiler import CompileOptions, SvgInput, compile_deck
from deck_master.compiler.native import normalize_notes, notes_records
from deck_master.export_sanitize import pptx
from deck_master.handoff import check_handoff
from deck_master.models import bump_revision, sha256_bytes, validate_schema, require_writer
from deck_master.operations import OperationError
from deck_master.pipeline import artifact, produce, output_matches_current_svg
from deck_master.store import Store
from rc3_fixtures import historical_snapshot

TEXTS = ['用户原注：54/120 = 45%\r\n\r\n不是节约率\n末段甲\n',
         '提示：先说明版本未知\n\n\n末段乙', '完整口播：交接完成不等于解决\n末段丙']


def _project(tmp_path, font=None, *, notes=TEXTS, ids=('p-a', 'p-b', 'p-c')):
    root = tmp_path / 'project'
    pages = [{'schema_version': 'deck_page_package.v2', 'page_id': pid,
              'customer_visible': {'title': pid, 'body_blocks': []},
              'visual_spec': {'intent': 'explicit synthetic notes file regression', 'reference_mode': 'new_design'},
              'speaker_notes': text} for pid, text in zip(ids, notes, strict=True)]
    service.create(root, brief='Synthetic notes file regression', draft={'pages': pages},
                   presentation_mode='read_alone')
    store = Store(root)
    doc = store.load_document()
    updated = bump_revision(doc, {'operation_id': 'seed-synthetic-svg', 'kind': 'artifact_adoption',
                                  'description': 'Synthetic initial SVGs; no model production claim', 'read_set': []})
    inputs = []
    for entry, page in zip(updated['pages'], pages, strict=True):
        path = tmp_path / (entry['page_id'] + '.svg')
        text = f'<text x="40" y="100" font-family="{font}" font-size="32">{page["page_id"]}</text>' if font else ''
        path.write_text('<svg viewBox="0 0 960 540"><rect width="960" height="540" fill="#fff"/>' + text + '</svg>')
        entry['svg'] = artifact(store, path, 'svg', page_id=entry['page_id'])
        inputs.append(SvgInput(entry['page_id'], path, page['speaker_notes']))
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='seed-synthetic-svg')
    return root, store, inputs


def _attach_compiled(store, inputs, tmp_path):
    compiled = compile_deck(inputs, CompileOptions(), tmp_path / 'compiled')
    doc = store.load_document()
    updated = bump_revision(doc, {'operation_id': 'seed-actual-compiled-output', 'kind': 'artifact_adoption',
                                  'description': 'Actual native bytes seeded for export isolation', 'read_set': []})
    updated['outputs']['pptx'] = artifact(store, compiled.pptx_path, 'pptx')
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='seed-actual-compiled-output')
    return compiled


@pytest.mark.render
def test_real_pipeline_to_export_preserves_full_notes_and_dependencies(tmp_path, resolvable_font_family):
    root, store, _ = _project(tmp_path, resolvable_font_family)
    from test_page_visual_gate import _start, _accept, _svg_envelope
    from page_visual_helpers import pass_page_review
    # This is a current production pipeline, so construct all real formal
    # stages using explicitly bounded L1 provider/reviewer fixtures. The lower
    # compiler-only tests below continue to claim only their actual bytes.
    original_svgs = {entry['page_id']: store.read_object_bytes(store.read_object_json(entry['svg'])['file'])
        for entry in store.load_document()['pages']}
    for pid in ('p-a', 'p-b', 'p-c'):
        task, svg = _start(root, pid, svg_builder=lambda page: original_svgs[page['page_id']])
        _accept(root, task, _svg_envelope(pid), file_name='page.svg', data=svg)
        pending = service.continue_project(root)
        assert pending['pending_tasks'][0]['review_stage'] == 'page_visual'
        pass_page_review(root, pending['pending_tasks'][0])
    before = store.load_document()
    report = produce(root)
    assert report['status'] == 'pass' and report['desktop_editing'] == 'not_evaluated'
    doc = store.load_document()
    assert doc['compatibility']['minimum_writer'] == 'method-catalog.v1'
    assert output_matches_current_svg(store, doc)
    assert [e['svg'] for e in doc['pages']] == [e['svg'] for e in before['pages']]
    assert doc['tasks'] == before['tasks']
    for entry in doc['pages']:
        preview = store.read_object_json(entry['svg_preview'])
        assert all(d['kind'] != 'content' for d in preview['dependencies'])
    result = exports.create(root, output_dir=tmp_path / 'public')
    data, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
    assert [r['notes'] for r in notes_records(data)] == [normalize_notes(t) for t in TEXTS]
    assert [r['page_id'] for r in notes_records(data)] == ['p-a', 'p-b', 'p-c']
    trace = json.loads(store.read_object_bytes(store.read_object_json(doc['outputs']['trace'])['file']))
    assert [p['notes_sha256'] for p in trace['pages']] == [sha256_bytes(normalize_notes(t).encode()) for t in TEXTS]
    output = store.read_object_json(doc['outputs']['pptx'])
    assert [d['identity'] for d in output['dependencies'] if d['kind'] == 'content'] == ['page_notes:p-a', 'page_notes:p-b', 'page_notes:p-c']
    changed = copy.deepcopy(doc)
    page = store.read_object_json(changed['pages'][0]['page'])
    page['speaker_notes'] += '\n新末段'
    changed['pages'][0]['page'] = store.put_json_object(page)
    assert not output_matches_current_svg(store, changed)
    same = copy.deepcopy(doc)
    page = store.read_object_json(same['pages'][0]['page'])
    page['speaker_notes'] = normalize_notes(page['speaker_notes'])
    same['pages'][0]['page'] = store.put_json_object(page)
    assert output_matches_current_svg(store, same)


@pytest.mark.parametrize('selection', ['pptx_notes', 'separate_script', 'both', 'pure_reading'])
def test_public_container_choice_full_readback_and_exact_replay(tmp_path, selection):
    root, store, inputs = _project(tmp_path)
    compiled = _attach_compiled(store, inputs, tmp_path)
    original = compiled.pptx_path.read_bytes()
    result = exports.create(root, notes_policy=selection, output_dir=tmp_path / 'public')
    policy = result['manifest']['notes_policy']
    assert policy['selection'] == selection and policy['source'] == 'explicit'
    data, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
    expected = [normalize_notes(t) for t in TEXTS] if selection in ('pptx_notes', 'both') else ['', '', '']
    assert [r['notes'] for r in notes_records(data)] == expected
    assert ('speaker-script.json' in result['files']) == (selection in ('separate_script', 'both'))
    assert ('speaker-script.txt' in result['files']) == (selection in ('separate_script', 'both'))
    if policy['separate_script']:
        script, _ = exports.download(root, export_id=result['export_id'], filename='speaker-script.json')
        script = json.loads(script)
        assert [p['notes'] for p in script['pages']] == [normalize_notes(t) for t in TEXTS]
        assert [p['page_id'] for p in script['pages']] == ['p-a', 'p-b', 'p-c']
        assert '.deckmaster/' not in json.dumps(script)
        readable, _ = exports.download(root, export_id=result['export_id'], filename='speaker-script.txt')
        assert all(normalize_notes(t) in readable.decode('utf-8') for t in TEXTS)
    assert exports.create(root, export_id=result['export_id'], notes_policy=selection)['manifest'] == result['manifest']
    other = 'pure_reading' if selection != 'pure_reading' else 'both'
    with pytest.raises(OperationError, match='another fixed request'):
        exports.create(root, export_id=result['export_id'], notes_policy=other)
    assert compiled.pptx_path.read_bytes() == original
    validate_schema('export_manifest', result['manifest'])
    handoff = check_handoff(root, file_path=tmp_path / 'public' / 'deck.pptx', notes_policy=selection)
    assert handoff['candidate_relation'] == 'sanitized_copy'
    assert not any('candidate_' in gap for gap in handoff['gaps'])


def test_notes_privacy_rejects_actual_text_but_pure_reading_cleans_hidden_metadata(tmp_path):
    root, store, inputs = _project(tmp_path, notes=['合法原注', '/Users/customer/secret', 'https://example.org/?token=private'])
    compiled = _attach_compiled(store, inputs, tmp_path)
    original = compiled.pptx_path.read_bytes()
    for selection in ('pptx_notes', 'separate_script', 'both'):
        with pytest.raises(OperationError) as failure:
            exports.create(root, notes_policy=selection)
        assert failure.value.payload()['error']['code'] == 'export_privacy_blocked'
    reading = exports.create(root, notes_policy='pure_reading')
    data, _ = exports.download(root, export_id=reading['export_id'], filename='deck.pptx')
    assert [r['notes'] for r in notes_records(data)] == ['', '', '']
    assert compiled.pptx_path.read_bytes() == original
    assert b'/Users/customer/secret' not in data


def test_sanitizer_rebuilds_only_trusted_notes_body_and_rejects_wrong_page_pairing(tmp_path):
    _, _, inputs = _project(tmp_path)
    compiled = compile_deck(inputs, CompileOptions(), tmp_path / 'compiled')
    deck = Presentation(compiled.pptx_path)
    deck.core_properties.author = '/Users/private-author'
    body = deck.slides[0].notes_slide.notes_text_frame
    body._txBody.getparent().xpath('./p:nvSpPr/p:cNvPr')[0].set('descr', '/Users/private-notes-metadata')
    stream = io.BytesIO(); deck.save(stream)
    raw = stream.getvalue()
    records = [{'page_id': item.page_id, 'notes': item.speaker_notes} for item in inputs]
    clean = pptx(raw, 'deck', notes=records)
    with zipfile.ZipFile(io.BytesIO(clean)) as archive:
        assert not any(n.startswith('docProps/') for n in archive.namelist())
        assert all(b'/Users/private' not in archive.read(n) for n in archive.namelist())
    assert [r['notes'] for r in notes_records(clean)] == [normalize_notes(t) for t in TEXTS]
    with pytest.raises(OperationError, match='stable identity'):
        pptx(raw, 'deck', notes=list(reversed(records)))
    with pytest.raises(OperationError):
        pptx(raw, 'deck', notes=[{**records[0], 'notes': records[0]['notes'][:-1]}, *records[1:]])


def test_order_and_none_preserve_original_notes_without_generating_script(tmp_path):
    root, store, inputs = _project(tmp_path, ids=('p-c', 'p-a', 'p-b'))
    _attach_compiled(store, inputs, tmp_path)
    result = exports.create(root)
    data, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
    assert [r['page_id'] for r in notes_records(data)] == ['p-c', 'p-a', 'p-b']
    assert [r['notes'] for r in notes_records(data)] == [normalize_notes(t) for t in TEXTS]
    assert 'speaker-script.json' not in result['files']
    assert result['manifest']['notes_policy']['speech_mode'] == 'none'


def test_separate_script_is_projection_without_compiling_or_creating_pptx(tmp_path):
    root, store, _ = _project(tmp_path)
    before = store.load_document()
    result = exports.create(root, notes_policy='separate_script')
    assert 'speaker-script.json' in result['files'] and 'deck.pptx' not in result['files']
    assert store.load_document() == before


def test_compiler_refuses_duplicate_ids_and_preserves_empty_notes(tmp_path):
    _, _, inputs = _project(tmp_path, notes=['', '', ''])
    with pytest.raises(ValueError, match='unique'):
        compile_deck([inputs[0], inputs[0]], CompileOptions(), tmp_path / 'rejected')
    result = compile_deck(inputs, CompileOptions(), tmp_path / 'compiled')
    assert [r['notes'] for r in notes_records(result.pptx_path)] == ['', '', '']


@pytest.mark.parametrize('mode,speech_source,effective_mode', [
    ('cues', 'provided', 'cues'), ('verbatim', 'provided', 'verbatim'), ('verbatim', 'default', 'cues')])
@pytest.mark.parametrize('output', ['pptx_notes', 'separate_script', 'both'])
def test_selected_real_plan_effective_policy_controls_containers(tmp_path, mode, speech_source, effective_mode, output):
    from deck_master.content_plan import build_v2_plan
    root, store, inputs = _project(tmp_path)
    _attach_compiled(store, inputs, tmp_path)
    doc = store.load_document()
    order = [p['page_id'] for p in doc['pages']]
    plan_input = {'schema_version': 'content_plan_input.v2',
                  'chapters': [{'chapter_id': 'c1', 'title': 'Synthetic notes', 'page_ids': order}],
                  'page_order': order,
                  'page_roles': [{'page_id': pid, 'role': 'synthetic file regression',
                                  'audience_question': '这一页是什么', 'intended_takeaway': '逐页配对',
                                  'narrative_node_ids': [], 'source_links': [], 'required_qualifications': [],
                                  'content_profile': 'simple', 'density_rationale': 'synthetic',
                                  'timing_seconds': None, 'transition_to_next': None} for pid in order],
                  'page_budget': {'exact': 3, 'target': 3, 'max': 3},
                  'time_budget': {'total_seconds': None, 'speaking_seconds': None,
                                  'interaction_seconds': None, 'constraint': 'unspecified'},
                  'speech_policy': {'mode': mode, 'source': speech_source, 'output': output},
                  'issues': [], 'tradeoff_issue_ids': []}
    material = store.put_blob(b'Synthetic fixture; no human confirmation claimed.', ext='txt')
    evidence = {'kind': 'imported_unknown', 'assurance': 'unverified', 'source_ref': material,
                'source_locator': 'synthetic', 'actor_id': None, 'transport_receipt_ref': None,
                'presentation_ref': None}
    plan_ref = build_v2_plan(store, doc, {'task_id': None, 'intent': 'import_draft',
                                        'dispatch_revision': doc['revision_id']},
                             plan_input, doc['pages'], narrative_ref=None, mode='imported_mapping',
                             imported_manuscript={'source_ref': material, 'claimed_confirmation_evidence': evidence,
                                                  'preserve_mainline': True, 'unresolved_mapping_items': []})
    updated = bump_revision(doc, {'operation_id': 'seed-plan', 'kind': 'content_update',
                                  'description': 'Validated synthetic mapping; not formal user import', 'read_set': []})
    updated['content_plan'] = plan_ref
    require_writer(updated, 'rc3.v1')
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='seed-plan')
    before = store.load_document()
    result = exports.create(root)
    policy = result['manifest']['notes_policy']
    assert policy['source'] == 'plan' and policy['speech_mode'] == effective_mode and policy['selection'] == output
    assert policy['plan_sha256'] == plan_ref['sha256']
    data, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
    assert [r['notes'] for r in notes_records(data)] == ([normalize_notes(t) for t in TEXTS] if output != 'separate_script' else ['', '', ''])
    assert ('speaker-script.json' in result['files']) == (output != 'pptx_notes')
    assert store.load_document() == before


@pytest.mark.parametrize('canary', ['/home/private/secret', '/workspace/work/private-source.txt',
                                   'https://example.org/?TOKEN=private', 'Bearer ABCDEFGHIJKLMNOP'])
@pytest.mark.parametrize('selection', ['pptx_notes', 'separate_script'])
def test_each_recognizable_notes_private_value_rejected_without_changing_original(tmp_path, canary, selection):
    root, store, inputs = _project(tmp_path, notes=[canary, '合法中文', '末段'])
    compiled = _attach_compiled(store, inputs, tmp_path)
    original = compiled.pptx_path.read_bytes()
    before = store.load_document()
    with pytest.raises(OperationError) as failure:
        exports.create(root, notes_policy=selection)
    assert failure.value.payload()['error']['code'] == 'export_privacy_blocked'
    assert store.load_document() == before and compiled.pptx_path.read_bytes() == original


def _editorial_fixture(tmp_path):
    """Formal v1 editorial ingress; explicit synthetic goals, not a user approval."""
    root = tmp_path / 'project'
    pages = [{'schema_version': 'deck_page_package.v2', 'page_id': pid,
              'customer_visible': {'title': pid, 'body_blocks': []},
              'visual_spec': {'intent': 'Synthetic stable notes identity', 'reference_mode': 'new_design'},
              'speaker_notes': text} for pid, text in zip(('p-a', 'p-b', 'p-c'), TEXTS, strict=True)]
    goals = [{'goal_id': 'g-' + p['page_id'], 'page_id': p['page_id'], 'purpose': p['page_id'],
              'source_links': [], 'unresolved_facts': ['Synthetic fixture; no customer facts.']} for p in pages]
    plan = {'schema_version': 'content_plan_input.v1', 'input_summary': 'Synthetic editorial notes regression.',
            'chapters': [{'chapter_id': 'c1', 'title': 'Synthetic', 'goal_ids': [g['goal_id'] for g in goals]}],
            'goals': goals, 'unresolved_facts': ['No actual model author or human acceptance claimed.']}
    service.create(root, brief='Synthetic formal editorial notes regression', project_format='workbench.v3',
                   draft={'pages': pages, 'content_plan': plan})
    from test_content_ops import prepare_editorial_project
    prepare_editorial_project(root, review=False)
    return root, Store(root)


def _editorial_change(root, store, action, selected=(), **extra):
    from deck_master import content_ops
    doc = store.load_document()
    entries = {e['page_id']: e for e in doc['pages']}
    value = {'schema_version': 'content_operation_input.v1', 'project_id': doc['project_id'],
             'base_revision': doc['revision_id'], 'content_plan_ref': doc.get('content_plan'),
             'action': action, 'instruction': 'Explicit synthetic notes identity operation: ' + action,
             'targets': [{'page_id': pid, 'page_ref': entries[pid]['page']} for pid in selected], **extra}
    planned = content_ops.plan(root, input=value)
    return content_ops.commit(root, plan_id=planned['plan_id'], base_revision=doc['revision_id'],
                              operation_id=str(uuid.uuid4()))['operation_result']


def _compile_selected_export(root, store, tmp_path):
    doc = store.load_document()
    inputs = []
    for entry in doc['pages']:
        page = store.read_object_json(entry['page'])
        path = tmp_path / (entry['page_id'] + '.svg')
        path.write_text('<svg viewBox="0 0 960 540"><rect width="960" height="540" fill="#fff"/></svg>')
        inputs.append(SvgInput(entry['page_id'], path, page.get('speaker_notes', '')))
    _attach_compiled(store, inputs, tmp_path)
    result = exports.create(root, notes_policy='both')
    actual, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
    return result, notes_records(actual)


@pytest.mark.parametrize('action', ['reorder', 'remove', 'split', 'merge'])
def test_formal_editorial_identity_changes_then_actual_compile_and_export(tmp_path, action):
    from deck_master import content_ops, tasks
    from test_content_ops import derived_page, v2_revision_envelope
    root, store = _editorial_fixture(tmp_path)
    original = copy.deepcopy(store.load_document())
    if action == 'reorder':
        _editorial_change(root, store, action, page_order=['p-c', 'p-a', 'p-b'])
    elif action == 'remove':
        _editorial_change(root, store, action, ('p-b',))
    else:
        selected = ['p-a'] if action == 'split' else ['p-a', 'p-b']
        dispatch = _editorial_change(root, store, action, selected)
        task = tasks._lookup_task(store.load_document(), dispatch['task_ids'][0], store)
        service.task_start(root, task_id=task['task_id'], execution_ref='synthetic-notes-editorial-host',
                           supported_protocols=[task['protocol_version']], capabilities=task['required_capabilities'])
        doc = store.load_document()
        template = store.read_object_json(next(e['page'] for e in doc['pages'] if e['page_id'] == 'p-a'))
        mapping = [('p-a1', '用户原注：54/120 = 45%\n不是节约率'), ('p-a2', '末段甲\n')] if action == 'split' else [
            ('p-ab', '合并口播：54/120 = 45%，不是节约率。\n先说明版本未知。')]
        new_pages = []
        for pid, notes in mapping:
            page = derived_page(template, pid)
            page['customer_visible']['title'] = 'Synthetic explicitly mapped ' + pid
            page['speaker_notes'] = notes
            page['content_mapping']['speech_segments'][0]['end_codepoint'] = len(notes)
            new_pages.append(page)
        outline = copy.deepcopy(store.read_object_json(doc['content_plan'])['input'])
        order = [e['page_id'] for e in doc['pages'] if e['page_id'] not in selected] + [p['page_id'] for p in new_pages]
        role = outline['page_roles'][0]
        outline['page_roles'] = [g for g in outline['page_roles'] if g['page_id'] not in selected] + [
            {**copy.deepcopy(role), 'page_id': p['page_id'], 'role': 'Explicitly mapped synthetic notes',
                'source_links': [], 'narrative_node_ids': [], 'required_qualifications': []} for p in new_pages]
        outline['page_order'] = order
        outline['chapters'] = [{'chapter_id': 'mapped', 'title': 'Explicit synthetic mapping', 'page_ids': order}]
        payload = {'kind': 'compose', 'content_update': {'input_digest': task['input_digest'],
                    'upsert_pages': new_pages, 'remove_page_ids': selected, 'page_order': order,
                    'impact_summary': 'Explicit supplied notes mapping; never copy all source notes to each derived page.'},
                   'content_plan': outline}
        unknown = copy.deepcopy(payload)
        unknown['content_update']['page_order'][-1] = 'unknown-notes-mapping'
        unchanged = copy.deepcopy(store.load_document())
        with pytest.raises((OperationError, service.ServiceError, ValueError)):
            service.accept_result(root, task_id=task['task_id'], operation_id=task['operation_id'],
                                  produced_against=task['produced_against'], result_payload=v2_revision_envelope(store, task, unknown))
        assert store.load_document() == unchanged
        service.accept_result(root, task_id=task['task_id'], operation_id=task['operation_id'],
                              produced_against=task['produced_against'], result_payload=v2_revision_envelope(store, task, payload))
        for pid, notes in mapping:
            lineage = content_ops.lineage(root, page_id=pid)['records'][0]['derivation']
            assert [p['page_id'] for p in lineage['source_pages']] == selected
            assert notes != normalize_notes(TEXTS[0])
    doc = store.load_document()
    expected = [(e['page_id'], normalize_notes(store.read_object_json(e['page']).get('speaker_notes', '')))
                for e in doc['pages']]
    result, actual = _compile_selected_export(root, store, tmp_path)
    assert [(r['page_id'], r['notes']) for r in actual] == expected
    assert [p['page_id'] for p in result['manifest']['notes_policy']['pages']] == [e['page_id'] for e in doc['pages']]
    script, _ = exports.download(root, export_id=result['export_id'], filename='speaker-script.json')
    assert [(r['page_id'], r['notes']) for r in json.loads(script)['pages']] == expected
    assert store.load_document(original['revision_id'])['pages'] == original['pages']
    if action == 'remove':
        assert not any(r['page_id'] == 'p-b' or r['notes'] == normalize_notes(TEXTS[1]) for r in actual)
    if action in ('split', 'merge'):
        assert not any(r['page_id'] in selected for r in actual)


def test_legacy_empty_notes_handoff_does_not_claim_or_bypass_new_notes_guarantee(tmp_path):
    fixture = historical_snapshot(tmp_path, scenario='Actual old zero-slide predicate exemption, no notes promise',
        pages=[{'schema_version': 'deck_page_package.v2', 'page_id': pid,
            'customer_visible': {'title': pid, 'body_blocks': []},
            'visual_spec': {'intent': 'Explicit old predicate facts', 'reference_mode': 'new_design'}}
            for pid in ('p-a', 'p-b', 'p-c')])
    root, store = fixture['project'], Store(fixture['project'])
    candidate = tmp_path / 'old-synthetic-predicate.pptx'
    # Historical delivery-predicate fixture uses a zero-slide PPTX; it does
    # not claim real native compile/readback. New notes-aware outputs must
    # never inherit this narrow legacy provenance-only exemption.
    Presentation().save(candidate)
    doc = store.load_document()
    updated = bump_revision(doc, {'operation_id': 'legacy-output', 'kind': 'artifact_adoption',
                                  'description': 'Legacy synthetic predicate output', 'read_set': []})
    updated['outputs']['pptx'] = artifact(store, candidate, 'pptx')
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='legacy-output')
    old = check_handoff(root, file_path=candidate)
    assert old['candidate_relation'] == 'immutable_original'
    assert 'candidate_notes_policy_mismatch' not in old['gaps']
    explicit = check_handoff(root, file_path=candidate, notes_policy='pptx_notes')
    assert 'candidate_notes_policy_mismatch' in explicit['gaps']
    doc = store.load_document()
    updated = bump_revision(doc, {'operation_id': 'notes-aware-output', 'kind': 'artifact_adoption',
                                  'description': 'Register real notes consumption to reject incomplete file', 'read_set': []})
    updated['outputs']['pptx'] = artifact(store, candidate, 'pptx', dependencies=[{
        'kind': 'content', 'identity': 'page_notes:p-a', 'sha256': sha256_bytes(b'')}])
    require_writer(updated, 'rc3.v1')
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='notes-aware-output')
    assert 'candidate_notes_policy_mismatch' in check_handoff(root, file_path=candidate)['gaps']


@pytest.mark.parametrize('selection', ['pptx_notes', 'both', 'pure_reading', 'separate_script'])
@pytest.mark.parametrize('case', ['correct', 'wrong_order', 'wrong_count'])
def test_empty_notes_new_identity_promise_checked_in_all_final_containers(tmp_path, selection, case):
    root, store, inputs = _project(tmp_path, notes=['', '', ''])
    compiled_inputs = inputs if case == 'correct' else list(reversed(inputs)) if case == 'wrong_order' else inputs[:2]
    compiled = compile_deck(compiled_inputs, CompileOptions(), tmp_path / 'compiled')
    doc = store.load_document()
    deps = [{'kind': 'svg', 'identity': e['page_id'], 'sha256': e['svg']['sha256']} for e in doc['pages']]
    deps += [{'kind': 'content', 'identity': 'page_notes:' + e['page_id'], 'sha256': sha256_bytes(b'')}
             for e in doc['pages']]
    updated = bump_revision(doc, {'operation_id': 'identity-promise', 'kind': 'artifact_adoption',
                                  'description': 'Synthetic adversarial native file, actual ordered dependencies', 'read_set': []})
    updated['outputs']['pptx'] = artifact(store, compiled.pptx_path, 'pptx', dependencies=deps)
    require_writer(updated, 'rc3.v1')
    store.commit_change(base_revision=doc['revision_id'], document=updated, operation_id='identity-promise')
    original = compiled.pptx_path.read_bytes()
    if case != 'correct':
        with pytest.raises(OperationError) as failure:
            exports.create(root, notes_policy=selection)
        assert failure.value.payload()['error']['code'] == 'export_privacy_blocked'
        report = check_handoff(root, file_path=compiled.pptx_path, notes_policy=selection)
        assert any(gap in report['gaps'] for gap in ('candidate_notes_policy_mismatch', 'candidate_notes_page_identity_mismatch'))
    else:
        result = exports.create(root, notes_policy=selection)
        data, _ = exports.download(root, export_id=result['export_id'], filename='deck.pptx')
        assert [(r['page_id'], r['notes']) for r in notes_records(data)] == [('p-a', ''), ('p-b', ''), ('p-c', '')]
        candidate = tmp_path / 'final.pptx'; candidate.write_bytes(data)
        report = check_handoff(root, file_path=candidate, notes_policy=selection)
        assert report['candidate_relation'] == 'sanitized_copy'
        assert not any(gap.startswith('candidate_') for gap in report['gaps'])
    assert compiled.pptx_path.read_bytes() == original


@pytest.mark.parametrize('selection', ['pptx_notes', 'both', 'pure_reading', 'separate_script'])
@pytest.mark.parametrize('compiled_output', [False, True])
def test_private_value_in_public_page_identity_rejected_in_all_metadata(tmp_path, selection, compiled_output):
    canary = 'sk-SYNTHETIC_CANARY_NOT_A_REAL_SECRET'
    root, store, inputs = _project(tmp_path, ids=(canary, 'p-b', 'p-c'))
    if compiled_output:
        _attach_compiled(store, inputs, tmp_path)
    before = store.load_document()
    with pytest.raises(OperationError) as failure:
        exports.create(root, notes_policy=selection, output_dir=tmp_path / 'public')
    assert failure.value.payload()['error']['code'] == 'export_privacy_blocked'
    assert not (tmp_path / 'public').exists()
    assert store.load_document() == before
    assert store.read_object_json(before['pages'][0]['page'])['page_id'] == canary
