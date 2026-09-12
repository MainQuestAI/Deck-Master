"""Approved business notes survive delivery; unbound or internal notes do not."""
import hashlib
import json
import sys
import copy
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from build.native_content import build_native_content_lock
from production.page_package import PageContent, build_page_package
from quality.customer_visible_safety import evaluate_customer_visible_safety_gate


def sample(tmp_path, notes="客户当前值和目标值待确认；未完成订单单列，不计为已达标。"):
    root = tmp_path / "run"
    revision = "a" * 32
    package = build_page_package(run_id="run", content=PageContent(
        page_id="P001", order=1, title="试点统计口径", body_blocks=[], speaker_notes=notes,
        visual_spec={"page_role": "cover"}), status="ready_for_build")
    lock = build_native_content_lock(package)
    state = {"page_packages/P001.json": package,
             "high_density_build/content_locks/P001.content_lock.json": lock}
    hashes = {}
    for relative, data in state.items():
        path = root / "build/revisions" / revision / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, ensure_ascii=False))
        hashes[relative] = hashlib.sha256(path.read_bytes()).hexdigest()
    (root / "build/revisions" / revision / "revision_manifest.json").write_text(json.dumps({"full_snapshot": True, "files": hashes}))
    (root / "build/current_revision.json").write_text(json.dumps({"revision_id": revision}))
    deck = Presentation()
    for field in ("author", "last_modified_by", "comments", "keywords", "category"):
        setattr(deck.core_properties, field, "")
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 2000000, 400000).text = "试点统计口径"
    slide.notes_slide.notes_text_frame.text = notes
    artifact = root / "build/deck.pptx"
    deck.save(artifact)
    result = {"run_id": "run", "engine_id": "deck_native", "status": "compiled", "build_revision": revision,
              "pages": [{"page_id": "P001", "order": 1}],
              "outputs": {"deck_pptx": {"path": "build/deck.pptx", "sha256": hashlib.sha256(artifact.read_bytes()).hexdigest()}}}
    (root / "build/native_compile_result.json").write_text(json.dumps(result))
    return root, artifact, revision, result, lock


def findings(root, artifact, **kwargs):
    return evaluate_customer_visible_safety_gate("run", artifact, run_dir=root, **kwargs)["findings"]


def test_current_locked_business_notes_are_not_internal_leak(tmp_path):
    root, artifact, *_ = sample(tmp_path)
    assert findings(root, artifact) == []


@pytest.mark.parametrize("notes", ["内部提示：下一版补案例", "来源 /Users/private/client.txt", "来源 /Volumes/private/client.txt", "文件 C:/private/client.txt", "来源 /custom/private.txt", "执行 python3 scripts/build.py", "执行 deck-master build", "TODO 追加数据", "SCR 制作指引"])
def test_matching_lock_never_exempts_internal_notes(tmp_path, notes):
    root, artifact, *_ = sample(tmp_path, notes)
    assert findings(root, artifact)


@pytest.mark.parametrize("damage", ["revision", "artifact_hash", "page_order", "lock_hash", "notes_changed", "missing_compile", "draft_package", "absolute_ref", "run_identity", "snapshot_hash"])
def test_notes_require_current_immutable_content_and_artifact_binding(tmp_path, damage):
    root, artifact, revision, result, lock = sample(tmp_path)
    if damage == "revision": result["build_revision"] = "b" * 32
    if damage == "artifact_hash": result["outputs"]["deck_pptx"]["sha256"] = "0" * 64
    if damage == "page_order": result["pages"][0]["page_id"] = "P002"
    if damage == "absolute_ref": result["outputs"]["deck_pptx"]["path"] = str(artifact)
    if damage == "run_identity": result["run_id"] = "another-run"
    if damage == "snapshot_hash":
        (root / "build/revisions" / revision / "page_packages/P001.json").write_text("{}")
    if damage in {"lock_hash", "draft_package"}:
        folder = root / "build/revisions" / revision
        relative = "page_packages/P001.json" if damage == "draft_package" else "high_density_build/content_locks/P001.content_lock.json"
        path = folder / relative
        data = json.loads(path.read_text())
        if damage == "draft_package": data["status"] = "draft"
        else: data["page_package_sha256"] = "0" * 64
        path.write_text(json.dumps(data))
        manifest = json.loads((folder / "revision_manifest.json").read_text())
        manifest["files"][relative] = hashlib.sha256(path.read_bytes()).hexdigest()
        (folder / "revision_manifest.json").write_text(json.dumps(manifest))
    if damage == "notes_changed":
        deck = Presentation(artifact)
        deck.slides[0].notes_slide.notes_text_frame.text = "客户目标已经达成。"
        deck.save(artifact)
        result["outputs"]["deck_pptx"]["sha256"] = hashlib.sha256(artifact.read_bytes()).hexdigest()
    (root / "build/native_compile_result.json").write_text(json.dumps(result))
    if damage == "missing_compile": (root / "build/native_compile_result.json").unlink()
    assert findings(root, artifact)


def test_custom_forbidden_terms_still_scan_approved_notes(tmp_path):
    root, artifact, *_ = sample(tmp_path, "本地客户暗号：测试材料。")
    assert findings(root, artifact, forbidden_terms=["客户暗号"])


def test_public_url_in_locked_business_notes_is_preserved(tmp_path):
    root, artifact, *_ = sample(tmp_path, "统计方法参考 https://example.com/public/method；客户当前值待确认。")
    assert findings(root, artifact) == []


@pytest.mark.parametrize("kind", ["textbox", "table"])
def test_extra_note_shape_cannot_hide_unlocked_content(tmp_path, kind):
    root, artifact, _, result, _ = sample(tmp_path)
    deck = Presentation(artifact)
    if kind == "table":
        extra = deck.slides[0].shapes.add_table(1, 1, 0, 0, 2000000, 400000)
        extra.table.cell(0, 0).text = "不在批准备注中的客户数据"
    else:
        extra = deck.slides[0].shapes.add_textbox(0, 0, 2000000, 400000)
        extra.text = "不在批准备注中的客户数据"
    deck.slides[0].notes_slide.shapes._spTree.append(copy.deepcopy(extra._element))
    extra._element.getparent().remove(extra._element)
    deck.save(artifact)
    result["outputs"]["deck_pptx"]["sha256"] = hashlib.sha256(artifact.read_bytes()).hexdigest()
    (root / "build/native_compile_result.json").write_text(json.dumps(result))
    assert findings(root, artifact)


@pytest.mark.parametrize("damage", ["second_body", "split_runs", "accessibility"])
def test_complete_note_xml_is_checked_without_run_boundary_bypass(tmp_path, damage):
    notes = "SCR 制作指引" if damage == "split_runs" else "客户当前值和目标值待确认。"
    root, artifact, _, result, _ = sample(tmp_path, notes)
    deck = Presentation(artifact)
    note = deck.slides[0].notes_slide
    if damage == "second_body":
        body = note._element.xpath(".//p:sp[p:nvSpPr/p:nvPr/p:ph[@type='body']]")[0]
        extra = copy.deepcopy(body)
        extra.xpath(".//a:t")[0].text = "客户目标已经达成，收益提升35%。"
        note.shapes._spTree.append(extra)
    elif damage == "split_runs":
        paragraph = note.notes_text_frame.paragraphs[0]
        paragraph.clear()
        paragraph.add_run().text = "S"
        paragraph.add_run().text = "CR 制作指引"
    else:
        note._element.xpath(".//p:cNvPr")[0].set("descr", "来源 /Users/private/client.txt")
    deck.save(artifact)
    result["outputs"]["deck_pptx"]["sha256"] = hashlib.sha256(artifact.read_bytes()).hexdigest()
    (root / "build/native_compile_result.json").write_text(json.dumps(result))
    assert findings(root, artifact)
