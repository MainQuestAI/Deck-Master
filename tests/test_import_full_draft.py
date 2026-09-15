"""Full-draft import: agent-written page packages entering the production chain.

Covers the WP3 integration rules: validated one-pass reception, stable page
identity, add/delete/reorder reflected in the active set, interrupted-import
recovery, downstream pruning, and the stage fast path that drives a
package-complete run straight to the builder.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "scripts"))

from build.manifest import build_manifest_v2  # noqa: E402
from high_density.engine import prepare_high_density, run_high_density  # noqa: E402
from narrative.judgment_builder import build_judgments  # noqa: E402
from runtime.orchestration import import_plan  # noqa: E402
from runtime.run_state import create_run, read_json, write_json  # noqa: E402
from runtime.run_state_resolver import resolve_run_state  # noqa: E402

BACKEND_OK = {
    "name": "ppt-master",
    "production_capable": True,
    "contract_versions": ["deck_page_package.v1", "deck_build_manifest.v2"],
}


def _make_run(tmp_path: Path) -> Path:
    return create_run(
        tmp_path,
        {
            "run_id": "full-draft-run",
            "project_name": "Full Draft Run",
            "business_goal": "Deliver an imported complete draft",
            "run_mode": "fixture",
        },
        run_id="full-draft-run",
    )


def _beat(index: int, role: str = "solution") -> dict:
    return {
        "beat_id": f"beat_{index:02d}_{role}",
        "page_title": f"Page {index}",
        "role": role,
        "content_goal": f"Goal {index}",
        "reuse_query": f"query {index}",
        "generation_brief": f"Brief {index}",
    }


def _package(beat_id: str, title: str, body: str) -> dict:
    return {
        "beat_id": beat_id,
        "customer_visible": {
            "title": title,
            "body_blocks": [{"type": "text", "text": body}],
        },
    }


def _full_draft(beats: list[dict], packages: list[dict]) -> dict:
    return {
        "narrative_plan": {"title": "Full Draft", "density": "high", "beats": beats},
        "page_packages": packages,
    }


def _write_input(tmp_path: Path, payload: dict, name: str = "draft.json") -> Path:
    path = tmp_path / name
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _package_files(run_dir: Path) -> set[str]:
    directory = run_dir / "page_packages"
    if not directory.is_dir():
        return set()
    return {path.name for path in directory.glob("*.json") if path.name != "index.json"}


def _draft_v1() -> tuple[list[dict], list[dict]]:
    beats = [_beat(1, "opener"), _beat(2, "problem"), _beat(3, "solution")]
    packages = [
        _package("beat_01_opener", "开场标题", "开场正文"),
        _package("beat_02_problem", "问题标题", "问题正文"),
        _package("beat_03_solution", "方案标题", "方案正文"),
    ]
    return beats, packages


def test_full_draft_import_writes_packages_index_and_mode(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    payload = _full_draft(beats, packages)

    result = import_plan(run_dir, _write_input(tmp_path, payload), source="agent")

    assert result["status"] == "imported"
    assert result["mode"] == "full_draft"
    assert result["page_packages"] == 3
    narrative = read_json(run_dir / "narrative_plan.json")
    assert narrative["target_pages"] == 3
    assert [beat["order"] for beat in narrative["beats"]] == [1, 2, 3]
    written = read_json(run_dir / "page_packages" / "beat_01_opener.json")
    assert written["status"] == "ready_for_build"
    assert written["beat_id"] == "beat_01_opener"
    assert written["run_id"] == "full-draft-run"
    assert written["customer_visible"]["title"] == "开场标题"
    assert len(written["source_fingerprint"]) == 64
    index = read_json(run_dir / "page_packages" / "index.json")
    assert index["page_count"] == 3
    assert [entry["page_id"] for entry in index["pages"]] == [
        "beat_01_opener",
        "beat_02_problem",
        "beat_03_solution",
    ]
    events = (run_dir / "events.jsonl").read_text(encoding="utf-8")
    assert '"mode": "full_draft"' in events


def test_imported_full_draft_builds_without_mbb_and_keeps_nested_copy(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats = [_beat(1, "architecture"), _beat(2, "solution")]
    packages = [
        {
            "beat_id": beats[0]["beat_id"],
            "customer_visible": {
                "title": "职责与数据流",
                "body_blocks": [
                    {
                        "title": "业务工作台",
                        "responsibilities": ["接收任务", "确认结果"],
                        "connections": [{"target": "智能编排", "relation": "提交与回写"}],
                    },
                    {"title": "智能编排", "status": "建设中", "unit": "双向读写"},
                ],
                "footnotes": ["来源：合成输入"],
            },
            "visual_spec": {"page_type": "architecture", "page_role": "architecture"},
        },
        {
            "beat_id": beats[1]["beat_id"],
            "customer_visible": {
                "title": "方案选择",
                "body_blocks": [
                    {"title": "优先方案", "text": "先连接高频资料，再扩展审批回写。"},
                    {"title": "选择理由", "text": "先验证资料命中与责任闭环，降低跨系统改造风险。"},
                ],
            },
            "visual_spec": {"page_type": "comparison", "page_role": "solution"},
        },
    ]
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")
    blueprint_template = (REPO_ROOT / "tests/fixtures/high_density/blueprint.svg").read_bytes()
    blueprint_dir = run_dir / "high_density_build/blueprints"
    blueprint_dir.mkdir(parents=True, exist_ok=True)
    for beat in beats:
        (blueprint_dir / f"{beat['beat_id']}.svg").write_bytes(blueprint_template)

    prepared = prepare_high_density(run_dir)
    result = run_high_density(run_dir)

    assert prepared["status"] == "prepared"
    assert result["status"] == "completed", result
    assert not (run_dir / "high_density_build/mbb/mbb_plan.json").exists()
    assert not (run_dir / "deck_brief.json").exists()
    assert not (run_dir / "claim_map.json").exists()
    write_json(run_dir / "context_manifest.json", {"run_id": run_dir.name, "sources": []})
    assert resolve_run_state(run_dir, run_mode="fixture")["stage"] == "ready_for_client_export"
    presentation = Presentation(run_dir / "high_density_build/pptx/deck_high_density.pptx")
    visible = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    for expected in ("接收任务", "确认结果", "提交与回写", "建设中", "双向读写", "选择理由"):
        assert expected in visible


def test_json_plan_without_packages_reports_plan_mode(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, _ = _draft_v1()
    payload = {"narrative_plan": {"title": "Plan only", "density": "high", "beats": beats}}

    result = import_plan(run_dir, _write_input(tmp_path, payload), source="agent")

    assert result["mode"] == "plan"
    assert "page_packages" not in result
    assert not (run_dir / "page_packages").is_dir()


def test_missing_page_package_rejected_without_writes(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    payload = _full_draft(beats, packages[:2])  # beat_03 missing

    with pytest.raises(Exception, match="missing page packages for beats"):
        import_plan(run_dir, _write_input(tmp_path, payload), source="agent")

    assert not (run_dir / "page_packages").is_dir()
    assert not (run_dir / "narrative_plan.json").exists()


def test_unknown_beat_and_incomplete_body_rejected(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()

    unknown = _full_draft(beats, packages + [_package("beat_99_extra", "多出来的页", "正文")])
    with pytest.raises(Exception, match="unknown beat_id"):
        import_plan(run_dir, _write_input(tmp_path, unknown), source="agent")

    beats2, _ = _draft_v1()
    no_title = _full_draft(
        beats2,
        [
            _package("beat_01_opener", "", "正文"),
            _package("beat_02_problem", "问题标题", "问题正文"),
            _package("beat_03_solution", "方案标题", "方案正文"),
        ],
    )
    with pytest.raises(Exception, match="no complete customer-visible content"):
        import_plan(run_dir, _write_input(tmp_path, no_title), source="agent")

    metadata_only = _full_draft(
        beats2,
        [
            {
                "beat_id": "beat_01_opener",
                "customer_visible": {"title": "开场标题", "body_blocks": [{"type": "text", "text": ""}]},
            },
            *packages[1:],
        ],
    )
    with pytest.raises(Exception, match="no complete customer-visible content"):
        import_plan(run_dir, _write_input(tmp_path, metadata_only, name="metadata-only.json"), source="agent")

    assert not (run_dir / "page_packages").is_dir()


def test_duplicate_beat_package_rejected(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    duplicated = packages + [_package("beat_01_opener", "重复页", "重复正文")]

    with pytest.raises(Exception, match="duplicates beat_id"):
        import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, duplicated)), source="agent")

    assert not (run_dir / "page_packages").is_dir()


def test_citations_get_globally_unique_evidence_ids(tmp_path: Path) -> None:
    """Pages citing sources without ids must not collide on the builder's
    cross-page evidence ledger (found by the fixture build exercise)."""
    run_dir = _make_run(tmp_path)
    write_json(
        run_dir / "context_manifest.json",
        {"schema_version": "deck_context_manifest.v1", "sources": [{"source_id": "material.md", "path": "material.md"}]},
    )
    beats, packages = _draft_v1()
    for package in packages:
        package["citations"] = [{"source": "material.md"}]
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")

    ids: set[str] = set()
    for name in sorted(_package_files(run_dir)):
        package = read_json(run_dir / "page_packages" / name)
        citations = package["citations"]
        assert citations and citations[0]["evidence_id"]
        ids.add(citations[0]["evidence_id"])
    assert len(ids) == 3


def test_citations_require_exact_registered_source_identity(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    write_json(
        run_dir / "context_manifest.json",
        {"schema_version": "deck_context_manifest.v1", "sources": [{"source_id": "src_10", "path": "materials/source-10.md"}]},
    )
    beats, packages = _draft_v1()
    packages[0]["citations"] = [{"source_id": "src_1"}]

    with pytest.raises(Exception, match="citation#1"):
        import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")


def test_evidence_judgment_rejects_prefix_and_ambiguous_source_refs() -> None:
    base = {
        "request": {"run_id": "r1", "business_goal": "提高转化"},
        "deck_brief": {"run_id": "r1", "core_points": ["方案"]},
    }
    claim_map = {"claims": [{"claim_id": "c1", "claim": "结论", "evidence_refs": ["src_1"]}]}

    prefix = build_judgments(
        base["request"],
        base["deck_brief"],
        claim_map,
        {"sources": [{"source_id": "src_10"}]},
    )
    prefix_evidence = next(item for item in prefix["judgments"] if item["judgment_id"] == "judgment_evidence_sufficiency")
    assert prefix_evidence["statement"].startswith("0/1")

    ambiguous = build_judgments(
        base["request"],
        base["deck_brief"],
        claim_map,
        {"sources": [{"source_id": "src_1"}, {"source_id": "src_1"}]},
    )
    ambiguous_evidence = next(item for item in ambiguous["judgments"] if item["judgment_id"] == "judgment_evidence_sufficiency")
    assert ambiguous_evidence["statement"].startswith("0/1")


def test_add_delete_reorder_roundtrip(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")

    v2_beats = [_beat(3, "solution"), _beat(1, "opener"), _beat(4, "case")]
    v2_packages = [
        _package("beat_03_solution", "方案标题改", "方案正文改"),
        _package("beat_01_opener", "开场标题", "开场正文"),
        _package("beat_04_case", "案例标题", "案例正文"),
    ]
    result = import_plan(run_dir, _write_input(tmp_path, _full_draft(v2_beats, v2_packages), name="v2.json"), source="agent")

    assert result["mode"] == "full_draft"
    assert result["removed_pages"] == ["beat_02_problem"]
    assert result["new_pages"] == ["beat_04_case"]
    assert result["reordered"] is True
    assert _package_files(run_dir) == {
        "beat_01_opener.json",
        "beat_03_solution.json",
        "beat_04_case.json",
    }
    index = read_json(run_dir / "page_packages" / "index.json")
    assert [entry["page_id"] for entry in index["pages"]] == [
        "beat_03_solution",
        "beat_01_opener",
        "beat_04_case",
    ]
    narrative = read_json(run_dir / "narrative_plan.json")
    assert [beat["beat_id"] for beat in narrative["beats"]] == [
        "beat_03_solution",
        "beat_01_opener",
        "beat_04_case",
    ]
    assert narrative["target_pages"] == 3
    tasks = read_json(run_dir / "page_tasks.json")["tasks"]
    assert [task["order"] for task in tasks] == [1, 2, 3]
    # page identity survived: beat_01 keeps its original page_id from v1
    kept = read_json(run_dir / "page_packages" / "beat_01_opener.json")
    assert kept["page_id"] == "beat_01_opener"
    backup = Path(result["backup_dir"])
    assert (backup / "page_packages" / "beat_02_problem.json").exists()


def test_interrupted_import_restores_previous_state(tmp_path: Path, monkeypatch) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    v1_narrative = read_json(run_dir / "narrative_plan.json")

    v2_beats = [_beat(1, "opener"), _beat(2, "problem")]
    v2_packages = [
        _package("beat_01_opener", "开场标题", "开场正文"),
        _package("beat_02_problem", "问题标题", "问题正文"),
    ]

    import runtime.orchestration as orchestration_module

    real_write_json = orchestration_module.write_json

    def failing_write_json(path, payload):
        if Path(path).name == "page_tasks.json" and not getattr(failing_write_json, "armed_done", False):
            failing_write_json.armed_done = True
            raise RuntimeError("simulated crash between writes")
        return real_write_json(path, payload)

    monkeypatch.setattr(orchestration_module, "write_json", failing_write_json)

    with pytest.raises(RuntimeError, match="simulated crash"):
        import_plan(run_dir, _write_input(tmp_path, _full_draft(v2_beats, v2_packages), name="v2.json"), source="agent")

    monkeypatch.undo()
    assert read_json(run_dir / "narrative_plan.json") == v1_narrative
    assert _package_files(run_dir) == {
        "beat_01_opener.json",
        "beat_02_problem.json",
        "beat_03_solution.json",
    }
    index = read_json(run_dir / "page_packages" / "index.json")
    assert index["page_count"] == 3

    recovery = import_plan(run_dir, _write_input(tmp_path, _full_draft(v2_beats, v2_packages), name="v2.json"), source="agent")
    assert recovery["mode"] == "full_draft"
    assert _package_files(run_dir) == {"beat_01_opener.json", "beat_02_problem.json"}


def test_downstream_pruned_on_delete_and_change(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    write_json(
        run_dir / "sourcing_plan.json",
        {
            "run_id": "full-draft-run",
            "pages": [
                {"page_id": "beat_01_opener", "beat_id": "beat_01_opener", "decision": "generate"},
                {"page_id": "beat_02_problem", "beat_id": "beat_02_problem", "decision": "generate"},
                {"page_id": "beat_03_solution", "beat_id": "beat_03_solution", "decision": "generate"},
            ],
        },
    )
    write_json(
        run_dir / "preview_manifest.json",
        {
            "run_id": "full-draft-run",
            "pages": [
                {"page_id": "beat_01_opener", "beat_id": "beat_01_opener", "order": 1},
                {"page_id": "beat_02_problem", "beat_id": "beat_02_problem", "order": 2},
                {"page_id": "beat_03_solution", "beat_id": "beat_03_solution", "order": 3},
            ],
        },
    )

    v2_beats = [_beat(1, "opener"), _beat(3, "solution")]
    v2_packages = [
        _package("beat_01_opener", "开场标题", "开场正文"),
        _package("beat_03_solution", "方案标题改", "方案正文改"),
    ]
    result = import_plan(run_dir, _write_input(tmp_path, _full_draft(v2_beats, v2_packages), name="v2.json"), source="agent")

    assert result["changed_pages"] == ["beat_03_solution"]
    assert result["downstream"].get("sourcing_plan.json") == "pruned"
    sourcing = read_json(run_dir / "sourcing_plan.json")
    assert [page["page_id"] for page in sourcing["pages"]] == ["beat_01_opener", "beat_03_solution"]
    preview = read_json(run_dir / "preview_manifest.json")
    assert [page["page_id"] for page in preview["pages"]] == ["beat_01_opener"]


def test_material_change_ignores_spoofed_source_fingerprint_and_invalidates_build(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    old = read_json(run_dir / "page_packages" / "beat_03_solution.json")
    (run_dir / "build").mkdir(exist_ok=True)
    write_json(run_dir / "build" / "build_manifest.json", {"status": "completed"})
    (run_dir / "render_results").mkdir(exist_ok=True)
    write_json(run_dir / "render_results" / "render_result.json", {"status": "completed"})

    changed = _package("beat_03_solution", "方案标题", "已经改变的正文")
    changed["source_fingerprint"] = old["source_fingerprint"]
    v2 = [packages[0], packages[1], changed]
    result = import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, v2), name="v2.json"), source="agent")

    assert result["changed_pages"] == ["beat_03_solution"]
    assert not (run_dir / "build").exists()
    assert not (run_dir / "render_results").exists()


def test_visual_spec_change_invalidates_derived_output(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    (run_dir / "high_density_build").mkdir(exist_ok=True)
    write_json(run_dir / "high_density_build" / "status.json", {"status": "completed"})

    v2 = [dict(item) for item in packages]
    v2[2] = dict(v2[2])
    v2[2]["visual_spec"] = {"diagram": "three-stage operating model"}
    result = import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, v2), name="visual-v2.json"), source="agent")

    assert result["changed_pages"] == ["beat_03_solution"]
    assert (run_dir / "high_density_build").exists()
    assert not (run_dir / "high_density_build/status.json").exists()


def test_downstream_invalidation_failure_rolls_back_entire_import(tmp_path: Path, monkeypatch) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    old_package = read_json(run_dir / "page_packages" / "beat_03_solution.json")
    (run_dir / "build").mkdir(exist_ok=True)
    write_json(run_dir / "build" / "build_manifest.json", {"status": "completed"})
    (run_dir / "generation_tasks").mkdir(exist_ok=True)
    write_json(run_dir / "generation_tasks" / "index.json", {"tasks": [{"page_id": "beat_03_solution"}]})
    (run_dir / "generation_tasks" / "agent-note.txt").write_text("preserve me", encoding="utf-8")

    import runtime.orchestration as orchestration_module

    real_remove = orchestration_module._remove_path

    def fail_once(path: Path) -> None:
        if Path(path).name == "build" and not getattr(fail_once, "failed", False):
            fail_once.failed = True
            raise OSError("simulated invalidation failure")
        real_remove(path)

    monkeypatch.setattr(orchestration_module, "_remove_path", fail_once)
    changed = [packages[0], packages[1], _package("beat_03_solution", "新标题", "新正文")]

    with pytest.raises(OSError, match="simulated invalidation failure"):
        import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, changed), name="v2.json"), source="agent")

    assert read_json(run_dir / "page_packages" / "beat_03_solution.json") == old_package
    assert read_json(run_dir / "build" / "build_manifest.json")["status"] == "completed"
    assert (run_dir / "generation_tasks" / "agent-note.txt").read_text(encoding="utf-8") == "preserve me"
    assert not (run_dir / "overrides" / "plan_import_in_progress.json").exists()


def test_reorder_invalidates_deck_level_downstream_state(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages), name="v1.json"), source="agent")
    write_json(run_dir / "sourcing_plan.json", {"pages": [{"page_id": item["beat_id"]} for item in beats]})
    (run_dir / "quality_reports").mkdir(exist_ok=True)
    write_json(run_dir / "quality_reports" / "old.json", {"status": "pass"})

    order = [beats[2], beats[0], beats[1]]
    by_beat = {item["beat_id"]: item for item in packages}
    result = import_plan(
        run_dir,
        _write_input(tmp_path, _full_draft(order, [by_beat[item["beat_id"]] for item in order]), name="reordered.json"),
        source="agent",
    )

    assert result["reordered"] is True
    assert (run_dir / "sourcing_plan.json").exists()
    assert [page["page_id"] for page in read_json(run_dir / "sourcing_plan.json")["pages"]] == [
        "beat_03_solution", "beat_01_opener", "beat_02_problem"
    ]
    assert not (run_dir / "quality_reports").exists()


def test_stage_fast_path_drives_package_run_to_build(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    for name in ("context_manifest.json", "deck_brief.json", "claim_map.json"):
        write_json(run_dir / name, {})
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")

    state = resolve_run_state(run_dir, run_mode="fixture")

    assert state["stage"] == "needs_build"
    assert "--profile high-density" in state["next_command"]

    # A deleted package file must block production instead of rebuilding the
    # deliverable without the page.
    (run_dir / "page_packages" / "beat_02_problem.json").unlink()
    blocked = resolve_run_state(run_dir, run_mode="fixture")
    assert blocked["stage"] == "blocked_packages"
    assert "re-import the full draft" in blocked["next_command"] or any(
        "page packages are missing" in entry.get("reason", "") for entry in blocked["blocked_actions"]
    )


def test_completed_high_density_package_run_is_ready_without_external_standard_backend(
    tmp_path: Path, monkeypatch
) -> None:
    run_dir = _make_run(tmp_path)
    request = read_json(run_dir / "request.json")
    request.update({"run_mode": "production", "builder_profile": "high_density"})
    write_json(run_dir / "request.json", request)
    for name in ("context_manifest.json", "deck_brief.json", "claim_map.json"):
        write_json(run_dir / name, {})
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")
    (run_dir / "build").mkdir(exist_ok=True)
    write_json(run_dir / "build" / "build_manifest.json", {"schema_version": "deck_build_manifest.v2", "pages": []})
    (run_dir / "render_results").mkdir(exist_ok=True)
    write_json(
        run_dir / "render_results" / "render_result.json",
        {"schema_version": "deck_render_result.v2", "source_mode": "production", "artifact_path": "deck.pptx"},
    )
    monkeypatch.setattr("runtime.run_state_resolver.builder_backend_status", lambda: {"production_capable": False})
    monkeypatch.setattr(
        "runtime.run_state_resolver.resolve_workspace_for_run",
        lambda **_kwargs: {
            "blocked": False,
            "reasons": [],
            "workspace_required": True,
            "workspace_valid": True,
            "resolved_workspace": str(tmp_path),
        },
    )

    state = resolve_run_state(run_dir, run_mode="production")

    assert state["stage"] == "ready_for_client_export"


def test_incomplete_import_marker_blocks_until_retry(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    marker = run_dir / "overrides" / "plan_import_in_progress.json"
    write_json(marker, {"status": "writing", "backup_dir": str(run_dir / "overrides" / "missing")})

    state = resolve_run_state(run_dir, run_mode="fixture")

    assert state["stage"] == "blocked_packages"
    assert "import-plan" in state["next_command"]


def test_legacy_pipeline_unaffected_without_packages(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    for name in ("context_manifest.json", "deck_brief.json", "claim_map.json"):
        write_json(run_dir / name, {})
    beats, _ = _draft_v1()
    write_json(
        run_dir / "narrative_plan.json",
        {"run_id": "full-draft-run", "title": "Plan only", "target_pages": 3, "density": "high", "beats": beats},
    )
    write_json(run_dir / "page_tasks.json", {"run_id": "full-draft-run", "tasks": []})

    state = resolve_run_state(run_dir, run_mode="fixture")

    assert state["stage"] == "needs_sourcing"


def test_build_manifest_v2_carries_imported_content(tmp_path: Path) -> None:
    run_dir = _make_run(tmp_path)
    beats, packages = _draft_v1()
    import_plan(run_dir, _write_input(tmp_path, _full_draft(beats, packages)), source="agent")

    from build.manifest import package_sha256, whitelist_project
    from production.page_package import PagePackageIndex

    loaded = PagePackageIndex(run_dir).list_packages()
    manifest = build_manifest_v2(
        run_id="full-draft-run",
        packages=loaded,
        builder_backend=BACKEND_OK,
        output_profile="production_pptx",
        builder_profile="high_density",
    )
    assert [page["order"] for page in manifest["pages"]] == [1, 2, 3]
    assert manifest["pages"][0]["page_id"] == "beat_01_opener"
    # The build input is bound to the exact imported package file, and the
    # received body text is part of the customer payload projection — not just
    # an "imported" status.
    changed_package = read_json(run_dir / "page_packages" / "beat_03_solution.json")
    assert manifest["pages"][2]["page_package_sha256"] == package_sha256(changed_package)
    projected = whitelist_project(changed_package)
    assert projected["customer_visible"]["title"] == "方案标题"
    assert any("方案正文" in str(block) for block in projected["customer_visible"]["body_blocks"])
