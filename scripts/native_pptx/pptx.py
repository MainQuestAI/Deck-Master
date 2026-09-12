from __future__ import annotations

import base64
import copy
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from collections.abc import Callable as CallableType
from typing import Any, Callable
from xml.etree import ElementTree

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .canvas import NATIVE_CANVAS as _NATIVE_CANVAS
from .contracts import ContractError, assert_v2, read_json, sha256_file, sha256_json, utc_now, write_json
from .svg_pipeline import svg_path, join_cjk_text_lines
from .svg_paint import parse_node_paint, parse_svg_paint
from .svg_native import SvgNativeError, commands_to_svg_path, format_svg_native_error, parse_svg_native
from .visibility import validate_visibility_policy, visible_text_violation
from .visual import VisualMetricsError, compute_visual_metrics, write_visual_metrics

PPTX_DIR = Path("high_density_build/pptx")
TRACE_DIR = Path("high_density_build/traces")
READBACK_DIR = Path("high_density_build/readback")
PREVIEW_DIR = Path("high_density_build/previews")
CANVAS_WIDTH = 1672
CANVAS_HEIGHT = 941
SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5
PPTX_BBOX_TOLERANCE_PT = 0.75
SVG_NS = "http://www.w3.org/2000/svg"


def _slide_width() -> float:
    return 40 / 3 if _NATIVE_CANVAS.get() else SLIDE_WIDTH_IN


class PptxEditabilityError(ContractError):
    pass


def pptx_path(root: Path) -> Path:
    return root / PPTX_DIR / "deck_high_density.pptx"


def trace_path(root: Path) -> Path:
    return root / TRACE_DIR / "pptx_trace.json"


def readback_path(root: Path) -> Path:
    return root / READBACK_DIR / "readback_report.json"


def pptx_preview_path(root: Path, page_id: str) -> Path:
    return root / PREVIEW_DIR / f"{page_id}.pptx.png"


def _rgb(value: Any, default: str = "18212b") -> RGBColor:
    text = str(value or "").lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", text):
        text = default
    return RGBColor.from_string(text.upper())


def _inches(value: float, total: float) -> float:
    if _NATIVE_CANVAS.get():
        return float(value) / CANVAS_WIDTH * (40 / 3)
    return float(value) / total * (SLIDE_WIDTH_IN if total == CANVAS_WIDTH else SLIDE_HEIGHT_IN)


def _canvas_points(value: float) -> Pt:
    return Pt(float(value) * _slide_width() * 72 / CANVAS_WIDTH)


def _remove_children(parent: Any, names: set[str]) -> None:
    for child in list(parent):
        if str(child.tag).split("}")[-1] in names:
            parent.remove(child)


def _append_color(parent: Any, color: str, alpha: float = 1.0) -> None:
    color_node = OxmlElement("a:srgbClr")
    color_node.set("val", str(color).lstrip("#").upper())
    alpha_node = OxmlElement("a:alpha")
    alpha_node.set("val", str(max(0, min(100000, int(round(alpha * 100000))))))
    color_node.append(alpha_node)
    parent.append(color_node)


def _append_gradient(parent: Any, gradient: dict[str, Any], alpha: float) -> None:
    gradient_fill = OxmlElement("a:gradFill")
    gradient_fill.set("rotWithShape", "0")
    stop_list = OxmlElement("a:gsLst")
    for stop in gradient.get("stops") or []:
        stop_node = OxmlElement("a:gs")
        stop_node.set("pos", str(max(0, min(100000, int(round(float(stop["offset"]) * 100000))))))
        _append_color(stop_node, str(stop["color"]), alpha * float(stop.get("opacity", 1)))
        stop_list.append(stop_node)
    gradient_fill.append(stop_list)
    if gradient.get("gradient_type") == "linear":
        linear = OxmlElement("a:lin")
        linear.set("ang", str(int(round(float(gradient.get("angle") or 0) * 60000)) % 21600000))
        linear.set("scaled", "1")
        gradient_fill.append(linear)
    else:
        path = OxmlElement("a:path")
        path.set("path", "circle")
        coordinates = gradient.get("coordinates") or {}
        fill_to = OxmlElement("a:fillToRect")
        fill_to.set("l", str(int(round(float(coordinates.get("fx", 0.5)) * 100000))))
        fill_to.set("t", str(int(round(float(coordinates.get("fy", 0.5)) * 100000))))
        fill_to.set("r", str(int(round((1 - float(coordinates.get("fx", 0.5))) * 100000))))
        fill_to.set("b", str(int(round((1 - float(coordinates.get("fy", 0.5))) * 100000))))
        path.append(fill_to)
        gradient_fill.append(path)
    parent.append(gradient_fill)


def _append_effect(parent: Any, effect: dict[str, Any], alpha: float) -> None:
    effects = OxmlElement("a:effectLst")
    color = str(effect.get("color") or "#000000")
    opacity = alpha * float(effect.get("opacity") or 0)
    std_deviation = float(effect.get("std_deviation") or 0)
    if effect.get("kind") == "shadow":
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", str(int(round(_canvas_points(std_deviation) if _NATIVE_CANVAS.get() else std_deviation * 12700))))
        distance = math.hypot(float(effect.get("dx") or 0), float(effect.get("dy") or 0))
        shadow.set("dist", str(int(round(_canvas_points(distance) if _NATIVE_CANVAS.get() else distance * 12700))))
        shadow.set("dir", str(int(round(math.degrees(math.atan2(float(effect.get("dy") or 0), float(effect.get("dx") or 0))) % 360 * 60000))))
        _append_color(shadow, color, opacity)
        effects.append(shadow)
    else:
        glow = OxmlElement("a:glow")
        glow.set("rad", str(int(round(_canvas_points(std_deviation) if _NATIVE_CANVAS.get() else std_deviation * 12700))))
        _append_color(glow, color, opacity)
        effects.append(glow)
    parent.append(effects)


def _paint_trace(paint: dict[str, Any]) -> dict[str, Any]:
    result = {"kind": str(paint.get("kind") or "none"), "fidelity": str(paint.get("fidelity") or "native")}
    if paint.get("kind") == "solid":
        result["color"] = str(paint.get("color") or "")
    if paint.get("kind") == "gradient":
        result.update({"gradient_id": str(paint.get("gradient_id") or ""), "gradient_type": str(paint.get("gradient_type") or ""), "coordinates": dict(paint.get("coordinates") or {}), "stop_count": len(paint.get("stops") or [])})
    return result


def _line_cap_value(style: dict[str, Any]) -> str:
    value = str(style.get("stroke-linecap", style.get("stroke_linecap", style.get("linecap", "butt"))) or "butt").lower()
    return value if value in {"butt", "round", "square"} else "butt"


def _line_join_value(style: dict[str, Any]) -> str:
    value = str(style.get("stroke-linejoin", style.get("stroke_linejoin", style.get("linejoin", "miter"))) or "miter").lower()
    return value if value in {"round", "bevel", "miter"} else "miter"


def _apply_line_style(shape: Any, style: dict[str, Any], stroke_opacity: float = 1.0) -> None:
    line = shape.line._get_or_add_ln()
    line.set("cap", {"butt": "flat", "round": "rnd", "square": "sq"}[_line_cap_value(style)])
    _remove_children(line, {"round", "bevel", "miter"})
    line.append(OxmlElement(f"a:{_line_join_value(style)}"))
    pairs = style.get("stroke_dash_pairs") or []
    if pairs:
        _remove_children(line, {"prstDash", "custDash"})
        dash = OxmlElement("a:custDash")
        for length, gap in pairs:
            stop = OxmlElement("a:ds")
            stop.set("d", str(round(length * 100000)))
            stop.set("sp", str(round(gap * 100000)))
            dash.append(stop)
        join = next(child for child in line if str(child.tag).split("}")[-1] in {"round", "bevel", "miter"})
        line.insert(list(line).index(join), dash)
    has_gradient_fill = any(str(child.tag).split("}")[-1] == "gradFill" for child in line)
    if stroke_opacity < 1.0 and not has_gradient_fill:
        solid_fill = OxmlElement("a:solidFill")
        _append_color(solid_fill, str(style.get("stroke") or "#d5dde5"), stroke_opacity)
        _remove_children(line, {"solidFill", "gradFill", "noFill", "pattFill", "grpFill"})
        line.append(solid_fill)


def _set_shape_fill(shape: Any, style: dict[str, Any], paint: dict[str, Any] | None = None, trace_entry: dict[str, Any] | None = None) -> None:
    paint = paint or {
        "fill": {"kind": "solid", "color": str(style.get("fill") or "#18212b")} if str(style.get("fill") or "").lower() not in {"", "none", "transparent"} else {"kind": "none"},
        "stroke": {"kind": "solid", "color": str(style.get("stroke") or "#d5dde5")} if str(style.get("stroke") or "").lower() not in {"", "none", "transparent"} else {"kind": "none"},
        "opacity": float(style.get("opacity", 1)),
        "fill_opacity": 1.0,
        "stroke_opacity": 1.0,
        "effect": None,
    }
    overall_opacity = float(paint.get("opacity", 1))
    fill = paint.get("fill") or {"kind": "none"}
    if hasattr(shape, "fill"):
        if fill.get("kind") == "gradient":
            shape.fill.background()
            _remove_children(shape._element.spPr, {"solidFill", "gradFill", "noFill", "blipFill", "pattFill", "grpFill"})
            _append_gradient(shape._element.spPr, fill, overall_opacity * float(paint.get("fill_opacity", 1)))
        elif fill.get("kind") == "solid":
            shape.fill.solid()
            solid_fill = shape._element.spPr.find(qn("a:solidFill"))
            for child in list(solid_fill):
                solid_fill.remove(child)
            _append_color(solid_fill, str(fill.get("color") or "#18212b"),
                          overall_opacity * float(paint.get("fill_opacity", 1)))
        else:
            shape.fill.background()
    stroke = paint.get("stroke") or {"kind": "none"}
    if float(style.get("stroke_width", 1)) == 0:
        stroke = {"kind": "none"}
    if stroke.get("kind") == "gradient":
        line = shape.line._get_or_add_ln()
        _remove_children(line, {"solidFill", "gradFill", "noFill", "pattFill", "grpFill"})
        _append_gradient(line, stroke, overall_opacity * float(paint.get("stroke_opacity", 1)))
        shape.line.width = _canvas_points(float(style.get("stroke_width", 1)))
        _apply_line_style(shape, style, overall_opacity * float(paint.get("stroke_opacity", 1)))
    elif stroke.get("kind") == "solid":
        shape.line.color.rgb = _rgb(stroke.get("color"), "d5dde5")
        shape.line.width = _canvas_points(float(style.get("stroke_width", 1)))
        _apply_line_style(shape, {**style, "stroke": stroke.get("color")}, overall_opacity * float(paint.get("stroke_opacity", 1)))
    else:
        shape.line.fill.background()
    effect = paint.get("effect")
    if effect:
        sp_pr = shape._element.spPr
        _remove_children(sp_pr, {"effectLst"})
        _append_effect(sp_pr, effect, overall_opacity)
    if trace_entry is not None:
        trace_entry["paint"] = {"fill": _paint_trace(fill), "stroke": _paint_trace(stroke)}
        trace_entry["stroke_linecap"] = _line_cap_value(style)
        trace_entry["stroke_linejoin"] = _line_join_value(style)
        trace_entry["stroke_opacity"] = overall_opacity * float(paint.get("stroke_opacity", 1))
        trace_entry["fill_rule"] = str(style.get("fill-rule", style.get("fill_rule", "nonzero")) or "nonzero")
        if effect:
            trace_entry["effect"] = {"type": str(effect.get("kind") or ""), **{key: value for key, value in effect.items() if key not in {"kind"}}}


def _remove_theme_effects(shape: Any) -> None:
    """Remove theme effect references that make LibreOffice add shadows."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _add_text(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    lines = element.get("_text_lines") or []
    reversed_baselines = any(
        previous and following and previous[0].get("position") and following[0].get("position")
        and float(following[0]["position"]["y"]) <= float(previous[0]["position"]["y"])
        for previous, following in zip(lines, lines[1:])
    )
    if _NATIVE_CANVAS.get() and (reversed_baselines or any(run.get("absolute_x") for line in lines for run in line[1:])):
        # A text frame cannot position independent SVG text chunks horizontally.
        # Keep the logical element as an editable group and emit each positioned
        # chunk as a text object. In particular, LO ignores custom tab positions.
        group = slide.shapes.add_group_shape()
        group.name = str(element["element_id"])
        children: list[dict[str, Any]] = []
        chunks: list[list[dict[str, Any]]] = []
        for line in lines:
            for run in line:
                if not chunks or run is line[0] or run.get("absolute_x"):
                    chunks.append([])
                chunks[-1].append(run)
        for index, chunk in enumerate(chunks):
            child = dict(element)
            child["element_id"] = f"{element['element_id']}.text_chunk_{index}"
            child["text"] = "".join(run["text"] for run in chunk)
            child["_text_lines"] = [chunk]
            position = chunk[0].get("position") or {}
            if position.get("anchor") == "start":
                bbox = element["bbox"]
                x = float(position["x"])
                size = float(str(chunk[0]["style"].get("font_size") or 16).removesuffix("px"))
                y = float(position["y"]) - size
                width = float(bbox["x"]) + float(bbox["w"]) - x
                height = float(bbox["y"]) + float(bbox["h"]) - y
                if width <= 0 or height <= 0:
                    raise PptxEditabilityError("positioned text chunk is outside its declared bounds")
                child["bbox"] = {"x": x, "y": y, "w": width, "h": height}
            _add_text(group, child, children)
        # Preserve the SVG logical text container without scaling its children.
        # Group extents normally shrink to child bounds; matching the child
        # coordinate system to the declared container is an identity mapping.
        bbox = element["bbox"]
        xfrm = group._element.grpSpPr.xfrm
        left = Inches(_inches(float(bbox["x"]), CANVAS_WIDTH))
        top = Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT))
        width = Inches(_inches(float(bbox["w"]), CANVAS_WIDTH))
        height = Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT))
        xfrm.off.x = xfrm.chOff.x = left
        xfrm.off.y = xfrm.chOff.y = top
        xfrm.ext.cx = xfrm.chExt.cx = width
        xfrm.ext.cy = xfrm.chExt.cy = height
        trace.append({"element_id": element["element_id"], "object_type": "group",
                      "shape_name": group.name, "bbox": dict(bbox),
                      "priority": element.get("priority", ""), "component_id": element.get("component_id", ""),
                      "z_order": element.get("z_index", 0), "children": children,
                      "child_element_ids": [child["element_id"] for child in children],
                      "fidelity": "native_group"})
        return
    bbox = element["bbox"]
    shape = slide.shapes.add_textbox(Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    # LibreOffice can materialize the theme's default text-box outline unless
    # both the fill and line are explicitly disabled.
    shape.fill.background()
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = not (_NATIVE_CANVAS.get() and element.get("_text_lines"))
    if _NATIVE_CANVAS.get() and element.get("_text_lines"):
        # The template spAutoFit makes LibreOffice reflow even wrap=none.
        frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    style = element.get("style") or {}
    canvas_to_slide = (_slide_width() * 96) / CANVAS_WIDTH
    lines = element.get("_text_lines") or [[{"text": str(element.get("text") or ""), "style": style, "paint": element.get("_paint") or {}}]]
    run_trace: list[dict[str, Any]] = []
    for line_index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(0)
        if _NATIVE_CANVAS.get() and line and line[0].get("position"):
            position = line[0]["position"]
            left = float(position["x"]) - float(bbox["x"])
            width = float(bbox["w"])
            anchor = position["anchor"]
            ppr = paragraph._p.get_or_add_pPr()
            unit = 72 / 96 * canvas_to_slide
            if anchor == "middle":
                ppr.set("algn", "ctr")
                ppr.set("marL", str(round(max(0, 2 * left - width) * unit * 12700)))
                ppr.set("marR", str(round(max(0, width - 2 * left) * unit * 12700)))
            elif anchor == "end":
                ppr.set("algn", "r")
                ppr.set("marR", str(round(max(0, width - left) * unit * 12700)))
            elif anchor == "start":
                ppr.set("algn", "l")
                ppr.set("marL", str(round(max(0, left) * unit * 12700)))
                ppr.set("indent", str(round(min(0, left) * unit * 12700)))
            else:
                raise PptxEditabilityError(f"unsupported SVG text-anchor: {anchor}")
            if line_index == 0:
                size = float(str(line[0]["style"].get("font_size") or 16).removesuffix("px"))
                frame.margin_top = Pt(max(0, float(position["y"]) - size - float(bbox["y"])) * unit)
            if line_index > 0 and lines[line_index - 1][0].get("position"):
                # Paragraph spacing positions this paragraph relative to the
                # preceding baseline, not the following paragraph.
                step = float(position["y"]) - float(lines[line_index - 1][0]["position"]["y"])
                if step <= 0:
                    raise PptxEditabilityError("SVG text lines must have increasing baselines")
                paragraph.line_spacing = Pt(step * unit)
        for run_spec in line:
            run = paragraph.add_run()
            run.text = str(run_spec.get("text") or "")
            run_style = run_spec.get("style") or style
            font = run.font
            font.name = str(run_style.get("font_family") or "Arial")
            font_size = float(str(run_style.get("font_size") or style.get("font_size") or 16).removesuffix("px"))
            font.size = Pt(font_size * 72 / 96 * canvas_to_slide)
            font.bold = str(run_style.get("font_weight") or "400") in {"600", "700", "bold", "Bold"}
            font.italic = str(run_style.get("font_style") or "normal").lower() == "italic"
            font.color.rgb = _rgb(run_style.get("fill"), "18212b")
            paint = run_spec.get("paint") or element.get("_paint") or {}
            text_paint = paint.get("fill") or {"kind": "solid", "color": str(run_style.get("fill") or "#18212b"), "fidelity": "native"}
            run_properties = run._r.get_or_add_rPr()
            _remove_children(run_properties, {"solidFill", "gradFill", "noFill", "blipFill", "pattFill", "grpFill"})
            if text_paint.get("kind") == "gradient":
                _append_gradient(run_properties, text_paint, float(paint.get("opacity", 1)) * float(paint.get("fill_opacity", 1)))
            elif text_paint.get("kind") == "solid":
                solid = OxmlElement("a:solidFill")
                _append_color(solid, str(text_paint.get("color") or "#18212b"), float(paint.get("opacity", 1)) * float(paint.get("fill_opacity", 1)))
                run_properties.append(solid)
            run_trace.append({"text": run.text, "font_family": font.name, "font_size_px": font_size, "font_weight": str(run_style.get("font_weight") or "400"), "font_style": str(run_style.get("font_style") or "normal"), "paint": _paint_trace(text_paint)})
    trace_entry = {"element_id": element["element_id"], "object_type": "text", "shape_name": shape.name, "bbox": bbox, "text": element.get("text", ""), "text_ref": element.get("text_ref", ""), "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0), "runs": run_trace}
    paint = element.get("_paint") or {}
    text_paint = paint.get("fill") or {"kind": "solid", "color": str(style.get("fill") or "#18212b"), "fidelity": "native"}
    if paint.get("effect"):
        _remove_children(shape._element.spPr, {"effectLst"})
        _append_effect(shape._element.spPr, paint["effect"], float(paint.get("opacity", 1)))
    trace_entry["paint"] = {"fill": _paint_trace(text_paint), "stroke": _paint_trace(paint.get("stroke") or {"kind": "none"})}
    if paint.get("effect"):
        trace_entry["effect"] = {"type": str(paint["effect"].get("kind") or ""), **{key: value for key, value in paint["effect"].items() if key not in {"kind"}}}
    if _NATIVE_CANVAS.get():
        trace_entry["text"] = shape.text  # trace records declared SVG paragraph wraps
    trace.append(trace_entry)


def _add_rect(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    style = element.get("style") or {}
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if float(style.get("radius") or 0) > 0 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    trace_entry = {"element_id": element["element_id"], "object_type": "shape", "shape_name": shape.name, "bbox": bbox, "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0)}
    _set_shape_fill(shape, style, element.get("_paint"), trace_entry)
    trace.append(trace_entry)


def _add_line(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["x"]) + float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]) + float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    style = element.get("style") or {}
    trace_entry = {"element_id": element["element_id"], "object_type": "line", "shape_name": shape.name, "bbox": bbox, "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0)}
    _set_shape_fill(shape, style, element.get("_paint"), trace_entry)
    shape.line.width = _canvas_points(float(style.get("stroke_width") or 2))
    trace.append(trace_entry)


_PATH_TOKEN_RE = re.compile(r"([AaCcHhLlMmQqSsTtVvZz])|([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?)")


def _path_points(path_data: str, *, allow_curves: bool = True) -> tuple[list[tuple[float, float]], bool]:
    matches = list(_PATH_TOKEN_RE.finditer(path_data))
    cursor = 0
    for match in matches:
        if path_data[cursor : match.start()].strip(" ,\t\r\n"):
            raise PptxEditabilityError("path contains unsupported separators or syntax")
        cursor = match.end()
    if path_data[cursor:].strip(" ,\t\r\n"):
        raise PptxEditabilityError("path contains unsupported separators or syntax")
    tokens = [(match.group(1) or match.group(2)) for match in matches]
    if not tokens:
        raise PptxEditabilityError("path is empty")
    points: list[tuple[float, float]] = []
    current = (0.0, 0.0)
    start = current
    command = ""
    closed = False
    index = 0

    def number() -> float:
        nonlocal index
        if index >= len(tokens) or re.fullmatch(r"[A-Za-z]", tokens[index]):
            raise PptxEditabilityError("path command has incomplete coordinates")
        value = float(tokens[index])
        index += 1
        return value

    def point(x: float, y: float, relative: bool) -> tuple[float, float]:
        return (x + current[0], y + current[1]) if relative else (x, y)

    while index < len(tokens):
        token = tokens[index]
        if re.fullmatch(r"[A-Za-z]", token):
            command = token
            index += 1
            if command in "Zz":
                if len(points) < 2:
                    raise PptxEditabilityError("closed path needs at least two points")
                closed = True
                current = start
                command = ""
                continue
            if command.upper() not in {"M", "L", "H", "V", "C", "Q"}:
                raise PptxEditabilityError(f"unsupported path command: {command}")
            if command.upper() in {"C", "Q"} and not allow_curves:
                raise PptxEditabilityError(f"curved paths require native reconstruction: {command}")
        if not command:
            raise PptxEditabilityError("path coordinates are missing a command")
        relative = command.islower()
        upper = command.upper()
        if upper in {"M", "L"}:
            target = point(number(), number(), relative)
            current = target
            if upper == "M" and not points:
                start = target
            points.append(target)
            if upper == "M":
                command = "l" if relative else "L"
        elif upper == "H":
            x = number() + current[0] if relative else number()
            current = (x, current[1])
            points.append(current)
        elif upper == "V":
            y = number() + current[1] if relative else number()
            current = (current[0], y)
            points.append(current)
        elif upper == "C":
            c1 = point(number(), number(), relative)
            c2 = point(number(), number(), relative)
            target = point(number(), number(), relative)
            p0 = current
            for step in range(1, 9):
                t = step / 8
                inv = 1 - t
                points.append((inv**3 * p0[0] + 3 * inv**2 * t * c1[0] + 3 * inv * t**2 * c2[0] + t**3 * target[0], inv**3 * p0[1] + 3 * inv**2 * t * c1[1] + 3 * inv * t**2 * c2[1] + t**3 * target[1]))
            current = target
        elif upper == "Q":
            control = point(number(), number(), relative)
            target = point(number(), number(), relative)
            p0 = current
            for step in range(1, 7):
                t = step / 6
                inv = 1 - t
                points.append((inv**2 * p0[0] + 2 * inv * t * control[0] + t**2 * target[0], inv**2 * p0[1] + 2 * inv * t * control[1] + t**2 * target[1]))
            current = target
    if len(points) < 2:
        raise PptxEditabilityError("path needs at least two points")
    return points, closed


def _add_freeform(slide: Any, element: dict[str, Any], points: list[tuple[float, float]], *, closed: bool, trace: list[dict[str, Any]]) -> None:
    x_scale = float(Inches(_slide_width())) / CANVAS_WIDTH
    y_scale = x_scale if _NATIVE_CANVAS.get() else float(Inches(SLIDE_HEIGHT_IN)) / CANVAS_HEIGHT
    emu_points = [(x * x_scale, y * y_scale) for x, y in points]
    builder = slide.shapes.build_freeform(start_x=emu_points[0][0], start_y=emu_points[0][1], scale=1.0)
    builder.add_line_segments(emu_points[1:], close=closed)
    shape = builder.convert_to_shape()
    shape.name = str(element["element_id"])
    # python-pptx sizes the temporary freeform from endpoints only. Curves
    # can extend beyond those endpoints, so reset the shape frame to the
    # normalized SVG bbox before installing the native custom geometry.
    shape.left = Inches(_inches(float(element["bbox"]["x"]), CANVAS_WIDTH))
    shape.top = Inches(_inches(float(element["bbox"]["y"]), CANVAS_HEIGHT))
    shape.width = Inches(_inches(float(element["bbox"]["w"]), CANVAS_WIDTH))
    shape.height = Inches(_inches(float(element["bbox"]["h"]), CANVAS_HEIGHT))
    _remove_theme_effects(shape)
    trace_entry = {"element_id": element["element_id"], "object_type": "freeform", "shape_name": shape.name, "bbox": element["bbox"], "closed": closed, "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0)}
    _set_shape_fill(shape, element.get("style") or {}, element.get("_paint"), trace_entry)
    trace.append(trace_entry)


def _path_command_points(commands: list[dict[str, Any]]) -> list[tuple[float, float]]:
    points: list[tuple[float, float]] = []
    for command in commands:
        if command.get("op") in {"M", "L", "C"}:
            points.append((float(command["x"]), float(command["y"])))
    return points


def _set_custom_geometry(shape: Any, commands: list[dict[str, Any]], bbox: dict[str, Any]) -> None:
    """Replace the temporary builder geometry with native move/line/cubic XML."""
    sp_pr = shape._element.spPr
    _remove_children(sp_pr, {"prstGeom", "custGeom"})
    custom = OxmlElement("a:custGeom")
    for name in ("avLst", "gdLst", "ahLst", "cxnLst"):
        custom.append(OxmlElement(f"a:{name}"))
    rect = OxmlElement("a:rect")
    rect.set("l", "l")
    rect.set("t", "t")
    rect.set("r", "r")
    rect.set("b", "b")
    custom.append(rect)
    path_list = OxmlElement("a:pathLst")
    path = OxmlElement("a:path")
    path.set("w", "100000")
    path.set("h", "100000")
    origin_x, origin_y = float(bbox["x"]), float(bbox["y"])
    width, height = max(float(bbox["w"]), 1e-6), max(float(bbox["h"]), 1e-6)

    def pt(x: float, y: float) -> Any:
        point = OxmlElement("a:pt")
        point.set("x", str(int(round((x - origin_x) / width * 100000))))
        point.set("y", str(int(round((y - origin_y) / height * 100000))))
        return point

    for command in commands:
        op = command["op"]
        if op == "M":
            node = OxmlElement("a:moveTo")
            node.append(pt(float(command["x"]), float(command["y"])))
            path.append(node)
        elif op == "L":
            node = OxmlElement("a:lnTo")
            node.append(pt(float(command["x"]), float(command["y"])))
            path.append(node)
        elif op == "C":
            node = OxmlElement("a:cubicBezTo")
            node.append(pt(float(command["x1"]), float(command["y1"])))
            node.append(pt(float(command["x2"]), float(command["y2"])))
            node.append(pt(float(command["x"]), float(command["y"])))
            path.append(node)
        elif op == "Z":
            path.append(OxmlElement("a:close"))
    path_list.append(path)
    custom.append(path_list)
    sp_pr.append(custom)


def _add_native_path(slide: Any, element: dict[str, Any], commands: list[dict[str, Any]], *, trace: list[dict[str, Any]]) -> None:
    points = _path_command_points(commands)
    if not points:
        raise PptxEditabilityError(f"native SVG path has no endpoints: {element.get('element_id')}")
    x_scale = float(Inches(_slide_width())) / CANVAS_WIDTH
    y_scale = x_scale if _NATIVE_CANVAS.get() else float(Inches(SLIDE_HEIGHT_IN)) / CANVAS_HEIGHT
    emu_points = [(x * x_scale, y * y_scale) for x, y in points]
    builder = slide.shapes.build_freeform(start_x=emu_points[0][0], start_y=emu_points[0][1], scale=1.0)
    if len(emu_points) > 1:
        builder.add_line_segments(emu_points[1:], close=False)
    shape = builder.convert_to_shape()
    shape.name = str(element["element_id"])
    # The temporary builder frame is based on endpoints. Native cubic
    # controls may extend beyond those endpoints, so anchor the shape to the
    # normalized SVG bbox before writing the custom geometry.
    shape.left = Inches(_inches(float(element["bbox"]["x"]), CANVAS_WIDTH))
    shape.top = Inches(_inches(float(element["bbox"]["y"]), CANVAS_HEIGHT))
    shape.width = Inches(_inches(float(element["bbox"]["w"]), CANVAS_WIDTH))
    shape.height = Inches(_inches(float(element["bbox"]["h"]), CANVAS_HEIGHT))
    _remove_theme_effects(shape)
    _set_custom_geometry(shape, commands, element["bbox"])
    trace_entry = {
        "element_id": element["element_id"],
        "object_type": "freeform",
        "shape_name": shape.name,
        "bbox": element["bbox"],
        "closed": any(command.get("op") == "Z" for command in commands),
        "priority": element.get("priority", ""),
        "component_id": element.get("component_id", ""),
        "z_order": element.get("z_index", 0),
        "group_id": str(element.get("group_id") or ""),
        "visual_id": str(element.get("visual_id") or ""),
        "normalized_path": commands_to_svg_path(commands),
        "path_commands": [dict(command) for command in commands],
        "fidelity": "native_normalized",
    }
    _set_shape_fill(shape, element.get("style") or {}, element.get("_paint"), trace_entry)
    trace.append(trace_entry)


def _add_path(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]], *, allow_curves: bool) -> None:
    if element.get("_native_commands") is not None:
        _add_native_path(slide, element, list(element["_native_commands"]), trace=trace)
        return
    points, closed = _path_points(str(element.get("path") or ""), allow_curves=allow_curves)
    _add_freeform(slide, element, points, closed=closed, trace=trace)


def _rectangle_path_bbox(path_data: str) -> dict[str, float] | None:
    """Return a rectangle bbox for the lossless SVG rectangle-path subset."""
    try:
        points, closed = _path_points(path_data, allow_curves=False)
    except PptxEditabilityError:
        return None
    if not closed or len(points) != 4:
        return None
    xs = sorted({round(point[0], 6) for point in points})
    ys = sorted({round(point[1], 6) for point in points})
    if len(xs) != 2 or len(ys) != 2:
        return None
    expected = {(xs[0], ys[0]), (xs[1], ys[0]), (xs[1], ys[1]), (xs[0], ys[1])}
    actual = {(round(point[0], 6), round(point[1], 6)) for point in points}
    if actual != expected or xs[1] <= xs[0] or ys[1] <= ys[0]:
        return None
    return {"x": xs[0], "y": ys[0], "w": xs[1] - xs[0], "h": ys[1] - ys[0]}


def _add_polygon(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    points = [(float(point[0]), float(point[1])) for point in element.get("points") or []]
    if len(points) < 3:
        raise PptxEditabilityError(f"polygon needs at least three points in {element.get('element_id')}")
    _add_freeform(slide, element, points, closed=True, trace=trace)


def _add_ellipse(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    trace_entry = {"element_id": element["element_id"], "object_type": "shape", "shape_name": shape.name, "bbox": bbox, "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0)}
    _set_shape_fill(shape, element.get("style") or {}, element.get("_paint"), trace_entry)
    trace.append(trace_entry)


def _add_image(slide: Any, element: dict[str, Any], asset_paths: dict[str, Path], trace: list[dict[str, Any]]) -> None:
    asset_id = str(element.get("asset_ref") or "")
    asset_path = asset_paths.get(asset_id)
    if asset_path is None or not asset_path.is_file():
        raise PptxEditabilityError(f"registered image asset is unavailable: {asset_id}")
    bbox = element["bbox"]
    shape = slide.shapes.add_picture(str(asset_path), Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), width=Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), height=Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    trace.append({"element_id": element["element_id"], "object_type": "registered_asset", "shape_name": shape.name, "asset_id": asset_id, "bbox": bbox, "priority": element.get("priority", ""), "component_id": element.get("component_id", ""), "z_order": element.get("z_index", 0)})


def _bounds(node: Any) -> dict[str, float]:
    raw = str(node.get("data-pptx-bounds") or "")
    if raw:
        try:
            x, y, w, h = [float(value) for value in raw.split(",")]
            declared = {"x": x, "y": y, "w": w, "h": h}
        except ValueError as exc:
            raise PptxEditabilityError(f"invalid SVG bbox for {node.get('id')}") from exc
    else:
        declared = None
    tag = str(node.tag).split("}")[-1]
    try:
        if tag == "rect":
            return {"x": float(node.get("x") or 0), "y": float(node.get("y") or 0), "w": float(node.get("width") or 0), "h": float(node.get("height") or 0)}
        if tag == "line":
            x1, y1, x2, y2 = (float(node.get(name)) for name in ("x1", "y1", "x2", "y2"))
            return {"x": min(x1, x2), "y": min(y1, y2), "w": abs(x2 - x1), "h": abs(y2 - y1)}
        if tag == "circle":
            cx, cy, radius = (float(node.get(name)) for name in ("cx", "cy", "r"))
            return {"x": cx - radius, "y": cy - radius, "w": radius * 2, "h": radius * 2}
        if tag == "ellipse":
            cx, cy, rx, ry = (float(node.get(name)) for name in ("cx", "cy", "rx", "ry"))
            return {"x": cx - rx, "y": cy - ry, "w": rx * 2, "h": ry * 2}
        if tag in {"polygon", "polyline"}:
            values = [float(value) for value in re.findall(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", str(node.get("points") or ""))]
            if len(values) < 4 or len(values) % 2:
                raise ValueError("invalid point list")
            xs, ys = values[0::2], values[1::2]
            return {"x": min(xs), "y": min(ys), "w": max(xs) - min(xs), "h": max(ys) - min(ys)}
        if tag == "path":
            points, _ = _path_points(str(node.get("d") or ""), allow_curves=True)
            xs, ys = zip(*points)
            return {"x": min(xs), "y": min(ys), "w": max(xs) - min(xs), "h": max(ys) - min(ys)}
        if tag == "image":
            return {"x": float(node.get("x") or 0), "y": float(node.get("y") or 0), "w": float(node.get("width") or 0), "h": float(node.get("height") or 0)}
    except (TypeError, ValueError) as exc:
        raise PptxEditabilityError(f"invalid SVG geometry for {node.get('id')}") from exc
    if declared is not None:
        return declared
    raise PptxEditabilityError(f"SVG element has no geometry bounds: {node.get('id')}")


def _svg_style(node: Any, parent: dict[str, Any] | None = None) -> dict[str, Any]:
    style = dict(parent or {})
    defaults = {
        "fill": "#000000",
        "stroke": "none",
        "stroke_width": "0",
        "opacity": "1",
        "font_family": "Arial",
        "font_size": "16px",
        "font_weight": "400",
        "font_style": "normal",
        "radius": "0",
    }
    for key, value in defaults.items():
        style.setdefault(key, value)
    for attribute in ("fill", "stroke", "stroke-width", "opacity", "font-family", "font-size", "font-weight", "font-style", "rx"):
        if node.get(attribute) is not None:
            key = {"stroke-width": "stroke_width", "font-family": "font_family", "font-size": "font_size", "font-weight": "font_weight", "font-style": "font_style", "rx": "radius"}.get(attribute, attribute)
            style[key] = str(node.get(attribute))
    return style


def _computed_run_paint(node: Any, parent: Any, registry: dict[str, Any]) -> dict[str, Any]:
    inherited = {}
    for attribute in ("fill", "stroke", "opacity", "fill-opacity", "stroke-opacity", "filter"):
        value = node.get(attribute) if node.get(attribute) is not None else parent.get(attribute)
        if value is not None:
            inherited[attribute] = str(value)
    synthetic = ElementTree.Element("tspan", inherited)
    return parse_node_paint(synthetic, registry)


def _svg_text_lines(node: Any, registry: dict[str, Any], *, preserve_positions: bool = True) -> list[list[dict[str, Any]]]:
    parent_style = _svg_style(node)
    parent_paint = parse_node_paint(node, registry)
    lines: list[list[dict[str, Any]]] = [[]]
    current_x = float(node.get("x") or 0)
    current_y = float(node.get("y") or 0)
    anchor = str(node.get("text-anchor") or "start")
    direct_text = str(node.text or "")
    if direct_text:
        lines[0].append({"text": direct_text, "style": parent_style, "paint": parent_paint, "position": {"x": current_x, "y": current_y, "anchor": anchor}})
    for child in list(node):
        if str(child.tag).split("}")[-1] != "tspan":
            continue
        try:
            dy = float(str(child.get("dy") or 0).removesuffix("px"))
        except ValueError as exc:
            raise PptxEditabilityError(f"invalid tspan line offset on {node.get('id')}") from exc
        if lines[-1] and (child.get("y") is not None or abs(dy) > 0.01):
            lines.append([])
        current_x = float(child.get("x") if child.get("x") is not None else current_x) + float(child.get("dx") or 0)
        current_y = float(child.get("y") if child.get("y") is not None else current_y) + dy
        position = {"x": current_x, "y": current_y, "anchor": str(child.get("text-anchor") or anchor)}
        lines[-1].append({"text": "".join(child.itertext()), "style": _svg_style(child, parent_style), "paint": _computed_run_paint(child, node, registry), "position": position, "absolute_x": child.get("x") is not None})
        if child.tail:
            lines[-1].append({"text": str(child.tail), "style": parent_style, "paint": parent_paint})
    declared = str(node.get("data-pptx-text") or "")
    flattened = "".join(run["text"] for line in lines for run in line)
    line_joined = "\n".join("".join(run["text"] for run in line) for line in lines)
    if line_joined == declared:
        return lines
    if flattened == declared:
        # The lock binds characters; SVG tspans bind visual line breaks.
        return lines if preserve_positions else [[run for line in lines for run in line]]
    if " ".join(line_joined.split()) == " ".join(declared.split()):
        return lines
    if join_cjk_text_lines(line_joined.splitlines()) == join_cjk_text_lines(declared.splitlines()):
        return lines
    raise PptxEditabilityError(f"visible SVG text does not match data-pptx-text on {node.get('id')}")


def _node_element(node: Any, scene_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    element_id = str(node.get("id") or "")
    tag = str(node.tag).split("}")[-1]
    kind_by_tag = {"text": "text", "rect": "rect", "circle": "ellipse", "ellipse": "ellipse", "line": "line", "path": "path", "polygon": "polygon", "polyline": "path", "image": "image"}
    kind = kind_by_tag.get(tag)
    if kind is None:
        raise PptxEditabilityError(f"unsupported SVG element {tag} id={element_id}")
    if not element_id or element_id not in scene_by_id:
        raise PptxEditabilityError(f"SVG element is not registered in page_scene.v2: {element_id or '<anonymous>'}")
    source = scene_by_id[element_id]
    element = {
        "element_id": element_id,
        "component_id": str(source.get("component_id") or f"component.{element_id}"),
        "priority": str(source.get("priority") or ""),
        "role": str(source.get("role") or ""),
        "text_ref": str(source.get("text_ref") or ""),
        "asset_ref": str(source.get("asset_ref") or ""),
        "asset_sha256": str(source.get("asset_sha256") or ""),
        "kind": kind,
        "bbox": _bounds(node),
        "z_index": int(node.get("data-pptx-z") or 0),
        "style": _svg_style(node),
    }
    if tag == "text":
        element["text"] = str(node.get("data-pptx-text") or "".join(node.itertext()))
    elif tag in {"path", "polyline"}:
        element["path"] = str(node.get("d") or "") if tag == "path" else "M " + " L ".join(str(point) for point in str(node.get("points") or "").split())
    elif tag == "polygon":
        points: list[list[float]] = []
        for pair in str(node.get("points") or "").split():
            x, y = pair.split(",", 1)
            points.append([float(x), float(y)])
        element["points"] = points
    elif tag == "image":
        element["asset_ref"] = str(node.get("data-pptx-asset-id") or source.get("asset_ref") or "")
    elif tag == "circle":
        if any(node.get(name) is None for name in ("cx", "cy", "r")):
            raise PptxEditabilityError(f"circle geometry is incomplete: {element_id}")
    elif tag == "ellipse":
        if any(node.get(name) is None for name in ("cx", "cy", "rx", "ry")):
            raise PptxEditabilityError(f"ellipse geometry is incomplete: {element_id}")
    return element


def _native_style(style: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in style.items():
        result[key.replace("-", "_")] = value
    if "font_size" in result:
        result["font_size"] = f"{float(result['font_size']):g}px"
    if "stroke_width" in result:
        result["stroke_width"] = float(result["stroke_width"])
    raw_dash = str(result.get("stroke_dasharray") or "none")
    if raw_dash != "none":
        lengths = [float(v.removesuffix("px")) for v in re.split(r"[\s,]+", raw_dash.strip())]
        if len(lengths) % 2:
            lengths *= 2
        width = float(result.get("stroke_width") or 1)
        # DrawingML custom dash lengths are percentages of the line width.
        # Ratios remain invariant under subsequent uniform canvas scaling.
        if sum(lengths) > 0 and width > 0:
            result["stroke_dash_pairs"] = [(lengths[i] / width, lengths[i + 1] / width) for i in range(0, len(lengths), 2)]
    if "radius" not in result:
        result["radius"] = 0
    return result


def _declared_native_bbox(node: Any, fallback: dict[str, Any]) -> dict[str, float]:
    raw = str(node.get("data-pptx-bounds") or "") if node is not None else ""
    values = raw.split(",")
    if len(values) == 4:
        try:
            x, y, w, h = (float(value) for value in values)
            if w >= 0 and h >= 0:
                return {"x": x, "y": y, "w": w, "h": h}
        except ValueError:
            pass
    return dict(fallback)


def _svg_elements(root: Path, scene: dict[str, Any], asset_paths: dict[str, Path] | None = None) -> list[dict[str, Any]]:
    page_id = str(scene.get("page_id") or "")
    try:
        svg_root = ElementTree.fromstring(root.read_text(encoding="utf-8"))
    except (OSError, ElementTree.ParseError) as exc:
        raise PptxEditabilityError(f"approved SVG cannot be parsed: {exc}") from exc
    try:
        native_document = parse_svg_native(svg_root)
    except SvgNativeError as exc:
        raise PptxEditabilityError(f"approved SVG native normalization failed: {format_svg_native_error(exc, page_id=page_id)}") from exc
    scene_by_id = {str(item.get("element_id")): item for item in scene.get("elements", [])}
    scene_by_visual = {str(item.get("visual_id") or ""): item for item in scene.get("visual_registry") or [] if item.get("visual_id")}
    elements: list[dict[str, Any]] = []
    for native in native_document.get("elements") or []:
        element_id = str(native.get("element_id") or "")
        source_id = str(native.get("source_element_id") or "")
        group_id = str(native.get("group_id") or "")
        source = scene_by_id.get(element_id) or scene_by_id.get(source_id) or scene_by_id.get(group_id)
        if source is None and native.get("visual_id"):
            registry = scene_by_visual.get(str(native.get("visual_id")))
            if registry:
                source = scene_by_id.get(str(registry.get("element_id") or ""))
        if source is None:
            raise PptxEditabilityError(f"SVG element is not registered in page_scene.v2: {element_id or source_id or '<anonymous>'}; visual_id={native.get('visual_id') or '<none>'}")
        tag = str(native.get("normalized_tag") or native.get("tag") or "")
        kind = {"text": "text", "rect": "rect", "ellipse": "ellipse", "line": "line", "path": "path", "image": "image"}.get(tag)
        if kind is None:
            raise PptxEditabilityError(f"unsupported normalized SVG element {tag} id={element_id}")
        element = {
            "element_id": element_id,
            "component_id": str(source.get("component_id") or f"component.{element_id}"),
            "priority": str(source.get("priority") or "P2"),
            "role": str(source.get("role") or ""),
            "text_ref": str(source.get("text_ref") or ""),
            "asset_ref": str(source.get("asset_ref") or native.get("asset_ref") or ""),
            "asset_sha256": str(source.get("asset_sha256") or ""),
            "kind": kind,
            # Text has no intrinsic SVG bounds, so its locked layout box is
            # authoritative. Geometric native elements use the normalized
            # parser bbox so a stale data-pptx-bounds cannot hide a transform
            # or path mutation.
            "bbox": _declared_native_bbox(native.get("node"), native.get("bbox") or {}) if tag == "text" else dict(native.get("bbox") or {}),
            "z_index": int(native.get("z_index") or source.get("z_index") or 0),
            "style": _native_style(native.get("style") or {}),
            "_paint": native.get("paint") or {},
            "_native_commands": native.get("commands"),
            "group_id": group_id,
            "visual_id": str(native.get("visual_id") or ""),
        }
        if kind == "text":
            svg_text = str(native.get("text") or "")
            expected_text = str(source.get("text") or "")
            if svg_text != expected_text:
                raise PptxEditabilityError(f"SVG text drift for {element_id}: content-lock text does not match")
            element["text"] = svg_text
            if native.get("node") is not None and native.get("node").get("font-size") is not None and native.get("node").get("fill") is not None:
                try:
                    element["_text_lines"] = _svg_text_lines(native.get("node"), {"gradients": {}, "effects": {}}, preserve_positions=_NATIVE_CANVAS.get())
                except (ContractError, ValueError):
                    if _NATIVE_CANVAS.get():
                        # Replacing rejected positioned runs with declared copy
                        # silently loses both geometry and style in production.
                        raise
                    element["_text_lines"] = [[{"text": svg_text, "style": element["style"], "paint": element["_paint"]}]]
            else:
                element["_text_lines"] = [[{"text": svg_text, "style": element["style"], "paint": element["_paint"]}]]
        if kind == "image":
            asset_ref = str(element.get("asset_ref") or "")
            asset_path = (asset_paths or {}).get(asset_ref)
            expected_sha = str(source.get("asset_sha256") or "")
            if asset_path is None or not asset_path.is_file():
                raise PptxEditabilityError(f"registered image asset is unavailable: {asset_ref}")
            if expected_sha and sha256_file(asset_path) != expected_sha:
                raise PptxEditabilityError(f"registered image asset hash mismatch: {asset_ref}")
        elements.append(element)
    if not elements:
        raise PptxEditabilityError("approved SVG contains no visible elements")
    seen = {item["element_id"] for item in elements}
    missing = [str(item.get("element_id")) for item in scene.get("elements", []) if item.get("kind") != "image" and item.get("priority") in {"P0", "P1"} and str(item.get("element_id")) not in seen]
    if missing:
        raise PptxEditabilityError(f"approved SVG is missing P0/P1 elements: {', '.join(missing)}")
    return elements


def _emit_element(container: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    kind = element.get("kind")
    if kind == "text":
        _add_text(container, element, trace)
    elif kind == "rect":
        _add_rect(container, element, trace)
    elif kind == "line":
        _add_line(container, element, trace)
    elif kind == "path":
        if element.get("_native_commands") is None:
            rectangle_bbox = _rectangle_path_bbox(str(element.get("path") or ""))
        else:
            rectangle_bbox = None
        if rectangle_bbox is not None:
            rectangle = dict(element)
            rectangle["kind"] = "rect"
            rectangle["bbox"] = rectangle_bbox
            _add_rect(container, rectangle, trace)
        else:
            _add_path(container, element, trace, allow_curves=True)
    elif kind == "polygon":
        _add_polygon(container, element, trace)
    elif kind == "ellipse":
        _add_ellipse(container, element, trace)
    elif kind == "image":
        _add_image(container, element, element.get("_asset_paths") or {}, trace)
    else:  # pragma: no cover
        raise PptxEditabilityError(f"unsupported scene element kind: {kind}")


def _trace_bbox(items: list[dict[str, Any]]) -> dict[str, float]:
    if not items:
        return {"x": 0.0, "y": 0.0, "w": 0.0, "h": 0.0}
    left = min(float(item["bbox"]["x"]) for item in items)
    top = min(float(item["bbox"]["y"]) for item in items)
    right = max(float(item["bbox"]["x"]) + float(item["bbox"]["w"]) for item in items)
    bottom = max(float(item["bbox"]["y"]) + float(item["bbox"]["h"]) for item in items)
    return {"x": left, "y": top, "w": right - left, "h": bottom - top}


def _flatten_trace_entries(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    flattened: list[dict[str, Any]] = []
    for entry in entries:
        flattened.append(entry)
        children = entry.get("children") or []
        if isinstance(children, list):
            flattened.extend(_flatten_trace_entries([child for child in children if isinstance(child, dict)]))
    return flattened


def compile_pptx(root: Path, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]], *, asset_paths_by_page: dict[str, dict[str, Path]] | None = None, validate_approved: Callable[[Path, dict[str, Any], dict[str, Any], dict[str, Path]], None] | None = None, svg_paths: dict[str, Path] | None = None, output_root: Path | None = None, canvas_mode: str = "legacy") -> tuple[Path, Path]:
    """Compile with an isolated per-call canvas policy; legacy callers keep their mapping."""
    if canvas_mode not in {"native", "legacy"}:
        raise PptxEditabilityError("unknown compiler canvas mode")
    token = _NATIVE_CANVAS.set(canvas_mode == "native")
    try:
        return _compile_pptx(root, scenes, locks, asset_paths_by_page=asset_paths_by_page,
                             validate_approved=validate_approved, svg_paths=svg_paths, output_root=output_root)
    finally:
        _NATIVE_CANVAS.reset(token)


def _contain_elements(elements: list[dict[str, Any]], svg_file: Path) -> list[dict[str, Any]]:
    """Map SVG viewBox uniformly into a strict 16:9 drawing plane.

    The normalized plane keeps the legacy horizontal unit for the existing
    shape emitters, but uses 1672*9/16 vertically. Geometry, paths, fonts,
    strokes and effects share one source-to-slide scale.
    """
    document = ElementTree.parse(svg_file).getroot()
    try:
        x, y, width, height = [float(value) for value in re.split(r"[\s,]+", str(document.get("viewBox") or "").strip())]
    except ValueError as exc:
        raise PptxEditabilityError("native SVG requires a four-number viewBox") from exc
    if not all(math.isfinite(value) for value in (x, y, width, height)) or width <= 0 or height <= 0:
        raise PptxEditabilityError("native SVG viewBox dimensions must be positive and finite")
    target_height = CANVAS_WIDTH * 9 / 16
    scale = min(CANVAS_WIDTH / width, target_height / height)
    dx, dy = (CANVAS_WIDTH - width * scale) / 2 - x * scale, (target_height - height * scale) / 2 - y * scale
    mapped = copy.deepcopy(elements)
    scaled_styles: set[int] = set()
    def scale_style(style: dict[str, Any]) -> None:
        if id(style) in scaled_styles:
            return
        scaled_styles.add(id(style))
        for key in ("font_size", "stroke_width", "radius"):
            if key in style:
                value = float(str(style[key]).removesuffix("px")) * scale
                style[key] = f"{value:g}px" if key == "font_size" else value
    for element in mapped:
        bbox = element["bbox"]
        element["bbox"] = {"x": float(bbox["x"]) * scale + dx, "y": float(bbox["y"]) * scale + dy,
                           "w": float(bbox["w"]) * scale, "h": float(bbox["h"]) * scale}
        scale_style(element.get("style") or {})
        for line in element.get("_text_lines") or []:
            for run in line:
                scale_style(run.get("style") or {})
                if run.get("position"):
                    run["position"]["x"] = float(run["position"]["x"]) * scale + dx
                    run["position"]["y"] = float(run["position"]["y"]) * scale + dy
        for command in element.get("_native_commands") or []:
            for key in ("x", "x1", "x2", "y", "y1", "y2"):
                if key in command:
                    command[key] = float(command[key]) * scale + (dx if key.startswith("x") else dy)
        effect = (element.get("_paint") or {}).get("effect") or {}
        for key in ("std_deviation", "dx", "dy", "radius"):
            if key in effect:
                effect[key] = float(effect[key]) * scale
    return mapped


def _compile_pptx(root: Path, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]], *, asset_paths_by_page: dict[str, dict[str, Path]] | None = None, validate_approved: Callable | None = None, svg_paths: dict[str, Path] | None = None, output_root: Path | None = None) -> tuple[Path, Path]:
    if not scenes:
        raise PptxEditabilityError("cannot compile an empty high-density deck")
    run_id = str(scenes[0].get("run_id") or "")
    for scene in scenes:
        try:
            assert_v2("page_scene", scene)
            page_id = str(scene.get("page_id") or "")
            lock = locks.get(page_id, {})
            assert_v2("content_lock", lock)
        except ContractError as exc:
            raise PptxEditabilityError(f"PPTX compiler input contract is invalid: {exc}") from exc
        if str(scene.get("run_id") or "") != run_id:
            raise PptxEditabilityError("PPTX compiler scene run_ids are inconsistent")
        if str(lock.get("run_id") or "") != run_id or str(lock.get("page_id") or "") != page_id:
            raise PptxEditabilityError(f"PPTX compiler content lock identity is inconsistent on page {page_id}")
    presentation = Presentation()
    # python-pptx's bundled template contains its author's metadata. It is
    # neither customer content nor evidence about who approved this deck.
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.comments = ""
    presentation.slide_width = Inches(_slide_width())
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = presentation.slide_layouts[6]
    trace_pages: list[dict[str, Any]] = []
    for scene in scenes:
        page_id = str(scene["page_id"])
        svg_file = (svg_paths or {}).get(page_id) or svg_path(root, page_id)
        if not svg_file.exists():
            raise PptxEditabilityError(f"approved SVG is missing on page {page_id}")
        # SC-1.1 ND-01: business-level validation is injected by the adapter
        # (high-density engine / native run adapter); the compiler stays pure.
        if validate_approved is None:
            raise PptxEditabilityError(
                "compile_pptx requires validate_approved (adapter-injected approved-SVG validator); "
                "refusing to compile without the business validation pass"
            )
        try:
            validate_approved(svg_file, scene, locks[page_id], (asset_paths_by_page or {}).get(page_id, {}))
        except ContractError as exc:
            raise PptxEditabilityError(f"approved SVG validation failed on page {page_id}: {exc}") from exc
        slide = presentation.slides.add_slide(blank_layout)
        # python-pptx otherwise rescans every existing XML id for each new
        # shape. High-density SVG traces need the library's sequential-id
        # cache to keep compilation linear in the number of elements.
        slide.shapes._cached_max_shape_id = slide.shapes._spTree.max_shape_id
        trace: list[dict[str, Any]] = []
        elements = _svg_elements(svg_file, scene, (asset_paths_by_page or {}).get(page_id, {}))
        if _NATIVE_CANVAS.get():
            elements = _contain_elements(elements, svg_file)
        index = 0
        while index < len(elements):
            group_id = str(elements[index].get("group_id") or "")
            if not group_id:
                element = dict(elements[index])
                element["_asset_paths"] = (asset_paths_by_page or {}).get(page_id, {})
                _emit_element(slide, element, trace)
                index += 1
                continue
            group_items: list[dict[str, Any]] = []
            while index < len(elements) and str(elements[index].get("group_id") or "") == group_id:
                element = dict(elements[index])
                element["_asset_paths"] = (asset_paths_by_page or {}).get(page_id, {})
                group_items.append(element)
                index += 1
            group = slide.shapes.add_group_shape()
            group.name = group_id
            child_trace: list[dict[str, Any]] = []
            for element in group_items:
                _emit_element(group, element, child_trace)
            if _NATIVE_CANVAS.get():
                group._element.recalculate_extents()
            group_entry = {
                "element_id": group_id,
                "object_type": "group",
                "shape_name": group.name,
                "bbox": _trace_bbox(child_trace),
                "priority": min((str(item.get("priority") or "P2") for item in group_items), key=lambda value: {"P0": 0, "P1": 1, "P2": 2}.get(value, 3)),
                "component_id": str(group_items[0].get("component_id") or f"component.{group_id}"),
                "z_order": min(int(item.get("z_index") or 0) for item in group_items),
                "group_id": group_id,
                "visual_id": str(group_items[0].get("visual_id") or group_id),
                "children": child_trace,
                "child_element_ids": [str(item.get("element_id") or "") for item in child_trace],
                "fidelity": "native_group",
            }
            trace.append(group_entry)
        notes = str((locks.get(page_id) or {}).get("speaker_notes") or "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        trace_pages.append({"page_id": page_id, "svg_sha256": sha256_file(svg_file), "elements": trace, "speaker_notes_present": bool(notes)})
    output = Path(output_root) / "deck.pptx" if output_root is not None else pptx_path(root)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    trace_file = Path(output_root) / "pptx_trace.json" if output_root is not None else trace_path(root)
    svg_hashes = {str(page["page_id"]): str(page["svg_sha256"]) for page in trace_pages}
    trace_payload = {
        "schema_version": "deck_svg_to_drawingml_trace.v1",
        "run_id": run_id,
        "page_id": "deck",
        "svg_sha256": sha256_json(svg_hashes),
        "pptx_sha256": sha256_file(output),
        "pages": trace_pages,
        "elements": [item for page in trace_pages for item in _flatten_trace_entries(page["elements"])],
        "coverage": {"p0_p1": "pass", "visual_registry": "pass", "groups": "pass"},
        "created_at": utc_now(),
    }
    assert_v2("svg_to_drawingml_trace", trace_payload)
    write_json(trace_file, trace_payload)
    return output, trace_file


def _render_pptx_pages(root: Path, pptx: Path, pages: list[tuple[str, int]]) -> dict[str, Path]:
    """Legacy preview projection over the shared bounded native renderer."""
    from .render import RenderError, render_deck
    from PIL import Image

    if not pages:
        return {}
    page_count = len(Presentation(pptx).slides)
    if any(index < 0 or index >= page_count for _, index in pages):
        raise PptxEditabilityError("requested render page is outside the PPTX page set")
    with tempfile.TemporaryDirectory(prefix="deck-master-pptx-render-") as directory:
        temporary_output = Path(directory) / "render"
        try:
            rendered = render_deck(pptx, [f"SLIDE_{index + 1:04d}" for index in range(page_count)], temporary_output)
        except RenderError as exc:
            raise PptxEditabilityError(str(exc)) from exc
        outputs: dict[str, Path] = {}
        for page_id, page_index in pages:
            source = Path(rendered["pages"][page_index]["path"])
            output = pptx_preview_path(root, page_id)
            output.parent.mkdir(parents=True, exist_ok=True)
            with Image.open(source) as image:
                image.convert("RGB").resize((CANVAS_WIDTH, CANVAS_HEIGHT), Image.Resampling.LANCZOS).save(output, format="PNG")
            outputs[page_id] = output
        return outputs


def _render_pptx_page(root: Path, pptx: Path, page_id: str, page_index: int) -> Path:
    return _render_pptx_pages(root, pptx, [(page_id, page_index)])[page_id]


def _drawingml_paint_inventory(presentation: Presentation) -> dict[str, Any]:
    gradients = 0
    gradient_stops = 0
    effects = 0
    effect_types: list[str] = []
    alpha_values: list[int] = []
    elements: list[dict[str, Any]] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            shape_gradients = 0
            shape_effects = 0
            for node in shape._element.iter():
                local = str(node.tag).split("}")[-1]
                if local == "gradFill":
                    shape_gradients += 1
                    gradient_stops += sum(1 for child in node.iter() if str(child.tag).split("}")[-1] == "gs")
                elif local in {"outerShdw", "glow"}:
                    shape_effects += 1
                    effect_types.append("shadow" if local == "outerShdw" else "glow")
                elif local == "alpha" and node.get("val") is not None:
                    try:
                        alpha_values.append(int(node.get("val") or 0))
                    except ValueError:
                        pass
            gradients += shape_gradients
            effects += shape_effects
            if shape_gradients or shape_effects:
                elements.append({"shape_name": str(shape.name or ""), "gradients": shape_gradients, "effects": shape_effects})
    return {"gradient_count": gradients, "gradient_stop_count": gradient_stops, "effect_count": effects, "effect_types": sorted(effect_types), "gradient_stop_alpha_values": alpha_values, "elements": elements}


def _expected_shape_type(object_type: str) -> Any | None:
    return {
        "text": MSO_SHAPE_TYPE.TEXT_BOX,
        "line": MSO_SHAPE_TYPE.LINE,
        "freeform": MSO_SHAPE_TYPE.FREEFORM,
        "registered_asset": MSO_SHAPE_TYPE.PICTURE,
        "shape": MSO_SHAPE_TYPE.AUTO_SHAPE,
        "group": MSO_SHAPE_TYPE.GROUP,
    }.get(object_type)


def _flatten_shape_objects(shapes: Any) -> list[Any]:
    flattened: list[Any] = []
    for shape in shapes:
        flattened.append(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(_flatten_shape_objects(shape.shapes))
    return flattened


def _shape_canvas_bbox(shape: Any, presentation: Presentation) -> dict[str, float]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = _flatten_shape_objects(shape.shapes)
        if not children:
            return {"x": 0.0, "y": 0.0, "w": 0.0, "h": 0.0}
        boxes = [_shape_canvas_bbox(child, presentation) for child in children]
        left = min(box["x"] for box in boxes)
        top = min(box["y"] for box in boxes)
        right = max(box["x"] + box["w"] for box in boxes)
        bottom = max(box["y"] + box["h"] for box in boxes)
        return {"x": left, "y": top, "w": right - left, "h": bottom - top}
    return {
        "x": float(shape.left) / float(presentation.slide_width) * CANVAS_WIDTH,
        "y": float(shape.top) / float(presentation.slide_height) * CANVAS_HEIGHT,
        "w": float(shape.width) / float(presentation.slide_width) * CANVAS_WIDTH,
        "h": float(shape.height) / float(presentation.slide_height) * CANVAS_HEIGHT,
    }


def _image_relationship_count(slide: Any) -> int:
    return sum(1 for relationship in slide.part.rels.values() if str(relationship.reltype).endswith("/image"))


def readback_pptx(root: Path, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]], output: Path) -> Path:
    for scene in scenes:
        try:
            assert_v2("page_scene", scene)
            lock = locks.get(str(scene.get("page_id") or ""), {})
            assert_v2("content_lock", lock)
            validate_visibility_policy(lock.get("visibility_policy") or {}, page_id=str(scene.get("page_id") or ""))
        except ContractError as exc:
            raise PptxEditabilityError(f"PPTX readback input contract is invalid: {exc}") from exc
    presentation = Presentation(output)
    try:
        trace_payload = read_json(trace_path(root))
        assert_v2("svg_to_drawingml_trace", trace_payload)
    except ContractError as exc:
        raise PptxEditabilityError("SVG-to-DrawingML trace is missing before readback") from exc
    run_id = str(scenes[0]["run_id"]) if scenes else ""
    if str(trace_payload.get("run_id") or "") != run_id:
        raise PptxEditabilityError("SVG-to-DrawingML trace run_id is stale")
    if str(trace_payload.get("pptx_sha256") or "") != sha256_file(output):
        raise PptxEditabilityError("SVG-to-DrawingML trace is stale for the PPTX")
    expected_svg_hashes = {str(scene["page_id"]): sha256_file(svg_path(root, str(scene["page_id"]))) for scene in scenes}
    if str(trace_payload.get("svg_sha256") or "") != sha256_json(expected_svg_hashes):
        raise PptxEditabilityError("SVG-to-DrawingML trace is stale for the approved SVG set")
    trace_pages = {str(page.get("page_id") or ""): page for page in trace_payload.get("pages", []) if isinstance(page, dict)}
    if set(trace_pages) != set(expected_svg_hashes) or any(trace_pages[page_id].get("svg_sha256") != svg_hash for page_id, svg_hash in expected_svg_hashes.items()):
        raise PptxEditabilityError("SVG-to-DrawingML trace page coverage is stale")
    expected_text = [
        str(element.get("text") or "")
        for page in trace_payload.get("pages", [])
        for element in page.get("elements", [])
        if element.get("object_type") == "text" and element.get("priority") in {"P0", "P1"}
    ]
    actual_text: list[str] = []
    missing_elements: list[str] = []
    text_mismatches: list[dict[str, str]] = []
    visibility_violations: list[dict[str, str]] = []
    geometry_errors: list[str] = []
    geometry_mismatches: list[dict[str, Any]] = []
    shape_type_mismatches: list[dict[str, str]] = []
    z_order_mismatches: list[dict[str, Any]] = []
    unexpected_elements: list[dict[str, str]] = []
    group_readback_pages: list[dict[str, Any]] = []
    notes_count = 0
    expected_trace_elements = [element for element in trace_payload.get("elements", []) if isinstance(element, dict)]
    expected_image_count = sum(1 for element in expected_trace_elements if element.get("object_type") == "registered_asset")
    actual_image_count = sum(1 for slide in presentation.slides for shape in _flatten_shape_objects(slide.shapes) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    expected_media_relationships = expected_image_count
    actual_media_relationships = sum(_image_relationship_count(slide) for slide in presentation.slides)
    expected_notes_count = sum(1 for scene in scenes if str((locks.get(str(scene["page_id"])) or {}).get("speaker_notes") or "").strip())
    pages: list[dict[str, Any]] = []
    pptx_geometries: dict[str, dict[str, dict[str, float]]] = {}
    for slide_index, slide in enumerate(presentation.slides):
        all_shapes = _flatten_shape_objects(slide.shapes)
        actual_text.extend(shape.text for shape in all_shapes if hasattr(shape, "text") and shape.text)
        if slide_index < len(scenes):
            page_id_for_visibility = str(scenes[slide_index].get("page_id") or "")
            policy = (locks.get(page_id_for_visibility) or {}).get("visibility_policy") or {}
            for shape in all_shapes:
                value = str(getattr(shape, "text", "") or "")
                violation = visible_text_violation(policy, value, page_id=page_id_for_visibility)
                if violation:
                    visibility_violations.append({"page_id": page_id_for_visibility, "element_id": str(shape.name or ""), "category": violation, "text": value})
        if slide.notes_slide.notes_text_frame.text.strip():
            notes_count += 1
        for shape in all_shapes:
            name = str(shape.name or "")
            if name and hasattr(shape, "left") and (shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height):
                geometry_errors.append(name)
        if slide_index >= len(scenes):
            continue
        page_trace_elements = [
            item
            for item in _flatten_trace_entries(trace_pages.get(str(scenes[slide_index]["page_id"]), {}).get("elements", []))
            if isinstance(item, dict)
        ]
        expected_elements = {
            str(element.get("element_id")): element
            for element in page_trace_elements
        }
        actual_elements = {str(shape.name): shape for shape in all_shapes if shape.name}
        expected_order = [str(item.get("element_id") or "") for item in trace_pages.get(str(scenes[slide_index]["page_id"]), {}).get("elements", []) if isinstance(item, dict)]
        # Include unnamed shapes in the order inventory so an unregistered
        # object cannot bypass the trace/readback gate.
        actual_order = [str(shape.name or "") for shape in slide.shapes]
        if expected_order != actual_order:
            z_order_mismatches.append({"page_id": str(scenes[slide_index]["page_id"]), "expected": expected_order, "actual": actual_order})
        for name in (str(shape.name or "") for shape in all_shapes):
            if name not in expected_elements:
                unexpected_elements.append({"page_id": str(scenes[slide_index]["page_id"]), "element_id": name})
        page_missing: list[str] = []
        page_mismatch: list[dict[str, Any]] = []
        page_shape_type_mismatches: list[dict[str, str]] = []
        page_geometry: dict[str, dict[str, float]] = {}
        expected_groups = [item for item in page_trace_elements if item.get("object_type") == "group"]
        actual_groups = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP]
        group_mismatches: list[dict[str, Any]] = []
        for group_entry in expected_groups:
            group_id = str(group_entry.get("element_id") or "")
            group_shape = actual_elements.get(group_id)
            expected_children = [item for item in _flatten_trace_entries(group_entry.get("children") or []) if isinstance(item, dict)]
            actual_children = _flatten_shape_objects(group_shape.shapes) if group_shape is not None and group_shape.shape_type == MSO_SHAPE_TYPE.GROUP else []
            expected_child_ids = [str(item.get("element_id") or "") for item in expected_children]
            actual_child_ids = [str(shape.name or "") for shape in actual_children]
            if group_shape is None or group_shape.shape_type != MSO_SHAPE_TYPE.GROUP or expected_child_ids != actual_child_ids:
                group_mismatches.append({"group_id": group_id, "expected_child_ids": expected_child_ids, "actual_child_ids": actual_child_ids, "actual_group_count": len(actual_groups)})
        group_readback_pages.append({"page_id": str(scenes[slide_index]["page_id"]), "expected_group_count": len(expected_groups), "actual_group_count": len(actual_groups), "groups": [{"group_id": str(item.get("element_id") or ""), "child_element_ids": [str(child.get("element_id") or "") for child in _flatten_trace_entries(item.get("children") or []) if isinstance(child, dict)]} for item in expected_groups], "mismatches": group_mismatches, "status": "pass" if not group_mismatches else "failed"})
        for element_id, element in expected_elements.items():
            shape = actual_elements.get(element_id)
            if shape is None:
                missing_elements.append(element_id)
                page_missing.append(element_id)
                continue
            expected_type = _expected_shape_type(str(element.get("object_type") or ""))
            if expected_type is not None and shape.shape_type != expected_type:
                mismatch = {"element_id": element_id, "expected": str(expected_type), "actual": str(shape.shape_type)}
                shape_type_mismatches.append(mismatch)
                page_shape_type_mismatches.append(mismatch)
            if element.get("object_type") == "text" and element.get("priority") in {"P0", "P1"}:
                expected_text_value = str(element.get("text") or "")
                actual_text_value = str(getattr(shape, "text", "") or "")
                if " ".join(actual_text_value.split()) != " ".join(expected_text_value.split()):
                    text_mismatches.append({"element_id": element_id, "expected": expected_text_value, "actual": actual_text_value})
            expected = element["bbox"]
            actual = _shape_canvas_bbox(shape, presentation)
            page_geometry[element_id] = actual
            deltas = {key: abs(actual[key] - float(expected[key])) for key in ("x", "y", "w", "h")}
            px_to_pt = SLIDE_WIDTH_IN * 72 / CANVAS_WIDTH
            if max(deltas.values()) * px_to_pt > PPTX_BBOX_TOLERANCE_PT:
                entry = {"element_id": element_id, "expected": expected, "actual": actual, "delta": deltas}
                geometry_mismatches.append(entry)
                page_mismatch.append(entry)
        page_id = str(scenes[slide_index]["page_id"])
        pptx_geometries[page_id] = page_geometry
        pages.append({"page_id": page_id, "missing_elements": page_missing, "geometry_mismatches": page_mismatch, "shape_type_mismatches": page_shape_type_mismatches, "shape_count": len(slide.shapes), "shape_inventory": [{"name": str(shape.name or ""), "shape_type": str(shape.shape_type), "z_order": index} for index, shape in enumerate(all_shapes)], "media_relationship_count": _image_relationship_count(slide), "notes_present": bool(slide.notes_slide.notes_text_frame.text.strip())})
    normalized_actual = "\n".join(actual_text)
    missing_text = [item["expected"] for item in text_mismatches]
    traced_ids = {str(item.get("element_id") or "") for item in expected_trace_elements}
    required_trace_ids = {str(element.get("element_id") or "") for scene in scenes for element in scene.get("elements", []) if element.get("priority") in {"P0", "P1"}}
    required_trace_ids.update(str(visual.get("svg_group_id") or visual.get("visual_id") or "") for scene in scenes for visual in scene.get("visual_registry") or [])
    required_trace_ids.update(str(child_id) for scene in scenes for visual in scene.get("visual_registry") or [] for child_id in visual.get("child_element_ids") or [])
    missing_trace_ids = sorted(required_trace_ids - traced_ids)
    trace_status = "pass" if not missing_trace_ids else "failed"
    expected_gradient_count = sum(1 for item in trace_payload.get("elements", []) if isinstance(item, dict) and any(str(paint.get("kind") or "") == "gradient" for paint in (item.get("paint") or {}).values() if isinstance(paint, dict)))
    expected_gradient_stop_count = sum(int(paint.get("stop_count") or 0) for item in trace_payload.get("elements", []) if isinstance(item, dict) for paint in (item.get("paint") or {}).values() if isinstance(paint, dict) and str(paint.get("kind") or "") == "gradient")
    expected_effect_count = sum(1 for item in trace_payload.get("elements", []) if isinstance(item, dict) and item.get("effect"))
    expected_effect_types = sorted(str(item.get("effect", {}).get("type") or "") for item in trace_payload.get("elements", []) if isinstance(item, dict) and item.get("effect"))
    paint_inventory = _drawingml_paint_inventory(presentation)
    paint_status = "pass" if paint_inventory["gradient_count"] == expected_gradient_count and paint_inventory["gradient_stop_count"] == expected_gradient_stop_count and paint_inventory["effect_count"] == expected_effect_count and paint_inventory["effect_types"] == expected_effect_types else "failed"
    visual_parity: dict[str, Any] = {"status": "pass", "pages": []}
    try:
        pptx_previews = _render_pptx_pages(root, output, [(str(scene["page_id"]), index) for index, scene in enumerate(scenes)])
    except (PptxEditabilityError, OSError) as exc:
        visual_parity["status"] = "failed"
        visual_parity["pages"] = [{"page_id": scene["page_id"], "status": "failed", "error": str(exc)} for scene in scenes]
    else:
        for scene in scenes:
            try:
                page_id = str(scene["page_id"])
                pptx_preview = pptx_previews[page_id]
                svg_preview = root / PREVIEW_DIR / f"{page_id}.png"
                metrics = compute_visual_metrics(root, scene, svg_preview, pptx_preview, comparison="svg_vs_pptx", candidate_geometry=pptx_geometries.get(page_id))
                metrics_file = write_visual_metrics(root, scene, metrics, comparison="svg_vs_pptx")
                visual_parity["pages"].append({"page_id": page_id, "metrics_path": str(metrics_file.relative_to(root)), "status": metrics["status"]})
                if metrics["status"] != "pass":
                    visual_parity["status"] = "failed"
            except (PptxEditabilityError, VisualMetricsError, OSError) as exc:
                visual_parity["status"] = "failed"
                visual_parity["pages"].append({"page_id": scene["page_id"], "status": "failed", "error": str(exc)})
    group_readback = {"pages": group_readback_pages, "expected_group_count": sum(int(page.get("expected_group_count") or 0) for page in group_readback_pages), "actual_group_count": sum(int(page.get("actual_group_count") or 0) for page in group_readback_pages), "status": "pass" if all(page.get("status") == "pass" for page in group_readback_pages) else "failed"}
    report = {
        "schema_version": "deck_pptx_readback.v2",
        "run_id": str(scenes[0]["run_id"]) if scenes else "",
        "pptx_sha256": sha256_file(output),
        "slide_count": len(presentation.slides),
        "expected_slide_count": len(scenes),
        "speaker_notes_count": notes_count,
        "expected_speaker_notes_count": expected_notes_count,
        "pages": pages,
        "text_readback": {"expected_p0_p1": len(expected_text), "missing": missing_text, "mismatches": text_mismatches, "visibility_violations": visibility_violations},
        "geometry": {"out_of_bounds": geometry_errors, "mismatches": geometry_mismatches},
        "trace_coverage": {"status": trace_status, "p0_p1": "pass" if not missing_trace_ids else "failed", "missing_element_ids": missing_trace_ids},
        "drawingml_paint": {**paint_inventory, "expected_gradient_count": expected_gradient_count, "expected_gradient_stop_count": expected_gradient_stop_count, "expected_effect_count": expected_effect_count, "expected_effect_types": expected_effect_types, "status": paint_status},
        "shape_readback": {"shape_type_mismatches": shape_type_mismatches, "z_order_mismatches": z_order_mismatches, "unexpected_elements": unexpected_elements, "status": "pass" if not shape_type_mismatches and not z_order_mismatches and not unexpected_elements else "failed"},
        "group_readback": group_readback,
        "media_relationships": {"expected": expected_media_relationships, "actual": actual_media_relationships, "status": "pass" if expected_media_relationships == actual_media_relationships else "failed"},
        "visual_parity": visual_parity,
        "local_visual_parity": visual_parity,
        "editable_object_count": sum(len(slide.shapes) for slide in presentation.slides),
        "image_shape_count": actual_image_count,
        "expected_image_count": expected_image_count,
        "status": "pass" if len(presentation.slides) == len(scenes) and notes_count == expected_notes_count and actual_image_count == expected_image_count and expected_media_relationships == actual_media_relationships and not missing_text and not text_mismatches and not visibility_violations and not missing_elements and not geometry_errors and not geometry_mismatches and not shape_type_mismatches and not z_order_mismatches and not unexpected_elements and trace_status == "pass" and paint_status == "pass" and group_readback["status"] == "pass" and visual_parity["status"] == "pass" else "failed",
        "created_at": utc_now(),
    }
    assert_v2("pptx_readback", report)
    if report["status"] != "pass":
        raise PptxEditabilityError(f"PPTX readback failed: {report}")
    path = readback_path(root)
    write_json(path, report)
    return path


__all__ = ["PptxEditabilityError", "compile_pptx", "pptx_path", "pptx_preview_path", "readback_path", "readback_pptx", "trace_path"]
