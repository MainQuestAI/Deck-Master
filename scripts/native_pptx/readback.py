"""Native structural readback independent of HD planning and preview paths.

Visual review is a separate required gate. This report never claims that a
rendered page was compared with its blueprint or approved SVG.
"""
from __future__ import annotations
from pathlib import Path
from typing import Any
import unicodedata

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .contracts import assert_v2, read_json, sha256_file, sha256_json, utc_now, write_json
from .pptx import _flatten_shape_objects, _flatten_trace_entries, _expected_shape_type, _drawingml_paint_inventory
from .visibility import validate_visibility_policy, visible_text_violation


def readback_native(result: Any, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]]) -> dict[str, Any]:
    from .api import NativeCompileError

    errors: list[str] = []
    trace = read_json(result.trace_path)
    assert_v2('svg_to_drawingml_trace', trace)
    deck = Presentation(result.pptx_path)
    ids = [str(scene.get('page_id') or '') for scene in scenes]
    if len(set(ids)) != len(ids) or [page.get('page_id') for page in trace['pages']] != ids or len(deck.slides) != len(ids):
        errors.append('slide/page/trace order mismatch')
    if trace['pptx_sha256'] != sha256_file(result.pptx_path):
        errors.append('trace PPTX hash is stale')
    if any(not path.is_file() or sha256_file(path) != result.input_sha256.get(page_id) for page_id, path in result.svg_paths.items()):
        errors.append('approved SVG inputs changed since compilation')
    if trace['svg_sha256'] != sha256_json(result.input_sha256):
        errors.append('trace SVG input hash is stale')
    if result.canvas_mode == 'native' and deck.slide_width * 9 != deck.slide_height * 16:
        errors.append('native slide is not exactly 16:9')
    pages: list[dict[str, Any]] = []
    for scene, page, slide in zip(scenes, trace['pages'], deck.slides):
        page_id = str(scene['page_id'])
        lock = locks[page_id]
        assert_v2('page_scene', scene)
        assert_v2('content_lock', lock)
        if scene['run_id'] != trace['run_id'] or lock['run_id'] != trace['run_id'] or lock['page_id'] != page_id:
            errors.append(f'{page_id}: scene/lock/run identity mismatch')
        validate_visibility_policy(lock['visibility_policy'], page_id=page_id)
        entries = _flatten_trace_entries(page['elements'])
        shapes = _flatten_shape_objects(slide.shapes)
        expected = {entry['element_id']: entry for entry in entries}
        actual = {shape.name: shape for shape in shapes}
        if set(actual) != set(expected) or len(actual) != len(shapes):
            errors.append(f'{page_id}: object inventory mismatch')
        if [entry['element_id'] for entry in page['elements']] != [shape.name for shape in slide.shapes]:
            errors.append(f'{page_id}: z-order mismatch')
        required = {str(element.get('element_id')) for element in scene['elements'] if element.get('priority') in {'P0', 'P1'}}
        if required - set(actual):
            errors.append(f'{page_id}: required P0/P1 object missing')
        for name, entry in expected.items():
            shape = actual.get(name)
            if shape is None:
                continue
            expected_type = _expected_shape_type(entry['object_type'])
            if expected_type is not None and shape.shape_type != expected_type:
                errors.append(f'{page_id}/{name}: shape type mismatch')
            if entry['object_type'] == 'group':
                children = _flatten_shape_objects(shape.shapes) if shape.shape_type == MSO_SHAPE_TYPE.GROUP else []
                if [child.name for child in children] != [child['element_id'] for child in _flatten_trace_entries(entry.get('children', []))]:
                    errors.append(f'{page_id}/{name}: group child mismatch')
            text = str(getattr(shape, 'text', '') or '')
            if entry['object_type'] == 'text' and unicodedata.normalize('NFC', text) != unicodedata.normalize('NFC', str(entry.get('text', ''))):
                errors.append(f'{page_id}/{name}: text mismatch')
            if visible_text_violation(lock['visibility_policy'], text, page_id=page_id):
                errors.append(f'{page_id}/{name}: visibility violation')
            plane_height = 1672 * 9 / 16 if result.canvas_mode == 'native' else 941
            bbox = {'x': shape.left / deck.slide_width * 1672, 'y': shape.top / deck.slide_height * plane_height,
                    'w': shape.width / deck.slide_width * 1672, 'h': shape.height / deck.slide_height * plane_height}
            if any(abs(bbox[key] - float(entry['bbox'][key])) * (40/3*72/1672) > .75 for key in bbox):
                errors.append(f'{page_id}/{name}: geometry mismatch')
            # Allow 1 EMU rounding at a page boundary only.
            if min(shape.left, shape.top) < -1 or shape.left + shape.width > deck.slide_width + 1 or shape.top + shape.height > deck.slide_height + 1:
                errors.append(f'{page_id}/{name}: out of bounds')
        if slide.notes_slide.notes_text_frame.text.strip() != str(lock.get('speaker_notes') or '').strip():
            errors.append(f'{page_id}: speaker notes mismatch')
        pages.append({'page_id': page_id, 'text_objects': sum(entry['object_type'] == 'text' for entry in entries),
                      'image_objects': sum(entry['object_type'] == 'image' for entry in entries),
                      'shape_objects': sum(entry['object_type'] not in {'text', 'image', 'group'} for entry in entries),
                      'group_objects': sum(entry['object_type'] == 'group' for entry in entries)})
    inventory = _drawingml_paint_inventory(deck)
    entries = trace['elements']
    expected_gradients = sum(1 for entry in entries if any(paint.get('kind') == 'gradient' for paint in (entry.get('paint') or {}).values() if isinstance(paint, dict)))
    expected_stops = sum(paint.get('stop_count', 0) for entry in entries for paint in (entry.get('paint') or {}).values() if isinstance(paint, dict) and paint.get('kind') == 'gradient')
    expected_effects = sorted(entry['effect']['type'] for entry in entries if entry.get('effect'))
    if inventory['gradient_count'] != expected_gradients or inventory['gradient_stop_count'] != expected_stops or inventory['effect_types'] != expected_effects:
        errors.append('DrawingML paint inventory mismatch')
    if errors:
        raise NativeCompileError('; '.join(errors), code='NDC_READBACK_FAILED', recovery='repair and recompile the affected pages')
    report = {'schema_version': 'deck_native_readback.v1', 'status': 'pass', 'scope': 'structural',
              'visual_review': 'not_run', 'pptx_sha256': sha256_file(result.pptx_path), 'pages': pages,
              'created_at': utc_now(), 'drawingml_paint': inventory}
    path = Path(result.output_root or result.pptx_path.parent) / 'native_readback.json'
    write_json(path, report)
    return {**report, 'report_path': str(path)}
