"""SC-1.1 ND-02: real native runtime capability probe.

Replaces the external PPT Master binding assertion with a probe of what the
default route actually needs: the kernel modules import, python-pptx is
present, the SVG subset pieces load, and (when a renderer is configured)
the renderer/字体 environment responds. No HOME probing, no external
product directory checks, no environment-variable shortcuts — ``installed``
or ``contract_declared`` never equals ``verified``.
"""

from __future__ import annotations

import json
import hashlib
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

PROBE_SCHEMA_VERSION = "deck_native_runtime_probe.v1"
FONTS_ENV = "DECK_MASTER_NATIVE_FONTS_DIR"


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _engine_fingerprint() -> str:

    digest = hashlib.sha256()
    package_dir = Path(__file__).resolve().parent
    for module in sorted(package_dir.glob("*.py")):
        digest.update(module.name.encode("utf-8"))
        digest.update(hashlib.sha256(module.read_bytes()).digest())
    return digest.hexdigest()[:16]


def probe_native_runtime() -> dict[str, Any]:
    """Probe the built-in native runtime for real capability evidence."""

    checks: dict[str, Any] = {}
    failures: list[str] = []
    engine_fp = _engine_fingerprint()

    try:
        import pptx  # noqa: F401

        checks["python_pptx"] = {"status": "verified", "version": str(getattr(pptx, "__version__", ""))}
    except Exception as exc:  # noqa: BLE001
        checks["python_pptx"] = {"status": "blocked", "error": str(exc)}
        failures.append("python-pptx is unavailable: " + str(exc))

    for module_name in ("native_pptx.pptx", "native_pptx.svg_pipeline", "native_pptx.svg_native", "native_pptx.svg_paint", "native_pptx.visibility"):
        try:
            __import__(module_name)
            checks[module_name] = {"status": "verified"}
        except Exception as exc:  # noqa: BLE001
            checks[module_name] = {"status": "blocked", "error": str(exc)}
            failures.append(f"{module_name} failed to import: {exc}")

    from .canvas import CANVAS_HEIGHT, CANVAS_WIDTH

    checks["canvas"] = {"status": "verified", "width": CANVAS_WIDTH, "height": CANVAS_HEIGHT}

    fonts_value = str(__import__("os").environ.get(FONTS_ENV, "")).strip()
    fonts_dir = Path(fonts_value) if fonts_value else None
    font_verified = False
    font_detail: dict[str, Any] = {"status": "unverified", "source": str(FONTS_ENV) if fonts_value else "fontconfig", "requested_family": "Noto Sans SC"}
    candidates: list[Path] = []
    if fonts_dir is not None and fonts_dir.is_dir():
        candidates = sorted(list(fonts_dir.glob("*.ttf")) + list(fonts_dir.glob("*.otf")) + list(fonts_dir.glob("*.ttc")))
    if fonts_dir is None:
        import shutil
        import subprocess
        fc_match = shutil.which("fc-match")
        if fc_match:
            try:
                resolved = subprocess.run([fc_match, "-f", "%{file}", "Noto Sans SC"], capture_output=True, text=True, timeout=5, check=True)
                candidates = [Path(resolved.stdout.strip())]
            except Exception as exc:
                font_detail["error"] = str(exc)
    for font_file in candidates:
        try:
            observed = _verify_chinese_font(font_file)
            font_detail.update(observed)
            font_verified = observed["status"] == "verified"
            if font_verified:
                break
        except Exception as exc:
            font_detail["error"] = str(exc)
    if font_verified:
        checks["fonts"] = font_detail
    else:
        # A non-empty directory is NOT font evidence — a real font file must
        # load. Unverified fonts degrade (readback text fidelity) but do not
        # block compilation.
        checks["fonts"] = font_detail
        failures.append("fonts unverified: no verified Chinese glyph rendering from an actual font file")

    # SC-1.1 review: "verified" needs real evidence, not import success.
    # compilable: run a minimal SVG->PPTX compile through the kernel.
    compilable_check = _probe_compilable()
    checks["compile_smoke"] = compilable_check
    if compilable_check["status"] != "verified":
        failures.append("native compile smoke failed: " + str(compilable_check.get("error", "")))

    render_check = _probe_renderable(engine_fp + ":" + str(font_detail.get("font_sha256") or "unverified"))
    checks["render_smoke"] = render_check
    if render_check["status"] != "verified":
        failures.append("native render unavailable: " + str(render_check.get("error", "")))
    svg_render_check = _probe_svg_renderer(engine_fp)
    checks["rsvg_convert"] = svg_render_check
    if svg_render_check["status"] != "verified":
        failures.append("SVG raster probe failed: " + str(svg_render_check.get("error", "")))
    engine_version = engine_fp
    hard_failures = [f for f in failures if "failed to import" in f or "python-pptx" in f or "compile smoke failed" in f or "SVG raster probe failed" in f]
    status = "ready" if not failures else ("blocked" if hard_failures else "degraded_ready")
    observed_fingerprint = hashlib.sha256(json.dumps({
        "engine": engine_fp, "font": font_detail.get("font_sha256", font_detail.get("status")),
        "render": render_check.get("toolchain_fingerprint", render_check.get("status")),
        "svg_render": svg_render_check.get("toolchain_fingerprint", svg_render_check.get("status")),
    }, sort_keys=True).encode("utf-8")).hexdigest()[:16]
    return {
        "schema_version": PROBE_SCHEMA_VERSION,
        "probe_id": f"native_probe_{observed_fingerprint}",
        "engine_id": "deck_native",
        "engine_version": engine_version,
        "checked_at": _utc_now(),
        "checks": checks,
        "required_missing": failures if status == "blocked" else [],
        "optional_unavailable": [f for f in failures if status != "blocked"],
        "status": status,
        "renderer": {"configured": render_check["status"] == "verified", "status": render_check["status"]},
        "capabilities": {"compile": compilable_check["status"], "render": render_check["status"], "fonts": checks["fonts"]["status"], "svg_render": svg_render_check["status"]},
    }


def native_runtime_ready(probe: dict[str, Any] | None = None) -> bool:
    probe = probe or probe_native_runtime()
    return str(probe.get("status") or "") in {"ready", "degraded_ready"}


def public_probe_summary(probe: dict[str, Any]) -> dict[str, Any]:
    """Redacted public view: no absolute paths, no machine specifics."""

    return {
        "schema_version": PROBE_SCHEMA_VERSION,
        "engine_id": probe.get("engine_id", "deck_native"),
        "probe_id": probe.get("probe_id", ""),
        "status": probe.get("status", ""),
        "checks": {name: {"status": item.get("status", "")} for name, item in (probe.get("checks") or {}).items()},
        "required_missing": [name for name, check in (probe.get("checks") or {}).items() if check.get("status") == "blocked" and name not in {"fonts", "render_smoke"}],
        "optional_unavailable": [name for name in ("fonts", "render_smoke") if (probe.get("checks", {}).get(name) or {}).get("status") != "verified"],
        "renderer": probe.get("renderer", {}),
    }


def _probe_compilable(output_root: Path | None = None) -> dict[str, Any]:
    """Emit a real SVG-derived native slide and reopen its OOXML output.

    This is a synthetic kernel smoke, not customer production/UAT evidence.
    It exercises the same native normalization and shape emitters as the
    compiler without inventing business approvals or an MBB workflow.
    """

    import tempfile

    try:
        with tempfile.TemporaryDirectory(prefix="native_probe_") as tmp:
            root = output_root or Path(tmp)
            root.mkdir(parents=True, exist_ok=True)
            from .svg_pipeline import CANVAS_HEIGHT, CANVAS_WIDTH, validate_svg

            svg_path = root / "probe.svg"
            svg_path.write_text(
                f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {CANVAS_WIDTH} {CANVAS_HEIGHT}" '
                'data-pptx-page-role="content" width="1672" height="941">'
                '<rect id="p.bg" x="0" y="0" width="1672" height="941" fill="#ffffff" stroke="none" '
                'stroke-width="0" data-pptx-bounds="0,0,1672,941"/>'
                '<text id="p.t" x="80" y="120" font-family="Noto Sans SC" font-size="28" font-weight="bold" '
                'fill="#18212b" data-pptx-bounds="80,90,400,50" data-pptx-text="Native 中文 123" data-pptx-text-ref="p.t">Native 中文 123</text>'
                "</svg>",
                encoding="utf-8",
            )
            result = validate_svg(svg_path, page_id="PROBE")
            if not result.get("valid"):
                return {"status": "blocked", "error": "kernel validate_svg rejected the probe SVG"}
            from xml.etree import ElementTree

            from .svg_native import parse_svg_native
            from .svg_paint import parse_svg_paint

            document = ElementTree.fromstring(svg_path.read_text(encoding="utf-8"))
            native = parse_svg_native(document)
            paint = parse_svg_paint(document)
            if not native.get("elements") or not paint:
                return {"status": "blocked", "error": "kernel parse returned no elements/paint"}
            from pptx import Presentation
            from .pptx import _svg_elements, _emit_element
            from .contracts import sha256_file

            scene = {"page_id": "PROBE", "elements": [
                {"element_id": "p.bg", "kind": "rect", "priority": "P2"},
                {"element_id": "p.t", "kind": "text", "priority": "P0", "text": "Native 中文 123", "text_ref": "p.t"},
            ]}
            elements = _svg_elements(svg_path, scene)
            deck = Presentation()
            from pptx.util import Inches
            deck.slide_width = Inches(40 / 3)
            deck.slide_height = Inches(7.5)
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            trace: list[dict[str, Any]] = []
            for element in elements:
                _emit_element(slide, element, trace)
            pptx = root / "probe.pptx"
            deck.save(pptx)
            reread = Presentation(pptx)
            if len(reread.slides) != 1 or len(reread.slides[0].shapes) != len(elements):
                raise RuntimeError("compile smoke PPTX object count mismatch")
            if not any(getattr(shape, "text", "") == "Native 中文 123" for shape in reread.slides[0].shapes):
                raise RuntimeError("compile smoke text is missing after reopening PPTX")
            return {"status": "verified", "evidence": "svg_to_drawingml_save_reopen",
                    "elements": len(elements), "pptx_sha256": sha256_file(pptx), "synthetic_probe": True}
    except Exception as exc:  # noqa: BLE001 - probe reports the failure
        return {"status": "blocked", "error": str(exc)}


def _probe_renderable(engine_fp: str) -> dict[str, Any]:
    import shutil
    paths = tuple(shutil.which(name) for name in ("soffice", "pdftoppm"))
    if not all(paths):
        return {"status": "unverified", "error": "soffice and pdftoppm are required"}
    try:
        identity = tuple((str(path), Path(path).stat().st_mtime_ns, Path(path).stat().st_size) for path in paths)
        return dict(_cached_render_smoke(engine_fp, identity))
    except Exception as exc:
        return {"status": "blocked", "error": str(exc)}


from functools import lru_cache


@lru_cache(maxsize=4)
def _cached_render_smoke(engine_fp: str, executables: tuple) -> dict[str, Any]:
    """Reuse observed evidence only while compiler and executable identities hold."""
    import tempfile
    from .render import render_deck
    with tempfile.TemporaryDirectory(prefix="native_render_probe_") as directory:
        root = Path(directory)
        compiled = _probe_compilable(root)
        if compiled["status"] != "verified":
            # Failed probes raise and are not cached.
            raise RuntimeError(str(compiled.get("error", "compile probe failed")))
        rendered = render_deck(root / "probe.pptx", ["PROBE"], root / "render", timeout_seconds=30)
        return {"status": "verified", "evidence": "soffice_pdf_pdftoppm_decode",
                "toolchain_fingerprint": hashlib.sha256(json.dumps(executables).encode()).hexdigest(),
                "page_count": len(rendered["pages"]), "png_sha256": rendered["pages"][0]["sha256"], "synthetic_probe": True}


def _verify_chinese_font(font_file: Path) -> dict[str, Any]:
    """Rasterize actual CJK glyphs; loading a Latin font alone is insufficient."""
    from PIL import ImageFont
    from .contracts import sha256_file

    sample = "原生编译中文测试 123.45% −"
    loaded = ImageFont.truetype(str(font_file), 24)
    missing = loaded.getmask("\U0010ffff")
    missing_signature = (missing.size, bytes(missing))
    missing_chars = []
    for character in sorted(set(sample) - {" "}):
        glyph = loaded.getmask(character)
        if not glyph.getbbox() or (glyph.size, bytes(glyph)) == missing_signature:
            missing_chars.append(character)
    if missing_chars:
        return {"status": "unverified", "font": font_file.name, "family": loaded.getname()[0],
                "error": "font lacks required glyphs: " + "".join(missing_chars)}
    raster = loaded.getmask(sample)
    return {"status": "verified", "font": font_file.name, "family": loaded.getname()[0],
            "font_sha256": sha256_file(font_file), "fallback_used": loaded.getname()[0] != "Noto Sans SC",
            "sample": sample, "raster_sha256": hashlib.sha256(bytes(raster)).hexdigest(),
            "scope": "chinese_latin_sample_glyphs"}


def _probe_svg_renderer(engine_fp: str) -> dict[str, Any]:
    import shutil
    executable = shutil.which("rsvg-convert")
    if not executable:
        return {"status": "blocked", "error": "rsvg-convert is required for SVG validation and parity"}
    try:
        path = Path(executable)
        identity = (str(path), path.stat().st_mtime_ns, path.stat().st_size)
        return dict(_cached_svg_render_smoke(engine_fp, identity))
    except Exception as exc:
        return {"status": "blocked", "error": str(exc)}


@lru_cache(maxsize=4)
def _cached_svg_render_smoke(engine_fp: str, executable: tuple) -> dict[str, Any]:
    import subprocess
    import tempfile
    from PIL import Image
    from .contracts import sha256_file
    with tempfile.TemporaryDirectory(prefix="native_svg_render_probe_") as directory:
        root = Path(directory)
        svg = root / "probe.svg"
        png = root / "probe.png"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 16 9"><rect width="16" height="9" fill="#1278ab"/></svg>', encoding="utf-8")
        subprocess.run([executable[0], "-w", "160", "-h", "90", "-o", str(png), str(svg)], capture_output=True, text=True, timeout=30, check=True)
        with Image.open(png) as image:
            image.load()
            if image.size != (160, 90) or image.convert("RGB").getpixel((80, 45)) != (18, 120, 171):
                raise RuntimeError("SVG raster probe produced incorrect dimensions or pixels")
        return {"status": "verified", "evidence": "rsvg_png_pixel_decode", "toolchain_fingerprint": hashlib.sha256(json.dumps(executable).encode()).hexdigest(), "png_sha256": sha256_file(png), "synthetic_probe": True}
