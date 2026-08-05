from __future__ import annotations

import base64
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .contracts import ContractError, assert_v2, read_json, sha256_file, sha256_json, utc_now, write_json
from .svg import svg_path, validate_svg, wrap_text, _estimated_width
from .visual import VisualMetricsError, compute_visual_metrics, write_visual_metrics

PPTX_DIR = Path("high_density_build/pptx")
TRACE_DIR = Path("high_density_build/traces")
READBACK_DIR = Path("high_density_build/readback")
PREVIEW_DIR = Path("high_density_build/previews")
CANVAS_WIDTH = 1672
CANVAS_HEIGHT = 941
SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5
PPTX_BBOX_TOLERANCE_PT = 0.75
SVG_NS = "http://www.w3.org/2000/svg"


class PptxEditabilityError(ContractError):
    pass


def pptx_path(root: Path) -> Path:
    return root / PPTX_DIR / "deck_high_density.pptx"


def trace_path(root: Path) -> Path:
    return root / TRACE_DIR / "pptx_trace.json"


def readback_path(root: Path) -> Path:
    return root / READBACK_DIR / "readback_report.json"


def pptx_preview_path(root: Path, page_id: str) -> Path:
    return root / PREVIEW_DIR / f"{page_id}.pptx.png"


def _rgb(value: Any, default: str = "18212b") -> RGBColor:
    text = str(value or "").lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", text):
        text = default
    return RGBColor.from_string(text.upper())


def _inches(value: float, total: float) -> float:
    return float(value) / total * (SLIDE_WIDTH_IN if total == CANVAS_WIDTH else SLIDE_HEIGHT_IN)


def _set_shape_fill(shape: Any, style: dict[str, Any]) -> None:
    fill = str(style.get("fill") or "").lower()
    if fill in {"", "none", "transparent"} or fill.startswith("url("):
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        if float(style.get("opacity") or 1) < 1:
            shape.fill.transparency = max(0, min(100, int((1 - float(style.get("opacity") or 1)) * 100)))
    stroke = str(style.get("stroke") or "").lower()
    if stroke in {"", "none", "transparent"}:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(stroke, "d5dde5")
        shape.line.width = Pt(float(style.get("stroke_width") or 1))


def _remove_theme_effects(shape: Any) -> None:
    """Remove theme effect references that make LibreOffice add shadows."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _fit_text_size(element: dict[str, Any]) -> float:
    bbox = element["bbox"]
    fit = element.get("text_fit") or {}
    size = float(fit.get("preferred_size_px") or 18)
    min_size = float(fit.get("min_size_px") or max(9, size * 0.7))
    max_lines = int(fit.get("max_lines") or 1)
    line_height = float(fit.get("line_height") or 1.18)
    text = str(element.get("text") or "")
    lines = wrap_text(text, float(bbox["w"]), size)
    while len(lines) > max_lines and size > min_size:
        size = max(min_size, size - 1)
        lines = wrap_text(text, float(bbox["w"]), size)
    if len(lines) > max_lines or len(lines) * size * line_height > float(bbox["h"]) + 0.01 or any(_estimated_width(line, size) > float(bbox["w"]) + 0.01 for line in lines):
        raise PptxEditabilityError(f"text does not fit in {element.get('element_id')}")
    return size


def _add_text(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    shape = slide.shapes.add_textbox(Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = str(element.get("text") or "")
    style = element.get("style") or {}
    font = run.font
    font.name = str(style.get("font_family") or "Arial")
    # SVG coordinates use a 1672px canvas; convert its pixel typography into
    # the physical slide width before LibreOffice renders it back to pixels.
    canvas_to_slide = (SLIDE_WIDTH_IN * 96) / CANVAS_WIDTH
    font.size = Pt(_fit_text_size(element) * 72 / 96 * canvas_to_slide)
    font.bold = str(style.get("font_weight") or "400") in {"600", "700", "bold", "Bold"}
    font.color.rgb = _rgb(style.get("fill"), "18212b")
    paragraph.space_after = Pt(0)
    trace.append({"element_id": element["element_id"], "object_type": "text", "shape_name": shape.name, "bbox": bbox, "text": element.get("text", ""), "text_ref": element.get("text_ref", "")})


def _add_rect(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    style = element.get("style") or {}
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if float(style.get("radius") or 0) > 0 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    _set_shape_fill(shape, style)
    trace.append({"element_id": element["element_id"], "object_type": "shape", "shape_name": shape.name, "bbox": bbox})


def _add_line(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["x"]) + float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]) + float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    style = element.get("style") or {}
    shape.line.color.rgb = _rgb(style.get("stroke"), "657485")
    shape.line.width = Pt(float(style.get("stroke_width") or 2))
    trace.append({"element_id": element["element_id"], "object_type": "line", "shape_name": shape.name, "bbox": bbox})


_PATH_TOKEN_RE = re.compile(r"([AaCcHhLlMmQqSsTtVvZz])|([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?)")


def _path_points(path_data: str, *, allow_curves: bool = True) -> tuple[list[tuple[float, float]], bool]:
    matches = list(_PATH_TOKEN_RE.finditer(path_data))
    cursor = 0
    for match in matches:
        if path_data[cursor : match.start()].strip(" ,\t\r\n"):
            raise PptxEditabilityError("path contains unsupported separators or syntax")
        cursor = match.end()
    if path_data[cursor:].strip(" ,\t\r\n"):
        raise PptxEditabilityError("path contains unsupported separators or syntax")
    tokens = [(match.group(1) or match.group(2)) for match in matches]
    if not tokens:
        raise PptxEditabilityError("path is empty")
    points: list[tuple[float, float]] = []
    current = (0.0, 0.0)
    start = current
    command = ""
    closed = False
    index = 0

    def number() -> float:
        nonlocal index
        if index >= len(tokens) or re.fullmatch(r"[A-Za-z]", tokens[index]):
            raise PptxEditabilityError("path command has incomplete coordinates")
        value = float(tokens[index])
        index += 1
        return value

    def point(x: float, y: float, relative: bool) -> tuple[float, float]:
        return (x + current[0], y + current[1]) if relative else (x, y)

    while index < len(tokens):
        token = tokens[index]
        if re.fullmatch(r"[A-Za-z]", token):
            command = token
            index += 1
            if command in "Zz":
                if len(points) < 2:
                    raise PptxEditabilityError("closed path needs at least two points")
                closed = True
                current = start
                command = ""
                continue
            if command.upper() not in {"M", "L", "H", "V", "C", "Q"}:
                raise PptxEditabilityError(f"unsupported path command: {command}")
            if command.upper() in {"C", "Q"} and not allow_curves:
                raise PptxEditabilityError(f"curved paths require native reconstruction: {command}")
        if not command:
            raise PptxEditabilityError("path coordinates are missing a command")
        relative = command.islower()
        upper = command.upper()
        if upper in {"M", "L"}:
            target = point(number(), number(), relative)
            current = target
            if upper == "M" and not points:
                start = target
            points.append(target)
            if upper == "M":
                command = "l" if relative else "L"
        elif upper == "H":
            x = number() + current[0] if relative else number()
            current = (x, current[1])
            points.append(current)
        elif upper == "V":
            y = number() + current[1] if relative else number()
            current = (current[0], y)
            points.append(current)
        elif upper == "C":
            c1 = point(number(), number(), relative)
            c2 = point(number(), number(), relative)
            target = point(number(), number(), relative)
            p0 = current
            for step in range(1, 9):
                t = step / 8
                inv = 1 - t
                points.append((inv**3 * p0[0] + 3 * inv**2 * t * c1[0] + 3 * inv * t**2 * c2[0] + t**3 * target[0], inv**3 * p0[1] + 3 * inv**2 * t * c1[1] + 3 * inv * t**2 * c2[1] + t**3 * target[1]))
            current = target
        elif upper == "Q":
            control = point(number(), number(), relative)
            target = point(number(), number(), relative)
            p0 = current
            for step in range(1, 7):
                t = step / 6
                inv = 1 - t
                points.append((inv**2 * p0[0] + 2 * inv * t * control[0] + t**2 * target[0], inv**2 * p0[1] + 2 * inv * t * control[1] + t**2 * target[1]))
            current = target
    if len(points) < 2:
        raise PptxEditabilityError("path needs at least two points")
    return points, closed


def _add_freeform(slide: Any, element: dict[str, Any], points: list[tuple[float, float]], *, closed: bool, trace: list[dict[str, Any]]) -> None:
    x_scale = float(Inches(SLIDE_WIDTH_IN)) / CANVAS_WIDTH
    y_scale = float(Inches(SLIDE_HEIGHT_IN)) / CANVAS_HEIGHT
    emu_points = [(x * x_scale, y * y_scale) for x, y in points]
    builder = slide.shapes.build_freeform(start_x=emu_points[0][0], start_y=emu_points[0][1], scale=1.0)
    builder.add_line_segments(emu_points[1:], close=closed)
    shape = builder.convert_to_shape()
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    _set_shape_fill(shape, element.get("style") or {})
    trace.append({"element_id": element["element_id"], "object_type": "freeform", "shape_name": shape.name, "bbox": element["bbox"], "closed": closed})


def _add_path(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]], *, allow_curves: bool) -> None:
    points, closed = _path_points(str(element.get("path") or ""), allow_curves=allow_curves)
    _add_freeform(slide, element, points, closed=closed, trace=trace)


def _add_polygon(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    points = [(float(point[0]), float(point[1])) for point in element.get("points") or []]
    if len(points) < 3:
        raise PptxEditabilityError(f"polygon needs at least three points in {element.get('element_id')}")
    _add_freeform(slide, element, points, closed=True, trace=trace)


def _add_ellipse(slide: Any, element: dict[str, Any], trace: list[dict[str, Any]]) -> None:
    bbox = element["bbox"]
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    _remove_theme_effects(shape)
    _set_shape_fill(shape, element.get("style") or {})
    trace.append({"element_id": element["element_id"], "object_type": "shape", "shape_name": shape.name, "bbox": bbox})


def _add_image(slide: Any, element: dict[str, Any], asset_paths: dict[str, Path], trace: list[dict[str, Any]]) -> None:
    asset_id = str(element.get("asset_ref") or "")
    asset_path = asset_paths.get(asset_id)
    if asset_path is None or not asset_path.is_file():
        raise PptxEditabilityError(f"registered image asset is unavailable: {asset_id}")
    bbox = element["bbox"]
    shape = slide.shapes.add_picture(str(asset_path), Inches(_inches(float(bbox["x"]), CANVAS_WIDTH)), Inches(_inches(float(bbox["y"]), CANVAS_HEIGHT)), width=Inches(_inches(float(bbox["w"]), CANVAS_WIDTH)), height=Inches(_inches(float(bbox["h"]), CANVAS_HEIGHT)))
    shape.name = str(element["element_id"])
    trace.append({"element_id": element["element_id"], "object_type": "registered_asset", "shape_name": shape.name, "asset_id": asset_id, "bbox": bbox})


def _bounds(node: Any) -> dict[str, float]:
    raw = str(node.get("data-pptx-bounds") or "")
    if raw:
        try:
            x, y, w, h = [float(value) for value in raw.split(",")]
            declared = {"x": x, "y": y, "w": w, "h": h}
        except ValueError as exc:
            raise PptxEditabilityError(f"invalid SVG bbox for {node.get('id')}") from exc
    else:
        declared = None
    tag = str(node.tag).split("}")[-1]
    try:
        if tag == "rect":
            return {"x": float(node.get("x") or 0), "y": float(node.get("y") or 0), "w": float(node.get("width") or 0), "h": float(node.get("height") or 0)}
        if tag == "line":
            x1, y1, x2, y2 = (float(node.get(name)) for name in ("x1", "y1", "x2", "y2"))
            return {"x": min(x1, x2), "y": min(y1, y2), "w": abs(x2 - x1), "h": abs(y2 - y1)}
        if tag == "circle":
            cx, cy, radius = (float(node.get(name)) for name in ("cx", "cy", "r"))
            return {"x": cx - radius, "y": cy - radius, "w": radius * 2, "h": radius * 2}
        if tag == "ellipse":
            cx, cy, rx, ry = (float(node.get(name)) for name in ("cx", "cy", "rx", "ry"))
            return {"x": cx - rx, "y": cy - ry, "w": rx * 2, "h": ry * 2}
        if tag in {"polygon", "polyline"}:
            values = [float(value) for value in re.findall(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", str(node.get("points") or ""))]
            if len(values) < 4 or len(values) % 2:
                raise ValueError("invalid point list")
            xs, ys = values[0::2], values[1::2]
            return {"x": min(xs), "y": min(ys), "w": max(xs) - min(xs), "h": max(ys) - min(ys)}
        if tag == "path":
            points, _ = _path_points(str(node.get("d") or ""), allow_curves=True)
            xs, ys = zip(*points)
            return {"x": min(xs), "y": min(ys), "w": max(xs) - min(xs), "h": max(ys) - min(ys)}
        if tag == "image":
            return {"x": float(node.get("x") or 0), "y": float(node.get("y") or 0), "w": float(node.get("width") or 0), "h": float(node.get("height") or 0)}
    except (TypeError, ValueError) as exc:
        raise PptxEditabilityError(f"invalid SVG geometry for {node.get('id')}") from exc
    if declared is not None:
        return declared
    raise PptxEditabilityError(f"SVG element has no geometry bounds: {node.get('id')}")


def _node_element(node: Any, scene_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    element_id = str(node.get("id") or "")
    tag = str(node.tag).split("}")[-1]
    kind_by_tag = {"text": "text", "rect": "rect", "circle": "ellipse", "ellipse": "ellipse", "line": "line", "path": "path", "polygon": "polygon", "polyline": "path", "image": "image"}
    kind = kind_by_tag.get(tag)
    if kind is None:
        raise PptxEditabilityError(f"unsupported SVG element {tag} id={element_id}")
    if not element_id or element_id not in scene_by_id:
        raise PptxEditabilityError(f"SVG element is not registered in page_scene.v2: {element_id or '<anonymous>'}")
    source = scene_by_id[element_id]
    element = dict(source)
    element["kind"] = kind
    element["bbox"] = _bounds(node)
    style = dict(source.get("style") or {})
    for key in ("fill", "stroke", "stroke-width", "opacity", "font-family", "font-size", "font-weight", "rx"):
        if node.get(key) is not None:
            style[key.replace("-", "_")] = node.get(key)
    if "stroke_width" not in style and node.get("stroke-width") is not None:
        style["stroke_width"] = node.get("stroke-width")
    if tag == "text":
        element["text"] = str(node.get("data-pptx-text") or "".join(node.itertext()))
        if node.get("x") is not None:
            style["x"] = node.get("x")
    elif tag in {"path", "polyline"}:
        element["path"] = str(node.get("d") or "") if tag == "path" else "M " + " L ".join(str(point) for point in str(node.get("points") or "").split())
    elif tag == "polygon":
        points: list[list[float]] = []
        for pair in str(node.get("points") or "").split():
            x, y = pair.split(",", 1)
            points.append([float(x), float(y)])
        element["points"] = points
    elif tag == "image":
        element["asset_ref"] = str(node.get("data-pptx-asset-id") or source.get("asset_ref") or "")
    elif tag == "circle":
        if any(node.get(name) is None for name in ("cx", "cy", "r")):
            raise PptxEditabilityError(f"circle geometry is incomplete: {element_id}")
    elif tag == "ellipse":
        if any(node.get(name) is None for name in ("cx", "cy", "rx", "ry")):
            raise PptxEditabilityError(f"ellipse geometry is incomplete: {element_id}")
    element["style"] = style
    return element


def _svg_elements(root: Path, scene: dict[str, Any], asset_paths: dict[str, Path] | None = None) -> list[dict[str, Any]]:
    try:
        svg_root = ElementTree.fromstring(root.read_text(encoding="utf-8"))
    except (OSError, ElementTree.ParseError) as exc:
        raise PptxEditabilityError(f"approved SVG cannot be parsed: {exc}") from exc
    scene_by_id = {str(item.get("element_id")): item for item in scene.get("elements", [])}
    elements: list[dict[str, Any]] = []
    for node in svg_root.iter():
        tag = str(node.tag).split("}")[-1]
        if tag in {"svg", "g", "defs", "linearGradient", "radialGradient", "stop", "symbol", "tspan", "title", "desc", "metadata"}:
            continue
        element = _node_element(node, scene_by_id)
        if element.get("kind") == "text":
            svg_text = str(element.get("text") or "")
            expected_text = str(scene_by_id[element["element_id"]].get("text") or "")
            if svg_text != expected_text:
                raise PptxEditabilityError(f"SVG text drift for {element['element_id']}: content-lock text does not match")
        if element.get("kind") == "image":
            asset_ref = str(element.get("asset_ref") or "")
            asset_path = (asset_paths or {}).get(asset_ref)
            expected_sha = str(scene_by_id[element["element_id"]].get("asset_sha256") or "")
            if asset_path is None or not asset_path.is_file():
                raise PptxEditabilityError(f"registered image asset is unavailable: {asset_ref}")
            if expected_sha and sha256_file(asset_path) != expected_sha:
                raise PptxEditabilityError(f"registered image asset hash mismatch: {asset_ref}")
        elements.append(element)
    if not elements:
        raise PptxEditabilityError("approved SVG contains no visible elements")
    seen = {item["element_id"] for item in elements}
    missing = [str(item.get("element_id")) for item in scene.get("elements", []) if item.get("kind") != "image" and item.get("priority") in {"P0", "P1"} and str(item.get("element_id")) not in seen]
    if missing:
        raise PptxEditabilityError(f"approved SVG is missing P0/P1 elements: {', '.join(missing)}")
    return elements


def compile_pptx(root: Path, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]], *, asset_paths_by_page: dict[str, dict[str, Path]] | None = None) -> tuple[Path, Path]:
    if not scenes:
        raise PptxEditabilityError("cannot compile an empty high-density deck")
    run_id = str(scenes[0].get("run_id") or "")
    for scene in scenes:
        try:
            assert_v2("page_scene", scene)
            page_id = str(scene.get("page_id") or "")
            lock = locks.get(page_id, {})
            assert_v2("content_lock", lock)
        except ContractError as exc:
            raise PptxEditabilityError(f"PPTX compiler input contract is invalid: {exc}") from exc
        if str(scene.get("run_id") or "") != run_id:
            raise PptxEditabilityError("PPTX compiler scene run_ids are inconsistent")
        if str(lock.get("run_id") or "") != run_id or str(lock.get("page_id") or "") != page_id:
            raise PptxEditabilityError(f"PPTX compiler content lock identity is inconsistent on page {page_id}")
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = presentation.slide_layouts[6]
    trace_pages: list[dict[str, Any]] = []
    for scene in scenes:
        page_id = str(scene["page_id"])
        svg_file = svg_path(root, page_id)
        if not svg_file.exists():
            raise PptxEditabilityError(f"approved SVG is missing on page {page_id}")
        try:
            validate_svg(svg_file, page_id=page_id)
        except ContractError as exc:
            raise PptxEditabilityError(f"approved SVG validation failed on page {page_id}: {exc}") from exc
        slide = presentation.slides.add_slide(blank_layout)
        trace: list[dict[str, Any]] = []
        elements = _svg_elements(svg_file, scene, (asset_paths_by_page or {}).get(page_id, {}))
        allow_curves = str(scene.get("source") or "") != "fixture_auto"
        for element in elements:
            kind = element.get("kind")
            if kind == "text":
                _add_text(slide, element, trace)
            elif kind == "rect":
                _add_rect(slide, element, trace)
            elif kind == "line":
                _add_line(slide, element, trace)
            elif kind == "path":
                _add_path(slide, element, trace, allow_curves=allow_curves)
            elif kind == "polygon":
                _add_polygon(slide, element, trace)
            elif kind == "ellipse":
                _add_ellipse(slide, element, trace)
            elif kind == "image":
                _add_image(slide, element, (asset_paths_by_page or {}).get(page_id, {}), trace)
            else:  # pragma: no cover
                raise PptxEditabilityError(f"unsupported scene element kind: {kind}")
        notes = str((locks.get(page_id) or {}).get("speaker_notes") or "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        trace_pages.append({"page_id": page_id, "svg_sha256": sha256_file(svg_file), "elements": trace, "speaker_notes_present": bool(notes)})
    output = pptx_path(root)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    trace_file = trace_path(root)
    svg_hashes = {str(page["page_id"]): str(page["svg_sha256"]) for page in trace_pages}
    trace_payload = {
        "schema_version": "deck_svg_to_drawingml_trace.v1",
        "run_id": run_id,
        "page_id": "deck",
        "svg_sha256": sha256_json(svg_hashes),
        "pptx_sha256": sha256_file(output),
        "pages": trace_pages,
        "elements": [item for page in trace_pages for item in page["elements"]],
        "coverage": {"p0_p1": "pass"},
        "created_at": utc_now(),
    }
    assert_v2("svg_to_drawingml_trace", trace_payload)
    write_json(trace_file, trace_payload)
    return output, trace_file


def _render_pptx_page(root: Path, pptx: Path, page_id: str, page_index: int) -> Path:
    soffice = shutil.which("soffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise PptxEditabilityError("soffice and pdftoppm are required for PPTX render parity")
    with tempfile.TemporaryDirectory(prefix="deck-master-pptx-render-") as directory:
        temp = Path(directory)
        result = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(pptx)], capture_output=True, text=True)
        pdf = temp / f"{pptx.stem}.pdf"
        if result.returncode != 0 or not pdf.exists():
            raise PptxEditabilityError(result.stderr.strip() or "LibreOffice failed to render PPTX")
        prefix = temp / "page"
        result = subprocess.run([pdftoppm, "-png", "-f", str(page_index + 1), "-singlefile", "-r", "144", str(pdf), str(prefix)], capture_output=True, text=True)
        if result.returncode != 0 or not (temp / "page.png").exists():
            raise PptxEditabilityError(result.stderr.strip() or "pdftoppm failed to render PPTX")
        output = pptx_preview_path(root, page_id)
        output.parent.mkdir(parents=True, exist_ok=True)
        from PIL import Image

        with Image.open(temp / "page.png") as image:
            image.convert("RGB").resize((CANVAS_WIDTH, CANVAS_HEIGHT), Image.Resampling.LANCZOS).save(output, format="PNG")
        return output


def readback_pptx(root: Path, scenes: list[dict[str, Any]], locks: dict[str, dict[str, Any]], output: Path) -> Path:
    for scene in scenes:
        try:
            assert_v2("page_scene", scene)
            assert_v2("content_lock", locks.get(str(scene.get("page_id") or ""), {}))
        except ContractError as exc:
            raise PptxEditabilityError(f"PPTX readback input contract is invalid: {exc}") from exc
    presentation = Presentation(output)
    try:
        trace_payload = read_json(trace_path(root))
        assert_v2("svg_to_drawingml_trace", trace_payload)
    except ContractError as exc:
        raise PptxEditabilityError("SVG-to-DrawingML trace is missing before readback") from exc
    run_id = str(scenes[0]["run_id"]) if scenes else ""
    if str(trace_payload.get("run_id") or "") != run_id:
        raise PptxEditabilityError("SVG-to-DrawingML trace run_id is stale")
    if str(trace_payload.get("pptx_sha256") or "") != sha256_file(output):
        raise PptxEditabilityError("SVG-to-DrawingML trace is stale for the PPTX")
    expected_svg_hashes = {str(scene["page_id"]): sha256_file(svg_path(root, str(scene["page_id"]))) for scene in scenes}
    if str(trace_payload.get("svg_sha256") or "") != sha256_json(expected_svg_hashes):
        raise PptxEditabilityError("SVG-to-DrawingML trace is stale for the approved SVG set")
    trace_pages = {str(page.get("page_id") or ""): page for page in trace_payload.get("pages", []) if isinstance(page, dict)}
    if set(trace_pages) != set(expected_svg_hashes) or any(trace_pages[page_id].get("svg_sha256") != svg_hash for page_id, svg_hash in expected_svg_hashes.items()):
        raise PptxEditabilityError("SVG-to-DrawingML trace page coverage is stale")
    expected_text: list[str] = []
    for scene in scenes:
        expected_text.extend(str(element.get("text") or "") for element in scene.get("elements", []) if element.get("kind") == "text" and element.get("priority") in {"P0", "P1"})
    actual_text: list[str] = []
    missing_elements: list[str] = []
    text_mismatches: list[dict[str, str]] = []
    geometry_errors: list[str] = []
    geometry_mismatches: list[dict[str, Any]] = []
    notes_count = 0
    expected_image_count = sum(1 for scene in scenes for element in scene.get("elements", []) if element.get("kind") == "image")
    actual_image_count = sum(1 for slide in presentation.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    expected_notes_count = sum(1 for scene in scenes if str((locks.get(str(scene["page_id"])) or {}).get("speaker_notes") or "").strip())
    pages: list[dict[str, Any]] = []
    pptx_geometries: dict[str, dict[str, dict[str, float]]] = {}
    for slide_index, slide in enumerate(presentation.slides):
        actual_text.extend(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text)
        if slide.notes_slide.notes_text_frame.text.strip():
            notes_count += 1
        for shape in slide.shapes:
            name = str(shape.name or "")
            if name and hasattr(shape, "left") and (shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height):
                geometry_errors.append(name)
        if slide_index >= len(scenes):
            continue
        expected_elements = {str(element.get("element_id")): element for element in scenes[slide_index].get("elements", [])}
        actual_elements = {str(shape.name): shape for shape in slide.shapes if shape.name}
        page_missing: list[str] = []
        page_mismatch: list[dict[str, Any]] = []
        page_geometry: dict[str, dict[str, float]] = {}
        for element_id, element in expected_elements.items():
            shape = actual_elements.get(element_id)
            if shape is None:
                missing_elements.append(element_id)
                page_missing.append(element_id)
                continue
            if element.get("kind") == "text" and element.get("priority") in {"P0", "P1"}:
                expected_text_value = str(element.get("text") or "")
                actual_text_value = str(getattr(shape, "text", "") or "")
                if actual_text_value != expected_text_value:
                    text_mismatches.append({"element_id": element_id, "expected": expected_text_value, "actual": actual_text_value})
            expected = element["bbox"]
            actual = {"x": float(shape.left) / float(presentation.slide_width) * CANVAS_WIDTH, "y": float(shape.top) / float(presentation.slide_height) * CANVAS_HEIGHT, "w": float(shape.width) / float(presentation.slide_width) * CANVAS_WIDTH, "h": float(shape.height) / float(presentation.slide_height) * CANVAS_HEIGHT}
            page_geometry[element_id] = actual
            deltas = {key: abs(actual[key] - float(expected[key])) for key in ("x", "y", "w", "h")}
            px_to_pt = SLIDE_WIDTH_IN * 72 / CANVAS_WIDTH
            if max(deltas.values()) * px_to_pt > PPTX_BBOX_TOLERANCE_PT:
                entry = {"element_id": element_id, "expected": expected, "actual": actual, "delta": deltas}
                geometry_mismatches.append(entry)
                page_mismatch.append(entry)
        page_id = str(scenes[slide_index]["page_id"])
        pptx_geometries[page_id] = page_geometry
        pages.append({"page_id": page_id, "missing_elements": page_missing, "geometry_mismatches": page_mismatch, "shape_count": len(slide.shapes), "shape_inventory": [{"name": str(shape.name or ""), "shape_type": str(shape.shape_type), "z_order": index} for index, shape in enumerate(slide.shapes)], "notes_present": bool(slide.notes_slide.notes_text_frame.text.strip())})
    normalized_actual = "\n".join(actual_text)
    missing_text = [item["expected"] for item in text_mismatches]
    traced_ids = {str(item.get("element_id") or "") for item in trace_payload.get("elements", []) if isinstance(item, dict)}
    required_trace_ids = {str(element.get("element_id") or "") for scene in scenes for element in scene.get("elements", []) if element.get("priority") in {"P0", "P1"}}
    missing_trace_ids = sorted(required_trace_ids - traced_ids)
    trace_status = "pass" if not missing_trace_ids else "failed"
    visual_parity: dict[str, Any] = {"status": "pass", "pages": []}
    for index, scene in enumerate(scenes):
        try:
            pptx_preview = _render_pptx_page(root, output, str(scene["page_id"]), index)
            svg_preview = root / PREVIEW_DIR / f"{scene['page_id']}.png"
            metrics = compute_visual_metrics(root, scene, svg_preview, pptx_preview, comparison="svg_vs_pptx", candidate_geometry=pptx_geometries.get(str(scene["page_id"])))
            metrics_file = write_visual_metrics(root, scene, metrics, comparison="svg_vs_pptx")
            visual_parity["pages"].append({"page_id": scene["page_id"], "metrics_path": str(metrics_file.relative_to(root)), "status": metrics["status"]})
            if metrics["status"] != "pass":
                visual_parity["status"] = "failed"
        except (PptxEditabilityError, VisualMetricsError, OSError) as exc:
            visual_parity["status"] = "failed"
            visual_parity["pages"].append({"page_id": scene["page_id"], "status": "failed", "error": str(exc)})
    report = {
        "schema_version": "deck_pptx_readback.v2",
        "run_id": str(scenes[0]["run_id"]) if scenes else "",
        "pptx_sha256": sha256_file(output),
        "slide_count": len(presentation.slides),
        "expected_slide_count": len(scenes),
        "speaker_notes_count": notes_count,
        "expected_speaker_notes_count": expected_notes_count,
        "pages": pages,
        "text_readback": {"expected_p0_p1": len(expected_text), "missing": missing_text, "mismatches": text_mismatches},
        "geometry": {"out_of_bounds": geometry_errors, "mismatches": geometry_mismatches},
        "trace_coverage": {"status": trace_status, "p0_p1": "pass" if not missing_trace_ids else "failed", "missing_element_ids": missing_trace_ids},
        "visual_parity": visual_parity,
        "editable_object_count": sum(len(slide.shapes) for slide in presentation.slides),
        "image_shape_count": actual_image_count,
        "expected_image_count": expected_image_count,
        "status": "pass" if len(presentation.slides) == len(scenes) and notes_count == expected_notes_count and actual_image_count == expected_image_count and not missing_text and not text_mismatches and not missing_elements and not geometry_errors and not geometry_mismatches and trace_status == "pass" and visual_parity["status"] == "pass" else "failed",
        "created_at": utc_now(),
    }
    assert_v2("pptx_readback", report)
    if report["status"] != "pass":
        raise PptxEditabilityError(f"PPTX readback failed: {report}")
    path = readback_path(root)
    write_json(path, report)
    return path


__all__ = ["PptxEditabilityError", "compile_pptx", "pptx_path", "pptx_preview_path", "readback_path", "readback_pptx", "trace_path"]
