"""Strict read-only extraction/review reports; import never elevates identity."""

from __future__ import annotations
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import os
import tempfile
import uuid
import zipfile

from .models import ModelError, canonical_json_bytes, sha256_bytes, validate_review_semantics, writer_at_least
from .store import Store
from . import collaboration, eligibility

MAX_BYTES = 64 * 1024 * 1024


def _snapshot(paths):
    files = []
    for path in paths:
        if not path.exists():
            raise ModelError("source", "review source does not exist")
        entries = list(path.rglob("*")) if path.is_dir() else [path]
        if any(p.is_symlink() for p in entries):
            raise ModelError("source", "review directory contains symlinks; select explicit files instead")
        files.extend(sorted(p for p in entries if p.is_file()))
    if sum(p.stat().st_size for p in files) > MAX_BYTES:
        raise ModelError("source", "review source exceeds the bounded read size")
    return {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}


def evaluate(source, *, persistence="none", report_output=None, source_sha256=None, scope=None):
    if persistence != "none":
        raise ModelError("persistence", "external evaluation is strictly read-only; import is a separate explicit action")
    paths = [Path(p).expanduser().resolve() for p in (source if isinstance(source, list) else [source])]
    scope = scope or {}
    if not isinstance(scope, dict) or set(scope) - {"page_ids", "slide_numbers"}:
        raise ModelError("scope", "scope supports explicit page_ids or slide_numbers")
    for key in scope:
        values = scope[key]
        if not isinstance(values, list) or not values or len(set(values)) != len(values):
            raise ModelError("scope/" + key, "selection must be a nonempty unique array")
        if key == "slide_numbers" and any(type(v) is not int or v < 1 for v in values):
            raise ModelError("scope/slide_numbers", "slide numbers must be positive integers")
        if key == "page_ids" and any(not isinstance(v, str) or not v for v in values):
            raise ModelError("scope/page_ids", "Page IDs must be nonempty strings")
    before = _snapshot(paths)
    if source_sha256 and (len(before) != 1 or next(iter(before.values())) != source_sha256):
        raise ModelError("source_sha256", "source bytes differ from the selected fixed file")
    manuscripts = []
    for path in paths:
        if path.is_dir() and (path / ".deckmaster/current.json").is_file():
            store = Store(path)
            doc = store.load_document()
            if "slide_numbers" in scope:
                raise ModelError("scope", "project snapshots use page_ids, not slide_numbers")
            selected = set(scope.get("page_ids") or [e["page_id"] for e in doc["pages"]])
            if selected - {e["page_id"] for e in doc["pages"]}:
                raise ModelError("scope/page_ids", "selected Page is not in the current snapshot")
            manuscripts.append(
                {
                    "kind": "project_snapshot",
                    "project_id": doc["project_id"],
                    "revision_id": doc["revision_id"],
                    "task": doc.get("task"),
                    "pages": [store.read_object_json(e["page"]) for e in doc["pages"] if e["page_id"] in selected],
                    "narrative": store.read_object_json(doc["narrative"]) if doc.get("narrative") else None,
                    "content_plan": store.read_object_json(doc["content_plan"]) if doc.get("content_plan") else None,
                }
            )
            continue
        files = sorted(p for p in path.rglob("*") if p.is_file()) if path.is_dir() else [path]
        for file in files:
            if file.suffix.lower() == ".pptx":
                with zipfile.ZipFile(file) as archive:
                    if sum(i.file_size for i in archive.infolist()) > MAX_BYTES:
                        raise ModelError("source", "presentation uncompressed size exceeds the review bound")
                from pptx import Presentation

                presentation = Presentation(file)
                if "page_ids" in scope:
                    raise ModelError("scope", "external presentations use slide_numbers, not Page IDs")
                selected = set(scope.get("slide_numbers") or range(1, len(presentation.slides) + 1))
                if selected - set(range(1, len(presentation.slides) + 1)):
                    raise ModelError("scope/slide_numbers", "selected slide is not in this presentation")
                slides = []
                for index, slide in enumerate(presentation.slides, 1):
                    if index not in selected:
                        continue
                    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
                    tables = [[[cell.text for cell in row.cells] for row in shape.table.rows] for shape in slide.shapes if shape.has_table]
                    slides.append(
                        {
                            "slide_number": index,
                            "texts": texts,
                            "tables": tables,
                            "speaker_notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "",
                        }
                    )
                manuscripts.append({"kind": "presentation", "source_name": file.name, "slides": slides})
            elif file.suffix.lower() in (".txt", ".md", ".json", ".csv", ".tsv"):
                if scope:
                    raise ModelError("scope", "plain manuscripts do not have Page/slide selectors")
                manuscripts.append({"kind": "manuscript", "source_name": file.name, "text": file.read_text(encoding="utf-8")})
    if not manuscripts:
        raise ModelError("source", "no supported manuscript or presentation was selected")
    after = _snapshot(paths)
    if after != before:
        raise ModelError("source", "source changed during read-only review; no fixed report is valid")
    report = {
        "schema_version": "external_review_report.v1",
        "status": "needs_reviewer",
        "persistence": "none",
        "scope": scope,
        "source_hashes": [{"source_uri": Path(path).as_uri(), "sha256": digest} for path, digest in before.items()],
        "manuscript": manuscripts,
        "evidence_level": "L1",
        "reviewer": {"type": "tool", "id": "deck-master-read-only-extractor"},
        "professional_evidence": "not_evaluated",
        "desktop_editing": "not_evaluated",
        "source_unchanged": True,
        "limitations": [
            "Extraction and byte integrity are mechanical evidence; semantic/professional/visual review has not run",
            "An unverified report or reviewer label cannot become an independent/human receipt",
        ],
    }
    if report_output:
        output = Path(report_output).expanduser().resolve()
        if any(output == p or (p.is_dir() and output.is_relative_to(p)) for p in paths):
            raise ModelError("report_output", "report must be outside every reviewed source directory/file")
        output.parent.mkdir(parents=True, exist_ok=True)
        data = canonical_json_bytes(report)
        with tempfile.NamedTemporaryFile(dir=output.parent, delete=False) as temporary:
            temporary.write(data)
            temporary.flush()
            os.fsync(temporary.fileno())
            temp = Path(temporary.name)
        try:
            os.link(temp, output)
        finally:
            temp.unlink()
        report["report_output"] = {"name": output.name, "sha256": sha256_bytes(data)}
    return report


def import_report(project, *, input, base_revision, operation_id):
    """Persist an explicit, scoped low-assurance review through the common CAS.

    The original report and bindings remain immutable evidence. All imported
    records stay tool/needs_review; human/independent claims in a payload do not
    acquire authority by being imported.
    """
    store = Store(Path(project).expanduser())
    digest = collaboration.write_digest("review.import", input)
    with store._locked():
        doc = store.load_document()
        replay = collaboration.replay_or_conflict(store, doc, operation_id, "review.import", digest)
        if replay:
            return replay
        if doc["revision_id"] != base_revision:
            from .store import ConflictError

            raise ConflictError("base_revision", "project changed before report import")
        if not writer_at_least(doc, "collaboration.v1"):
            raise ModelError("migration_required", "report import requires explicit collaboration migration")
        bindings = input.get("source_bindings") or []
        report = input.get("external_report") or {}
        if report.get("schema_version") != "external_review_report.v1" or not bindings:
            raise ModelError("external_report", "provide a versioned report and explicit current source bindings")
        subjects = []
        current = {e["page_id"]: e["page"] for e in doc["pages"]}
        for binding in bindings:
            page_id = binding.get("page_id")
            if page_id not in current or binding.get("page_ref") != current[page_id]:
                raise ModelError("source_bindings", "bind the exact current Page; historical/foreign refs are refused")
            if not any(
                s.get("source_uri") == binding.get("source_uri") and s.get("sha256") == binding.get("source_sha256")
                for s in report.get("source_hashes") or []
            ):
                raise ModelError("source_bindings", "external source identity is not in the report")
            subjects.append(current[page_id])
        eligibility.guard(
            store, doc, "review.record", context={"targets": {"page_ids": [b["page_id"] for b in bindings], "layers": ["review"]}}
        )
        record = {
            "schema_version": "deck_review.v2",
            "review_id": "import-" + uuid.uuid4().hex,
            "review_stage": "final",
            "kind": "content",
            "status": "needs_review",
            "subjects": subjects,
            "dependencies": [
                {"kind": "content", "identity": "page:" + b["page_id"], "sha256": current[b["page_id"]]["sha256"]} for b in bindings
            ],
            "reviewer": {
                "type": "tool",
                "id": "external-report-import",
                "execution_ref": None,
                "independence_confirmed": False,
                "evidence_ref": None,
                "assurance": "unverified",
            },
            "observations": ["Imported external report has unverified source-to-Page bindings; semantic review still required"],
            "findings": [],
            "created_at": datetime.now(timezone.utc).isoformat(),
            "replaces": None,
        }
        validate_review_semantics(record)
        raw_ref = store.put_json_object({"report": report, "source_bindings": bindings})
        ref = store.put_json_object(record)
        updated = {**doc, "reviews": [*doc["reviews"], ref]}
        return collaboration.commit_collaboration_change(
            store,
            doc,
            operation_id=operation_id,
            action="review.import",
            digest=digest,
            description="Explicit low-assurance external report import",
            updated=updated,
            response={
                "review_refs": [ref],
                "original_report_ref": raw_ref,
                "assurance": "unverified",
                "next_action": "complete_semantic_review",
            },
        )
