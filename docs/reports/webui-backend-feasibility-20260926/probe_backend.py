"""Audit current backend affordances with synthetic projects, never real Host calls.

Run from repository root with PYTHONPATH=src and --out <result.json>.
Expected gaps are observations, not feature acceptance. All services stop on exit.
"""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path
import subprocess
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from unittest.mock import patch as mock_patch

from deck_master import editing, service, tasks, view
from deck_master.models import bump_revision
from deck_master.store import Store
from deck_master.web import WorkbenchServer


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[3]
    observations = []

    def record(key, finding, facts, verified=True):
        observations.append({'id': key, 'finding': finding, 'facts': facts, 'verified': verified})
        args.out.write_text(json.dumps({'status': 'incomplete', 'observations': observations}, ensure_ascii=False, indent=2)+'\n')
        assert verified, key

    def pages(count):
        return [{'schema_version': 'deck_page_package.v2', 'page_id': f'p{i:02d}',
                 'customer_visible': {'title': f'合成核验页 {i}', 'body_blocks': []},
                 'visual_spec': {'intent': 'Backend feasibility synthetic fixture',
                                 'reference_mode': 'new_design'}} for i in range(1, count+1)]

    def create(root, name, count=2):
        project = root / name
        service.create(project, brief='后端核验：纯合成，不是业务产物', draft={'pages': pages(count)})
        return project, Store(project)

    def repair_result(page):
        return {'kind': 'repair', 'files': [], 'pages': [page], 'artifact_specs': [],
                'reviews': [], 'usage_events': [], 'notes': 'Synthetic audit result; no Host/model invoked'}

    with tempfile.TemporaryDirectory(prefix='deck-master-feasibility-') as directory:
        root = Path(directory)
        project, store = create(root, 'read-http', 24)
        baseline = store.load_document()
        old_revision = baseline['revision_id']
        page = store.read_object_json(baseline['pages'][0]['page'])
        page['customer_visible']['title'] = '已保存的新标题'
        saved = editing.edit_page(project, page=page, base_revision=old_revision,
                                  page_hash=baseline['pages'][0]['page']['sha256'], operation_id='audit-edit-1')
        replay = editing.edit_page(project, page=page, base_revision=old_revision,
                                   page_hash=baseline['pages'][0]['page']['sha256'], operation_id='audit-edit-1')
        old_view = view.project_view(project, revision=old_revision)
        current_view = view.project_view(project)
        record('P01', '24-page current read and core historical read work',
               {'page_count': current_view['page_count'], 'old_title': old_view['pages'][0]['title'],
                'current_title': current_view['pages'][0]['title'], 'replay_status': replay['status']},
               current_view['page_count'] == 24 and old_view['pages'][0]['title'] != current_view['pages'][0]['title']
               and replay['status'] == 'already_applied')

        server = WorkbenchServer(project)
        base_url = server.start().rstrip('/')

        def http(path, payload=None):
            headers = {}
            data = None
            if payload is not None:
                headers = {'Content-Type': 'application/json', 'Origin': base_url,
                           'X-Deck-Token': token}
                data = json.dumps(payload).encode()
            request = urllib.request.Request(base_url + path, data=data, headers=headers)
            try:
                with urllib.request.urlopen(request, timeout=10) as response:
                    body = response.read()
                    return response.status, json.loads(body) if 'application/json' in response.headers.get('Content-Type', '') else body.decode()
            except urllib.error.HTTPError as exc:
                return exc.code, json.loads(exc.read())

        try:
            token = http('/api/session')[1]['token']
            code, old_http = http('/api/view?revision=' + old_revision)
            invalid_code, invalid_http = http('/api/view?revision=missing-audit-revision')
            single_code, single = http('/api/pages/p01?revision=' + old_revision)
            record('P02', 'HTTP revision queries are ignored, including nonexistent revisions',
                   {'http_status': code, 'matches_current': old_http['revision_id'] == saved['revision_id'],
                    'requested_revision': old_http.get('requested_revision'), 'invalid_revision_status': invalid_code,
                    'invalid_returns_current': invalid_http['revision_id'] == saved['revision_id'],
                    'single_page_status': single_code, 'single_page_title': single['page']['customer_visible']['title']},
                   code == invalid_code == single_code == 200 and old_http['revision_id'] == saved['revision_id']
                   and single['page']['customer_visible']['title'] == page['customer_visible']['title'])
            paths = ['/api/view/summary', '/v2/', '/api/projects', '/api/inputs', '/api/operations/audit-edit-1']
            codes = {path: http(path)[0] for path in paths}
            record('P03', 'Required/proposed read surfaces are absent', codes, all(v == 404 for v in codes.values()))
            record('P04', 'Existing view omits task facts and sources',
                   {'view_keys': sorted(current_view), 'core_has_task': 'task' in baseline, 'core_has_sources': 'sources' in baseline},
                   'task' not in current_view and 'sources' not in current_view)

            payload = {'page_id': 'p01', 'instruction': '保留正文，检查右上区域', 'operation_id': 'client-fixed-id',
                       'artifact_ref': baseline['pages'][0]['page'],
                       'selection': {'type': 'rect', 'x': 0.1, 'y': 0.2, 'width': 0.3, 'height': 0.4}}
            status, feedback = http('/api/feedback', payload)
            stored = tasks._lookup_task(store.load_document(), feedback['task_id'], store)
            second_status, second = http('/api/feedback', {**payload, 'instruction': '使用同一操作ID更换要求'})
            record('P05', 'Feedback drops structured selection and ignores client operation identity',
                   {'http_statuses': [status, second_status], 'selection_stored': 'selection' in stored,
                    'operation_matches_client': stored['operation_id'] == payload['operation_id'],
                    'same_client_id_created_distinct_tasks': feedback['task_id'] != second['task_id']},
                   status == second_status == 200 and 'selection' not in stored
                   and stored['operation_id'] != payload['operation_id'] and feedback['task_id'] != second['task_id'])
            before = len(store.load_document()['tasks'])
            batch_status, batch_error = http('/api/feedback', {'operation_id': 'batch-audit', 'requests': [payload]})
            record('P06', 'Batch feedback shape is unsupported; rejected without adding a task',
                   {'status': batch_status, 'error': batch_error, 'task_count_unchanged': len(store.load_document()['tasks']) == before},
                   batch_status == 400 and len(store.load_document()['tasks']) == before)
            service.task_start(project, task_id=feedback['task_id'], execution_ref='synthetic-audit-claim')
            projected = next(t for t in http('/api/tasks')[1]['tasks'] if t['task_id'] == feedback['task_id'])
            record('P07', 'Real claim state exists; HTTP task projection omits execution identity/result references',
                   {'status': projected['status'], 'projected_fields': sorted(projected)},
                   projected['status'] == 'running' and 'execution_ref' not in projected)

            doc = store.load_document()
            aged = bump_revision(doc, {'operation_id': 'audit-age-task', 'kind': 'task_update', 'description': 'Synthetic clock fixture', 'read_set': []})
            for index, ref in enumerate(aged['tasks']):
                obj = store.read_object_json(ref)
                if obj['task_id'] == feedback['task_id']:
                    obj['created_at'] = obj['updated_at'] = '2000-01-01T00:00:00Z'
                    aged['tasks'][index] = store.put_json_object(obj)
            store.commit_change(base_revision=doc['revision_id'], document=aged, operation_id='audit-age-task')
            continued = service.continue_project(project)
            aged_task = tasks._lookup_task(store.load_document(), feedback['task_id'], store)
            record('P08', 'Old running task is not timed out by continue',
                   {'status': aged_task['status'], 'updated_at': aged_task['updated_at'], 'continue_status': continued['status']},
                   aged_task['status'] == 'running' and aged_task['updated_at'].startswith('2000'))

            export_dir = root / 'review-export'
            no_ppt_error = None
            try:
                editing.export_project(project, output_dir=export_dir, purpose='review')
            except Exception as exc:
                no_ppt_error = type(exc).__name__ + ': ' + str(exc)
            record('P09a', 'Even review export refuses a content-only project',
                   {'error': no_ppt_error}, bool(no_ppt_error and 'no current PPT' in no_ppt_error))
            # Seed a clearly synthetic, genuine PPTX container solely to inspect
            # exporter packaging; no renderer or quality pass is invented.
            from pptx import Presentation
            from deck_master.pipeline import artifact as make_artifact
            fixture_ppt = root / 'synthetic-export-fixture.pptx'
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = 'Synthetic backend exporter fixture'
            presentation.save(fixture_ppt)
            doc = store.load_document()
            staged_doc = bump_revision(doc, {'operation_id': 'audit-export-fixture', 'kind': 'task_update',
                                            'description': 'Synthetic export-only fixture; not a produced deck', 'read_set': []})
            staged_doc['outputs']['pptx'] = make_artifact(store, fixture_ppt, 'pptx')
            store.commit_change(base_revision=doc['revision_id'], document=staged_doc, operation_id='audit-export-fixture')
            exported = editing.export_project(project, output_dir=export_dir, purpose='review')
            has_internal = (export_dir / 'project/.deckmaster/objects').is_dir()
            ref = exported['output_dir'] + '/delivery.json'
            download_status, _ = http('/api/file?path=' + urllib.parse.quote(ref) + '&sha256=' + '0'*64)
            engineering_error = revision_error = None
            try:
                editing.export_project(project, output_dir=root/'engineering', purpose='engineering')
            except Exception as exc:
                engineering_error = type(exc).__name__ + ': ' + str(exc)
            try:
                editing.export_project(project, output_dir=root/'old-export', purpose='review', revision=old_revision)
            except TypeError as exc:
                revision_error = str(exc)
            record('P09', 'Current review contains project objects; engineering/revision/download missing',
                   {'review_has_engineering_objects': has_internal, 'engineering_error': engineering_error,
                    'revision_error': revision_error, 'export_file_route_status': download_status},
                   has_internal and engineering_error is not None and revision_error is not None and download_status == 403)
        finally:
            server.stop()

        project, store = create(root, 'adoption')
        a = service.open_host_task(store, kind='repair', page_ids=['p01'], instruction='Synthetic first repair')
        b = service.open_host_task(store, kind='repair', page_ids=['p02'], instruction='Synthetic second repair')
        page_a = copy.deepcopy(pages(2)[0]); page_a['customer_visible']['title'] = 'p01 returned result'
        page_b = copy.deepcopy(pages(2)[1]); page_b['customer_visible']['title'] = 'p02 returned result'
        results = []
        for task, changed in [(b, page_b), (a, page_a)]:
            result = service.accept_result(project, task_id=task['task_id'], operation_id=task['operation_id'],
                                           produced_against=task['produced_against'], result_payload=repair_result(changed))
            results.append(result['status'])
        current = view.project_view(project)
        record('P10', 'Disjoint results accepted out of order, but acceptance immediately changes current',
               {'accept_statuses': results, 'current_titles': [p['title'] for p in current['pages']],
                'separate_user_candidate_adoption_called': False},
               results == ['accepted', 'accepted'] and current['pages'][0]['title'] == page_a['customer_visible']['title'])
        task = service.open_host_task(store, kind='repair', page_ids=['p01'], instruction='cancel late audit')
        service.task_cancel(project, task_id=task['task_id'])
        page_a['customer_visible']['title'] = 'must not be applied'
        refused = None
        try:
            service.accept_result(project, task_id=task['task_id'], operation_id=task['operation_id'],
                                  produced_against=task['produced_against'], result_payload=repair_result(page_a))
        except tasks.TaskConflict as exc:
            refused = str(exc)
        record('P11', 'Cancelled task rejects late product', {'refusal': refused,
               'current_title': view.project_view(project)['pages'][0]['title']},
               bool(refused) and view.project_view(project)['pages'][0]['title'] != page_a['customer_visible']['title'])

        project, store = create(root, 'prompt')
        dispatched = service.continue_project(project)['pending_tasks'][0]
        frozen = tasks._lookup_task(store.load_document(), dispatched['task_id'], store)
        requests = [obj for ref in frozen['inputs'] if isinstance(obj := store.read_object_json(ref), dict)
                    and obj.get('schema_version') == 'deck_blueprint_request.v1']
        record('P12', 'Prepared generation request already persists in Task inputs',
               {'request_count': len(requests), 'has_prompt': bool(requests and requests[0].get('prompt')),
                'task_kind': dispatched['kind']}, len(requests) == 1 and dispatched['kind'] == 'blueprint')
        staging = store.deck_root / 'staging' / dispatched['operation_id']; staging.mkdir(parents=True, exist_ok=True)
        from PIL import Image
        Image.new('RGB', (160, 90), '#a8b0b8').save(staging/'original.png')
        page_ref = store.load_document()['pages'][0]['page']
        envelope = {'kind': 'blueprint', 'files': [{'file_id': 'image', 'path': 'original.png', 'media_type': 'image/png'}],
                    'artifact_specs': [{'role': 'blueprint', 'page_id': 'p01', 'file_id': 'image', 'derived_from': [page_ref],
                                        'provenance': {'source_type': 'unknown', 'tool': 'synthetic-audit-no-model',
                                                       'invocation_ref': None, 'generated_from_page': page_ref}}],
                    'usage_events': [], 'notes': 'Pure protocol fixture, no actual generation or image quality evidence'}
        result = service.accept_result(project, task_id=dispatched['task_id'], operation_id=dispatched['operation_id'],
                                       produced_against=dispatched['produced_against'], result_payload=envelope)
        artifact = store.read_object_json(store.load_document()['pages'][0]['blueprint'])
        completed = tasks._lookup_task(store.load_document(), dispatched['task_id'], store)
        record('P15', 'Returned artifact references are not persisted into completed Task.result_refs',
               {'response_result_ref_count': len(result.get('result_refs') or []),
                'stored_task_result_ref_count': len(completed['result_refs'])},
               bool(result.get('result_refs')) and not completed['result_refs'])
        next_step = service.continue_project(project)
        record('P13', 'Blueprint accepted without submitted prompt; next work is same-page reconstruction',
               {'accept_status': result['status'], 'submitted_prompt_present': 'submitted_prompt' in artifact['provenance'],
                'next_action': next_step['next_action'], 'next_scope': next_step['pending_tasks'][0]['scope_pages']},
               result['status'] == 'accepted' and 'submitted_prompt' not in artifact['provenance']
               and next_step['next_action'] == 'codex_reconstruct_svg' and next_step['pending_tasks'][0]['scope_pages'] == ['p01'])

        project, store = create(root, 'inputs')
        before = store.load_document()
        response = service.inputs_update(project, patch={'task_patch': {'audience': '合成的新受众'}, 'reason': 'audit'},
                                         base_revision=before['revision_id'], operation_id='audit-inputs')
        after = store.load_document()
        pending = service.continue_project(project)['pending_tasks'][0]
        record('P14', 'Input update saves facts and dispatches impact/reconciliation work',
               {'status': response['status'], 'audience': after['task']['audience'],
                'content_unchanged': [p['page'] for p in before['pages']] == [p['page'] for p in after['pages']],
                'pending_kind': pending['kind'], 'pending_intent': pending.get('intent')},
               after['task']['audience'] == '合成的新受众' and pending['kind'] == 'compose' and pending.get('intent') == 'input_revision')

        project, store = create(root, 'unknown-result')
        task = service.open_host_task(store, kind='repair', page_ids=['p01'], instruction='audit result journal crash')
        changed = pages(2)[0]; changed['customer_visible']['title'] = '已提交但回执写入失败'
        envelope = repair_result(changed)
        failed = retry_error = None
        with mock_patch('deck_master.tasks.write_operation_journal', side_effect=OSError('synthetic failure after current commit')):
            try:
                service.accept_result(project, task_id=task['task_id'], operation_id=task['operation_id'],
                                      produced_against=task['produced_against'], result_payload=envelope)
            except OSError as exc:
                failed = str(exc)
        try:
            service.accept_result(project, task_id=task['task_id'], operation_id=task['operation_id'],
                                  produced_against=task['produced_against'], result_payload=envelope)
        except tasks.TaskConflict as exc:
            retry_error = str(exc)
        title = view.project_view(project)['pages'][0]['title']
        record('P16', 'Fault after current commit but before Host-result journal leaves retry conflicted',
               {'fault': failed, 'result_already_current': title == changed['customer_visible']['title'],
                'replay_error': retry_error},
               bool(failed and retry_error) and title == changed['customer_visible']['title'])

    result = {'scope': 'Real current Python core and local HTTP; synthetic input/results only',
              'production_acceptance': False, 'real_host_invoked': False, 'temporary_projects_cleaned': True,
              'source_commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=repo, text=True).strip(),
              'observations': observations, 'verified_observations': len(observations), 'status': 'passed'}
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'status': result['status'], 'observations': len(observations)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
