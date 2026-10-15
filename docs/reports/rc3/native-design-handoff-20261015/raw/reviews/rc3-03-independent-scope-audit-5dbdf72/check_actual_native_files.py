"""Independent readback and edit-copy proof over the retained actual p04 bytes."""
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from deck_master.compiler.svg import parse_svg
from deck_master.pipeline import readback

OUT=Path(__file__).resolve().parent
ROOT=Path('/workspace/work/rc3/native-image-pilot-03')
PPT=ROOT/'run-02-native-ppt'
SVG=ROOT/'run-02-native-svg-preview'
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
paths=[PPT/'actual-p04.svg',PPT/'fixed-page.json',PPT/'native/deck.pptx',
 PPT/'actual-readback.json',PPT/'rendered/page-1.png',PPT/'reader-identity.json',
 SVG/'actual-p04.svg',SVG/'actual-p04-preview.png',ROOT/'run-02-native/original-blueprint.png']
before={str(path):digest(path) for path in paths}
page=json.loads((PPT/'fixed-page.json').read_text())
parsed=parse_svg((PPT/'actual-p04.svg').read_bytes(),page_id='p04',assets={})
report=readback(PPT/'native/deck.pptx',[parsed],[page])
assert report['status']=='pass',report
assert report==json.loads((PPT/'actual-readback.json').read_text())
assert (PPT/'actual-p04.svg').read_bytes()==(SVG/'actual-p04.svg').read_bytes()
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(PPT/'native/deck.pptx') as archive:
 slide=ET.fromstring(archive.read('ppt/slides/slide1.xml'))
 counts={'native_shapes':len(slide.findall('.//p:sp',NS)),
  'text_runs':len(slide.findall('.//a:t',NS)),'raster_pictures':len(slide.findall('.//p:pic',NS))}
assert counts=={'native_shapes':59,'text_runs':35,'raster_pictures':0},counts
deck=Presentation(PPT/'native/deck.pptx')
text_shape=next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and shape.text)
original_name=text_shape.name
text_shape.text_frame.paragraphs[0].runs[0].text='AUDIT EDIT COPY ONLY'
proof=OUT/'actual-native-edit-proof-copy.pptx'
deck.save(proof)
reopened=Presentation(proof)
assert 'AUDIT EDIT COPY ONLY' in next(shape.text for shape in reopened.slides[0].shapes if shape.name==original_name)
assert {str(path):digest(path) for path in paths}==before
result={'status':'PASS','scope':'one actual p04 native engineering file only','reader_source_sha':'5dbdf72089d6625449a23a52ef3e23859c597b4e',
 'independent_readback':report,'counts':counts,'files':[{'path':key,'sha256':value} for key,value in before.items()],
 'native_edit_copy':{'path':str(proof),'sha256':digest(proof),'shape':original_name,'reopened_modified_text':True},
 'actual_original_files_unchanged':True,
 'visual_observation':'Auditor viewed both original model PNG and final native render. Main four-node chain stays forward, the two exception rows retain column ownership and qualifiers; final render explicitly displays the fixed main-chain applicability qualification.',
 'limitations':['No whole-deck delivery or risk coverage result.','No fair E2 comparison or completed full E3 protocol.','Programmatic native text edit and reopen is not human desktop editing.','Independent page_visual workflow remains Root-bound; this audit does not submit a review or mutate the active project.']}
(OUT/'actual-native-file-check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({key:value for key,value in result.items() if key not in ('files','independent_readback')},ensure_ascii=False,indent=2))
