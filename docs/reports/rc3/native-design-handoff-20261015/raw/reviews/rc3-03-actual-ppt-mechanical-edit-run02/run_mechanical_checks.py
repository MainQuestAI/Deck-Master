"""Actual native p04 artifact mutations on exclusive copies, fixed 5db reader.

This is headless engineering evidence, never desktop or model quality evidence.
No project Store writer, test fixture, review report or newer source is used.
"""
from __future__ import annotations
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import runpy
import shutil
import sys
import traceback
import xml.etree.ElementTree as ET
import zipfile

import deck_master
from deck_master import pipeline
from deck_master.compiler import CompileOptions,SvgInput,compile_deck
from deck_master.compiler.native import notes_records
from deck_master.compiler.svg import parse_svg
from pptx import Presentation

OUT=Path(__file__).resolve().parent
SOURCE=Path('/workspace/work/rc3/native-image-pilot-03')
ENGINEERING=SOURCE/'run-02-native-ppt'
INPUTS={'original.pptx':ENGINEERING/'native/deck.pptx',
 'fixed-page.json':ENGINEERING/'fixed-page.json','actual-svg-ir.json':ENGINEERING/'actual-svg-ir.json',
 'original-render.png':ENGINEERING/'rendered/page-1.png',
 'original.svg':SOURCE/'run-02-svg-native-author/p04-native-v1.svg'}
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
SLIDE='ppt/slides/slide1.xml'
SVG_NS='http://www.w3.org/2000/svg'

def digest(data):return hashlib.sha256(data).hexdigest()
def save(name,obj):
    with (OUT/name).open('x') as f:json.dump(obj,f,ensure_ascii=False,indent=2,sort_keys=True);f.write('\n')
def inventory():return {k:{'path':str(v),'bytes':v.stat().st_size,'sha256':digest(v.read_bytes())} for k,v in INPUTS.items()}
def root_of(path):
    with zipfile.ZipFile(path) as z:return ET.fromstring(z.read(SLIDE))
def texts(root):return [t.text or '' for t in root.findall('.//a:t',NS)]
def shapes(root,name):return [s for s in root.findall('.//p:sp',NS) if s.find('p:nvSpPr/p:cNvPr',NS).get('name')==name]
def one(root,name):
    found=shapes(root,name);assert len(found)==1,(name,len(found));return found[0]
def replace_slide(target,root):
    with zipfile.ZipFile(OUT/'inputs/original.pptx') as old,zipfile.ZipFile(target,'x',zipfile.ZIP_DEFLATED) as new:
        for meta in old.infolist():new.writestr(meta,ET.tostring(root,encoding='utf-8',xml_declaration=True) if meta.filename==SLIDE else old.read(meta.filename))
def geometry(native):
    off=native.find('p:spPr/a:xfrm/a:off',NS);ext=native.find('p:spPr/a:xfrm/a:ext',NS)
    return {'name':native.find('p:nvSpPr/p:cNvPr',NS).get('name'),'offset':dict(off.attrib),'extent':dict(ext.attrib),
            'text':''.join(texts(native))}
def report(path,ir,page,name):
    result=pipeline.readback(path,[ir],[page]);save(name,result);return result
def admission():
    helper=SOURCE/'prepare-pilot.py'
    assert digest(helper.read_bytes())=='4f51363e867cdc367416a4b665a7537db5f2d1fa1c343b73cab848187790e782'
    return runpy.run_path(str(helper))['installation'](
       Path('/workspace/work/rc3/candidate-03-5dbdf72/manifest.json'),'5dbdf72089d6625449a23a52ef3e23859c597b4e')

def case_a(ir,page):
    path=OUT/'A-python-object-text-edited.pptx'
    prs=Presentation(OUT/'inputs/original.pptx')
    matches=[s for s in prs.slides[0].shapes if s.name=='atom:p04:title']
    assert len(matches)==1 and matches[0].has_text_frame
    shape=matches[0];old=shape.text;run=shape.text_frame.paragraphs[0].runs[0]
    assert run.text==old
    new='机械编辑验证：'+old
    run.text=new
    prs.save(path)
    reopened=Presentation(path)
    again=next(s for s in reopened.slides[0].shapes if s.name=='atom:p04:title')
    assert again.text==new and again.text!=old
    assert len(reopened.slides)==1 and len(reopened.slides[0].shapes)==len(prs.slides[0].shapes)
    assert notes_records(path)==notes_records(OUT/'inputs/original.pptx')
    r=report(path,ir,page,'A-readback-against-unchanged-page.json')
    assert r['status']=='fail'  # Intentional copy edit must not become the fixed manuscript.
    result={'status':'PASS_headless_native_object_edit_save_reopen_only','object_name':shape.name,
      'old_text':old,'new_text':new,'reopened_text':again.text,'slide_count':len(reopened.slides),
      'shape_count':len(reopened.slides[0].shapes),'notes_unchanged':True,'pptx_sha256':digest(path.read_bytes()),
      'readback_against_unchanged_page':r['status'],'desktop_open_edit_save_reopen':'NOT_RUN'}
    save('A-edit-observation.json',result);return result

def case_b(ir,page):
    root=root_of(OUT/'inputs/original.pptx');original_text=texts(root)
    upper_id='atom:p04:cell:exceptions:unknown:action:display_text'
    lower_id='atom:p04:cell:exceptions:unsuitable:action:display_text'
    upper,lower=shapes(root,upper_id),shapes(root,lower_id)
    assert len(upper)==3 and len(lower)==2
    before=[geometry(s) for s in upper+lower]
    uy=min(int(s.find('p:spPr/a:xfrm/a:off',NS).get('y')) for s in upper)
    ly=min(int(s.find('p:spPr/a:xfrm/a:off',NS).get('y')) for s in lower)
    delta=ly-uy;assert delta==126*9525
    for group,amount in [(upper,delta),(lower,-delta)]:
        for s in group:
            off=s.find('p:spPr/a:xfrm/a:off',NS);off.set('y',str(int(off.get('y'))+amount))
    path=OUT/'B-native-cell-row-ownership-swapped.pptx';replace_slide(path,root)
    reopened=root_of(path)
    assert texts(reopened)==original_text and Counter(''.join(texts(reopened)))==Counter(''.join(original_text))
    r=report(path,ir,page,'B-actual-readback.json')
    findings=[f for f in r['findings'] if f['code']=='table_cell_ownership_mismatch']
    result={'status':'PASS_negative_detected' if r['status']=='fail' and {f.get('atom_id') for f in findings}=={upper_id,lower_id} else 'FAIL_checker_missed_or_mislocalized',
      'mutation':'swap all native text-carrier row positions for two action cells; preserve text, names, Page and source IR',
      'offset_emu':delta,'offset_px':delta/9525,'carriers_before':before,'carriers_after':[geometry(s) for s in upper+lower],
      'exact_text_sequence_unchanged':True,'character_multiset_unchanged':True,'character_count':len(''.join(original_text)),
      'readback_status':r['status'],'finding_codes':[f['code'] for f in r['findings']],
      'located_ownership_findings':findings,'pptx_sha256':digest(path.read_bytes())}
    save('B-ownership-observation.json',result);assert result['status']=='PASS_negative_detected';return result

def case_c(ir,page):
    root=root_of(OUT/'inputs/original.pptx');original_text=texts(root)
    line=one(root,'edge-collect-retrieve');head=one(root,'arrowhead-collect-retrieve')
    path=line.find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
    points=[child.find('a:pt',NS) for child in path if child.find('a:pt',NS) is not None]
    assert len(points)==2
    original_endpoints=[dict(p.attrib) for p in points]
    points[0].attrib,points[1].attrib=dict(points[1].attrib),dict(points[0].attrib)
    triangle=head.find('p:spPr/a:custGeom/a:pathLst/a:path',NS);width=int(triangle.get('w'))
    triangle_points=triangle.findall('.//a:pt',NS);original_triangle=[dict(p.attrib) for p in triangle_points]
    assert len(triangle_points)==3
    for p in triangle_points:p.set('x',str(width-int(p.get('x'))))
    target=OUT/'C-native-edge-and-visible-arrow-reversed.pptx';replace_slide(target,root)
    assert texts(root_of(target))==original_text
    r=report(target,ir,page,'C-actual-readback.json')
    findings=[f for f in r['findings'] if f['code']=='reversed_arrow' and f.get('edge_id')=='collect-retrieve']
    result={'status':'PASS_negative_detected' if r['status']=='fail' and findings else 'FAIL_checker_missed_reversed_edge',
      'mutation':'reverse actual native edge path endpoint order and mirror its real separate visible arrowhead; no text/Page/IR changes',
      'edge_id':'collect-retrieve','path_before':original_endpoints,'path_after':[dict(p.attrib) for p in points],
      'visible_arrowhead_before':original_triangle,'visible_arrowhead_after':[dict(p.attrib) for p in triangle_points],
      'exact_text_sequence_unchanged':True,'character_multiset_unchanged':True,'character_count':len(''.join(original_text)),
      'readback_status':r['status'],'finding_codes':[f['code'] for f in r['findings']],
      'located_direction_findings':findings,'pptx_sha256':digest(target.read_bytes())}
    save('C-direction-observation.json',result);assert result['status']=='PASS_negative_detected';return result

def case_d(ir,page):
    original=(OUT/'inputs/original.svg').read_bytes()
    assert parse_svg(original,page_id='p04',assets={})==ir
    root=ET.fromstring(original)
    group=next(g for g in root.findall('{'+SVG_NS+'}g') if g.get('id')=='main-chain')
    assert group.get('transform') is None
    group.set('transform','translate(0 2)')
    ET.register_namespace('',SVG_NS)
    svg=OUT/'D-real-svg-main-chain-moved-2px.svg'
    svg.write_bytes(ET.tostring(root,encoding='utf-8',xml_declaration=True))
    moved=parse_svg(svg.read_bytes(),page_id='p04',assets={})
    identity=lambda s:(s['id'],s['kind'],s.get('text'),s.get('atom_id'),s.get('node_ref'),s.get('edge_ref'))
    assert [identity(s) for s in moved['shapes']]==[identity(s) for s in ir['shapes']]
    assert [s['text'] for s in moved['shapes'] if s['kind']=='text']==[s['text'] for s in ir['shapes'] if s['kind']=='text']
    save('D-actual-transformed-svg-ir.json',moved)
    fonts=pipeline.resolve_fonts([moved])
    compiled=compile_deck([SvgInput('p04',svg,page.get('speaker_notes',''))],
         CompileOptions(width_px=1600,height_px=900,fonts=fonts),OUT/'D-real-recompile')
    r=report(compiled.pptx_path,moved,page,'D-actual-readback.json')
    result={'status':'PASS_layout_only_positive' if r['status']=='pass' else 'FAIL_layout_false_positive',
      'mutation':'actual original SVG main-chain group translated 2px down, then real installed parse and native compilation',
      'canvas':[1600,900],'translation_px':[0,2],'minimum_external_margin_px_at_least':48,
      'complete_Page_unchanged':True,'atom_node_edge_bindings_and_text_unchanged':True,
      'svg_sha256':digest(svg.read_bytes()),'pptx_sha256':digest(compiled.pptx_path.read_bytes()),
      'readback_status':r['status'],'findings':r['findings'],'compiler_diagnostics':compiled.diagnostics,
      'font_files':{k:{'path':v,'sha256':digest(Path(v).read_bytes())} for k,v in fonts.items()},
      'visual_quality':'NOT_EVALUATED','desktop_editing':'NOT_RUN'}
    save('D-layout-positive-observation.json',result);assert result['status']=='PASS_layout_only_positive';return result

def main():
    assert not (OUT/'inputs').exists();(OUT/'inputs').mkdir()
    save('command-and-scope.json',{'argv':sys.argv,'python':sys.executable,'package':deck_master.__file__,
       'source_phase':'5dbdf72089d6625449a23a52ef3e23859c597b4e','desktop':'NOT_RUN',
       'project_writes':False,'source_edits':False,'fixture_imports':False,'model_output_generation':False})
    identity=admission();save('installed-admission-before.json',identity)
    before=inventory();save('original-input-inventory-before.json',before)
    for name,path in INPUTS.items():shutil.copy2(path,OUT/'inputs'/name)
    page=json.loads((OUT/'inputs/fixed-page.json').read_text());ir=json.loads((OUT/'inputs/actual-svg-ir.json').read_text())
    baseline=report(OUT/'inputs/original.pptx',ir,page,'baseline-independent-actual-readback.json')
    assert baseline['status']=='pass',baseline
    results={}
    for name,fn in [('A',case_a),('B',case_b),('C',case_c),('D',case_d)]:
        try:results[name]=fn(ir,page)
        except Exception as e:
            results[name]={'status':'FAIL_OR_UNSUPPORTED','error':str(e),'traceback':traceback.format_exc()}
            save(name+'-failure.json',results[name])
    after=inventory();save('original-input-inventory-after.json',after);assert before==after
    installed_after=admission();save('installed-admission-after.json',installed_after);assert installed_after==identity
    summary={'evidence_kind':'actual_original_native_PPTX_headless_edit_and_semantic_geometry_pairs',
       'baseline_readback':'pass','original_inputs_unchanged':True,'fixed_installed_bytes_unchanged':True,
       'cases':{k:v['status'] for k,v in results.items()},'actual_project_writer_calls':0,
       'native_model_calls':0,'image_tool_calls':0,'PowerPoint_desktop':'NOT_RUN',
       'professional_visual_quality':'NOT_EVALUATED','fair_E3':'NOT_CLAIMED','card_completion':'NOT_CLAIMED'}
    save('summary.json',summary);print(json.dumps(summary,ensure_ascii=False,indent=2,sort_keys=True))
    assert all(v['status'].startswith('PASS') for v in results.values())

if __name__=='__main__':main()
