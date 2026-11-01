"""Native DrawingML emitter. Paint derived from B0 pptx::_append_color.

B0 2a866cf: keep explicit alpha=0; no business runtime imports.
The current IR contains only source-derived coordinates and visible text.
"""
from pathlib import Path
from PIL import ImageFont, ImageColor
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
import io
import hashlib
import copy


def _role_name(source, page_id):
    if source.get('atom_id') == f'atom:{page_id}:title':
        return '标题'
    if source.get('atom_id') == f'atom:{page_id}:subtitle':
        return '副标题'
    role = {'text': '正文', 'image': '图片', 'line': '连线',
            'path': '图形', 'polygon': '箭头', 'polyline': '连线',
            'rect': '底板', 'circle': '图形', 'ellipse': '图形'}[source['kind']]
    # The parser appends a run index even to a single authored text box.
    label = source['id'].removesuffix(':0') if source['kind'] == 'text' else source['id']
    return f"{source.get('group') + '-' if source.get('group') else ''}{role}-{label}"


def normalize_notes(text):
    """Normalize transport newlines only; never trim speech or blank paragraphs."""
    if not isinstance(text, str):
        raise ValueError('speaker_notes must be a string')
    return text.replace('\r\n', '\n').replace('\r', '\n')


def notes_records(data):
    """Read complete body-placeholder notes in actual presentation relationship order."""
    deck = Presentation(io.BytesIO(data) if isinstance(data, bytes) else data)
    rows = []
    for index, slide in enumerate(deck.slides, 1):
        text = ''
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                text = normalize_notes(frame.text)
        rows.append({'slide': index, 'page_id': slide.name or None, 'notes': text,
                     'notes_sha256': hashlib.sha256(text.encode('utf-8')).hexdigest()})
    return rows


def _write_notes(slide, text):
    """Use readable notes-page typography in every native compiler exit."""
    frame = slide.notes_slide.notes_text_frame
    frame.text = normalize_notes(text)
    # The stock notes master inherits a large slide-text font. It can hide
    # paragraphs in compatible notes-page views even with complete XML text.
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(12)


def attach_notes(data, records):
    """Attach trusted Page text to a notes-free presentation, including stable IDs."""
    deck = Presentation(io.BytesIO(data))
    if len(deck.slides) != len(records):
        raise ValueError('notes page count differs from actual slides')
    for slide, record in zip(deck.slides, records, strict=True):
        slide.name = record['page_id']
        text = normalize_notes(record['notes'])
        if text:
            _write_notes(slide, text)
    result = io.BytesIO()
    deck.save(result)
    return result.getvalue()


def paint(parent, color, alpha):
    for child in list(parent):
        if child.tag.rsplit('}',1)[-1] in ('solidFill','noFill','gradFill'):
            parent.remove(child)
    if isinstance(color,dict):
        fill=OxmlElement('a:gradFill');fill.set('rotWithShape','1');stops=OxmlElement('a:gsLst')
        for item in color['stops']:
            stop=OxmlElement('a:gs');stop.set('pos',str(round(item['offset']*100000)))
            holder=OxmlElement('a:rPr');paint(holder,item['color'],alpha*item['opacity'])
            stop.append(holder[0][0]);stops.append(stop)
        fill.append(stops)
        if color['kind']=='linear':
            direction=OxmlElement('a:lin');direction.set('ang',str(round(color['angle']*60000)));direction.set('scaled','1')
        else:
            direction=OxmlElement('a:path');direction.set('path','circle');rect=OxmlElement('a:fillToRect')
            for key in ('l','t','r','b'):rect.set(key,'50000')
            direction.append(rect)
        fill.append(direction);parent.append(fill);return
    if color == 'none':
        parent.append(OxmlElement('a:noFill')); return
    fill=OxmlElement('a:solidFill'); rgb=OxmlElement('a:srgbClr')
    rgba=ImageColor.getrgb(color)
    rgb.set('val',''.join(f'{v:02X}' for v in rgba[:3]))
    opacity=OxmlElement('a:alpha');opacity.set('val',str(round(alpha*(rgba[3]/255 if len(rgba)==4 else 1)*100000)))
    rgb.append(opacity);fill.append(rgb);parent.append(fill)


def _trim_table_text_slack(grid_shapes, text_boxes, page_id):
    """Trim only measured horizontal slack at actual drawn column boundaries.

    Readback still verifies the full native grid, row/column ownership and
    perimeter. No source label, font-size tolerance or padding exemption can
    make an out-of-cell glyph fit. Alignment anchors stay fixed in EMUs.
    """
    from ..readback_relations import _native_grid_shape
    native = [(source, _native_grid_shape(shape._element, frame=bool(source.get('table_frame'))))
              for source, shape in grid_shapes]
    grids = {}
    for source, frame in native:
        if not source.get('table_frame') or frame is None:
            continue
        left, top = frame['x'], frame['y']
        right, bottom = left + frame['width'], top + frame['height']
        epsilon = .02  # native zero-width line representation, never text slop
        vertical = [rule['x'] for original, rule in native if original['kind'] == 'line' and rule
                    and rule['width'] <= epsilon and left < rule['x'] < right
                    and abs(rule['y'] - top) <= epsilon and abs(rule['height'] - frame['height']) <= epsilon
                    and all(any(a-epsilon <= rule['x'] <= b+epsilon for a, b in frame['sides'][side])
                            for side in ('top', 'bottom'))]
        grids.setdefault(source['table_id'], set()).add((left, *sorted(vertical), right))
    for source, shape, glyph_width, anchor in text_boxes:
        if source.get('rotation'):
            continue
        identity = source.get('atom_id') or ''
        matches = [xs for table_id, choices in grids.items() if len(choices) == 1
                   and any(identity.startswith(f'atom:{page_id}:{kind}:{table_id}:') for kind in ('cell', 'column'))
                   for xs in choices]
        if len(matches) != 1:
            continue
        cells = [(round(left*9525), round(right*9525)) for left, right in zip(matches[0], matches[0][1:])
                 if (round(left*9525) <= anchor < round(right*9525) if source['anchor'] == 'start'
                     else round(left*9525) < anchor <= round(right*9525) if source['anchor'] == 'end'
                     else round(left*9525) < anchor < round(right*9525))]
        if len(cells) != 1:
            continue
        left, right = cells[0]
        maximum = {'middle': 2*min(anchor-left, right-anchor), 'end': anchor-left}.get(source['anchor'], right-anchor)
        if glyph_width <= maximum < shape.width:
            shape.width = maximum
            shape.left = anchor - ({'middle': maximum // 2, 'end': maximum}.get(source['anchor'], 0))


def emit(pages, width, height, fonts, output):
    pres=Presentation();pres.slide_width=round(width*9525);pres.slide_height=round(height*9525)
    for page in pages:
        slide=pres.slides.add_slide(pres.slide_layouts[6]); k=min(width/page['width'],height/page['height'])
        slide.name = page['page_id']
        if page.get('speaker_notes'):
            _write_notes(slide, page['speaker_notes'])
        ox=(width-page['width']*k)/2;oy=(height-page['height']*k)/2
        def X(x):return round((ox+x*k)*9525)
        def Y(y):return round((oy+y*k)*9525)
        def S(v):return round(v*k*9525)
        grid_shapes, text_boxes = [], []
        native_by_source = {}
        for s in page['shapes']:
            if s['kind']=='image':
                from PIL import Image
                x,y,w,h=s['x'],s['y'],s['width'],s['height']
                if s['fit']=='contain':
                    with Image.open(s['asset_path']) as raster:iw,ih=raster.size
                    ratio=min(w/iw,h/ih);x+=(w-iw*ratio)/2;y+=(h-ih*ratio)/2;w=iw*ratio;h=ih*ratio
                sh=slide.shapes.add_picture(s['asset_path'],X(x),Y(y),S(w),S(h));sh.name=s.get('atom_id') or s['id']
                sh._element.nvPicPr.cNvPr.set('title', _role_name(s, page['page_id']))
                native_by_source[id(s)] = sh
                if s['fit']=='cover':
                    with Image.open(s['asset_path']) as raster:iw,ih=raster.size
                    if iw/ih>w/h:sh.crop_left=sh.crop_right=(1-(w/h)/(iw/ih))/2
                    else:sh.crop_top=sh.crop_bottom=(1-(iw/ih)/(w/h))/2
                continue
            if s['kind']=='text':
                family=s['font_family'];key=family+(':bold' if s['bold'] else '')
                filename=fonts.get(key,fonts.get(family))
                if not filename:raise ValueError(f"{page['page_id']}/{s['id']}: font file not supplied: {key}")
                text_width = 0
                for text_run in s.get('runs', [s]):
                    run_key = text_run['font_family'] + (':bold' if text_run['bold'] else '')
                    run_file = fonts.get(run_key) or fonts.get(text_run['font_family'])
                    if not run_file:
                        raise ValueError(f"{page['page_id']}/{s['id']}: font file not supplied: {run_key}")
                    run_font = ImageFont.truetype(str(run_file), size=1000)
                    text_width += run_font.getlength(text_run['text']) / 1000 * text_run['font_size']
                text_width += max(0, len(s['text']) - 1) * s.get('letter_spacing', 0)
                w=s.get('text_width', text_width+s['font_size']*1.0)
                anchor=s['x']
                # The extra glyph of slack keeps Office from clipping text, but
                # should not itself create a false slide-boundary overflow.
                # Shrink only slack and recompute the box from the same anchor;
                # otherwise centered/right-aligned glyphs move with the box.
                if not s.get('rotation') and 0 <= anchor <= page['width']:
                    maximum = {'middle': 2*min(anchor, page['width']-anchor),
                               'end': anchor}.get(s['anchor'], page['width']-anchor)
                    if text_width <= maximum < w:
                        w = maximum
                x=anchor
                if s['anchor']=='middle':x-=w/2
                if s['anchor']=='end':x-=w
                y=s['y']-s['font_size']*.88;h=s.get('text_height', s['font_size']*1.5)
                # The IR anchor is already in final page coordinates; glyph rotation is
                # applied once through sh.rotation below, so the box stays at the anchor.
                sh=slide.shapes.add_textbox(X(x),Y(y),S(w),S(h))
                text_boxes.append((s, sh, S(text_width), X(anchor)))
                tf=sh.text_frame;tf.clear();tf.word_wrap='text_width' in s;tf.vertical_anchor=MSO_ANCHOR.TOP
                from pptx.enum.text import MSO_AUTO_SIZE
                tf.auto_size=MSO_AUTO_SIZE.NONE
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                p=tf.paragraphs[0];p.alignment={'middle':PP_ALIGN.CENTER,'end':PP_ALIGN.RIGHT}.get(s['anchor'],PP_ALIGN.LEFT)
                p.space_before=p.space_after=Pt(0)
                if s.get('bullet'):
                    props = p._p.get_or_add_pPr()
                    props.set('marL', str(S(s['font_size'])))
                    props.set('indent', str(-S(s['font_size'] * .75)))
                    bullet = OxmlElement('a:buChar')
                    bullet.set('char', '•')
                    props.append(bullet)
                for text_run in s.get('runs', [s]):
                    run = p.add_run()
                    run.text = text_run['text']
                    props = run._r.get_or_add_rPr()
                    # Fill properties must precede latin/ea/cs in a:rPr.
                    paint(props, text_run['fill'], text_run['opacity'] * text_run['fill_opacity'])
                    run.font.size = Pt(text_run['font_size'] * k * .75)
                    run.font.name = text_run['font_family']
                    run.font.bold = text_run['bold']
                    for name in ('ea', 'cs'):
                        element = OxmlElement('a:' + name)
                        element.set('typeface', text_run['font_family'])
                        props.append(element)
                # Letter spacing is written once by the api.py post-pass, which
                # keys on the shape name for both native and node emitters.
            elif s['kind'] in ('rect','circle','ellipse'):
                shape=MSO_SHAPE.OVAL if s['kind']!='rect' else MSO_SHAPE.ROUNDED_RECTANGLE if s.get('rx') else MSO_SHAPE.RECTANGLE
                sh=slide.shapes.add_shape(shape,X(s['x']),Y(s['y']),S(s['width']),S(s['height']))
                if shape==MSO_SHAPE.ROUNDED_RECTANGLE:
                    sh.adjustments[0]=min(.5,s['rx']/min(s['width'],s['height']))
            else:
                cmds=list(s.get('commands') or [{('moveTo' if i==0 else 'lineTo'):{'x':p[0],'y':p[1]}} for i,p in enumerate(s['points'])])
                if s['kind']=='polygon':cmds.append({'close':{}})
                pts=[next(iter(c.values())) for c in cmds if 'close' not in c]
                l=min(p['x'] for p in pts);t=min(p['y'] for p in pts);w=max(.01,max(p['x'] for p in pts)-l);h=max(.01,max(p['y'] for p in pts)-t)
                sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,X(l),Y(t),S(w),S(h))
                sp=sh._element.spPr
                for old in list(sp):
                    if old.tag.rsplit('}',1)[-1]=='prstGeom':sp.remove(old)
                geom=OxmlElement('a:custGeom')
                for tag in ('avLst','gdLst','ahLst','cxnLst'):geom.append(OxmlElement('a:'+tag))
                rect=OxmlElement('a:rect')
                for name,val in [('l','0'),('t','0'),('r','r'),('b','b')]:rect.set(name,val)
                geom.append(rect);lst=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w',str(max(1,S(w))));path.set('h',str(max(1,S(h))))
                for c in cmds:
                    op=next(iter(c));node=OxmlElement('a:'+('lnTo' if op=='lineTo' else op))
                    if op!='close':
                        point=OxmlElement('a:pt');point.set('x',str(S(c[op]['x']-l)));point.set('y',str(S(c[op]['y']-t)));node.append(point)
                    path.append(node)
                lst.append(path);geom.append(lst);sp.append(geom)
            native_by_source[id(s)] = sh
            if s.get('rotation'):sh.rotation=s['rotation']
            sh.name=s.get('atom_id') or s['id']
            identity = sh._element.nvSpPr.cNvPr
            identity.set('title', _role_name(s, page['page_id']))
            if s.get('atom_id') == f"atom:{page['page_id']}:title":
                placeholder = OxmlElement('p:ph')
                placeholder.set('type', 'title')
                sh._element.nvSpPr.nvPr.append(placeholder)
            for node in list(sh._element):
                if node.tag.rsplit('}',1)[-1]=='style':sh._element.remove(node)
            if s['kind']!='text':
                # CT_ShapeProperties order: geometry, fill, a:ln, effectLst.
                paint(sh._element.spPr,s['fill'],s['opacity']*s['fill_opacity'])
                line=sh._element.spPr.get_or_add_ln();line.set('w',str(S(s['stroke_width'])))
                paint(line,s['stroke'],s['opacity']*s['stroke_opacity'])
                caps={'butt':'flat','round':'rnd','square':'sq'}
                joins={'miter':'miter','round':'round','bevel':'bevel'}
                if s.get('stroke_linecap','butt') not in caps or s.get('stroke_linejoin','miter') not in joins:
                    raise ValueError(f"{page['page_id']}/{s['id']}: unsupported stroke cap/join")
                line.set('cap',caps[s.get('stroke_linecap','butt')])
                join=OxmlElement('a:'+joins[s.get('stroke_linejoin','miter')])
                if s.get('stroke_linejoin','miter')=='miter':join.set('lim',str(round(s.get('stroke_miterlimit',4)*100000)))
                line.append(join)
                sh._element.spPr.append(OxmlElement('a:effectLst'))
            else:
                sh._element.spPr.append(OxmlElement('a:effectLst'))
            if s['kind'] == 'line' or s.get('table_frame'):
                grid_shapes.append((s, sh))
        _trim_table_text_slack(grid_shapes, text_boxes, page['page_id'])
        _editable_tables(slide, page, X, Y, S, native_by_source)
        _module_groups(slide, page, native_by_source)
    pres.save(output)


def _editable_tables(slide, page, X, Y, S, native_by_source):
    """Replace explicitly marked SVG grids with editable PowerPoint tables."""
    for frame in page['shapes']:
        if not frame.get('native_table') or not frame.get('table_frame'):
            continue
        table_id = frame['table_id']
        left, top = frame['x'], frame['y']
        right, bottom = left + frame['width'], top + frame['height']
        members = [s for s in page['shapes'] if s.get('table_id') == table_id]
        xs = sorted({left, right, *(s['points'][0][0] for s in members
                      if s['kind'] == 'line' and s['points'][0][0] == s['points'][1][0])})
        ys = sorted({top, bottom, *(s['points'][0][1] for s in members
                      if s['kind'] == 'line' and s['points'][0][1] == s['points'][1][1])})
        shape = slide.shapes.add_table(len(ys) - 1, len(xs) - 1,
                                      X(left), Y(top), S(right - left), S(bottom - top))
        shape.name = table_id
        shape._element.nvGraphicFramePr.cNvPr.set('title', '表格-' + (frame['id'] if frame['id'] != 'rect' else table_id))
        table = shape.table
        table.first_row = True
        table.horz_banding = False
        for col, (a, b) in zip(table.columns, zip(xs, xs[1:])):
            col.width = S(b - a)
        for row, (a, b) in zip(table.rows, zip(ys, ys[1:])):
            row.height = S(b - a)
        originals = native_by_source
        for r in range(len(ys) - 1):
            for c in range(len(xs) - 1):
                cell = table.cell(r, c)
                cell.margin_left = cell.margin_right = S(12)
                cell.margin_top = cell.margin_bottom = S(8)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                props = cell._tc.get_or_add_tcPr()
                for side in ('L', 'R', 'T', 'B'):
                    line = OxmlElement('a:ln' + side)
                    line.set('w', str(S(frame['stroke_width'])))
                    paint(line, frame['stroke'], 1)
                    props.append(line)
                # CT_TableCellProperties places all border lines before fill.
                paint(props, '#eaf1f6' if r == 0 else '#ffffff', 1)
                texts = [s for s in members if s['kind'] == 'text'
                         and xs[c] <= s['x'] < xs[c + 1] and ys[r] <= s['y'] < ys[r + 1]]
                if texts:
                    texts.sort(key=lambda text: (text['y'], text['x']))
                    original = originals[id(texts[0])]
                    body = copy.deepcopy(original._element.txBody)
                    body.tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}txBody'
                    body.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr').set('wrap', 'square')
                    for text in texts[1:]:
                        for paragraph in originals[id(text)]._element.txBody.p_lst:
                            body.append(copy.deepcopy(paragraph))
                    cell._tc.replace(cell._tc.txBody, body)
        elements = [originals[id(s)]._element for s in members]
        first = min(slide.shapes._spTree.index(e) for e in elements)
        slide.shapes._spTree.remove(shape._element)
        slide.shapes._spTree.insert(first, shape._element)
        for element in elements:
            slide.shapes._spTree.remove(element)


def _module_groups(slide, page, native_by_source):
    groups = {}
    for source in page['shapes']:
        if source.get('group') and not source.get('table_id'):
            groups.setdefault(source['group'], []).append(native_by_source[id(source)])
    for name, members in groups.items():
        if members:
            position = min(slide.shapes._spTree.index(shape._element) for shape in members)
            group = slide.shapes.add_group_shape(members)
            group.name = name
            group._element.nvGrpSpPr.cNvPr.set('title', name)
            slide.shapes._spTree.remove(group._element)
            slide.shapes._spTree.insert(position, group._element)
