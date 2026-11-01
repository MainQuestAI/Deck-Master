"""S10/CQ-05 neutral native geometry fixtures; no model/human acceptance.

Every PPTX is genuinely compiled; corruptions edit only the named native
carrier. No actual benchmark input, image or private Judge material is used.
"""
from copy import deepcopy
from pathlib import Path
import json
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
from pptx import Presentation
import pytest

from deck_tools.compiler import CompileOptions, SvgInput, compile_deck
from deck_tools.compiler.svg import parse_svg
from deck_tools.render import executable, readback, render_deck, run
from hostenv import host_fonts, resolve_host_font

pytestmark = pytest.mark.render
PID = 'neutral-table'
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def atom(row, col):
    return f'atom:{PID}:cell:ledger:{row}:{col}:display_text'


def page_fixture():
    page = {'schema_version': 'deck_page_package.v2', 'page_id': PID,
        'customer_visible': {'title': 'Neutral ownership fixture', 'body_blocks': [{
            'id': 'ledger', 'type': 'table', 'columns': [
                {'id': 'key', 'label': 'Key'}, {'id': 'action', 'label': 'Action'},
                {'id': 'context', 'label': 'Context'}],
            'rows': [{'id': row, 'cells': [
                {'column_id': 'key', 'display_text': label},
                {'column_id': 'action', 'display_text': action},
                {'column_id': 'context', 'display_text': context}]} for row, label, action, context in (
                    ('alpha', 'Alpha', 'Wait 10 ms', 'Repeated note'),
                    ('beta', 'Beta', 'Retain basis', 'Repeated note'),
                    ('gamma', 'Gamma', 'Keep record', 'First cue. Retain basis.'))]}]},
        'visual_spec': {'intent': 'Neutral fixed row/column ownership', 'reference_mode': 'new_design'},
        'speaker_notes': 'ENGINEERING_ONLY fixed notes.\n\nKeep the ending.\n'}
    return page


def svg_fixture(page, *, transform=None, wrong_source=None):
    root = ET.Element('svg', {'xmlns': 'http://www.w3.org/2000/svg', 'viewBox': '0 0 1280 720'})
    ET.SubElement(root, 'rect', {'id': 'background', 'width': '1280', 'height': '720', 'fill': '#ffffff'})
    def text(parent, identity, value, x, y, anchor='start'):
        e = ET.SubElement(parent, 'text', {'id': identity.replace(':', '-'), 'data-atom-id': identity,
            'x': str(x), 'y': str(y), 'font-size': '24', 'font-family': resolve_host_font(),
            'fill': '#20242a', 'text-anchor': anchor})
        e.text = value
        return e
    text(root, f'atom:{PID}:title', page['customer_visible']['title'], 40, 80)
    group = ET.SubElement(root, 'g', {'transform': transform} if transform else {})
    ET.SubElement(group, 'rect', {'id': 'ledger-frame', 'data-table-id': 'ledger',
        'x': '40', 'y': '170', 'width': '1200', 'height': '360', 'rx': '5', 'fill': 'none',
        'stroke': '#8090a0', 'stroke-width': '1'})
    for i, x in enumerate((260, 700)):
        ET.SubElement(group, 'line', {'id': f'column-rule-{i}', 'x1': str(x), 'x2': str(x),
            'y1': '170', 'y2': '530', 'stroke': '#8090a0'})
    for i, y in enumerate((260, 350, 440)):
        ET.SubElement(group, 'line', {'id': f'row-rule-{i}', 'x1': '40', 'x2': '1240',
            'y1': str(y), 'y2': str(y), 'stroke': '#8090a0'})
    table = page['customer_visible']['body_blocks'][0]
    for column, x in zip(table['columns'], (150, 480, 970), strict=True):
        text(group, f"atom:{PID}:column:ledger:{column['id']}:label", column['label'], x, 225, 'middle')
    for index, row in enumerate(table['rows']):
        for cell, x in zip(row['cells'], (60, 280, 720), strict=True):
            y = 312 + index * 90
            if wrong_source == 'row' and row['id'] == 'alpha' and cell['column_id'] == 'action':
                y += 90
            if wrong_source == 'column' and row['id'] == 'alpha' and cell['column_id'] == 'action':
                x = 720
            identity = atom(row['id'], cell['column_id'])
            e = text(group, identity, cell['display_text'], x, y)
            if row['id'] == 'gamma' and cell['column_id'] == 'context':
                e.text = None
                for line, baseline in [('First cue. ', 468), ('Retain basis.', 502)]:
                    ET.SubElement(e, 'tspan', {'x': str(x), 'y': str(baseline)}).text = line
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def compile_fixture(tmp_path, *, name='neutral', **options):
    page = page_fixture()
    path = tmp_path / (name + '.svg')
    path.write_bytes(svg_fixture(page, **options))
    compiled = compile_deck([SvgInput(PID, path, page['speaker_notes'])],
                            CompileOptions(fonts=host_fonts()), tmp_path / (name + '-native'))
    source = parse_svg(path.read_bytes(), page_id=PID)
    return page, path, compiled, source


def report_for(path, source, page, output):
    report = readback(path, [source], [page])
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    return report


def native_texts(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read('ppt/slides/slide1.xml'))
    return [e.text or '' for e in root.findall('.//a:t', NS)]


def test_mixed_alignment_multiline_has_all_native_carriers_and_passes(tmp_path):
    page, _, compiled, source = compile_fixture(tmp_path)
    before = deepcopy(page)
    pres = Presentation(compiled.pptx_path)
    multiline = [s for s in pres.slides[0].shapes if s.name == atom('gamma', 'context')]
    assert len(multiline) == 2 and multiline[0].top != multiline[1].top
    report = report_for(compiled.pptx_path, source, page, tmp_path / 'readback.json')
    assert report['status'] == 'pass', report
    assert page == before and report['visual_review'] == report['desktop_editing'] == 'not_evaluated'


@pytest.mark.parametrize('transform', ['translate(12 24)', 'translate(28 0) scale(0.94)'])
def test_only_moving_complete_table_relationship_layout_remains_valid(tmp_path, transform):
    page, _, original, _ = compile_fixture(tmp_path, name='original')
    same_page, _, moved, source = compile_fixture(tmp_path, name='moved', transform=transform)
    assert page == same_page
    assert native_texts(original.pptx_path) == native_texts(moved.pptx_path)
    assert report_for(moved.pptx_path, source, same_page, tmp_path / 'moved-readback.json')['status'] == 'pass'


def corrupt_native(original, path, mutation):
    deck = Presentation(original)
    slide = deck.slides[0]
    def named(name): return [s for s in slide.shapes if s.name == name]
    target = named(atom('alpha', 'action'))[0]
    if mutation == 'wrong_row':
        target.top = named(atom('beta', 'action'))[0].top
    elif mutation == 'wrong_column':
        target.left = named(atom('alpha', 'context'))[0].left
    elif mutation == 'same_text_swap':
        left, right = named(atom('alpha', 'context'))[0], named(atom('beta', 'context'))[0]
        left.top, right.top = right.top, left.top
    elif mutation in ('first_line_wrong_row', 'last_line_wrong_column'):
        selected = named(atom('gamma', 'context'))[0 if mutation.startswith('first') else 1]
        if mutation.startswith('first'):
            selected.top = named(atom('beta', 'context'))[0].top
        else:
            selected.left = target.left
    elif mutation in ('missing_carrier', 'extra_carrier'):
        selected = named(atom('gamma', 'context'))[0]
        if mutation == 'missing_carrier':
            slide.shapes._spTree.remove(selected._element)
        else:
            extra = deepcopy(selected._element)
            extra.find('p:nvSpPr/p:cNvPr', NS).set('id', '9999')
            slide.shapes._spTree.append(extra)
    elif mutation == 'unknown_geometry':
        selected = named(atom('gamma', 'context'))[0]
        selected._element.spPr.remove(selected._element.spPr.xfrm)
    elif mutation == 'missing_header':
        slide.shapes._spTree.remove(named(f'atom:{PID}:column:ledger:context:label')[0]._element)
    elif mutation == 'missing_grid_line':
        slide.shapes._spTree.remove(named('column-rule-1')[0]._element)
    elif mutation == 'rotated_carrier':
        target.rotation = 90
    else:
        raise AssertionError(mutation)
    deck.save(path)


@pytest.mark.parametrize('mutation', ['wrong_row', 'wrong_column', 'same_text_swap',
    'first_line_wrong_row', 'last_line_wrong_column', 'missing_carrier', 'extra_carrier',
    'unknown_geometry', 'missing_header', 'missing_grid_line', 'rotated_carrier'])
def test_bad_or_insufficient_native_ownership_never_passes(tmp_path, mutation):
    page, _, compiled, source = compile_fixture(tmp_path)
    corrupted = tmp_path / (mutation + '.pptx')
    corrupt_native(compiled.pptx_path, corrupted, mutation)
    report = report_for(corrupted, source, page, tmp_path / (mutation + '-readback.json'))
    assert report['status'] == 'fail', report
    assert any(f['code'] in ('table_cell_ownership_mismatch', 'unverifiable_table_ownership') for f in report['findings']), report
    if mutation in ('first_line_wrong_row', 'last_line_wrong_column'):
        finding = next(f for f in report['findings'] if f.get('atom_id') == atom('gamma', 'context'))
        carriers = finding['native_carriers']
        actual = [s for s in Presentation(corrupted).slides[0].shapes if s.name == atom('gamma', 'context')]
        assert len(carriers) == len(actual) == len({c['native_shape_id'] for c in carriers}) == 2
        assert [(c['x'], c['y']) for c in carriers] == [(s.left / 9525, s.top / 9525) for s in actual]
    if mutation in ('missing_carrier', 'extra_carrier', 'unknown_geometry'):
        finding = next(f for f in report['findings'] if f.get('atom_id') == atom('gamma', 'context')
                       and f['code'] == 'unverifiable_table_ownership')
        assert finding['code'] == 'unverifiable_table_ownership'
        assert finding['expected_carrier_count'] == 2
        assert finding['actual_carrier_count'] == len(finding['native_shape_ids']) == {
            'missing_carrier': 1, 'extra_carrier': 3, 'unknown_geometry': 2}[mutation]
    if mutation in ('wrong_row', 'wrong_column', 'same_text_swap', 'first_line_wrong_row', 'last_line_wrong_column', 'missing_grid_line'):
        assert native_texts(compiled.pptx_path) == native_texts(corrupted)


@pytest.mark.parametrize('mutation', ['row', 'column'])
def test_same_text_wrong_source_cell_is_located_after_real_compile(tmp_path, mutation):
    page, _, compiled, source = compile_fixture(tmp_path, wrong_source=mutation)
    report = report_for(compiled.pptx_path, source, page, tmp_path / 'wrong-source-readback.json')
    assert any(f['code'] == 'table_cell_ownership_mismatch' and f.get('atom_id') == atom('alpha', 'action') for f in report['findings'])


def test_actual_mixed_alignment_and_same_text_swap_renders_are_retained(tmp_path):
    page, svg, compiled, source = compile_fixture(tmp_path)
    corrupted = tmp_path / 'same-text-swap.pptx'
    corrupt_native(compiled.pptx_path, corrupted, 'same_text_swap')
    svg_png = tmp_path / 'source.svg.png'
    run([executable('rsvg-convert'), '-o', str(svg_png), str(svg)])
    good = render_deck(compiled.pptx_path, tmp_path / 'good-render', fonts=host_fonts())
    bad = render_deck(corrupted, tmp_path / 'bad-render', fonts=host_fonts())
    assert len(good) == len(bad) == 1
    # Text strings repeat; moving differently owned but visually equal text can
    # leave the rendered pixels equal. Native ownership must still reject it.
    for p in [svg_png, good[0], bad[0]]:
        with Image.open(p) as image:
            assert image.width >= 1280 and image.height >= 720
    assert report_for(compiled.pptx_path, source, page, tmp_path / 'good.json')['status'] == 'pass'
    assert report_for(corrupted, source, page, tmp_path / 'bad.json')['status'] == 'fail'
