"""Synthetic historical PPTX, also used by the Rust subprocess tests."""

import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

from PIL import Image, PngImagePlugin
from pptx import Presentation
from pptx.util import Inches

from deck_tools.compiler.native import notes_records

PRIVATE_PATH = "/Users/" + "synthetic-author/private/source.md"


def make_pptx(path, variant="metadata"):
    deck = Presentation()
    deck.core_properties.author = PRIVATE_PATH
    image = io.BytesIO()
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Source", PRIVATE_PATH)
    Image.new("RGB", (12, 12), "navy").save(image, format="PNG", pnginfo=metadata)
    for index in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        # The second historical slide has no stable page ID.
        slide.name = "historical-1" if index == 1 else ""
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.text = f"Historical slide {index}"
        props = shape._element.find(".//{*}cNvPr")
        props.set("name", f"atom:p{index}:title")
        props.set("descr", PRIVATE_PATH)
        props.set("title", PRIVATE_PATH if variant == "private-title" and index == 2 else "Title")
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "10"
        table.cell(0, 1).text = "20"
        slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(6), Inches(1))
        notes = f"Talk {index}\n\nKeep notes.\n"
        if variant == "private-notes" and index == 2:
            notes = f"Talk {index}\n\nSource: {PRIVATE_PATH}\n"
        slide.notes_slide.notes_text_frame.text = notes
    deck.save(path)


def inspect_pptx(path):
    deck = Presentation(path)
    slides = []
    for slide in deck.slides:
        slides.append({
            "text": [shape.text for shape in slide.shapes if shape.has_text_frame],
            "tables": [[[cell.text for cell in row.cells] for row in shape.table.rows]
                       for shape in slide.shapes if shape.has_table],
            "images": [hashlib.sha256(Image.open(io.BytesIO(shape.image.blob)).tobytes()).hexdigest()
                       for shape in slide.shapes if shape.shape_type == 13],
        })
    with zipfile.ZipFile(path) as archive:
        payload = b"".join(archive.read(name) for name in archive.namelist())
    return {"slides": slides, "notes": notes_records(Path(path)),
            "private": PRIVATE_PATH.encode() in payload, "atoms": b"atom:" in payload}


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make_pptx(sys.argv[2], sys.argv[3])
    else:
        print(json.dumps(inspect_pptx(sys.argv[2])))
