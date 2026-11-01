"""Actual native compile/render proof for tight tables; no model effect claim."""
import xml.etree.ElementTree as ET
import subprocess

from PIL import ImageFont
from pptx import Presentation
import pytest

from deck_tools.compiler import CompileOptions, SvgInput, compile_deck
from deck_tools.compiler.svg import parse_svg
from deck_tools.render import readback, render_deck
from hostenv import host_fonts, resolve_host_font
from test_native_table_geometry import PID, atom, page_fixture, svg_fixture

pytestmark = pytest.mark.render


def rendered_words(png):
    data = subprocess.run(['pdftotext','-bbox',str(png.parent/'deck.pdf'),'-'],
                          check=True,capture_output=True).stdout
    return [(word.text, [float(word.get(key)) for key in ('xMin','yMin','xMax','yMax')])
            for word in ET.fromstring(data).iter('{http://www.w3.org/1999/xhtml}word')]


@pytest.mark.parametrize('anchor', ['start', 'end'])
@pytest.mark.parametrize('multiline', [False, True])
def test_column_boundary_anchor_trims_only_empty_slack(tmp_path, anchor, multiline):
    page = page_fixture()
    root = ET.fromstring(svg_fixture(page))
    identity = atom('gamma', 'context') if multiline else atom('alpha', 'action')
    text = next(e for e in root.iter() if e.get('data-atom-id') == identity)
    font = ImageFont.truetype(str(host_fonts()[resolve_host_font()]), size=1000)
    left, right = (700, 1240) if multiline else (260, 700)
    # Fill the cell with visible glyphs at the ordinary body size; only the
    # emitter's extra 1em crosses the rule. No platform-specific font required.
    lines = list(text) if multiline else [text]
    value = 'boundaryglyph ' * 20
    while font.getlength(value)/1000*24 > right-left-10:
        value = value[:-1]
    value = value.rstrip()
    row, column = ('gamma', 'context') if multiline else ('alpha', 'action')
    table = page['customer_visible']['body_blocks'][0]
    cell = next(cell for item in table['rows'] if item['id'] == row
                for cell in item['cells'] if cell['column_id'] == column)
    cell['display_text'] = ' '.join([value]*len(lines))
    text.set('text-anchor', anchor)
    for index, line in enumerate(lines):
        line.text = value + (' ' if index < len(lines)-1 else '')
        line.set('x', str(left if anchor == 'start' else right))
    path = tmp_path/'boundary.svg'; path.write_bytes(ET.tostring(root))
    source = parse_svg(path.read_bytes(), page_id=PID)
    compiled = compile_deck([SvgInput(PID, path, page['speaker_notes'])],
                            CompileOptions(fonts=host_fonts()), tmp_path/'boundary')
    report = readback(compiled.pptx_path, [source], [page])
    assert report['status'] == 'pass', report
    carriers = [shape for shape in Presentation(compiled.pptx_path).slides[0].shapes
                if shape.name == identity]
    assert len(carriers) == len(lines)
    for carrier in carriers:
        position = carrier.left if anchor == 'start' else carrier.left+carrier.width
        assert position == round((left if anchor == 'start' else right)*9525)
        assert carrier.left >= round(left*9525)
        assert carrier.left+carrier.width <= round(right*9525)
    rendered = render_deck(compiled.pptx_path, tmp_path/'boundary-render', fonts=host_fonts())
    expected = {word for line in lines for word in line.text.split()}
    selected = [box for word, box in rendered_words(rendered[0]) if word in expected]
    assert selected
    assert min(box[0] for box in selected) >= left*.75-.075
    assert max(box[2] for box in selected) <= right*.75+.075


def test_table_can_share_slide_with_unbound_caption_text(tmp_path):
    page = page_fixture()
    root = ET.fromstring(svg_fixture(page))
    ET.SubElement(root, 'text', {'id':'unbound-caption','x':'80','y':'140','font-size':'20',
        'font-family':resolve_host_font(),'fill':'#20242a'}).text = 'Additional caption outside the table'
    path = tmp_path/'caption.svg'; path.write_bytes(ET.tostring(root))
    source = parse_svg(path.read_bytes(), page_id=PID)
    assert next(shape for shape in source['shapes'] if shape['kind']=='text'
                and shape['text']=='Additional caption outside the table')['atom_id'] is None
    compiled = compile_deck([SvgInput(PID,path,page['speaker_notes'])],
                            CompileOptions(fonts=host_fonts()),tmp_path/'compiled')
    assert readback(compiled.pptx_path,[source],[page])['status']=='pass'
    assert any(shape.text=='Additional caption outside the table'
               for shape in Presentation(compiled.pptx_path).slides[0].shapes if shape.has_text_frame)


def tight_table(tmp_path, anchor, gap, *, multiline=False, name='tight'):
    page = page_fixture()
    root = ET.fromstring(svg_fixture(page))
    identity = atom('gamma', 'context') if multiline else atom('alpha', 'action')
    text = next(e for e in root.iter() if e.get('data-atom-id') == identity)
    font = ImageFont.truetype(str(host_fonts()[resolve_host_font()]), size=1000)
    left, right = (700, 1240) if multiline else (260, 700)
    for line in list(text) if multiline else [text]:
        # Place visible glyphs, not trailing separator whitespace, at the
        # intended distance from the rule. Keep the source text unchanged.
        width = font.getlength(line.text.rstrip()) / 1000 * 24
        x = {'start': right-width-gap, 'middle': right-width/2-gap, 'end': left+width+gap}[anchor]
        line.set('x', str(x))
    text.set('text-anchor', anchor)
    path = tmp_path / (name + '.svg'); path.write_bytes(ET.tostring(root))
    compiled = compile_deck([SvgInput(PID, path, page['speaker_notes'])],
                            CompileOptions(fonts=host_fonts()), tmp_path / name)
    return page, parse_svg(path.read_bytes(), page_id=PID), compiled, identity


@pytest.mark.parametrize('anchor', ['start', 'middle', 'end'])
@pytest.mark.parametrize('multiline', [False, True])
def test_only_empty_slack_is_trimmed_with_alignment_anchor_preserved(tmp_path, anchor, multiline, monkeypatch):
    from deck_tools.compiler import native
    original_trim = native._trim_table_text_slack
    monkeypatch.setattr(native, '_trim_table_text_slack', lambda *args: None)
    page, source, original, identity = tight_table(tmp_path, anchor, 10, multiline=multiline, name='untrimmed')
    before = readback(original.pptx_path, [source], [page])
    assert any(item['code'] == 'table_cell_ownership_mismatch' for item in before['findings'])
    monkeypatch.setattr(native, '_trim_table_text_slack', original_trim)
    page, source, compiled, _ = tight_table(tmp_path, anchor, 10, multiline=multiline, name='trimmed')
    after = readback(compiled.pptx_path, [source], [page])
    assert after['status'] == 'pass', after
    old = [shape for shape in Presentation(original.pptx_path).slides[0].shapes if shape.name == identity]
    new = [shape for shape in Presentation(compiled.pptx_path).slides[0].shapes if shape.name == identity]
    assert len(old) == len(new) == (2 if multiline else 1)
    fraction = {'start': 0, 'middle': .5, 'end': 1}[anchor]
    for a, b in zip(old, new):
        assert b.width < a.width
        assert abs((a.left + fraction*a.width) - (b.left + fraction*b.width)) <= 1
        assert (a.top, a.height, a.text) == (b.top, b.height, b.text)
    # Both are genuinely rendered. Impress quantizes text layout to .01 mm;
    # verify actual word positions as well as exact native alignment anchors.
    first = render_deck(original.pptx_path, tmp_path / 'original-render', fonts=host_fonts())
    second = render_deck(compiled.pptx_path, tmp_path / 'trimmed-render', fonts=host_fonts())
    a, b = rendered_words(first[0]), rendered_words(second[0])
    assert len(a) == len(b)
    for (text, old), (current, new) in zip(a, b):
        assert text == current
        assert max(abs(x-y) for x, y in zip(old, new)) <= .075  # 0.1 px, renderer quantization only


@pytest.mark.parametrize('anchor', ['start', 'middle', 'end'])
@pytest.mark.parametrize('multiline', [False, True])
def test_true_glyph_overflow_is_still_refused(tmp_path, anchor, multiline):
    page, source, compiled, identity = tight_table(tmp_path, anchor, -6, multiline=multiline)
    report = readback(compiled.pptx_path, [source], [page])
    assert report['status'] == 'fail'
    assert any(item['code'] == 'table_cell_ownership_mismatch' and item['atom_id'] == identity
               for item in report['findings'])
    rendered = render_deck(compiled.pptx_path, tmp_path/'overflow-render', fonts=host_fonts())
    selected = [box for text, box in rendered_words(rendered[0])
                if text in (('First','cue.') if multiline else ('Wait','10','ms'))]
    assert selected
    left, right = (700,1240) if multiline else (260,700)
    if anchor == 'end':
        assert min(box[0] for box in selected) < left*.75
    else:
        assert max(box[2] for box in selected) > right*.75
