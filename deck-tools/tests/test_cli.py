import json
from pathlib import Path
import subprocess
import sys

import pytest

from deck_tools.sources import read_source


def invoke(tmp_path, path):
    job = tmp_path / "job.json"
    job.write_text(json.dumps({"path": str(path)}), encoding="utf-8")
    return subprocess.run(
        [sys.executable, "-m", "deck_tools", "read-source", str(job)],
        capture_output=True, text=True, check=False,
    )


@pytest.mark.parametrize("suffix", ["txt", "md", "json"])
def test_subprocess_keeps_text_locators_and_zero(tmp_path, suffix):
    text = '{"a/b~c":0}' if suffix == "json" else "合成材料\n末尾约束\n"
    source = tmp_path / f"source.{suffix}"
    source.write_text(text, encoding="utf-8")
    output = invoke(tmp_path, source)
    assert output.returncode == 0, output.stderr
    data = json.loads(output.stdout)
    assert data["status"] == "read"
    assert data["locators"]
    if suffix == "json":
        assert data["locators"][0] == {"locator": "/a~1b~0c", "kind": "value", "text": "0"}
    else:
        assert data["text"] == text
    assert source.read_text(encoding="utf-8") == text


@pytest.mark.parametrize("suffix,data", [("json", b"{broken"), ("txt", b"\xff"), ("docx", b"not a zip"), ("docx", b"PK")])
def test_reader_failure_has_no_success_json(tmp_path, suffix, data):
    source = tmp_path / f"source.{suffix}"
    source.write_bytes(data)
    output = invoke(tmp_path, source)
    assert output.returncode == 2
    assert not output.stdout
    assert output.stderr


def test_pptx_table_and_notes_are_read_without_project_kernel(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(4), Inches(1)).table
    table.cell(0, 0).text = "合成项"
    table.cell(0, 1).text = "0"
    slide.notes_slide.notes_text_frame.text = "合成讲稿"
    source = tmp_path / "source.pptx"
    presentation.save(source)
    output = invoke(tmp_path, source)
    assert output.returncode == 0, output.stderr
    result = json.loads(output.stdout)
    assert "合成项" in result["text"] and "合成讲稿" in result["text"]
    assert result["locators"][0]["locator"] == "slide-1"


def test_missing_pdf_dependency_preserves_needs_tool(tmp_path, monkeypatch):
    monkeypatch.setattr("deck_tools.sources.shutil.which", lambda name: None)
    source = tmp_path / "source.pdf"
    source.write_bytes(b"synthetic pdf input")
    result = read_source(source)
    assert result.status == "needs_tool"
    assert "pdftotext" in result.detail


def test_image_reports_visual_work_and_unknown_format_reports_missing_reader(tmp_path):
    source = tmp_path / "source.svg"
    source.write_text('<svg xmlns="http://www.w3.org/2000/svg"/>', encoding="utf-8")
    assert json.loads(invoke(tmp_path, source).stdout)["status"] == "pending_visual"
    unknown = tmp_path / "source.xyz"
    unknown.write_text("synthetic unknown format", encoding="utf-8")
    assert json.loads(invoke(tmp_path, unknown).stdout)["status"] == "needs_tool"


def test_non_object_job_is_input_error(tmp_path):
    job = tmp_path / "job.json"
    job.write_text("null", encoding="utf-8")
    output = subprocess.run([sys.executable, "-m", "deck_tools", "read-source", str(job)], capture_output=True, text=True, check=False)
    assert output.returncode == 2 and not output.stdout
