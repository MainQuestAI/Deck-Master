"""Native cells must agree with content, even when the SVG is self-consistent."""

import json

from pptx import Presentation
import pytest

from deck_tools.compiler import CompileOptions, SvgInput, compile_deck
from deck_tools.render import readback
from hostenv import host_fonts, resolve_host_font


def table_case(tmp_path, variant="correct", count=1):
    blocks = []
    markup = ['<text x="30" y="40" data-atom-id="atom:p1:title">Comparison</text>']
    for index in range(count):
        table_id = f"t{index}"
        left = 30 + 600 * index
        blocks.append({"id": table_id, "type": "table", "columns": [
            {"id": "name", "label": "Option"}, {"id": "value", "label": "Value"}],
            "rows": [{"id": row, "cells": [{"column_id": "name", "display_text": row.upper()},
                     {"column_id": "value", "display_text": value}]} for row, value in
                     [("a", "" if variant == "empty" else "10"), ("b", "10" if variant == "rich-repeat" else "20")]]})
        markup.append(f'<rect x="{left}" y="90" width="500" height="240" fill="none" stroke="#333" data-table-id="{table_id}" data-native-table="true"/>')
        markup.extend(f'<line x1="{left}" y1="{y}" x2="{left + 500}" y2="{y}" stroke="#333" data-table-id="{table_id}"/>' for y in (160, 240))
        markup.append(f'<line x1="{left + 270}" y1="90" x2="{left + 270}" y2="330" stroke="#333" data-table-id="{table_id}"/>')
        for column, x, label in [("name", left + 20, "Option"), ("value", left + 290, "Value")]:
            markup.append(f'<text x="{x}" y="130" data-table-id="{table_id}" data-atom-id="atom:p1:column:{table_id}:{column}:label">{label}</text>')
        for row, y, value in [("a", 200, "10"), ("b", 280, "20")]:
            markup.append(f'<text x="{left + 20}" y="{y}" data-table-id="{table_id}" data-atom-id="atom:p1:cell:{table_id}:{row}:name:display_text">{row.upper()}</text>')
            if variant == "empty" and row == "a":
                continue
            bound_row = ("b" if row == "a" else "a") if variant == "wrong-row" else row
            binding = f'data-atom-id="atom:p1:cell:{table_id}:{bound_row}:value:display_text"'
            if variant == "missing-binding" and row == "a":
                binding = ""
            if variant == "swapped":
                value = "20" if row == "a" else "10"
            if variant == "rich-repeat":
                value = f'<tspan x="{left + 290}" y="{y}" font-weight="700">1</tspan><tspan x="{left + 305}" y="{y}">0</tspan>'
            markup.append(f'<text x="{left + 290}" y="{y}" data-table-id="{table_id}" {binding}>{value}</text>')
    svg = tmp_path / "table.svg"
    svg.write_text(f'<svg viewBox="0 0 {600 * count} 400" font-family="{resolve_host_font()}" font-size="24">' + ''.join(markup) + '</svg>')
    result = compile_deck([SvgInput("p1", svg)], CompileOptions(width_px=600 * count, height_px=400,
                                                               fonts=host_fonts()), tmp_path / "compiled")
    sources = json.loads(result.manifest_path.read_text())["pages"]
    page = {"page_id": "p1", "customer_visible": {"title": "Comparison", "body_blocks": blocks}}
    return result.pptx_path, sources, [page]


@pytest.mark.parametrize("variant,code,rows", [
    ("swapped", "table_cell_text_mismatch", {"a", "b"}),
    ("wrong-row", "table_cell_ownership_mismatch", {"a", "b"}),
    ("missing-binding", "unverifiable_table_ownership", {"a"}),
])
def test_native_content_reports_cells_not_global_word_presence(tmp_path, variant, code, rows):
    args = table_case(tmp_path, variant)
    findings = readback(*args)["findings"]
    assert len(findings) == len(rows), findings
    assert {item["row_id"] for item in findings} == rows
    assert all(item["code"] == code and item["table_id"] == "t0" and item["column_id"] == "value"
               and item["page_id"] == "p1" for item in findings)


def test_multiple_tables_and_rich_runs_preserve_identical_cell_text(tmp_path):
    assert readback(*table_case(tmp_path, "rich-repeat", count=2))["status"] == "pass"


def test_empty_content_cell_needs_no_text_carrier(tmp_path):
    assert readback(*table_case(tmp_path, "empty"))["status"] == "pass"


def test_ppt_cell_edit_is_located_once_and_svg_fidelity_stays_checked(tmp_path):
    path, sources, pages = table_case(tmp_path)
    assert readback(path, sources, pages)["status"] == "pass"
    deck = Presentation(path)
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    table.cell(1, 1).text = "99"
    changed = tmp_path / "changed.pptx"
    deck.save(changed)
    findings = readback(changed, sources, pages)["findings"]
    assert len(findings) == 1, findings
    assert findings[0]["row_id"] == "a" and findings[0]["column_id"] == "value"
    # A table not declared by content still gets SVG-to-PPT fidelity checking.
    pages[0]["customer_visible"]["body_blocks"] = []
    assert any(item["code"] == "svg_ppt_text_mismatch" for item in readback(changed, sources, pages)["findings"])
