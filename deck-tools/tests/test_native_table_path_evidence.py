"""S10-R1-F01: actual editable paths, never stale transform-box proof.

The neutral partial rule is the same in-viewport quarter-span mutation as the
independent R1 reproducer. All atom/name/ID/text/notes/xfrm remain unchanged.
"""
from copy import deepcopy
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
import pytest

from deck_tools.compiler import CompileOptions, SvgInput, compile_deck
from deck_tools.compiler.svg import parse_svg
from deck_tools.render import readback, render_deck
from hostenv import host_fonts
from test_native_table_geometry import compile_fixture, page_fixture, svg_fixture, NS, PID

pytestmark = pytest.mark.render


def change_paths(original, target, mutation):
    deck = Presentation(original)
    slide = deck.slides[0]
    rule = next(s for s in slide.shapes if s.name == 'column-rule-1')
    frame = next(s for s in slide.shapes if s.name == 'ledger-frame')
    selected = frame if mutation.startswith('frame_') else rule
    path = selected._element.find('p:spPr/a:custGeom/a:pathLst/a:path', NS)
    assert path is not None
    if mutation.startswith('frame_clip_'):
        w, h = int(path.get('w')), int(path.get('h'))
        if mutation == 'frame_clip_bottom_left':
            points = [(0, 0), (w, 0), (w, h), (220*9525, h), (0, 270*9525)]
        elif mutation == 'frame_clip_top_left':
            points = [(220*9525, 0), (w, 0), (w, h), (0, h), (0, 90*9525)]
        elif mutation == 'frame_clip_bottom_right':
            points = [(0, 0), (w, 0), (w, 270*9525), (660*9525, h), (0, h)]
        elif mutation == 'frame_clip_top_right':
            points = [(0, 0), (660*9525, 0), (w, 90*9525), (w, h), (0, h)]
        else:
            assert mutation in ('frame_clip_safe', 'frame_clip_safe_reversed')
            d = 20*9525
            points = [(d, 0), (w-d, 0), (w, d), (w, h-d), (w-d, h), (d, h), (0, h-d), (0, d)]
            if mutation.endswith('_reversed'):
                points.reverse()
        for child in list(path):
            path.remove(child)
        for i, (x, y) in enumerate(points):
            command = OxmlElement('a:moveTo' if i == 0 else 'a:lnTo')
            point = OxmlElement('a:pt'); point.set('x', str(x)); point.set('y', str(y))
            command.append(point); path.append(command)
        path.append(OxmlElement('a:close'))
    elif mutation in ('partial_rule', 'frame_partial'):
        if mutation == 'partial_rule':
            path.find('a:lnTo/a:pt', NS).set('y', str(int(path.get('h')) // 4))
        else:
            for pt in path.findall('.//a:pt', NS):
                pt.set('y', str(int(pt.get('y')) // 4))
    elif mutation in ('empty_rule', 'frame_empty'):
        for child in list(path):
            path.remove(child)
    elif mutation == 'collapsed_rule':
        end = path.find('a:lnTo/a:pt', NS)
        end.set('x', '0'); end.set('y', '0')
    elif mutation == 'bent_rule':
        end = path.find('a:lnTo/a:pt', NS)
        elbow = deepcopy(path.find('a:lnTo', NS))
        elbow[0].set('x', path.get('w')); elbow[0].set('y', str(int(path.get('h')) // 2))
        path.insert(1, elbow)
        assert end.get('y') == path.get('h')
    elif mutation == 'unsupported_curve':
        line = path.find('a:lnTo', NS)
        path.remove(line)
        curve = OxmlElement('a:cubicBezTo')
        path.append(curve)
        for fraction in (1, 2, 3):
            point = OxmlElement('a:pt')
            point.set('x', '0'); point.set('y', str(int(path.get('h')) * fraction // 3))
            curve.append(point)
    elif mutation == 'stroke_disabled':
        path.set('stroke', '0')
    elif mutation == 'frame_diamond':
        for child in list(path):
            path.remove(child)
        w, h = int(path.get('w')), int(path.get('h'))
        for i, (x, y) in enumerate(((w//2, 0), (w, h//2), (w//2, h), (0, h//2))):
            command = OxmlElement('a:moveTo' if i == 0 else 'a:lnTo')
            point = OxmlElement('a:pt'); point.set('x', str(x)); point.set('y', str(y))
            command.append(point); path.append(command)
        path.append(OxmlElement('a:close'))
    elif mutation == 'ambiguous_paths':
        path.getparent().append(deepcopy(path))
    elif mutation == 'frame_open':
        path.remove(path.find('a:close', NS))
    elif mutation == 'invisible_rule':
        line = rule._element.find('p:spPr/a:ln', NS)
        for child in list(line):
            if child.tag.rsplit('}', 1)[-1] in ('solidFill', 'gradFill', 'noFill'):
                line.remove(child)
        line.append(OxmlElement('a:noFill'))
    elif mutation in ('scaled_viewports', 'flipped_geometry'):
        for shape in slide.shapes:
            path = shape._element.find('p:spPr/a:custGeom/a:pathLst/a:path', NS)
            if path is None:
                continue
            if mutation == 'scaled_viewports':
                path.set('w', str(int(path.get('w')) * 3))
                path.set('h', str(int(path.get('h')) * 2))
                for pt in path.findall('.//a:pt', NS):
                    pt.set('x', str(int(pt.get('x')) * 3)); pt.set('y', str(int(pt.get('y')) * 2))
            else:
                transform = shape._element.find('p:spPr/a:xfrm', NS)
                transform.set('flipH', '1'); transform.set('flipV', 'true'); transform.set('rot', '10800000')
    else:
        raise AssertionError(mutation)
    deck.save(target)


def native_invariants(path):
    with zipfile.ZipFile(path) as archive:
        slide = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
        return {'text': [e.text for e in slide.findall('.//a:t', NS)],
                'identities': [ET.tostring(e) for e in slide.findall('.//p:cNvPr', NS)],
                'transforms': [ET.tostring(e) for e in slide.findall('.//a:xfrm', NS)],
                'notes': archive.read('ppt/notesSlides/notesSlide1.xml')}


@pytest.mark.parametrize('mutation', ['partial_rule', 'empty_rule', 'collapsed_rule',
    'frame_partial', 'frame_empty', 'frame_open', 'bent_rule', 'unsupported_curve',
    'ambiguous_paths', 'invisible_rule', 'stroke_disabled', 'frame_diamond'])
def test_native_grid_path_evidence_cannot_be_replaced_by_unchanged_xfrm(tmp_path, mutation):
    page, _, compiled, source = compile_fixture(tmp_path)
    target = tmp_path / (mutation + '.pptx')
    change_paths(compiled.pptx_path, target, mutation)
    assert native_invariants(compiled.pptx_path) == native_invariants(target)
    report = readback(target, [source], [page])
    assert report['status'] == 'fail', report
    assert any(f['code'] == 'unverifiable_table_ownership' and f.get('table_id') == 'ledger'
               for f in report['findings']), report


@pytest.mark.parametrize('mutation', ['scaled_viewports', 'flipped_geometry'])
def test_supported_native_path_scale_and_transform_preserve_real_grid(tmp_path, mutation):
    page, _, compiled, source = compile_fixture(tmp_path)
    target = tmp_path / (mutation + '.pptx')
    change_paths(compiled.pptx_path, target, mutation)
    if mutation == 'scaled_viewports':
        assert native_invariants(compiled.pptx_path) == native_invariants(target)
    report = readback(target, [source], [page])
    assert report['status'] == 'pass', report


def test_actual_unrounded_rect_frame_remains_supported(tmp_path):
    page = page_fixture()
    svg = ET.fromstring(svg_fixture(page))
    next(e for e in svg.iter() if e.get('id') == 'ledger-frame').attrib.pop('rx')
    path = tmp_path / 'plain-frame.svg'; path.write_bytes(ET.tostring(svg))
    compiled = compile_deck([SvgInput(PID, path, page['speaker_notes'])],
        CompileOptions(fonts=host_fonts()), tmp_path / 'plain-native')
    assert readback(compiled.pptx_path, [parse_svg(path.read_bytes(), page_id=PID)], [page])['status'] == 'pass'


def test_valid_header_only_rule_renders_incomplete_grid_and_is_unverified(tmp_path):
    page, _, compiled, source = compile_fixture(tmp_path)
    target = tmp_path / 'header-only-native-rule.pptx'
    change_paths(compiled.pptx_path, target, 'partial_rule')
    assert native_invariants(compiled.pptx_path) == native_invariants(target)
    good = render_deck(compiled.pptx_path, tmp_path / 'full-grid-render', fonts=host_fonts())
    bad = render_deck(target, tmp_path / 'partial-grid-render', fonts=host_fonts())
    assert len(good) == len(bad) == 1
    with Image.open(good[0]) as a, Image.open(bad[0]) as b:
        assert a.size == b.size and a.convert('RGB').tobytes() != b.convert('RGB').tobytes()
    assert readback(compiled.pptx_path, [source], [page])['status'] == 'pass'
    report = readback(target, [source], [page])
    assert report['status'] == 'fail' and any(f['code'] == 'unverifiable_table_ownership' for f in report['findings'])


def test_native_group_transform_is_not_flat_grid_evidence(tmp_path):
    page, _, compiled, source = compile_fixture(tmp_path)
    deck = Presentation(compiled.pptx_path)
    slide = deck.slides[0]
    rule = next(s for s in slide.shapes if s.name == 'column-rule-1')
    original = ET.tostring(rule._element)
    group = slide.shapes.add_group_shape([rule])
    group.left += 95250  # legal enclosing transform, local rule xfrm unchanged
    assert ET.tostring(rule._element) == original
    target = tmp_path / 'grouped-native-rule.pptx'; deck.save(target)
    assert native_invariants(target)['text'] == native_invariants(compiled.pptx_path)['text']
    report = readback(target, [source], [page])
    assert report['status'] == 'fail' and any(f['code'] == 'unverifiable_table_ownership' for f in report['findings'])


def test_unsupported_native_rule_rotation_is_unverified(tmp_path):
    page, _, compiled, source = compile_fixture(tmp_path)
    deck = Presentation(compiled.pptx_path)
    rule = next(s for s in deck.slides[0].shapes if s.name == 'column-rule-1')
    rule.rotation = 45
    target = tmp_path / 'oblique-native-rule.pptx'; deck.save(target)
    report = readback(target, [source], [page])
    assert report['status'] == 'fail' and any(f['code'] == 'unverifiable_table_ownership' for f in report['findings'])


@pytest.mark.parametrize('environment', ['slide_root_transform', 'shape_3d'])
def test_additional_native_coordinate_environment_is_unverified(tmp_path, environment):
    page, _, compiled, source = compile_fixture(tmp_path)
    deck = Presentation(compiled.pptx_path)
    slide = deck.slides[0]
    if environment == 'slide_root_transform':
        parent = slide._element.find('p:cSld/p:spTree/p:grpSpPr', NS)
        transform = OxmlElement('a:xfrm')
        for name, attrs in (('off', {'x': '95250', 'y': '0'}),
                            ('ext', {'cx': '12192000', 'cy': '6858000'}),
                            ('chOff', {'x': '0', 'y': '0'}),
                            ('chExt', {'cx': '12192000', 'cy': '6858000'})):
            child = OxmlElement('a:' + name)
            for key, value in attrs.items():
                child.set(key, value)
            transform.append(child)
        parent.append(transform)
    else:
        rule = next(s for s in slide.shapes if s.name == 'column-rule-1')
        shape3d = OxmlElement('a:sp3d'); shape3d.set('extrusionH', '190500')
        rule._element.find('p:spPr', NS).append(shape3d)
    target = tmp_path / (environment + '.pptx'); deck.save(target)
    assert len(Presentation(target).slides[0].shapes) == len(slide.shapes)
    assert native_invariants(target)['text'] == native_invariants(compiled.pptx_path)['text']
    report = readback(target, [source], [page])
    assert report['status'] == 'fail' and any(f['code'] == 'unverifiable_table_ownership' for f in report['findings'])


@pytest.mark.parametrize('corner', ['bottom_left', 'top_left', 'bottom_right', 'top_right'])
def test_real_convex_frame_perimeter_bounds_every_cell_and_header_carrier(tmp_path, corner):
    page, _, compiled, source = compile_fixture(tmp_path)
    target = tmp_path / ('clipped-' + corner + '.pptx')
    change_paths(compiled.pptx_path, target, 'frame_clip_' + corner)
    assert native_invariants(compiled.pptx_path) == native_invariants(target)
    report = readback(target, [source], [page])
    assert report['status'] == 'fail', report
    if corner.startswith('top_'):
        assert any(f['code'] == 'unverifiable_table_ownership' and f.get('table_id') == 'ledger'
                   for f in report['findings']), report
    else:
        atom_id = f"atom:{PID}:cell:ledger:gamma:{'key' if corner.endswith('left') else 'context'}:display_text"
        finding = next(f for f in report['findings'] if f.get('atom_id') == atom_id
                       and f['code'] == 'table_cell_ownership_mismatch')
        assert finding['native_frame_perimeter']
        if corner.endswith('right'):
            assert len(finding['native_carriers']) == 2  # include the last native line outside the sloped edge


@pytest.mark.parametrize('mutation', ['frame_clip_safe', 'frame_clip_safe_reversed'])
def test_true_convex_perimeter_with_all_carriers_inside_passes_in_both_windings(tmp_path, mutation):
    page, _, compiled, source = compile_fixture(tmp_path)
    target = tmp_path / (mutation + '.pptx')
    change_paths(compiled.pptx_path, target, mutation)
    assert native_invariants(compiled.pptx_path) == native_invariants(target)
    report = readback(target, [source], [page])
    assert report['status'] == 'pass', report
