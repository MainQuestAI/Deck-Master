"""T10/AC-K01 isolation: compile_deck depends only on the SVG bytes, approved
assets and explicit parameters; input bytes are consistency-checked; a failed
compile never publishes deck.pptx."""
import json
import os
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pytest
from PIL import ImageFont

import deck_tools.compiler.native as native_emit
from deck_tools.compiler import CompileOptions, SvgInput, compile_deck
from hostenv import host_fonts, resolve_host_font
from deck_tools.sanitize import pptx
from deck_tools.render import readback
from pptx import Presentation
import io

REPO_ROOT = Path(__file__).resolve().parents[2]


def test_editable_table_and_title_survive_delivery_cleanup(tmp_path):
    family = resolve_host_font()
    svg = tmp_path / 'editable.svg'
    svg.write_text(f'''<svg viewBox="0 0 600 400" font-family="{family}" font-size="24">
      <text id="标题" x="30" y="40" data-atom-id="atom:p1:title">方案比较</text>
      <text id="流程说明" x="30" y="75">比较后提交</text>
      <rect id="评分项" x="30" y="90" width="500" height="220" fill="none"
            stroke="#333333" data-table-id="评分项" data-native-table="true"/>
      <line x1="30" y1="160" x2="530" y2="160" stroke="#333333" data-table-id="评分项"/>
      <line x1="300" y1="90" x2="300" y2="310" stroke="#333333" data-table-id="评分项"/>
      <text x="50" y="130" data-table-id="评分项">评分项</text>
      <text x="320" y="130" data-table-id="评分项">比较项</text>
      <text x="50" y="200" data-table-id="评分项" data-text-width="230">
        <tspan x="50" y="200" font-weight="700" fill="#AA0000">方案</tspan><tspan x="110" y="200">说明</tspan>
      </text>
      <text x="50" y="250" data-table-id="评分项">补充说明</text>
      <text x="320" y="200" data-table-id="评分项">结果</text>
      <line id="流程连线" x1="30" y1="350" x2="530" y2="350" stroke="#333333"
            stroke-linecap="round" stroke-linejoin="bevel"/>
    </svg>''')
    result = compile_deck([SvgInput('p1', svg)],
                          CompileOptions(width_px=600, height_px=400, fonts=host_fonts()),
                          tmp_path / 'out')
    source = json.loads(result.manifest_path.read_text())['pages']
    assert readback(result.pptx_path, source, [{'page_id': 'p1', 'customer_visible': {'title': '方案比较'}}])['status'] == 'pass'
    changed = Presentation(result.pptx_path)
    next(shape.table for shape in changed.slides[0].shapes if shape.has_table).cell(1, 1).text = '不同结果'
    changed_path = tmp_path / 'changed.pptx'
    changed.save(changed_path)
    findings = readback(changed_path, source, [{'page_id': 'p1', 'customer_visible': {'title': '方案比较'}}])['findings']
    assert any(finding['code'] == 'svg_ppt_text_mismatch' and finding.get('element') == '评分项'
               for finding in findings)
    cleaned = pptx(result.pptx_path.read_bytes(), 'deck.pptx')
    delivered = pptx(cleaned, 'deck.pptx')
    slide = Presentation(io.BytesIO(delivered)).slides[0]
    assert slide.shapes.title.text == '方案比较'
    assert slide.slide_layout.name
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    assert table.first_row
    assert table.cell(1, 0).text == '方案说明\n补充说明'
    runs = table.cell(1, 0).text_frame.paragraphs[0].runs
    assert runs[0].font.bold
    assert str(runs[0].font.color.rgb) == 'AA0000'
    assert not runs[1].font.bold
    assert [shape.name for shape in slide.shapes] == [shape.name for shape in Presentation(io.BytesIO(cleaned)).slides[0].shapes]
    assert '表格-评分项' in [shape.name for shape in slide.shapes]
    assert '连线-流程连线' in [shape.name for shape in slide.shapes]
    assert '正文-流程说明' in [shape.name for shape in slide.shapes]
    ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(io.BytesIO(delivered)) as archive:
        xml = archive.read('ppt/slides/slide1.xml')
    root = ET.fromstring(xml)
    graphic = root.find('.//p:graphicFrame', ns)
    assert graphic is not None
    for props in graphic.findall('.//a:tcPr', ns):
        assert [child.tag.rsplit('}', 1)[-1] for child in props] == [
            'lnL', 'lnR', 'lnT', 'lnB', 'solidFill']
    shape = next(sp for sp in root.findall('.//p:sp', ns)
                 if sp.find('p:nvSpPr/p:cNvPr', ns).get('name') == '连线-流程连线')
    line = shape.find('p:spPr/a:ln', ns)
    assert line.get('cap') == 'rnd'
    assert line.find('a:bevel', ns) is not None


def test_native_table_marker_without_identity_keeps_drawn_shape_readback(tmp_path):
    family = resolve_host_font()
    svg = tmp_path / 'drawn.svg'
    svg.write_text(f'''<svg viewBox="0 0 600 400" font-family="{family}" font-size="24">
      <rect x="30" y="90" width="500" height="220" fill="none"
            stroke="#333333" data-native-table="true"/>
      <text x="50" y="130" data-atom-id="atom:p1:title">方案比较</text>
    </svg>''')
    result = compile_deck([SvgInput('p1', svg)],
                          CompileOptions(width_px=600, height_px=400, fonts=host_fonts()),
                          tmp_path / 'out')
    source = json.loads(result.manifest_path.read_text())['pages']
    assert readback(result.pptx_path, source, [
        {'page_id': 'p1', 'customer_visible': {'title': '方案比较'}}])['status'] == 'pass'


def test_compile_isolated_from_home_cwd_and_repo(monkeypatch, tmp_path):
    # Only SVG bytes + explicit parameters: no HOME/config/env defaults, cwd in
    # an empty directory, and the manifest leaks no repository absolute paths.
    home = tmp_path / 'empty-home'
    home.mkdir()
    for key in ('HOME', 'USERPROFILE', 'XDG_CONFIG_HOME', 'XDG_DATA_HOME',
                'DECK_MASTER_NODE_EXECUTABLE', 'DECK_MASTER_ARTIFACT_MODULE'):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv('HOME', str(home))
    monkeypatch.chdir(tmp_path)
    svg = tmp_path / 'isolated.svg'
    svg.write_text('<svg viewBox="0 0 960 720"><rect width="960" height="720" fill="#102030"/></svg>')
    result = compile_deck([SvgInput('p1', svg)],
                          CompileOptions(width_px=960, height_px=720), tmp_path / 'out')
    assert result.pptx_path.is_file()
    manifest_text = result.manifest_path.read_text()
    assert str(REPO_ROOT) not in manifest_text
    assert json.loads(manifest_text)['pages'][0]['page_id'] == 'p1'


def test_input_bytes_changed_mid_compile_is_rejected(monkeypatch, tmp_path):
    # Compile reads the input up front; if the bytes change during compilation
    # the result is rejected instead of registering a mixed input.
    svg = tmp_path / 'changing.svg'
    svg.write_text('<svg viewBox="0 0 960 720"><rect width="960" height="720"/></svg>')
    original_emit = native_emit.emit

    def mutating_emit(pages, width, height, fonts, output):
        svg.write_text('<svg viewBox="0 0 960 720"><rect width="400" height="200"/></svg>')
        return original_emit(pages, width, height, fonts, output)

    monkeypatch.setattr(native_emit, 'emit', mutating_emit)
    with pytest.raises(ValueError, match='input changed during compilation'):
        compile_deck([SvgInput('p1', svg)], CompileOptions(width_px=960, height_px=720),
                     tmp_path / 'out')
    assert not (tmp_path / 'out' / 'deck.pptx').exists()


def test_failed_compile_publishes_no_pptx(tmp_path):
    # A failed emit (missing required font file) must not leave deck.pptx behind.
    svg = tmp_path / 'text.svg'
    svg.write_text('<svg viewBox="0 0 960 720"><text x="50" y="200" font-family="Noto Sans SC">文字</text></svg>')
    with pytest.raises(ValueError, match='font file not supplied'):
        compile_deck([SvgInput('p1', svg)], CompileOptions(width_px=960, height_px=720, fonts={}),
                     tmp_path / 'out')
    assert not (tmp_path / 'out' / 'deck.pptx').exists()


def test_text_box_slack_does_not_cause_false_slide_overflow(tmp_path):
    family = resolve_host_font()
    fonts = host_fonts()
    size = 16
    glyph_width = ImageFont.truetype(fonts[family], size=1000).getlength('A') / 1000 * size
    x = 200 - glyph_width - 2
    svg = tmp_path / 'near-edge.svg'
    svg.write_text(f'<svg viewBox="0 0 200 100"><text x="{x}" y="50" '
                   f'font-family="{family}" font-size="{size}">A</text></svg>')
    result = compile_deck([SvgInput('p1', svg)],
                          CompileOptions(width_px=200, height_px=100, fonts=fonts),
                          tmp_path / 'near-edge-out')
    ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(result.pptx_path) as archive:
        slide = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
    shape = slide.find('.//p:sp/p:spPr/a:xfrm', ns)
    off = shape.find('a:off', ns)
    ext = shape.find('a:ext', ns)
    assert int(off.get('x')) + int(ext.get('cx')) <= 200 * 9525
