"""W11 browser QA over a synthetic project with a real compiled/rendered PPTX.

Quality records are explicit test inputs, not Host or professional approval.
No model calls. All projects, exports and screenshots remain under --out.
"""

from __future__ import annotations
import argparse
import copy
import hashlib
import html
import io
import json
from pathlib import Path
import uuid
import zipfile
from urllib.parse import urlencode

from playwright.sync_api import sync_playwright, expect
from pptx import Presentation
from deck_master import service, editing, ui_journal
from deck_master.compiler import CompileOptions, SvgInput, compile_deck
from deck_master.compiler.svg import parse_svg
from deck_master.pipeline import artifact, render_deck, resolve_fonts, readback
from deck_master.models import bump_revision
from deck_master.local_state import write_json
from deck_master.samples import FACTORY_VERSION
from deck_master.store import Store
from deck_master.web import WorkbenchServer


def advance(store, description="synthetic background state"):
    doc = store.load_document()
    updated = bump_revision(
        copy.deepcopy(doc), {"operation_id": str(uuid.uuid4()), "kind": "policy_update", "description": description, "read_set": []}
    )
    updated["policy"]["user_stop"] = True
    store.commit_change(base_revision=doc["revision_id"], document=updated, operation_id=updated["change"]["operation_id"])
    return updated


def fixture(root, font="Arial"):
    project = root / "synthetic-project"
    page = {
        "schema_version": "deck_page_package.v2",
        "page_id": "p1",
        "customer_visible": {"title": "W11 file verification", "body_blocks": []},
        "visual_spec": {"intent": "Explicit synthetic export test", "reference_mode": "new_design"},
    }
    plan = {
        "schema_version": "content_plan_input.v1",
        "input_summary": "Synthetic test only",
        "chapters": [{"chapter_id": "c1", "title": "Synthetic", "goal_ids": ["g1"]}],
        "goals": [
            {
                "goal_id": "g1",
                "page_id": "p1",
                "purpose": "Verify fixed downloads",
                "source_links": [],
                "unresolved_facts": ["Synthetic verification fixture; no customer source evidence"],
            }
        ],
        "unresolved_facts": [],
    }
    service.create(
        project,
        brief="Synthetic W11 browser test, no customer facts",
        project_format="workbench.v3",
        draft={"pages": [page], "content_plan": plan},
    )
    write_json(
        project / ".deckmaster/workbench/sample.json",
        {"format": FACTORY_VERSION, "readonly": False, "evidence_level": "synthetic", "model_calls": 0},
    )
    store = Store(project)
    old = store.current_revision_id()
    svg = root / "page.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#f5f5f5"/><text x="120" y="220" font-family="Arial" font-size="44" fill="#154d3e">W11 file verification</text></svg>'.replace(
            "Arial", html.escape(font, quote=True)
        )
    )
    parsed = parse_svg(svg.read_bytes(), page_id="p1", assets={})
    fonts = resolve_fonts([parsed])
    compiled = compile_deck([SvgInput("p1", svg)], CompileOptions(width_px=1280, height_px=720, fonts=fonts), root / "compiled")
    deck = Presentation(compiled.pptx_path)
    deck.core_properties.author = "/Users/synthetic/PRIVATE-CANARY"
    deck.slides[0].notes_slide.notes_text_frame.text = "/Users/synthetic/PRIVATE-CANARY"
    deck.save(compiled.pptx_path)
    images = render_deck(compiled.pptx_path, root / "rendered", fonts=fonts)
    report = readback(compiled.pptx_path, [parsed], [page])
    assert report["status"] == "pass", report
    report_path = root / "readback.json"
    report_path.write_text(json.dumps(report))
    doc = store.load_document()
    updated = bump_revision(
        copy.deepcopy(doc),
        {
            "operation_id": "synthetic-production",
            "kind": "artifact_adoption",
            "description": "Explicit test artifacts; real compiler/render/readback",
            "read_set": [],
        },
    )
    entry = updated["pages"][0]
    entry["svg"] = artifact(store, svg, "svg", page_id="p1")
    deps = [{"kind": "svg", "identity": "p1", "sha256": entry["svg"]["sha256"]}]
    for slot in ("blueprint", "svg_preview", "ppt_preview"):
        entry[slot] = artifact(store, images[0], slot, page_id="p1")
    updated["outputs"] = {
        "pptx": artifact(store, compiled.pptx_path, "pptx", dependencies=deps),
        "trace": artifact(store, compiled.manifest_path, "object_trace", dependencies=deps),
        "render_report": artifact(store, report_path, "render_report", dependencies=deps),
    }
    for kind in ("content", "blueprint_content", "blueprint_fidelity", "conversion", "readability", "privacy"):
        review = {
            "schema_version": "deck_review.v1",
            "review_id": "synthetic-" + kind,
            "kind": kind,
            "status": "pass",
            "subjects": [entry["page"], updated["outputs"]["pptx"]],
            "dependencies": [{"kind": "content", "identity": "page:p1", "sha256": entry["page"]["sha256"]}],
            "reviewer": {"type": "host_self", "id": "synthetic-test", "execution_ref": None, "independence_confirmed": False},
            "observations": ["Explicit synthetic gate input; not an actual Host or professional review."],
            "findings": [],
            "created_at": "2026-09-30T00:00:00Z",
            "replaces": None,
        }
        updated["reviews"].append(store.put_json_object(review))
    updated["policy"]["user_stop"] = True
    store.commit_change(base_revision=doc["revision_id"], document=updated, operation_id="synthetic-production")
    assert editing.check_summary(store, store.load_document())["status"] == "pass"
    return project, store, old


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--font", default="Arial", help="Explicit installed font; CI Linux uses DejaVu Sans")
    args = parser.parse_args()
    root = args.out.resolve()
    root.mkdir(parents=True, exist_ok=False)
    project, store, old = fixture(root, args.font)
    current = store.current_revision_id()
    info = ui_journal.project_info(project)
    server = WorkbenchServer(project)
    url = server.start()
    checks = {}
    errors = []
    downloads = []

    def route(revision):
        return (
            url
            + "v2/#"
            + urlencode({"project": info["project_identity"], "surface": "runs", "layer": "content", "revision": revision, "zoom": 1})
        )

    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            context = browser.new_context(viewport={"width": 1440, "height": 1050}, accept_downloads=True)
            page = context.new_page()
            page.on("pageerror", lambda err: errors.append(str(err)))

            def open_version(revision):
                page.goto(route(revision))
                expect(page.get_by_role("heading", name="版本与文件", exact=True)).to_be_visible()

            def download_result(label):
                row = page.locator(".export-result").filter(has=page.get_by_role("heading", name=label)).first
                expect(row.get_by_role("link", name="下载 ZIP", exact=True)).to_be_visible()
                with page.expect_download() as event:
                    row.get_by_role("link", name="下载 ZIP", exact=True).click()
                file = root / (str(len(downloads)) + "-" + event.value.suggested_filename)
                event.value.save_as(file)
                digest = hashlib.sha256(file.read_bytes()).hexdigest()
                expect(row).to_contain_text(digest)
                with zipfile.ZipFile(file) as archive:
                    manifest = json.loads(archive.read("manifest.json"))
                    for item in manifest["files"]:
                        assert hashlib.sha256(archive.read(item["path"])).hexdigest() == item["sha256"]
                    if manifest["purpose"] != "engineering":
                        assert not any(n.startswith("project/") for n in archive.namelist())
                        if "deck.pptx" in archive.namelist():
                            with zipfile.ZipFile(io.BytesIO(archive.read("deck.pptx"))) as ppt:
                                assert all(b"PRIVATE-CANARY" not in ppt.read(n) for n in ppt.namelist())
                downloads.append(
                    {"purpose": manifest["purpose"], "revision_id": manifest["revision_id"], "sha256": digest, "filename": file.name}
                )
                return manifest, row

            open_version(current)
            page.get_by_text("查看此版本的交付缺项与制作证据", exact=True).click()
            expect(page.get_by_text("工程检查：记录为通过。输入：按此快照记录。", exact=True)).to_be_visible()
            page.get_by_text("制作与可编辑性", exact=True).click()
            expect(page.get_by_text("专业使用评估：未评估", exact=True)).to_be_visible()
            checks["separate_engineering_and_professional_evidence"] = True
            for purpose, label in [("review", "审阅包"), ("delivery", "正式交付包"), ("engineering", "内部工程包")]:
                page.get_by_role("button", name="生成" + label, exact=True).click()
                manifest, row = download_result(label + " ·")
                assert manifest["purpose"] == purpose and manifest["revision_id"] == current
                checks[purpose + "_browser_download"] = True
            page.get_by_role("heading", name="版本与文件", exact=True).scroll_into_view_if_needed()
            page.screenshot(path=str(root / "desktop-1440.png"), full_page=True)
            page.set_viewport_size({"width": 1180, "height": 1000})
            page.screenshot(path=str(root / "desktop-1180.png"), full_page=True)
            assert page.evaluate("() => document.documentElement.scrollWidth <= innerWidth")
            checks["two_desktop_widths_no_horizontal_overflow"] = True
            page.set_viewport_size({"width": 1440, "height": 1050})
            # Switch through the real history picker, preserving current Document.
            page.get_by_label("阅读历史版本").select_option(old)
            page.get_by_role("button", name="读取所选版本", exact=True).click()
            expect(page.locator(".history-banner")).to_be_visible()
            assert store.current_revision_id() == current
            page.get_by_role("button", name="生成审阅包", exact=True).click()
            manifest, _ = download_result("审阅包 · R " + old[:8])
            assert not any(f["path"] == "deck.pptx" for f in manifest["files"])
            checks["historical_no_ppt_review"] = True
            page.get_by_role("button", name="生成正式交付包", exact=True).click()
            expect(page.get_by_text("所选版本尚不满足正式交付条件，请处理下面的页与检查项。", exact=True)).to_be_visible()
            checks["historical_delivery_keeps_structured_gaps"] = True
            page.get_by_role("button", name="预览恢复此版本", exact=True).click()
            expect(page.get_by_role("heading", name="确认历史恢复影响", exact=True)).to_be_visible()
            page.get_by_role("button", name="取消恢复", exact=True).click()
            assert store.current_revision_id() == current
            checks["restore_cancel_no_revision"] = True
            # Current advances after a plan; both versions survive rejection.
            page.get_by_role("button", name="预览恢复此版本", exact=True).click()
            expect(page.get_by_role("heading", name="确认历史恢复影响", exact=True)).to_be_visible()
            newer = advance(store)["revision_id"]
            page.get_by_role("button", name="确认恢复并创建新版本", exact=True).click()
            expect(page.get_by_role("heading", name="本次未提交，输入已保留", exact=True)).to_be_visible()
            assert store.current_revision_id() == newer
            checks["restore_conflict_preserves_both_versions"] = True
            page.get_by_role("button", name="关闭", exact=True).click()
            # Server commits, but browser loses the response. Reload and verify once.
            open_version(old)

            def lose_restore(request):
                request.fetch()
                request.abort()

            page.route("**/api/history/commit-restore", lose_restore, times=1)
            page.get_by_role("button", name="预览恢复此版本", exact=True).click()
            page.get_by_role("button", name="确认恢复并创建新版本", exact=True).click()
            expect(page.get_by_text("结果待核实；原请求已保留，可关闭后在顶部核实。", exact=True)).to_be_visible()
            committed = store.current_revision_id()
            assert committed != newer and store.load_document()["policy"]["user_stop"]
            page.reload()
            expect(page.get_by_role("button", name="核实保存结果", exact=True)).to_be_visible()
            page.get_by_role("button", name="核实保存结果", exact=True).click()
            expect(page.get_by_role("button", name="查看恢复后的版本", exact=True)).to_be_visible()
            assert store.current_revision_id() == committed
            checks["unknown_restore_recovers_across_reload_without_repeat"] = True
            open_version(committed)

            def lose_export(request):
                request.fetch()
                advance(store, "background after frozen export")
                request.abort()

            page.route("**/api/exports", lose_export, times=1)
            page.get_by_role("button", name="生成审阅包", exact=True).click()
            expect(page.get_by_text("生成结果待核实；原版本、用途与编号已保留，勿改成当前新版本重试。", exact=True)).to_be_visible()
            latest = store.current_revision_id()
            assert latest != committed
            page.reload()
            page.get_by_role("button", name="核实这个文件包", exact=True).first.click()
            manifest, row = download_result("审阅包 · R " + committed[:8])
            assert manifest["revision_id"] == committed
            assert "revision=" + committed in page.url and store.current_revision_id() == latest
            checks["unknown_export_recovers_frozen_id_after_reload"] = True
            checks["background_result_does_not_relabel_download"] = True
            # A real truncated HTTP response, followed by the same frozen URL.
            from deck_master.web import WorkbenchHandler

            send_bytes = WorkbenchHandler._send_bytes
            cut = {"sent": False, "requests": 0}

            def truncate_once(handler, body, media, **kwargs):
                if kwargs.get("download_name"):
                    cut["requests"] += 1
                if kwargs.get("download_name") and not cut["sent"]:
                    cut["sent"] = True
                    handler.send_response(200)
                    handler.send_header("Content-Type", media)
                    handler.send_header("Content-Length", str(len(body)))
                    handler.send_header("Content-Disposition", 'attachment; filename="interrupted.zip"')
                    handler.end_headers()
                    handler.wfile.write(body[:100])
                    handler.close_connection = True
                    return
                return send_bytes(handler, body, media, **kwargs)

            WorkbenchHandler._send_bytes = truncate_once
            try:
                with page.expect_download() as event:
                    row.get_by_role("link", name="下载 ZIP", exact=True).click()
                failure = event.value.failure()
                if failure is None:
                    probe = root / "interrupted-transfer.zip"
                    event.value.save_as(probe)
                    cut["first_download_sha256"] = hashlib.sha256(probe.read_bytes()).hexdigest()
                cut["browser_failure"] = failure
            finally:
                WorkbenchHandler._send_bytes = send_bytes
            retried, _ = download_result("审阅包 · R " + committed[:8])
            assert retried == manifest and cut["sent"]
            checks["truncated_download_retries_same_manifest"] = True
            (root / "interrupted-transfer.json").write_text(json.dumps(cut, indent=2) + "\n")
            assert not errors, errors
            browser.close()
        result = {
            "checks": checks,
            "downloads": downloads,
            "page_errors": errors,
            "model_calls": 0,
            "fixture": "real compiler/rendered PPT, synthetic quality records; no professional acceptance",
        }
        (root / "checks.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(result, ensure_ascii=False))
    finally:
        server.stop()


if __name__ == "__main__":
    main()
